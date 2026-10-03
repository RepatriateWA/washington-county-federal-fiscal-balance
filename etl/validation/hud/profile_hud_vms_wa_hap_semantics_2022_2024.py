"""
ETL-174B-5Y-2B
HUD VMS Washington HAP semantic profile, CY2022-CY2024.

PURPOSE
-------
Test the semantics and internal behavior of HUD VMS monthly HAP observations
for Washington PHAs before using VMS as an accounting control.

READ ONLY.

THIS SCRIPT:
    * reads official preserved HUD VMS quarterly workbooks,
    * reads Vms1 only,
    * selects PHA codes beginning with WA,
    * extracts monthly observations,
    * tests PHA-month uniqueness,
    * tests month coverage,
    * profiles HAP Total,
    * profiles selected HAP-related diagnostic fields,
    * compares HAP Total with selected component fields,
    * writes TXT/CSV/JSON validation artifacts.

THIS SCRIPT DOES NOT:
    * write to PostgreSQL,
    * alter source workbooks,
    * reconcile to USAspending,
    * classify AL 14.871 as additive or overlapping,
    * allocate dollars to counties,
    * infer household geography,
    * make a ledger decision.

Checkpoint:
    ETL-174B-5Y-2B
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# =============================================================================
# PATHS / CONTROLS
# =============================================================================

REPO_ROOT = Path(__file__).resolve().parents[3]

RAW_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "quarterly"
)

OUTPUT_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "validation"
    / "profile_wa_hap_semantics_2022_2024"
)

EXPECTED_YEARS = (2022, 2023, 2024)
EXPECTED_WORKBOOK_COUNT = 12
EXPECTED_MONTH_COUNT = 36

TARGET_SHEET = "Vms1"
WA_PREFIX = "WA"


# =============================================================================
# HEADER NORMALIZATION
# =============================================================================

def clean_text(value: Any) -> str:
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


def normalize_header(value: Any) -> str:
    text = clean_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


# =============================================================================
# TARGET FIELDS
# =============================================================================

# Core fields that must exist in every Vms1 workbook.

CORE_FIELDS = (
    "pha_code",
    "pha_name",
    "submission_month",
    "total_vouchers",
    "hap_total",
)


# Diagnostic fields discovered in ETL-174B-5Y-2A.
#
# These are NOT asserted to be additive components of HAP Total.
# Their purpose here is semantic inspection only.

DIAGNOSTIC_FIELDS = (
    "portable_voucher_paid_hap",
    "total_hap_for_portable_units_administered",
    "mainstream_hap",
    "hcv_all_voucher_hap_expenses_after_the_first_of_the_month",
    "project_based_vouchers_total_hap_expenses",
    "veterans_affair_supported_housing_vash_hap",
    "homeownership_hap",
    "tenant_protection_hap",
    "foster_youth_to_independence_hap",
    "housing_stability_vouchers_hap_expenses",
)


# These fields are selected for a deliberately NON-DEFINITIVE component sum.
#
# IMPORTANT:
# The resulting sum is a diagnostic only.
#
# We do NOT assume:
#
#     HAP Total = sum(DIAGNOSTIC_COMPONENT_FIELDS)
#
# Some fields may overlap, be subsets, or have different accounting semantics.

DIAGNOSTIC_COMPONENT_FIELDS = (
    "portable_voucher_paid_hap",
    "mainstream_hap",
    "veterans_affair_supported_housing_vash_hap",
    "homeownership_hap",
    "tenant_protection_hap",
    "foster_youth_to_independence_hap",
    "housing_stability_vouchers_hap_expenses",
)


# =============================================================================
# BASIC HELPERS
# =============================================================================

def numeric(value: Any) -> float | None:
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)

    text = clean_text(value)

    if not text:
        return None

    text = (
        text
        .replace("$", "")
        .replace(",", "")
        .strip()
    )

    try:
        return float(text)
    except ValueError:
        return None


def parse_month(value: Any) -> date | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = clean_text(value)

    if not text:
        return None

    candidates = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m/%d/%y",
    )

    for fmt in candidates:
        try:
            return datetime.strptime(
                text,
                fmt,
            ).date()
        except ValueError:
            pass

    return None


def month_key(value: date) -> str:
    return value.strftime("%Y-%m")


def money(value: float | None) -> str:
    if value is None:
        return "NULL"

    return f"${value:,.2f}"


def number(value: float | None) -> str:
    if value is None:
        return "NULL"

    return f"{value:,.2f}"


# =============================================================================
# WORKBOOK DISCOVERY
# =============================================================================

def discover_workbooks() -> list[Path]:
    paths: list[Path] = []

    for year in EXPECTED_YEARS:
        year_dir = RAW_ROOT / f"CY{year}"

        if not year_dir.exists():
            raise RuntimeError(
                f"Missing expected directory: {year_dir}"
            )

        paths.extend(
            sorted(year_dir.glob("*.xlsx"))
        )

    if len(paths) != EXPECTED_WORKBOOK_COUNT:
        raise RuntimeError(
            f"Expected {EXPECTED_WORKBOOK_COUNT} workbooks; "
            f"found {len(paths)}."
        )

    return paths


# =============================================================================
# WORKBOOK EXTRACTION
# =============================================================================

def extract_wa_rows(
    path: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:

    wb = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:
        if TARGET_SHEET not in wb.sheetnames:
            raise RuntimeError(
                f"{path.name}: missing sheet {TARGET_SHEET}"
            )

        ws = wb[TARGET_SHEET]

        rows = ws.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            raise RuntimeError(
                f"{path.name}: empty {TARGET_SHEET}"
            )

        normalized_headers = [
            normalize_header(value)
            for value in header_row
        ]

        header_positions: dict[str, int] = {}

        duplicate_headers = []

        for index, header in enumerate(
            normalized_headers
        ):
            if not header:
                continue

            if header in header_positions:
                duplicate_headers.append(header)
                continue

            header_positions[header] = index

        missing_core = [
            field
            for field in CORE_FIELDS
            if field not in header_positions
        ]

        if missing_core:
            raise RuntimeError(
                f"{path.name}: missing core fields: "
                f"{', '.join(missing_core)}"
            )

        present_diagnostics = [
            field
            for field in DIAGNOSTIC_FIELDS
            if field in header_positions
        ]

        missing_diagnostics = [
            field
            for field in DIAGNOSTIC_FIELDS
            if field not in header_positions
        ]

        extracted: list[dict[str, Any]] = []

        source_rows = 0
        wa_rows = 0
        blank_rows = 0
        invalid_month_rows = 0

        for worksheet_row_number, row in enumerate(
            rows,
            start=2,
        ):
            source_rows += 1

            pha_index = header_positions["pha_code"]

            if pha_index >= len(row):
                blank_rows += 1
                continue

            pha_code = clean_text(row[pha_index])

            if not pha_code:
                blank_rows += 1
                continue

            if not pha_code.upper().startswith(WA_PREFIX):
                continue

            wa_rows += 1

            submission_value = row[
                header_positions["submission_month"]
            ]

            submission_month = parse_month(
                submission_value
            )

            if submission_month is None:
                invalid_month_rows += 1

            record: dict[str, Any] = {
                "source_workbook": path.name,
                "source_row_number": worksheet_row_number,
                "pha_code": pha_code,
                "pha_name": clean_text(
                    row[
                        header_positions["pha_name"]
                    ]
                ),
                "submission_month": (
                    submission_month.isoformat()
                    if submission_month
                    else ""
                ),
                "submission_month_key": (
                    month_key(submission_month)
                    if submission_month
                    else ""
                ),
                "total_vouchers": numeric(
                    row[
                        header_positions["total_vouchers"]
                    ]
                ),
                "hap_total": numeric(
                    row[
                        header_positions["hap_total"]
                    ]
                ),
            }

            for field in DIAGNOSTIC_FIELDS:
                if field not in header_positions:
                    record[field] = None
                    continue

                position = header_positions[field]

                if position >= len(row):
                    record[field] = None
                else:
                    record[field] = numeric(
                        row[position]
                    )

            extracted.append(record)

        metadata = {
            "source_workbook": path.name,
            "source_rows_after_header": source_rows,
            "wa_rows": wa_rows,
            "blank_or_unkeyed_rows": blank_rows,
            "invalid_month_rows": invalid_month_rows,
            "duplicate_normalized_headers":
                sorted(set(duplicate_headers)),
            "missing_core_fields": missing_core,
            "present_diagnostic_fields":
                present_diagnostics,
            "missing_diagnostic_fields":
                missing_diagnostics,
        }

        return extracted, metadata

    finally:
        wb.close()


# =============================================================================
# VALIDATION
# =============================================================================

def analyze(
    records: list[dict[str, Any]],
) -> dict[str, Any]:

    result: dict[str, Any] = {}

    # -------------------------------------------------------------------------
    # Month coverage
    # -------------------------------------------------------------------------

    months = sorted(
        {
            record["submission_month_key"]
            for record in records
            if record["submission_month_key"]
        }
    )

    result["distinct_months"] = months
    result["distinct_month_count"] = len(months)

    months_by_year: dict[int, list[str]] = {}

    for year in EXPECTED_YEARS:
        prefix = f"{year}-"

        months_by_year[year] = [
            month
            for month in months
            if month.startswith(prefix)
        ]

    result["months_by_year"] = months_by_year

    # -------------------------------------------------------------------------
    # PHA coverage
    # -------------------------------------------------------------------------

    pha_codes = sorted(
        {
            record["pha_code"]
            for record in records
        }
    )

    result["pha_codes"] = pha_codes
    result["distinct_pha_count"] = len(pha_codes)

    pha_names: dict[str, set[str]] = defaultdict(set)

    for record in records:
        if record["pha_name"]:
            pha_names[
                record["pha_code"]
            ].add(
                record["pha_name"]
            )

    result["pha_name_variants"] = {
        code: sorted(names)
        for code, names in sorted(
            pha_names.items()
        )
    }

    # -------------------------------------------------------------------------
    # PHA-month uniqueness
    # -------------------------------------------------------------------------

    pha_month_counter = Counter(
        (
            record["pha_code"],
            record["submission_month_key"],
        )
        for record in records
        if record["submission_month_key"]
    )

    duplicate_pha_months = [
        {
            "pha_code": key[0],
            "month": key[1],
            "count": count,
        }
        for key, count in sorted(
            pha_month_counter.items()
        )
        if count > 1
    ]

    result[
        "duplicate_pha_months"
    ] = duplicate_pha_months

    result[
        "duplicate_pha_month_count"
    ] = len(duplicate_pha_months)

    # -------------------------------------------------------------------------
    # Observation counts by month
    # -------------------------------------------------------------------------

    observations_by_month = Counter(
        record["submission_month_key"]
        for record in records
        if record["submission_month_key"]
    )

    result["observations_by_month"] = dict(
        sorted(observations_by_month.items())
    )

    # -------------------------------------------------------------------------
    # Missing HAP Total
    # -------------------------------------------------------------------------

    missing_hap_total = [
        record
        for record in records
        if record["hap_total"] is None
    ]

    result["missing_hap_total_count"] = len(
        missing_hap_total
    )

    # -------------------------------------------------------------------------
    # HAP Total sign / distribution controls
    # -------------------------------------------------------------------------

    hap_values = [
        record["hap_total"]
        for record in records
        if record["hap_total"] is not None
    ]

    result["hap_total_numeric_count"] = len(
        hap_values
    )

    result["hap_total_zero_count"] = sum(
        1
        for value in hap_values
        if value == 0
    )

    result["hap_total_negative_count"] = sum(
        1
        for value in hap_values
        if value < 0
    )

    result["hap_total_positive_count"] = sum(
        1
        for value in hap_values
        if value > 0
    )

    if hap_values:
        result["hap_total_min"] = min(hap_values)
        result["hap_total_max"] = max(hap_values)
    else:
        result["hap_total_min"] = None
        result["hap_total_max"] = None

    # -------------------------------------------------------------------------
    # Monthly Washington HAP Total
    #
    # This is calculated as a DIAGNOSTIC.
    # It is NOT yet an accounting control.
    # -------------------------------------------------------------------------

    monthly_hap: dict[str, float] = defaultdict(float)

    for record in records:
        month = record["submission_month_key"]
        value = record["hap_total"]

        if month and value is not None:
            monthly_hap[month] += value

    result["monthly_wa_hap_total"] = dict(
        sorted(monthly_hap.items())
    )

    # -------------------------------------------------------------------------
    # Annual diagnostic HAP totals
    # -------------------------------------------------------------------------

    annual_hap: dict[int, float] = defaultdict(float)

    for month, value in monthly_hap.items():
        year = int(month[:4])
        annual_hap[year] += value

    result["annual_wa_hap_total"] = dict(
        sorted(annual_hap.items())
    )

    # -------------------------------------------------------------------------
    # Month-to-month behavior
    #
    # This helps distinguish monthly-flow behavior from cumulative/YTD behavior.
    # No formal semantic conclusion is made automatically.
    # -------------------------------------------------------------------------

    monthly_changes = []

    sorted_months = sorted(monthly_hap)

    for previous, current in zip(
        sorted_months,
        sorted_months[1:],
    ):
        previous_value = monthly_hap[previous]
        current_value = monthly_hap[current]

        difference = (
            current_value - previous_value
        )

        pct_change = None

        if previous_value != 0:
            pct_change = (
                difference
                / previous_value
            )

        monthly_changes.append(
            {
                "previous_month": previous,
                "current_month": current,
                "previous_value": previous_value,
                "current_value": current_value,
                "difference": difference,
                "pct_change": pct_change,
            }
        )

    result["monthly_changes"] = monthly_changes

    # -------------------------------------------------------------------------
    # Diagnostic component sum
    #
    # DO NOT interpret this as a reconciliation identity.
    # -------------------------------------------------------------------------

    component_diagnostics = []

    for record in records:
        available_values = []

        for field in DIAGNOSTIC_COMPONENT_FIELDS:
            value = record.get(field)

            if value is not None:
                available_values.append(value)

        diagnostic_sum = (
            sum(available_values)
            if available_values
            else None
        )

        hap_total = record["hap_total"]

        difference = None
        ratio = None

        if (
            diagnostic_sum is not None
            and hap_total is not None
        ):
            difference = (
                hap_total - diagnostic_sum
            )

            if hap_total != 0:
                ratio = (
                    diagnostic_sum
                    / hap_total
                )

        component_diagnostics.append(
            {
                "pha_code":
                    record["pha_code"],
                "month":
                    record[
                        "submission_month_key"
                    ],
                "hap_total":
                    hap_total,
                "diagnostic_component_sum":
                    diagnostic_sum,
                "hap_minus_diagnostic_sum":
                    difference,
                "diagnostic_sum_over_hap_total":
                    ratio,
            }
        )

    result[
        "component_diagnostics"
    ] = component_diagnostics

    # -------------------------------------------------------------------------
    # Field availability
    # -------------------------------------------------------------------------

    availability = {}

    for field in DIAGNOSTIC_FIELDS:
        values = [
            record.get(field)
            for record in records
        ]

        availability[field] = {
            "numeric_count": sum(
                value is not None
                for value in values
            ),
            "positive_count": sum(
                value is not None
                and value > 0
                for value in values
            ),
            "negative_count": sum(
                value is not None
                and value < 0
                for value in values
            ),
            "zero_count": sum(
                value == 0
                for value in values
                if value is not None
            ),
        }

    result[
        "diagnostic_field_availability"
    ] = availability

    return result


# =============================================================================
# OUTPUT
# =============================================================================

def write_csv(
    path: Path,
    records: list[dict[str, Any]],
) -> None:

    fields = [
        "source_workbook",
        "source_row_number",
        "pha_code",
        "pha_name",
        "submission_month",
        "submission_month_key",
        "total_vouchers",
        "hap_total",
        *DIAGNOSTIC_FIELDS,
    ]

    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=fields,
        )

        writer.writeheader()

        for record in records:
            writer.writerow(
                {
                    field: record.get(field)
                    for field in fields
                }
            )


def render_report(
    records: list[dict[str, Any]],
    workbook_metadata: list[dict[str, Any]],
    analysis: dict[str, Any],
) -> str:

    lines: list[str] = []

    def emit(text: str = "") -> None:
        lines.append(text)

    emit("=" * 100)
    emit("HUD VMS WASHINGTON HAP SEMANTIC PROFILE")
    emit("ETL-174B-5Y-2B")
    emit("=" * 100)
    emit()
    emit(
        "READ-ONLY diagnostic profile. "
        "Calculated totals below are exploratory controls, "
        "not ledger values."
    )
    emit()

    # -------------------------------------------------------------------------
    # Control
    # -------------------------------------------------------------------------

    emit("CONTROL SUMMARY")
    emit("-" * 100)
    emit(
        f"Workbooks expected                   : "
        f"{EXPECTED_WORKBOOK_COUNT}"
    )
    emit(
        f"Workbooks processed                  : "
        f"{len(workbook_metadata)}"
    )
    emit(
        f"Washington Vms1 observations         : "
        f"{len(records):,}"
    )
    emit(
        f"Distinct Washington PHA codes        : "
        f"{analysis['distinct_pha_count']:,}"
    )
    emit(
        f"Distinct submission months           : "
        f"{analysis['distinct_month_count']:,}"
    )
    emit(
        f"Expected submission months           : "
        f"{EXPECTED_MONTH_COUNT:,}"
    )
    emit(
        f"Duplicate PHA-month keys             : "
        f"{analysis['duplicate_pha_month_count']:,}"
    )
    emit(
        f"Missing HAP Total observations       : "
        f"{analysis['missing_hap_total_count']:,}"
    )
    emit()

    # -------------------------------------------------------------------------
    # Workbook extraction
    # -------------------------------------------------------------------------

    emit("=" * 100)
    emit("WORKBOOK EXTRACTION CONTROL")
    emit("=" * 100)

    for item in workbook_metadata:
        emit()
        emit(item["source_workbook"])
        emit(
            f"    source rows after header : "
            f"{item['source_rows_after_header']:,}"
        )
        emit(
            f"    Washington rows          : "
            f"{item['wa_rows']:,}"
        )
        emit(
            f"    invalid month rows        : "
            f"{item['invalid_month_rows']:,}"
        )

        missing = item[
            "missing_diagnostic_fields"
        ]

        if missing:
            emit(
                "    missing diagnostics      : "
                + ", ".join(missing)
            )
        else:
            emit(
                "    missing diagnostics      : NONE"
            )

    # -------------------------------------------------------------------------
    # Month coverage
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("MONTH COVERAGE")
    emit("=" * 100)
    emit()

    for year in EXPECTED_YEARS:
        months = analysis[
            "months_by_year"
        ][year]

        emit(
            f"CY{year}: "
            f"{len(months)} months"
        )

        for month in months:
            count = analysis[
                "observations_by_month"
            ].get(
                month,
                0,
            )

            emit(
                f"    {month}  "
                f"PHA observations={count:,}"
            )

    # -------------------------------------------------------------------------
    # PHA inventory
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("WASHINGTON PHA INVENTORY")
    emit("=" * 100)
    emit()

    for pha_code in analysis["pha_codes"]:
        names = analysis[
            "pha_name_variants"
        ].get(
            pha_code,
            [],
        )

        emit(
            f"{pha_code:<8} "
            f"{' | '.join(names)}"
        )

    # -------------------------------------------------------------------------
    # Duplicate keys
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("PHA-MONTH UNIQUENESS")
    emit("=" * 100)
    emit()

    if not analysis["duplicate_pha_months"]:
        emit(
            "PASS: no duplicate PHA Code x Submission Month keys."
        )
    else:
        emit(
            "REVIEW REQUIRED: duplicate PHA-month keys detected."
        )

        for item in analysis[
            "duplicate_pha_months"
        ]:
            emit(
                f"    {item['pha_code']} "
                f"{item['month']} "
                f"count={item['count']}"
            )

    # -------------------------------------------------------------------------
    # HAP Total semantics
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("HAP TOTAL FIELD PROFILE")
    emit("=" * 100)
    emit()

    emit(
        f"Numeric observations : "
        f"{analysis['hap_total_numeric_count']:,}"
    )
    emit(
        f"Positive             : "
        f"{analysis['hap_total_positive_count']:,}"
    )
    emit(
        f"Zero                 : "
        f"{analysis['hap_total_zero_count']:,}"
    )
    emit(
        f"Negative             : "
        f"{analysis['hap_total_negative_count']:,}"
    )
    emit(
        f"Minimum              : "
        f"{money(analysis['hap_total_min'])}"
    )
    emit(
        f"Maximum              : "
        f"{money(analysis['hap_total_max'])}"
    )

    # -------------------------------------------------------------------------
    # Monthly totals
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("WASHINGTON MONTHLY HAP TOTAL — DIAGNOSTIC ONLY")
    emit("=" * 100)
    emit()

    for month, value in analysis[
        "monthly_wa_hap_total"
    ].items():
        emit(
            f"{month}    {money(value)}"
        )

    # -------------------------------------------------------------------------
    # Annual diagnostic totals
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("WASHINGTON ANNUAL HAP TOTAL — DIAGNOSTIC ONLY")
    emit("=" * 100)
    emit()

    for year, value in analysis[
        "annual_wa_hap_total"
    ].items():
        emit(
            f"CY{year}    {money(value)}"
        )

    # -------------------------------------------------------------------------
    # Month-to-month behavior
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("MONTH-TO-MONTH HAP BEHAVIOR")
    emit("=" * 100)
    emit()

    for item in analysis["monthly_changes"]:
        pct = item["pct_change"]

        pct_text = (
            "NULL"
            if pct is None
            else f"{pct:+.4%}"
        )

        emit(
            f"{item['previous_month']} -> "
            f"{item['current_month']}    "
            f"{money(item['previous_value'])} -> "
            f"{money(item['current_value'])}    "
            f"change={money(item['difference'])}    "
            f"pct={pct_text}"
        )

    # -------------------------------------------------------------------------
    # Diagnostic field availability
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("DIAGNOSTIC FIELD AVAILABILITY")
    emit("=" * 100)
    emit()

    for field, stats in analysis[
        "diagnostic_field_availability"
    ].items():

        emit(field)
        emit(
            f"    numeric  : "
            f"{stats['numeric_count']:,}"
        )
        emit(
            f"    positive : "
            f"{stats['positive_count']:,}"
        )
        emit(
            f"    zero     : "
            f"{stats['zero_count']:,}"
        )
        emit(
            f"    negative : "
            f"{stats['negative_count']:,}"
        )

    # -------------------------------------------------------------------------
    # Methodological stop
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("METHODOLOGICAL STOP")
    emit("=" * 100)
    emit()
    emit(
        "HAP Total has NOT been certified as the WCFFBP accounting control."
    )
    emit(
        "Annual diagnostic sums are NOT ledger values."
    )
    emit(
        "Diagnostic HAP fields have NOT been assumed additive."
    )
    emit(
        "Vms2 / Emergency Housing Voucher values have NOT been added."
    )
    emit(
        "No USAspending reconciliation has been performed."
    )
    emit(
        "No county allocation has been performed."
    )
    emit(
        "No database writes occurred."
    )

    return "\n".join(lines) + "\n"


# =============================================================================
# MAIN
# =============================================================================

def main() -> int:

    print("=" * 88)
    print("HUD VMS WASHINGTON HAP SEMANTIC PROFILE")
    print("ETL-174B-5Y-2B")
    print("=" * 88)
    print()
    print(f"Input : {RAW_ROOT}")
    print(f"Output: {OUTPUT_ROOT}")
    print()
    print("Scope : Washington PHAs / Vms1 / CY2022-CY2024")
    print("Mode  : READ ONLY")
    print("DB    : NO DATABASE WRITES")
    print()

    try:
        paths = discover_workbooks()
    except RuntimeError as exc:
        print(f"ERROR: {exc}")
        return 1

    OUTPUT_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    all_records: list[dict[str, Any]] = []
    workbook_metadata: list[dict[str, Any]] = []

    for index, path in enumerate(
        paths,
        start=1,
    ):
        print(
            f"[{index:02}/{len(paths):02}] "
            f"{path.name}",
            flush=True,
        )

        try:
            records, metadata = extract_wa_rows(
                path
            )
        except Exception as exc:
            print(
                f"         ERROR: {exc}",
                flush=True,
            )
            return 1

        all_records.extend(records)
        workbook_metadata.append(metadata)

        print(
            f"         WA rows={len(records):,}",
            flush=True,
        )

    analysis = analyze(all_records)

    timestamp = datetime.now(
        timezone.utc
    ).strftime("%Y%m%dT%H%M%SZ")

    report_path = (
        OUTPUT_ROOT
        / (
            "profile_hud_vms_wa_hap_semantics_"
            f"2022_2024_{timestamp}.txt"
        )
    )

    csv_path = (
        OUTPUT_ROOT
        / (
            "hud_vms_wa_monthly_hap_diagnostic_"
            f"2022_2024_{timestamp}.csv"
        )
    )

    json_path = (
        OUTPUT_ROOT
        / (
            "profile_hud_vms_wa_hap_semantics_"
            f"2022_2024_{timestamp}.json"
        )
    )

    write_csv(
        csv_path,
        all_records,
    )

    report_path.write_text(
        render_report(
            all_records,
            workbook_metadata,
            analysis,
        ),
        encoding="utf-8",
    )

    json_payload = {
        "checkpoint":
            "ETL-174B-5Y-2B",
        "generated_utc":
            datetime.now(
                timezone.utc
            ).isoformat(),
        "scope":
            "Washington PHAs / Vms1 / CY2022-CY2024",
        "mode":
            "READ_ONLY_DIAGNOSTIC",
        "workbook_metadata":
            workbook_metadata,
        "analysis":
            analysis,
    }

    json_path.write_text(
        json.dumps(
            json_payload,
            indent=2,
            ensure_ascii=False,
            default=str,
        )
        + "\n",
        encoding="utf-8",
    )

    # -------------------------------------------------------------------------
    # Console control
    # -------------------------------------------------------------------------

    print()
    print("=" * 88)
    print("SEMANTIC PROFILE CONTROL")
    print("=" * 88)

    print(
        f"Workbooks processed        : "
        f"{len(workbook_metadata)}"
    )

    print(
        f"Washington observations    : "
        f"{len(all_records):,}"
    )

    print(
        f"Distinct WA PHAs           : "
        f"{analysis['distinct_pha_count']:,}"
    )

    print(
        f"Distinct months            : "
        f"{analysis['distinct_month_count']:,}"
    )

    print(
        f"Duplicate PHA-month keys   : "
        f"{analysis['duplicate_pha_month_count']:,}"
    )

    print(
        f"Missing HAP Total          : "
        f"{analysis['missing_hap_total_count']:,}"
    )

    print()

    for year in EXPECTED_YEARS:
        value = analysis[
            "annual_wa_hap_total"
        ].get(
            year
        )

        print(
            f"CY{year} diagnostic HAP    : "
            f"{money(value)}"
        )

    print()
    print(
        "Report : "
        f"{report_path.relative_to(REPO_ROOT)}"
    )
    print(
        "CSV    : "
        f"{csv_path.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON   : "
        f"{json_path.relative_to(REPO_ROOT)}"
    )
    print()

    structural_pass = (
        len(workbook_metadata)
        == EXPECTED_WORKBOOK_COUNT
        and analysis["distinct_month_count"]
        == EXPECTED_MONTH_COUNT
        and analysis[
            "duplicate_pha_month_count"
        ]
        == 0
        and analysis[
            "missing_hap_total_count"
        ]
        == 0
    )

    if structural_pass:
        print(
            "RESULT: PASS — STRUCTURAL / SEMANTIC "
            "DIAGNOSTIC CONTROLS"
        )
    else:
        print(
            "RESULT: REVIEW REQUIRED"
        )

    print()
    print(
        "No accounting-source certification was made."
    )
    print(
        "No USAspending reconciliation was performed."
    )
    print(
        "No county allocation or ledger decision was made."
    )

    return 0 if structural_pass else 2


if __name__ == "__main__":
    sys.exit(main())