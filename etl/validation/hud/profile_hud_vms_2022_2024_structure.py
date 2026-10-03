"""
ETL-174B-5Y-2A v0.2
Optimized structural profile of official HUD VMS quarterly workbooks,
CY2022-CY2024.

READ ONLY.

Optimization from v0.1:
    Each worksheet is traversed once for column profiling rather than once
    per column.

NO:
    financial aggregation,
    Washington filtering,
    USAspending reconciliation,
    accounting classification,
    database writes,
    ledger writes.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[3]

RAW_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "quarterly"
)

VALIDATION_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "validation"
    / "profile_2022_2024_structure"
)

EXPECTED_YEARS = (2022, 2023, 2024)
EXPECTED_WORKBOOK_COUNT = 12

MAX_HEADER_SCAN_ROWS = 25
MAX_SAMPLE_VALUES = 5
MAX_VALUE_LENGTH = 120


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def clean_text(value: Any) -> str:
    if value is None:
        return ""

    text = re.sub(r"\s+", " ", str(value)).strip()

    if len(text) > MAX_VALUE_LENGTH:
        text = text[: MAX_VALUE_LENGTH - 3] + "..."

    return text


def normalized_header(value: Any) -> str:
    text = clean_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def value_kind(value: Any) -> str:
    if value is None:
        return "NULL"

    if isinstance(value, bool):
        return "BOOLEAN"

    if isinstance(value, (int, float)):
        return "NUMERIC"

    if hasattr(value, "year") and hasattr(value, "month"):
        return "DATE_OR_DATETIME"

    return "TEXT"


def workbook_identity(path: Path) -> tuple[int | None, int | None]:
    year = None
    quarter = None

    match = re.search(r"CY(20\d{2})", str(path.parent), re.I)

    if match:
        year = int(match.group(1))

    name = path.name.lower().replace("_", "-")

    patterns = (
        (1, ("jan-mar",)),
        (2, ("apr-jun",)),
        (3, ("jul-sep", "jul-sept")),
        (4, ("oct-dec",)),
    )

    for candidate, terms in patterns:
        if any(term in name for term in terms):
            quarter = candidate
            break

    return year, quarter


def choose_header_from_rows(
    rows: list[tuple[Any, ...]],
) -> int | None:
    """
    Return one-based worksheet row number.

    Structural heuristic only.
    """
    candidates = []

    for row_number, row in enumerate(rows, start=1):
        populated = [
            value
            for value in row
            if clean_text(value)
        ]

        textual = [
            clean_text(value)
            for value in populated
            if isinstance(value, str)
        ]

        if len(populated) < 3 or len(textual) < 2:
            continue

        distinct_textual = len(
            {value.lower() for value in textual}
        )

        candidates.append(
            (
                len(textual),
                distinct_textual,
                len(populated),
                -row_number,
                row_number,
            )
        )

    if not candidates:
        return None

    candidates.sort(reverse=True)
    return candidates[0][-1]


def profile_sheet(ws) -> dict[str, Any]:
    """
    Profile worksheet with one principal traversal.

    First MAX_HEADER_SCAN_ROWS are buffered to discover a candidate header.
    Worksheet is then reopened through iter_rows() and traversed once for
    column statistics.
    """

    sheet: dict[str, Any] = {
        "sheet_name": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "candidate_header_row": None,
        "columns": [],
    }

    # ---------------------------------------------------------------
    # Small header discovery scan
    # ---------------------------------------------------------------

    initial_rows = list(
        ws.iter_rows(
            min_row=1,
            max_row=min(MAX_HEADER_SCAN_ROWS, ws.max_row),
            values_only=True,
        )
    )

    header_row = choose_header_from_rows(initial_rows)
    sheet["candidate_header_row"] = header_row

    if header_row is None:
        return sheet

    header_values = initial_rows[header_row - 1]

    active_columns = []

    for column_number, value in enumerate(
        header_values,
        start=1,
    ):
        header_text = clean_text(value)

        if not header_text:
            continue

        active_columns.append(
            {
                "column_number": column_number,
                "header_raw": header_text,
                "header_normalized": normalized_header(value),
                "non_null_count": 0,
                "type_counts": Counter(),
                "sample_values": [],
                "_sample_seen": set(),
            }
        )

    if not active_columns:
        return sheet

    by_position = {
        column["column_number"] - 1: column
        for column in active_columns
    }

    # ---------------------------------------------------------------
    # ONE traversal of data rows
    # ---------------------------------------------------------------

    for row in ws.iter_rows(
        min_row=header_row + 1,
        values_only=True,
    ):
        for zero_based_position, column in by_position.items():

            if zero_based_position >= len(row):
                value = None
            else:
                value = row[zero_based_position]

            kind = value_kind(value)
            column["type_counts"][kind] += 1

            rendered = clean_text(value)

            if not rendered:
                continue

            column["non_null_count"] += 1

            if (
                len(column["sample_values"])
                < MAX_SAMPLE_VALUES
                and rendered not in column["_sample_seen"]
            ):
                column["sample_values"].append(rendered)
                column["_sample_seen"].add(rendered)

    # Convert non-JSON structures.
    for column in active_columns:
        column["type_counts"] = dict(
            column["type_counts"]
        )
        del column["_sample_seen"]

    sheet["columns"] = active_columns

    return sheet


def profile_workbook(path: Path) -> dict[str, Any]:
    year, quarter = workbook_identity(path)

    result = {
        "path": str(path.relative_to(REPO_ROOT)),
        "filename": path.name,
        "calendar_year": year,
        "quarter": quarter,
        "byte_count": path.stat().st_size,
        "sha256": sha256_file(path),
        "sheet_names": [],
        "sheets": [],
    }

    wb = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:
        result["sheet_names"] = list(wb.sheetnames)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            result["sheets"].append(
                profile_sheet(ws)
            )

    finally:
        wb.close()

    return result


def schema_signature(
    sheet: dict[str, Any],
) -> tuple[str, ...]:
    return tuple(
        column["header_normalized"]
        for column in sheet["columns"]
    )


def render_report(
    profiles: list[dict[str, Any]],
) -> str:

    lines: list[str] = []

    def emit(value: str = "") -> None:
        lines.append(value)

    emit("=" * 100)
    emit("HUD VMS STRUCTURAL PROFILE")
    emit("ETL-174B-5Y-2A v0.2")
    emit("=" * 100)
    emit()
    emit(
        "READ-ONLY structural discovery. "
        "No financial values aggregated or interpreted."
    )
    emit()

    emit("CONTROL SUMMARY")
    emit("-" * 100)
    emit(
        f"Expected workbooks               : "
        f"{EXPECTED_WORKBOOK_COUNT}"
    )
    emit(
        f"Workbooks profiled               : "
        f"{len(profiles)}"
    )

    total_sheets = sum(
        len(profile["sheets"])
        for profile in profiles
    )

    emit(
        f"Worksheets profiled              : "
        f"{total_sheets}"
    )

    hashes = {
        profile["sha256"]
        for profile in profiles
    }

    emit(
        f"Distinct workbook hashes         : "
        f"{len(hashes)}"
    )
    emit()

    emit("=" * 100)
    emit("WORKBOOK INVENTORY")
    emit("=" * 100)

    for profile in profiles:
        emit()
        emit(
            f"CY{profile['calendar_year']} "
            f"Q{profile['quarter']} — "
            f"{profile['filename']}"
        )
        emit(
            f"Bytes   : {profile['byte_count']:,}"
        )
        emit(
            f"SHA-256 : {profile['sha256']}"
        )
        emit(
            f"Sheets  : "
            f"{', '.join(profile['sheet_names'])}"
        )

    emit()
    emit("=" * 100)
    emit("DETAILED STRUCTURE")
    emit("=" * 100)

    for profile in profiles:
        emit()
        emit("#" * 100)
        emit(
            f"CY{profile['calendar_year']} "
            f"Q{profile['quarter']} — "
            f"{profile['filename']}"
        )
        emit("#" * 100)

        for sheet in profile["sheets"]:
            emit()
            emit(
                f"SHEET: {sheet['sheet_name']}"
            )
            emit("-" * 100)
            emit(
                f"Dimensions           : "
                f"{sheet['max_row']} rows x "
                f"{sheet['max_column']} columns"
            )
            emit(
                f"Candidate header row : "
                f"{sheet['candidate_header_row']}"
            )
            emit(
                f"Profiled columns     : "
                f"{len(sheet['columns'])}"
            )
            emit()

            for column in sheet["columns"]:
                emit(
                    f"[{column['column_number']:>3}] "
                    f"{column['header_raw']}"
                )
                emit(
                    f"      normalized : "
                    f"{column['header_normalized']}"
                )
                emit(
                    f"      non-null   : "
                    f"{column['non_null_count']:,}"
                )

                type_summary = ", ".join(
                    f"{key}={value:,}"
                    for key, value in sorted(
                        column["type_counts"].items()
                    )
                )

                emit(
                    f"      types      : "
                    f"{type_summary}"
                )

                if column["sample_values"]:
                    emit("      samples    :")

                    for sample in column["sample_values"]:
                        emit(
                            f"          {sample}"
                        )

    # ---------------------------------------------------------------
    # Schema signatures
    # ---------------------------------------------------------------

    groups: dict[
        tuple[str, ...],
        list[str],
    ] = defaultdict(list)

    for profile in profiles:
        for sheet in profile["sheets"]:

            if not sheet["columns"]:
                continue

            signature = schema_signature(sheet)

            groups[signature].append(
                f"CY{profile['calendar_year']} "
                f"Q{profile['quarter']} / "
                f"{sheet['sheet_name']}"
            )

    emit()
    emit("=" * 100)
    emit("SCHEMA SIGNATURE COMPARISON")
    emit("=" * 100)
    emit()
    emit(
        f"Distinct structured schemas : "
        f"{len(groups)}"
    )
    emit()

    for number, (signature, labels) in enumerate(
        groups.items(),
        start=1,
    ):
        emit(
            f"SCHEMA {number}"
        )
        emit("-" * 100)

        for label in labels:
            emit(
                f"    {label}"
            )

        emit()
        emit(
            f"Column count: {len(signature)}"
        )

        for header in signature:
            emit(
                f"    {header}"
            )

        emit()

    # ---------------------------------------------------------------
    # Header vocabulary
    # ---------------------------------------------------------------

    occurrences: dict[
        str,
        list[str],
    ] = defaultdict(list)

    raw_labels: dict[
        str,
        set[str],
    ] = defaultdict(set)

    for profile in profiles:
        for sheet in profile["sheets"]:

            location = (
                f"CY{profile['calendar_year']} "
                f"Q{profile['quarter']} / "
                f"{sheet['sheet_name']}"
            )

            for column in sheet["columns"]:
                normalized = column[
                    "header_normalized"
                ]

                occurrences[normalized].append(
                    location
                )

                raw_labels[normalized].add(
                    column["header_raw"]
                )

    emit("=" * 100)
    emit("CROSS-WORKBOOK HEADER VOCABULARY")
    emit("=" * 100)
    emit()

    for normalized in sorted(occurrences):
        emit(normalized)
        emit(
            f"    occurrences : "
            f"{len(occurrences[normalized])}"
        )
        emit(
            f"    raw labels  : "
            f"{' | '.join(sorted(raw_labels[normalized]))}"
        )

    # ---------------------------------------------------------------
    # Semantic header discovery only
    # ---------------------------------------------------------------

    terms = (
        "pha",
        "housing",
        "agency",
        "state",
        "hap",
        "expense",
        "expenditure",
        "payment",
        "lease",
        "voucher",
        "mainstream",
        "emergency",
        "ehv",
        "vash",
        "port",
        "admin",
        "fee",
        "mtw",
        "reserve",
    )

    emit()
    emit("=" * 100)
    emit("SEMANTIC HEADER DISCOVERY")
    emit("=" * 100)
    emit()
    emit(
        "Header-name matches only. "
        "NOT accounting classifications."
    )
    emit()

    for term in terms:
        matches = [
            header
            for header in sorted(occurrences)
            if term in header
        ]

        emit(
            f"{term.upper():<15}: "
            f"{len(matches):>4}"
        )

        for header in matches:
            emit(
                f"    {header}"
            )

    emit()
    emit("=" * 100)
    emit("METHODOLOGICAL STOP")
    emit("=" * 100)
    emit()
    emit(
        "No financial column selected for aggregation."
    )
    emit(
        "No HAP/expenditure total calculated."
    )
    emit(
        "No Washington filter applied."
    )
    emit(
        "No USAspending reconciliation performed."
    )
    emit(
        "No database or ledger write occurred."
    )

    return "\n".join(lines) + "\n"


def main() -> int:

    print("=" * 88)
    print("HUD VMS STRUCTURAL PROFILE")
    print("ETL-174B-5Y-2A v0.2")
    print("=" * 88)
    print()
    print(f"Input : {RAW_ROOT}")
    print(f"Output: {VALIDATION_ROOT}")
    print()
    print("Mode  : READ ONLY / SINGLE-PASS PROFILE")
    print("DB    : NO DATABASE WRITES")
    print()

    paths: list[Path] = []

    for year in EXPECTED_YEARS:

        year_dir = RAW_ROOT / f"CY{year}"

        if not year_dir.exists():
            print(
                f"ERROR: missing directory: "
                f"{year_dir}"
            )
            return 1

        paths.extend(
            sorted(year_dir.glob("*.xlsx"))
        )

    if len(paths) != EXPECTED_WORKBOOK_COUNT:
        print(
            f"ERROR: expected "
            f"{EXPECTED_WORKBOOK_COUNT} workbooks; "
            f"found {len(paths)}."
        )
        return 1

    VALIDATION_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    profiles = []

    for index, path in enumerate(
        paths,
        start=1,
    ):
        print(
            f"[{index:02}/{len(paths):02}] "
            f"{path.name}",
            flush=True,
        )

        profiles.append(
            profile_workbook(path)
        )

        print(
            "         complete",
            flush=True,
        )

    timestamp = datetime.now(
        timezone.utc
    ).strftime("%Y%m%dT%H%M%SZ")

    report_path = (
        VALIDATION_ROOT
        / (
            "profile_hud_vms_2022_2024_"
            f"structure_v0_2_{timestamp}.txt"
        )
    )

    json_path = (
        VALIDATION_ROOT
        / (
            "profile_hud_vms_2022_2024_"
            f"structure_v0_2_{timestamp}.json"
        )
    )

    report_path.write_text(
        render_report(profiles),
        encoding="utf-8",
    )

    json_path.write_text(
        json.dumps(
            {
                "etl_checkpoint":
                    "ETL-174B-5Y-2A",
                "version": "0.2",
                "generated_utc":
                    datetime.now(
                        timezone.utc
                    ).isoformat(),
                "mode":
                    "READ_ONLY_SINGLE_PASS_STRUCTURE",
                "workbook_count":
                    len(profiles),
                "profiles":
                    profiles,
            },
            indent=2,
            ensure_ascii=False,
            default=str,
        )
        + "\n",
        encoding="utf-8",
    )

    hashes = {
        profile["sha256"]
        for profile in profiles
    }

    signatures = set()

    for profile in profiles:
        for sheet in profile["sheets"]:

            if sheet["columns"]:
                signatures.add(
                    schema_signature(sheet)
                )

    print()
    print("=" * 88)
    print("STRUCTURAL PROFILE CONTROL")
    print("=" * 88)
    print(
        f"Workbooks profiled       : "
        f"{len(profiles)}"
    )
    print(
        f"Distinct workbook hashes : "
        f"{len(hashes)}"
    )
    print(
        f"Schema signatures        : "
        f"{len(signatures)}"
    )
    print()
    print(
        "Report : "
        f"{report_path.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON   : "
        f"{json_path.relative_to(REPO_ROOT)}"
    )
    print()

    if (
        len(profiles) == EXPECTED_WORKBOOK_COUNT
        and len(hashes) == EXPECTED_WORKBOOK_COUNT
    ):
        print("RESULT: PASS")
        print()
        print(
            "No financial values were aggregated."
        )
        print(
            "No Washington filter was applied."
        )
        print(
            "No USAspending reconciliation was performed."
        )
        print(
            "No accounting or ledger decision was made."
        )
        return 0

    print("RESULT: FAIL")
    return 1


if __name__ == "__main__":
    sys.exit(main())