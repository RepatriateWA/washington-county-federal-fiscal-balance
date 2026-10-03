"""
ETL-174B-5Y-2D-2
HUD VMS HAP Total vs USAspending AL 14.871 obligation reconciliation.

SCOPE
-----
Washington
FFY2023: 2022-10 through 2023-09
FFY2024: 2023-10 through 2024-09

PURPOSE
-------
Put two independently validated monthly series on the same fiscal-period
time axis:

    HUD VMS HAP Total
    USAspending AL 14.871 federal_action_obligation

This is a descriptive reconciliation.

A difference is NOT treated as an error.

READ ONLY.

NO:
    accounting-source selection,
    obligation/expenditure equivalence assumption,
    county allocation,
    Vms2/EHV addition,
    database writes,
    ledger decision.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


CHECKPOINT = "ETL-174B-5Y-2D-2"

REPO_ROOT = Path(__file__).resolve().parents[3]

VMS_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "quarterly"
)

USASPENDING_PROFILE_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "validation"
    / "usaspending_14871_monthly_obligations"
)

OUTPUT_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "validation"
    / "reconcile_vms_hap_vs_usaspending_14871"
)

TARGET_MONTHS = {
    2023: [
        date(2022, 10, 1),
        date(2022, 11, 1),
        date(2022, 12, 1),
        date(2023, 1, 1),
        date(2023, 2, 1),
        date(2023, 3, 1),
        date(2023, 4, 1),
        date(2023, 5, 1),
        date(2023, 6, 1),
        date(2023, 7, 1),
        date(2023, 8, 1),
        date(2023, 9, 1),
    ],
    2024: [
        date(2023, 10, 1),
        date(2023, 11, 1),
        date(2023, 12, 1),
        date(2024, 1, 1),
        date(2024, 2, 1),
        date(2024, 3, 1),
        date(2024, 4, 1),
        date(2024, 5, 1),
        date(2024, 6, 1),
        date(2024, 7, 1),
        date(2024, 8, 1),
        date(2024, 9, 1),
    ],
}


def normalize_header(value) -> str:
    if value is None:
        return ""

    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def decimal_or_zero(value) -> Decimal:
    if value is None:
        return Decimal("0")

    if isinstance(value, Decimal):
        return value

    return Decimal(str(value))


def money(value: Decimal | None) -> str:
    if value is None:
        return "NULL"

    return f"${value:,.2f}"


def pct(value: Decimal | None) -> str:
    if value is None:
        return "NULL"

    return f"{value * Decimal('100'):.4f}%"


def month_start(value) -> date:
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)

    if isinstance(value, date):
        return date(value.year, value.month, 1)

    text = str(value).strip()

    # YYYY-MM-DD
    parsed = datetime.strptime(
        text[:10],
        "%Y-%m-%d",
    ).date()

    return date(
        parsed.year,
        parsed.month,
        1,
    )


def fiscal_year_for_month(month: date) -> int:
    if month.month >= 10:
        return month.year + 1

    return month.year


def discover_latest_usaspending_csv() -> Path:
    candidates = sorted(
        USASPENDING_PROFILE_ROOT.glob(
            "profile_14871_monthly_*.csv"
        )
    )

    if not candidates:
        raise FileNotFoundError(
            "No ETL-174B-5Y-2D-1 monthly CSV found under "
            f"{USASPENDING_PROFILE_ROOT}"
        )

    return candidates[-1]


def load_usaspending_monthly(path: Path):
    rows = {}

    with path.open(
        "r",
        encoding="utf-8",
        newline="",
    ) as handle:

        reader = csv.DictReader(handle)

        required = {
            "action_month",
            "fiscal_year",
            "transaction_count",
            "award_count",
            "recipient_count",
            "positive_obligation",
            "negative_obligation",
            "net_obligation",
        }

        missing = required - set(
            reader.fieldnames or []
        )

        if missing:
            raise RuntimeError(
                "USAspending CSV missing required columns: "
                + ", ".join(sorted(missing))
            )

        for raw in reader:
            month = month_start(
                raw["action_month"]
            )

            if month in rows:
                raise RuntimeError(
                    "Duplicate USAspending action month: "
                    f"{month}"
                )

            rows[month] = {
                "fiscal_year":
                    int(raw["fiscal_year"]),
                "transaction_count":
                    int(raw["transaction_count"]),
                "award_count":
                    int(raw["award_count"]),
                "recipient_count":
                    int(raw["recipient_count"]),
                "positive_obligation":
                    decimal_or_zero(
                        raw["positive_obligation"]
                    ),
                "negative_obligation":
                    decimal_or_zero(
                        raw["negative_obligation"]
                    ),
                "net_obligation":
                    decimal_or_zero(
                        raw["net_obligation"]
                    ),
            }

    return rows


def discover_vms_workbooks():
    workbooks = sorted(
        path
        for path in VMS_ROOT.rglob("*.xlsx")
        if not path.name.startswith("~$")
    )

    if not workbooks:
        raise FileNotFoundError(
            f"No VMS workbooks found under {VMS_ROOT}"
        )

    return workbooks


def find_column(headers, aliases):
    normalized = {
        normalize_header(value): index
        for index, value in enumerate(headers)
        if value is not None
    }

    for alias in aliases:
        if alias in normalized:
            return normalized[alias]

    return None


def load_vms_hap_monthly():
    """
    Extract reported HUD HAP Total directly from Vms1.

    This does NOT reconstruct the 33-component formula again.
    ETL-174B-5Y-2C.1 already certified reported HAP Total against
    that formula for 1,008/1,008 Washington PHA-month observations.
    """

    monthly = defaultdict(
        lambda: {
            "hap_total": Decimal("0"),
            "pha_codes": set(),
            "row_count": 0,
        }
    )

    workbook_controls = []

    for path in discover_vms_workbooks():

        wb = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        if "Vms1" not in wb.sheetnames:
            wb.close()
            continue

        ws = wb["Vms1"]

        iterator = ws.iter_rows(
            values_only=True
        )

        try:
            headers = next(iterator)
        except StopIteration:
            wb.close()
            continue

        pha_col = find_column(
            headers,
            (
                "pha_code",
                "pha_code_number",
            ),
        )

        month_col = find_column(
            headers,
            (
                "submission_month",
            ),
        )

        hap_col = find_column(
            headers,
            (
                "hap_total",
            ),
        )

        if (
            pha_col is None
            or month_col is None
            or hap_col is None
        ):
            wb.close()
            raise RuntimeError(
                f"Required Vms1 field missing in {path.name}: "
                f"PHA={pha_col}, month={month_col}, "
                f"HAP Total={hap_col}"
            )

        wa_rows = 0

        for row in iterator:

            pha_raw = row[pha_col]

            if pha_raw is None:
                continue

            pha_code = str(
                pha_raw
            ).strip().upper()

            if not pha_code.startswith("WA"):
                continue

            submission_month = row[
                month_col
            ]

            if submission_month is None:
                continue

            month = month_start(
                submission_month
            )

            # Restrict to the two fully observed FFYs.
            if month not in (
                set(TARGET_MONTHS[2023])
                | set(TARGET_MONTHS[2024])
            ):
                continue

            hap_value = decimal_or_zero(
                row[hap_col]
            )

            key = monthly[month]

            if pha_code in key["pha_codes"]:
                wb.close()
                raise RuntimeError(
                    "Duplicate VMS PHA-month encountered: "
                    f"{pha_code} {month}"
                )

            key["pha_codes"].add(
                pha_code
            )
            key["hap_total"] += hap_value
            key["row_count"] += 1
            wa_rows += 1

        workbook_controls.append(
            {
                "workbook": path.name,
                "target_wa_rows": wa_rows,
            }
        )

        wb.close()

    return monthly, workbook_controls


def build_reconciliation(
    vms_monthly,
    usa_monthly,
):
    results = []

    for fiscal_year in (2023, 2024):

        for month in TARGET_MONTHS[
            fiscal_year
        ]:

            if month not in vms_monthly:
                raise RuntimeError(
                    "Missing VMS target month: "
                    f"{month}"
                )

            if month not in usa_monthly:
                raise RuntimeError(
                    "Missing USAspending target month: "
                    f"{month}"
                )

            vms = vms_monthly[month]
            usa = usa_monthly[month]

            derived_fy = fiscal_year_for_month(
                month
            )

            if derived_fy != fiscal_year:
                raise RuntimeError(
                    "Internal fiscal-year mapping error "
                    f"for {month}"
                )

            if (
                usa["fiscal_year"]
                != fiscal_year
            ):
                raise RuntimeError(
                    "USAspending fiscal-year mismatch: "
                    f"{month} expected FY{fiscal_year}, "
                    f"found FY{usa['fiscal_year']}"
                )

            hap = vms["hap_total"]
            positive = usa[
                "positive_obligation"
            ]
            negative = usa[
                "negative_obligation"
            ]
            net = usa[
                "net_obligation"
            ]

            difference = net - hap

            net_over_hap = (
                net / hap
                if hap != 0
                else None
            )

            hap_over_net = (
                hap / net
                if net != 0
                else None
            )

            results.append(
                {
                    "fiscal_year":
                        fiscal_year,
                    "action_month":
                        month,
                    "vms_pha_count":
                        len(vms["pha_codes"]),
                    "vms_row_count":
                        vms["row_count"],
                    "vms_hap_total":
                        hap,
                    "usaspending_transaction_count":
                        usa[
                            "transaction_count"
                        ],
                    "usaspending_award_count":
                        usa[
                            "award_count"
                        ],
                    "usaspending_recipient_count":
                        usa[
                            "recipient_count"
                        ],
                    "usaspending_positive_obligation":
                        positive,
                    "usaspending_negative_obligation":
                        negative,
                    "usaspending_net_obligation":
                        net,
                    "net_minus_hap":
                        difference,
                    "net_over_hap":
                        net_over_hap,
                    "hap_over_net":
                        hap_over_net,
                }
            )

    return results


def summarize_fiscal_years(rows):
    summaries = []

    for fiscal_year in (2023, 2024):

        selected = [
            row
            for row in rows
            if row["fiscal_year"]
            == fiscal_year
        ]

        hap = sum(
            (
                row["vms_hap_total"]
                for row in selected
            ),
            Decimal("0"),
        )

        positive = sum(
            (
                row[
                    "usaspending_positive_obligation"
                ]
                for row in selected
            ),
            Decimal("0"),
        )

        negative = sum(
            (
                row[
                    "usaspending_negative_obligation"
                ]
                for row in selected
            ),
            Decimal("0"),
        )

        net = sum(
            (
                row[
                    "usaspending_net_obligation"
                ]
                for row in selected
            ),
            Decimal("0"),
        )

        difference = net - hap

        summaries.append(
            {
                "fiscal_year":
                    fiscal_year,
                "month_count":
                    len(selected),
                "vms_hap_total":
                    hap,
                "positive_obligation":
                    positive,
                "negative_obligation":
                    negative,
                "net_obligation":
                    net,
                "net_minus_hap":
                    difference,
                "net_over_hap":
                    (
                        net / hap
                        if hap != 0
                        else None
                    ),
                "hap_over_net":
                    (
                        hap / net
                        if net != 0
                        else None
                    ),
            }
        )

    return summaries


def write_csv(path, rows):
    fields = [
        "fiscal_year",
        "action_month",
        "vms_pha_count",
        "vms_row_count",
        "vms_hap_total",
        "usaspending_transaction_count",
        "usaspending_award_count",
        "usaspending_recipient_count",
        "usaspending_positive_obligation",
        "usaspending_negative_obligation",
        "usaspending_net_obligation",
        "net_minus_hap",
        "net_over_hap",
        "hap_over_net",
    ]

    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=fields,
        )

        writer.writeheader()

        for row in rows:
            writer.writerow(row)


def render_report(
    usa_path,
    rows,
    summaries,
    workbook_controls,
):
    lines = []

    def emit(text=""):
        lines.append(text)

    emit("=" * 108)
    emit(
        "HUD VMS HAP TOTAL vs USASPENDING "
        "AL 14.871 COMMON-PERIOD RECONCILIATION"
    )
    emit(CHECKPOINT)
    emit("=" * 108)
    emit()
    emit("Mode : READ ONLY")
    emit("Scope: Washington")
    emit("FFY2023 = Oct 2022 through Sep 2023")
    emit("FFY2024 = Oct 2023 through Sep 2024")
    emit()
    emit(
        "IMPORTANT: A difference is not treated as an error. "
        "HUD VMS HAP Total and USAspending federal_action_obligation "
        "represent different accounting stages."
    )

    emit()
    emit("INPUT CONTROL")
    emit("-" * 108)
    emit(
        "USAspending monthly source : "
        f"{usa_path.relative_to(REPO_ROOT)}"
    )
    emit(
        "HUD VMS source             : "
        f"{VMS_ROOT.relative_to(REPO_ROOT)}"
    )
    emit(
        "Target paired months       : 24"
    )

    emit()
    emit("VMS WORKBOOK TARGET-ROW CONTROL")
    emit("-" * 108)

    for item in workbook_controls:
        if item["target_wa_rows"] > 0:
            emit(
                f"{item['workbook']:<35} "
                f"WA target rows="
                f"{item['target_wa_rows']:,}"
            )

    emit()
    emit("MONTHLY RECONCILIATION")
    emit("-" * 108)

    for row in rows:
        emit(
            f"{row['action_month']} | "
            f"FY{row['fiscal_year']} | "
            f"PHAs={row['vms_pha_count']:>2} | "
            f"HAP={money(row['vms_hap_total'])} | "
            f"USA+={money(row['usaspending_positive_obligation'])} | "
            f"USA-={money(row['usaspending_negative_obligation'])} | "
            f"USA net={money(row['usaspending_net_obligation'])} | "
            f"net-HAP={money(row['net_minus_hap'])}"
        )

    emit()
    emit("FISCAL-YEAR RECONCILIATION")
    emit("-" * 108)

    for row in summaries:
        emit(f"FFY{row['fiscal_year']}")
        emit(
            f"    Months                         : "
            f"{row['month_count']}"
        )
        emit(
            f"    HUD VMS HAP Total              : "
            f"{money(row['vms_hap_total'])}"
        )
        emit(
            f"    USAspending positive obligation: "
            f"{money(row['positive_obligation'])}"
        )
        emit(
            f"    USAspending negative obligation: "
            f"{money(row['negative_obligation'])}"
        )
        emit(
            f"    USAspending net obligation     : "
            f"{money(row['net_obligation'])}"
        )
        emit(
            f"    Net obligation - HAP Total     : "
            f"{money(row['net_minus_hap'])}"
        )
        emit(
            f"    Net obligation / HAP Total     : "
            f"{pct(row['net_over_hap'])}"
        )
        emit(
            f"    HAP Total / net obligation     : "
            f"{pct(row['hap_over_net'])}"
        )
        emit()

    emit("POPULATION / COMPLETENESS CONTROL")
    emit("-" * 108)

    emit(
        f"Paired month observations : "
        f"{len(rows)}"
    )

    pha_counts = sorted(
        {
            row["vms_pha_count"]
            for row in rows
        }
    )

    emit(
        "Distinct monthly PHA counts: "
        + ", ".join(
            str(value)
            for value in pha_counts
        )
    )

    complete = (
        len(rows) == 24
        and pha_counts == [28]
        and all(
            row["vms_row_count"] == 28
            for row in rows
        )
        and all(
            summary["month_count"] == 12
            for summary in summaries
        )
    )

    emit(
        "Common-period completeness : "
        + (
            "PASS"
            if complete
            else "HOLD"
        )
    )

    emit()
    emit("INTERPRETIVE BOUNDARY")
    emit("-" * 108)
    emit(
        "This test measures the relationship between "
        "HUD-reported HAP expenditure/use and "
        "USAspending transaction obligations."
    )
    emit(
        "Equality is not expected or required for PASS."
    )
    emit(
        "A positive or negative residual is not classified "
        "as error, overlap, omission, or timing difference "
        "without additional evidence."
    )

    emit()
    emit("METHODOLOGICAL STOP")
    emit("-" * 108)
    emit(
        "No accounting-source selection was made."
    )
    emit(
        "No obligation/expenditure equivalence was assumed."
    )
    emit(
        "No Vms2/EHV values were added."
    )
    emit(
        "No county allocation was performed."
    )
    emit(
        "No database writes occurred."
    )
    emit(
        "No ledger decision was made."
    )

    emit()

    if complete:
        emit(
            "RESULT: PASS — FFY2023 AND FFY2024 "
            "COMMON-PERIOD RECONCILIATION COMPLETED"
        )
    else:
        emit(
            "RESULT: HOLD — COMMON-PERIOD POPULATION "
            "OR COMPLETENESS CONTROL FAILED"
        )

    return "\n".join(lines) + "\n", complete


def main():
    print("=" * 92)
    print(
        "HUD VMS HAP TOTAL vs USASPENDING "
        "AL 14.871 RECONCILIATION"
    )
    print(CHECKPOINT)
    print("=" * 92)
    print()
    print("Mode : READ ONLY")
    print("Scope: FFY2023 + FFY2024")
    print()

    OUTPUT_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    usa_path = (
        discover_latest_usaspending_csv()
    )

    print(
        "USAspending monthly source:",
        usa_path.name,
    )

    usa_monthly = load_usaspending_monthly(
        usa_path
    )

    vms_monthly, workbook_controls = (
        load_vms_hap_monthly()
    )

    rows = build_reconciliation(
        vms_monthly,
        usa_monthly,
    )

    summaries = summarize_fiscal_years(
        rows
    )

    report, complete = render_report(
        usa_path,
        rows,
        summaries,
        workbook_controls,
    )

    timestamp = datetime.now(
        timezone.utc
    ).strftime("%Y%m%dT%H%M%SZ")

    report_path = (
        OUTPUT_ROOT
        / f"reconcile_vms_hap_vs_14871_{timestamp}.txt"
    )

    csv_path = (
        OUTPUT_ROOT
        / f"reconcile_vms_hap_vs_14871_{timestamp}.csv"
    )

    json_path = (
        OUTPUT_ROOT
        / f"reconcile_vms_hap_vs_14871_{timestamp}.json"
    )

    report_path.write_text(
        report,
        encoding="utf-8",
    )

    write_csv(
        csv_path,
        rows,
    )

    json_payload = {
        "checkpoint": CHECKPOINT,
        "generated_utc":
            datetime.now(
                timezone.utc
            ).isoformat(),
        "scope": {
            "state": "Washington",
            "fiscal_years": [
                2023,
                2024,
            ],
            "paired_month_count":
                len(rows),
        },
        "usaspending_monthly_source":
            str(
                usa_path.relative_to(
                    REPO_ROOT
                )
            ),
        "fiscal_year_summary": [
            {
                key: (
                    str(value)
                    if isinstance(
                        value,
                        Decimal,
                    )
                    else value
                )
                for key, value
                in summary.items()
            }
            for summary in summaries
        ],
        "monthly_rows": [
            {
                key: (
                    value.isoformat()
                    if isinstance(
                        value,
                        date,
                    )
                    else (
                        str(value)
                        if isinstance(
                            value,
                            Decimal,
                        )
                        else value
                    )
                )
                for key, value
                in row.items()
            }
            for row in rows
        ],
        "common_period_complete":
            complete,
        "methodological_stop": {
            "accounting_source_selected":
                False,
            "obligation_expenditure_equivalence_assumed":
                False,
            "vms2_ehv_added":
                False,
            "county_allocation_performed":
                False,
            "database_writes":
                False,
            "ledger_decision":
                False,
        },
    }

    json_path.write_text(
        json.dumps(
            json_payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    print()

    for summary in summaries:
        print(
            f"FFY{summary['fiscal_year']}:"
        )
        print(
            "  VMS HAP Total          :",
            money(
                summary[
                    "vms_hap_total"
                ]
            ),
        )
        print(
            "  USAspending net        :",
            money(
                summary[
                    "net_obligation"
                ]
            ),
        )
        print(
            "  Net - HAP              :",
            money(
                summary[
                    "net_minus_hap"
                ]
            ),
        )
        print(
            "  USA net / HAP          :",
            pct(
                summary[
                    "net_over_hap"
                ]
            ),
        )
        print()

    print(
        "Paired months          :",
        len(rows),
    )

    print(
        "Completeness           :",
        "PASS" if complete else "HOLD",
    )

    print()
    print(
        "Report :",
        report_path.relative_to(
            REPO_ROOT
        ),
    )
    print(
        "CSV    :",
        csv_path.relative_to(
            REPO_ROOT
        ),
    )
    print(
        "JSON   :",
        json_path.relative_to(
            REPO_ROOT
        ),
    )
    print()

    if not complete:
        print(
            "RESULT: HOLD — COMMON-PERIOD "
            "POPULATION OR COMPLETENESS "
            "CONTROL FAILED"
        )
        return 2

    print(
        "RESULT: PASS — FFY2023 AND FFY2024 "
        "COMMON-PERIOD RECONCILIATION COMPLETED"
    )

    return 0


if __name__ == "__main__":
    sys.exit(main())