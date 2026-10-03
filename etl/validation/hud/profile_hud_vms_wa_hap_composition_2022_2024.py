"""
ETL-174B-5Y-2C
HUD VMS Washington HAP Total composition test, CY2022-CY2024.

PURPOSE
-------
Empirically determine how HUD's Vms1 "HAP Total" relates to the other
HAP-labelled fields in the official VMS workbooks.

READ ONLY.

CORE QUESTION
-------------
What, mathematically, is HUD's HAP Total?

METHOD
------
For every Washington PHA x submission-month observation:

    1. Extract HAP Total.
    2. Extract every Vms1 field whose normalized header contains "hap".
    3. Preserve every HAP-labelled field individually.
    4. Classify fields structurally using their HUD header text.
    5. Calculate several deliberately mechanical candidate sums.
    6. Test exact and near-exact equality with HAP Total.
    7. Search empirically for fields whose inclusion/exclusion explains
       residual differences.

IMPORTANT
---------
This script does NOT assume that every field containing "HAP" is additive.

Many HAP-labelled fields may be:
    * subsets,
    * memorandum items,
    * portability measures,
    * after-first-of-month measures,
    * unpaid/billed balances,
    * special-purpose voucher components,
    * overlapping classifications.

Therefore all candidate sums are DIAGNOSTIC ONLY.

NO:
    database writes,
    USAspending reconciliation,
    Vms2/EHV addition,
    county allocation,
    accounting-source certification,
    ledger decision.

Checkpoint:
    ETL-174B-5Y-2C
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# =============================================================================
# PATHS / CONTROLS
# =============================================================================

REPO_ROOT = Path(__file__).resolve().parents[3]

RAW_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "quarterly"
)

OUTPUT_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "validation"
    / "profile_wa_hap_composition_2022_2024"
)

EXPECTED_YEARS = (2022, 2023, 2024)
EXPECTED_WORKBOOK_COUNT = 12
EXPECTED_OBSERVATIONS = 1008
EXPECTED_MONTHS = 36

TARGET_SHEET = "Vms1"
WA_PREFIX = "WA"

ABSOLUTE_TOLERANCE = 0.01
NEAR_TOLERANCE = 1.00


# =============================================================================
# BASIC HELPERS
# =============================================================================

def clean_text(value: Any) -> str:
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


def normalize_header(value: Any) -> str:
    text = clean_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def numeric(value: Any) -> float | None:
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        if (
            isinstance(value, float)
            and math.isnan(value)
        ):
            return None

        return float(value)

    text = clean_text(value)

    if not text:
        return None

    text = (
        text
        .replace("$", "")
        .replace(",", "")
        .strip()
    )

    try:
        return float(text)
    except ValueError:
        return None


def parse_month(value: Any) -> date | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = clean_text(value)

    if not text:
        return None

    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m/%d/%y",
    )

    for fmt in formats:
        try:
            return datetime.strptime(
                text,
                fmt,
            ).date()
        except ValueError:
            pass

    return None


def month_key(value: date) -> str:
    return value.strftime("%Y-%m")


def money(value: float | None) -> str:
    if value is None:
        return "NULL"

    return f"${value:,.2f}"


def safe_ratio(
    numerator: float | None,
    denominator: float | None,
) -> float | None:

    if numerator is None:
        return None

    if denominator in (None, 0):
        return None

    return numerator / denominator


# =============================================================================
# FIELD CLASSIFICATION
# =============================================================================

def classify_hap_field(
    normalized: str,
) -> list[str]:
    """
    Structural classification based ONLY on HUD field name.

    A field can belong to multiple categories.

    These labels are diagnostic metadata, not accounting conclusions.
    """

    categories: list[str] = []

    if normalized == "hap_total":
        categories.append("HUD_TOTAL")

    if "portable" in normalized:
        categories.append("PORTABILITY")

    if "admin" in normalized:
        categories.append("ADMIN")

    if "unpaid" in normalized:
        categories.append("BALANCE_OR_RECEIVABLE")

    if "after_the_first_of_the_month" in normalized:
        categories.append("TIMING_SUBSET")

    if "mainstream" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "vash" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "foster_youth" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "housing_stability" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "tenant_protection" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "homeownership" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "rad" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "rental_assistance_demonstration" in normalized:
        categories.append("SPECIAL_PURPOSE")

    if "project_based" in normalized:
        categories.append("PROGRAM_SUBSET")

    if "total_hap_for_portable_units_administered" == normalized:
        categories.append("PORT_IN")

    if "portable_voucher_paid_hap" == normalized:
        categories.append("PORT_OUT_OR_PAID")

    if not categories:
        categories.append("UNCLASSIFIED_HAP")

    return categories


def candidate_sum_group(
    normalized: str,
) -> str:
    """
    Assign a conservative mechanical candidate-sum disposition.

    This does NOT assert accounting semantics.

    PRIMARY_CANDIDATE:
        HAP-labelled fields that do not look, from their names alone,
        like balances, timing subsets, admin fields, or portability
        memorandum measures.

    EXCLUDED_DIAGNOSTIC:
        fields that visibly look unsuitable for blind addition.

    HUD_TOTAL:
        HAP Total itself.
    """

    categories = set(
        classify_hap_field(normalized)
    )

    if "HUD_TOTAL" in categories:
        return "HUD_TOTAL"

    exclusion_categories = {
        "PORTABILITY",
        "ADMIN",
        "BALANCE_OR_RECEIVABLE",
        "TIMING_SUBSET",
    }

    if categories & exclusion_categories:
        return "EXCLUDED_DIAGNOSTIC"

    return "PRIMARY_CANDIDATE"


# =============================================================================
# WORKBOOK DISCOVERY
# =============================================================================

def discover_workbooks() -> list[Path]:
    paths: list[Path] = []

    for year in EXPECTED_YEARS:
        year_dir = RAW_ROOT / f"CY{year}"

        if not year_dir.exists():
            raise RuntimeError(
                f"Missing expected directory: {year_dir}"
            )

        paths.extend(
            sorted(year_dir.glob("*.xlsx"))
        )

    if len(paths) != EXPECTED_WORKBOOK_COUNT:
        raise RuntimeError(
            f"Expected {EXPECTED_WORKBOOK_COUNT} workbooks; "
            f"found {len(paths)}."
        )

    return paths


# =============================================================================
# EXTRACTION
# =============================================================================

def extract_workbook(
    path: Path,
) -> tuple[
    list[dict[str, Any]],
    dict[str, dict[str, Any]],
    dict[str, Any],
]:

    wb = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:
        if TARGET_SHEET not in wb.sheetnames:
            raise RuntimeError(
                f"{path.name}: missing {TARGET_SHEET}"
            )

        ws = wb[TARGET_SHEET]
        row_iterator = ws.iter_rows(
            values_only=True
        )

        try:
            raw_headers = next(row_iterator)
        except StopIteration:
            raise RuntimeError(
                f"{path.name}: empty sheet"
            )

        headers = [
            normalize_header(value)
            for value in raw_headers
        ]

        raw_header_by_normalized = {}

        positions = {}

        for position, (
            raw,
            normalized,
        ) in enumerate(
            zip(raw_headers, headers)
        ):
            if not normalized:
                continue

            if normalized in positions:
                # Preserve first occurrence.
                continue

            positions[normalized] = position

            raw_header_by_normalized[
                normalized
            ] = clean_text(raw)

        required = (
            "pha_code",
            "pha_name",
            "submission_month",
            "hap_total",
        )

        missing = [
            field
            for field in required
            if field not in positions
        ]

        if missing:
            raise RuntimeError(
                f"{path.name}: missing required fields: "
                + ", ".join(missing)
            )

        hap_fields = [
            header
            for header in positions
            if "hap" in header
        ]

        field_metadata = {}

        for field in hap_fields:
            field_metadata[field] = {
                "normalized_header": field,
                "raw_header":
                    raw_header_by_normalized[field],
                "categories":
                    classify_hap_field(field),
                "candidate_sum_group":
                    candidate_sum_group(field),
            }

        records = []

        source_rows = 0
        wa_rows = 0
        invalid_month_rows = 0

        for worksheet_row, row in enumerate(
            row_iterator,
            start=2,
        ):
            source_rows += 1

            pha_position = positions[
                "pha_code"
            ]

            if pha_position >= len(row):
                continue

            pha_code = clean_text(
                row[pha_position]
            )

            if not pha_code.upper().startswith(
                WA_PREFIX
            ):
                continue

            wa_rows += 1

            submission_position = positions[
                "submission_month"
            ]

            submission_value = (
                row[submission_position]
                if submission_position < len(row)
                else None
            )

            submission_month = parse_month(
                submission_value
            )

            if submission_month is None:
                invalid_month_rows += 1

            record = {
                "source_workbook":
                    path.name,
                "source_row_number":
                    worksheet_row,
                "pha_code":
                    pha_code,
                "pha_name":
                    clean_text(
                        row[
                            positions["pha_name"]
                        ]
                    ),
                "submission_month":
                    (
                        submission_month.isoformat()
                        if submission_month
                        else ""
                    ),
                "submission_month_key":
                    (
                        month_key(submission_month)
                        if submission_month
                        else ""
                    ),
                "hap_fields": {},
            }

            for field in hap_fields:
                position = positions[field]

                value = (
                    row[position]
                    if position < len(row)
                    else None
                )

                record["hap_fields"][field] = (
                    numeric(value)
                )

            records.append(record)

        metadata = {
            "source_workbook":
                path.name,
            "source_rows_after_header":
                source_rows,
            "wa_rows":
                wa_rows,
            "invalid_month_rows":
                invalid_month_rows,
            "hap_field_count":
                len(hap_fields),
            "hap_fields":
                hap_fields,
        }

        return (
            records,
            field_metadata,
            metadata,
        )

    finally:
        wb.close()


# =============================================================================
# FIELD INVENTORY
# =============================================================================

def consolidate_field_metadata(
    per_workbook_metadata: list[
        dict[str, dict[str, Any]]
    ],
) -> dict[str, dict[str, Any]]:

    consolidated: dict[
        str,
        dict[str, Any],
    ] = {}

    for metadata in per_workbook_metadata:
        for field, item in metadata.items():

            if field not in consolidated:
                consolidated[field] = {
                    "normalized_header":
                        field,
                    "raw_headers":
                        set(),
                    "categories":
                        set(),
                    "candidate_sum_group":
                        item[
                            "candidate_sum_group"
                        ],
                    "workbook_occurrences":
                        0,
                }

            target = consolidated[field]

            target["raw_headers"].add(
                item["raw_header"]
            )

            target["categories"].update(
                item["categories"]
            )

            target[
                "workbook_occurrences"
            ] += 1

    result = {}

    for field, item in sorted(
        consolidated.items()
    ):
        result[field] = {
            "normalized_header":
                field,
            "raw_headers":
                sorted(item["raw_headers"]),
            "categories":
                sorted(item["categories"]),
            "candidate_sum_group":
                item["candidate_sum_group"],
            "workbook_occurrences":
                item["workbook_occurrences"],
        }

    return result


# =============================================================================
# ANALYSIS
# =============================================================================

def analyze(
    records: list[dict[str, Any]],
    field_metadata: dict[
        str,
        dict[str, Any],
    ],
) -> dict[str, Any]:

    result: dict[str, Any] = {}

    # -------------------------------------------------------------------------
    # Core population controls
    # -------------------------------------------------------------------------

    months = sorted(
        {
            record["submission_month_key"]
            for record in records
            if record["submission_month_key"]
        }
    )

    phas = sorted(
        {
            record["pha_code"]
            for record in records
        }
    )

    result["observation_count"] = len(
        records
    )
    result["distinct_month_count"] = len(
        months
    )
    result["distinct_pha_count"] = len(
        phas
    )

    key_counter = Counter(
        (
            record["pha_code"],
            record["submission_month_key"],
        )
        for record in records
        if record["submission_month_key"]
    )

    result["duplicate_key_count"] = sum(
        1
        for count in key_counter.values()
        if count > 1
    )

    # -------------------------------------------------------------------------
    # Field-level statistics
    # -------------------------------------------------------------------------

    field_stats = {}

    for field in field_metadata:

        values = []

        for record in records:
            value = record[
                "hap_fields"
            ].get(field)

            if value is not None:
                values.append(value)

        field_stats[field] = {
            "numeric_count":
                len(values),
            "positive_count":
                sum(
                    value > 0
                    for value in values
                ),
            "zero_count":
                sum(
                    value == 0
                    for value in values
                ),
            "negative_count":
                sum(
                    value < 0
                    for value in values
                ),
            "sum":
                (
                    sum(values)
                    if values
                    else None
                ),
            "minimum":
                (
                    min(values)
                    if values
                    else None
                ),
            "maximum":
                (
                    max(values)
                    if values
                    else None
                ),
        }

    result["field_stats"] = field_stats

    # -------------------------------------------------------------------------
    # Candidate field groups
    # -------------------------------------------------------------------------

    primary_fields = [
        field
        for field, metadata
        in field_metadata.items()
        if metadata[
            "candidate_sum_group"
        ] == "PRIMARY_CANDIDATE"
    ]

    excluded_fields = [
        field
        for field, metadata
        in field_metadata.items()
        if metadata[
            "candidate_sum_group"
        ] == "EXCLUDED_DIAGNOSTIC"
    ]

    result["primary_candidate_fields"] = (
        primary_fields
    )
    result["excluded_diagnostic_fields"] = (
        excluded_fields
    )

    # -------------------------------------------------------------------------
    # Mechanical candidate sums
    # -------------------------------------------------------------------------

    observation_tests = []

    for record in records:

        hap_total = record[
            "hap_fields"
        ].get(
            "hap_total"
        )

        # All HAP-labelled fields except HAP Total.
        all_other_values = [
            value
            for field, value
            in record["hap_fields"].items()
            if (
                field != "hap_total"
                and value is not None
            )
        ]

        all_other_sum = sum(
            all_other_values
        )

        # Conservative primary candidate.
        primary_values = [
            record[
                "hap_fields"
            ].get(field)
            for field in primary_fields
        ]

        primary_sum = sum(
            value
            for value in primary_values
            if value is not None
        )

        # Excluded diagnostic group.
        excluded_values = [
            record[
                "hap_fields"
            ].get(field)
            for field in excluded_fields
        ]

        excluded_sum = sum(
            value
            for value in excluded_values
            if value is not None
        )

        primary_residual = (
            None
            if hap_total is None
            else hap_total - primary_sum
        )

        all_other_residual = (
            None
            if hap_total is None
            else hap_total - all_other_sum
        )

        observation_tests.append(
            {
                "pha_code":
                    record["pha_code"],
                "submission_month_key":
                    record[
                        "submission_month_key"
                    ],
                "hap_total":
                    hap_total,
                "primary_candidate_sum":
                    primary_sum,
                "primary_residual":
                    primary_residual,
                "all_other_hap_sum":
                    all_other_sum,
                "all_other_residual":
                    all_other_residual,
                "excluded_diagnostic_sum":
                    excluded_sum,
                "primary_over_hap_total":
                    safe_ratio(
                        primary_sum,
                        hap_total,
                    ),
                "all_other_over_hap_total":
                    safe_ratio(
                        all_other_sum,
                        hap_total,
                    ),
            }
        )

    result[
        "observation_tests"
    ] = observation_tests

    # -------------------------------------------------------------------------
    # Equality tests
    # -------------------------------------------------------------------------

    def equality_profile(
        residual_field: str,
    ) -> dict[str, Any]:

        residuals = [
            item[residual_field]
            for item in observation_tests
            if item[residual_field]
            is not None
        ]

        exact = [
            value
            for value in residuals
            if abs(value)
            <= ABSOLUTE_TOLERANCE
        ]

        near = [
            value
            for value in residuals
            if abs(value)
            <= NEAR_TOLERANCE
        ]

        return {
            "tested_count":
                len(residuals),
            "exact_count":
                len(exact),
            "near_count":
                len(near),
            "exact_rate":
                (
                    len(exact)
                    / len(residuals)
                    if residuals
                    else None
                ),
            "near_rate":
                (
                    len(near)
                    / len(residuals)
                    if residuals
                    else None
                ),
            "residual_sum":
                (
                    sum(residuals)
                    if residuals
                    else None
                ),
            "residual_min":
                (
                    min(residuals)
                    if residuals
                    else None
                ),
            "residual_max":
                (
                    max(residuals)
                    if residuals
                    else None
                ),
            "absolute_residual_sum":
                (
                    sum(
                        abs(value)
                        for value in residuals
                    )
                    if residuals
                    else None
                ),
        }

    result[
        "primary_equality"
    ] = equality_profile(
        "primary_residual"
    )

    result[
        "all_other_equality"
    ] = equality_profile(
        "all_other_residual"
    )

    # -------------------------------------------------------------------------
    # Residual association test
    #
    # For each HAP-labelled field, test whether adding or subtracting that
    # individual field from the PRIMARY residual improves absolute residual.
    #
    # This is empirical clue generation only.
    # -------------------------------------------------------------------------

    association = []

    for field in field_metadata:

        if field == "hap_total":
            continue

        baseline_abs = 0.0
        add_abs = 0.0
        subtract_abs = 0.0
        tested = 0

        exact_if_added = 0
        exact_if_subtracted = 0

        for record, test in zip(
            records,
            observation_tests,
        ):
            residual = test[
                "primary_residual"
            ]

            field_value = record[
                "hap_fields"
            ].get(field)

            if (
                residual is None
                or field_value is None
            ):
                continue

            tested += 1

            baseline_abs += abs(
                residual
            )

            added_residual = (
                residual - field_value
            )

            subtracted_residual = (
                residual + field_value
            )

            add_abs += abs(
                added_residual
            )

            subtract_abs += abs(
                subtracted_residual
            )

            if (
                abs(added_residual)
                <= ABSOLUTE_TOLERANCE
            ):
                exact_if_added += 1

            if (
                abs(subtracted_residual)
                <= ABSOLUTE_TOLERANCE
            ):
                exact_if_subtracted += 1

        if tested == 0:
            continue

        association.append(
            {
                "field":
                    field,
                "tested_count":
                    tested,
                "baseline_abs_residual":
                    baseline_abs,
                "abs_residual_if_added":
                    add_abs,
                "abs_residual_if_subtracted":
                    subtract_abs,
                "improvement_if_added":
                    baseline_abs - add_abs,
                "improvement_if_subtracted":
                    baseline_abs - subtract_abs,
                "exact_if_added":
                    exact_if_added,
                "exact_if_subtracted":
                    exact_if_subtracted,
            }
        )

    association.sort(
        key=lambda item: max(
            item[
                "improvement_if_added"
            ],
            item[
                "improvement_if_subtracted"
            ],
        ),
        reverse=True,
    )

    result[
        "residual_field_association"
    ] = association

    # -------------------------------------------------------------------------
    # Annual HAP Total control
    # -------------------------------------------------------------------------

    annual_hap_total = defaultdict(float)

    for record in records:
        month = record[
            "submission_month_key"
        ]

        value = record[
            "hap_fields"
        ].get(
            "hap_total"
        )

        if (
            month
            and value is not None
        ):
            annual_hap_total[
                int(month[:4])
            ] += value

    result[
        "annual_hap_total"
    ] = dict(
        sorted(
            annual_hap_total.items()
        )
    )

    return result


# =============================================================================
# OUTPUT HELPERS
# =============================================================================

def write_field_inventory_csv(
    path: Path,
    field_metadata: dict[
        str,
        dict[str, Any],
    ],
    analysis: dict[str, Any],
) -> None:

    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as handle:

        writer = csv.writer(handle)

        writer.writerow(
            [
                "normalized_header",
                "raw_headers",
                "categories",
                "candidate_sum_group",
                "workbook_occurrences",
                "numeric_count",
                "positive_count",
                "zero_count",
                "negative_count",
                "sum",
                "minimum",
                "maximum",
            ]
        )

        for field, metadata in sorted(
            field_metadata.items()
        ):
            stats = analysis[
                "field_stats"
            ][field]

            writer.writerow(
                [
                    field,
                    " | ".join(
                        metadata[
                            "raw_headers"
                        ]
                    ),
                    " | ".join(
                        metadata[
                            "categories"
                        ]
                    ),
                    metadata[
                        "candidate_sum_group"
                    ],
                    metadata[
                        "workbook_occurrences"
                    ],
                    stats["numeric_count"],
                    stats["positive_count"],
                    stats["zero_count"],
                    stats["negative_count"],
                    stats["sum"],
                    stats["minimum"],
                    stats["maximum"],
                ]
            )


def write_observation_csv(
    path: Path,
    analysis: dict[str, Any],
) -> None:

    fields = (
        "pha_code",
        "submission_month_key",
        "hap_total",
        "primary_candidate_sum",
        "primary_residual",
        "all_other_hap_sum",
        "all_other_residual",
        "excluded_diagnostic_sum",
        "primary_over_hap_total",
        "all_other_over_hap_total",
    )

    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=fields,
        )

        writer.writeheader()

        for item in analysis[
            "observation_tests"
        ]:
            writer.writerow(
                {
                    field: item.get(field)
                    for field in fields
                }
            )


def render_report(
    workbook_metadata: list[
        dict[str, Any]
    ],
    field_metadata: dict[
        str,
        dict[str, Any],
    ],
    analysis: dict[str, Any],
) -> str:

    lines: list[str] = []

    def emit(text: str = "") -> None:
        lines.append(text)

    emit("=" * 100)
    emit("HUD VMS WASHINGTON HAP TOTAL COMPOSITION TEST")
    emit("ETL-174B-5Y-2C")
    emit("=" * 100)
    emit()
    emit(
        "READ-ONLY empirical composition test. "
        "No candidate sum below is an accounting conclusion."
    )
    emit()

    # -------------------------------------------------------------------------
    # Population control
    # -------------------------------------------------------------------------

    emit("POPULATION CONTROL")
    emit("-" * 100)

    emit(
        f"Workbooks processed              : "
        f"{len(workbook_metadata)}"
    )
    emit(
        f"Washington PHA-month observations: "
        f"{analysis['observation_count']:,}"
    )
    emit(
        f"Distinct Washington PHAs         : "
        f"{analysis['distinct_pha_count']:,}"
    )
    emit(
        f"Distinct submission months       : "
        f"{analysis['distinct_month_count']:,}"
    )
    emit(
        f"Duplicate PHA-month keys         : "
        f"{analysis['duplicate_key_count']:,}"
    )
    emit(
        f"Distinct HAP-labelled fields     : "
        f"{len(field_metadata):,}"
    )
    emit()

    # -------------------------------------------------------------------------
    # Workbook HAP field counts
    # -------------------------------------------------------------------------

    emit("=" * 100)
    emit("WORKBOOK HAP FIELD INVENTORY")
    emit("=" * 100)

    for item in workbook_metadata:
        emit()
        emit(item["source_workbook"])
        emit(
            f"    WA rows          : "
            f"{item['wa_rows']:,}"
        )
        emit(
            f"    HAP fields       : "
            f"{item['hap_field_count']:,}"
        )
        emit(
            f"    invalid months   : "
            f"{item['invalid_month_rows']:,}"
        )

    # -------------------------------------------------------------------------
    # Consolidated fields
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("CONSOLIDATED HAP FIELD INVENTORY")
    emit("=" * 100)

    for field, metadata in sorted(
        field_metadata.items()
    ):
        stats = analysis[
            "field_stats"
        ][field]

        emit()
        emit(field)
        emit(
            "    raw header(s)      : "
            + " | ".join(
                metadata["raw_headers"]
            )
        )
        emit(
            "    categories         : "
            + ", ".join(
                metadata["categories"]
            )
        )
        emit(
            f"    candidate group    : "
            f"{metadata['candidate_sum_group']}"
        )
        emit(
            f"    workbook occurrence: "
            f"{metadata['workbook_occurrences']}"
        )
        emit(
            f"    numeric            : "
            f"{stats['numeric_count']:,}"
        )
        emit(
            f"    positive           : "
            f"{stats['positive_count']:,}"
        )
        emit(
            f"    zero               : "
            f"{stats['zero_count']:,}"
        )
        emit(
            f"    negative           : "
            f"{stats['negative_count']:,}"
        )
        emit(
            f"    sum                : "
            f"{money(stats['sum'])}"
        )

    # -------------------------------------------------------------------------
    # Candidate field groups
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("PRIMARY CANDIDATE FIELDS")
    emit("=" * 100)
    emit()

    for field in analysis[
        "primary_candidate_fields"
    ]:
        emit(
            f"    {field}"
        )

    emit()
    emit("=" * 100)
    emit("EXCLUDED DIAGNOSTIC FIELDS")
    emit("=" * 100)
    emit()

    for field in analysis[
        "excluded_diagnostic_fields"
    ]:
        emit(
            f"    {field}"
        )

    # -------------------------------------------------------------------------
    # Equality
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("CANDIDATE-SUM EQUALITY TESTS")
    emit("=" * 100)

    for label, key in (
        (
            "PRIMARY CANDIDATE SUM",
            "primary_equality",
        ),
        (
            "ALL OTHER HAP-LABELLED FIELDS",
            "all_other_equality",
        ),
    ):
        profile = analysis[key]

        emit()
        emit(label)
        emit("-" * 100)
        emit(
            f"Tested observations        : "
            f"{profile['tested_count']:,}"
        )
        emit(
            f"Exact within $0.01         : "
            f"{profile['exact_count']:,}"
        )
        emit(
            f"Near within $1.00          : "
            f"{profile['near_count']:,}"
        )

        exact_rate = profile[
            "exact_rate"
        ]

        emit(
            "Exact rate                 : "
            + (
                "NULL"
                if exact_rate is None
                else f"{exact_rate:.4%}"
            )
        )

        emit(
            f"Signed residual sum        : "
            f"{money(profile['residual_sum'])}"
        )
        emit(
            f"Absolute residual sum      : "
            f"{money(profile['absolute_residual_sum'])}"
        )
        emit(
            f"Minimum residual           : "
            f"{money(profile['residual_min'])}"
        )
        emit(
            f"Maximum residual           : "
            f"{money(profile['residual_max'])}"
        )

    # -------------------------------------------------------------------------
    # Residual association
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("PRIMARY-RESIDUAL FIELD ASSOCIATION")
    emit("=" * 100)
    emit()
    emit(
        "Fields ranked by whether adding/subtracting one field "
        "reduces the absolute PRIMARY residual."
    )
    emit(
        "This is clue generation, not automatic formula selection."
    )
    emit()

    for item in analysis[
        "residual_field_association"
    ]:
        emit(item["field"])
        emit(
            f"    tested                     : "
            f"{item['tested_count']:,}"
        )
        emit(
            f"    improvement if added       : "
            f"{money(item['improvement_if_added'])}"
        )
        emit(
            f"    improvement if subtracted  : "
            f"{money(item['improvement_if_subtracted'])}"
        )
        emit(
            f"    exact observations if added: "
            f"{item['exact_if_added']:,}"
        )
        emit(
            f"    exact if subtracted        : "
            f"{item['exact_if_subtracted']:,}"
        )

    # -------------------------------------------------------------------------
    # Annual control
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("HAP TOTAL ANNUAL CONTROL — DIAGNOSTIC ONLY")
    emit("=" * 100)
    emit()

    for year, value in analysis[
        "annual_hap_total"
    ].items():
        emit(
            f"CY{year}    {money(value)}"
        )

    # -------------------------------------------------------------------------
    # Methodological stop
    # -------------------------------------------------------------------------

    emit()
    emit("=" * 100)
    emit("METHODOLOGICAL STOP")
    emit("=" * 100)
    emit()
    emit(
        "No candidate formula has been automatically certified."
    )
    emit(
        "HAP Total remains a diagnostic HUD-reported measure."
    )
    emit(
        "No Vms2 / Emergency Housing Voucher values were added."
    )
    emit(
        "No USAspending reconciliation was performed."
    )
    emit(
        "No county allocation was performed."
    )
    emit(
        "No database writes occurred."
    )
    emit(
        "No ledger decision was made."
    )

    return "\n".join(lines) + "\n"


# =============================================================================
# MAIN
# =============================================================================

def main() -> int:

    print("=" * 88)
    print("HUD VMS WASHINGTON HAP TOTAL COMPOSITION TEST")
    print("ETL-174B-5Y-2C")
    print("=" * 88)
    print()
    print(f"Input : {RAW_ROOT}")
    print(f"Output: {OUTPUT_ROOT}")
    print()
    print("Scope : Washington PHAs / Vms1 / CY2022-CY2024")
    print("Mode  : READ ONLY")
    print("DB    : NO DATABASE WRITES")
    print()

    try:
        paths = discover_workbooks()
    except RuntimeError as exc:
        print(f"ERROR: {exc}")
        return 1

    OUTPUT_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    records: list[dict[str, Any]] = []
    workbook_metadata = []
    per_workbook_field_metadata = []

    for index, path in enumerate(
        paths,
        start=1,
    ):
        print(
            f"[{index:02}/{len(paths):02}] "
            f"{path.name}",
            flush=True,
        )

        try:
            (
                workbook_records,
                field_metadata,
                metadata,
            ) = extract_workbook(path)

        except Exception as exc:
            print(
                f"         ERROR: {exc}",
                flush=True,
            )
            return 1

        records.extend(
            workbook_records
        )

        workbook_metadata.append(
            metadata
        )

        per_workbook_field_metadata.append(
            field_metadata
        )

        print(
            f"         WA rows="
            f"{len(workbook_records):,}; "
            f"HAP fields="
            f"{metadata['hap_field_count']:,}",
            flush=True,
        )

    field_metadata = (
        consolidate_field_metadata(
            per_workbook_field_metadata
        )
    )

    analysis = analyze(
        records,
        field_metadata,
    )

    timestamp = datetime.now(
        timezone.utc
    ).strftime("%Y%m%dT%H%M%SZ")

    report_path = (
        OUTPUT_ROOT
        / (
            "profile_hud_vms_wa_hap_composition_"
            f"2022_2024_{timestamp}.txt"
        )
    )

    field_csv_path = (
        OUTPUT_ROOT
        / (
            "hud_vms_wa_hap_field_inventory_"
            f"2022_2024_{timestamp}.csv"
        )
    )

    observation_csv_path = (
        OUTPUT_ROOT
        / (
            "hud_vms_wa_hap_composition_observations_"
            f"2022_2024_{timestamp}.csv"
        )
    )

    json_path = (
        OUTPUT_ROOT
        / (
            "profile_hud_vms_wa_hap_composition_"
            f"2022_2024_{timestamp}.json"
        )
    )

    report_path.write_text(
        render_report(
            workbook_metadata,
            field_metadata,
            analysis,
        ),
        encoding="utf-8",
    )

    write_field_inventory_csv(
        field_csv_path,
        field_metadata,
        analysis,
    )

    write_observation_csv(
        observation_csv_path,
        analysis,
    )

    json_payload = {
        "checkpoint":
            "ETL-174B-5Y-2C",
        "generated_utc":
            datetime.now(
                timezone.utc
            ).isoformat(),
        "scope":
            "Washington PHAs / Vms1 / CY2022-CY2024",
        "mode":
            "READ_ONLY_EMPIRICAL_COMPOSITION_TEST",
        "workbook_metadata":
            workbook_metadata,
        "field_metadata":
            field_metadata,
        "analysis":
            analysis,
    }

    json_path.write_text(
        json.dumps(
            json_payload,
            indent=2,
            ensure_ascii=False,
            default=str,
        )
        + "\n",
        encoding="utf-8",
    )

    primary = analysis[
        "primary_equality"
    ]

    all_other = analysis[
        "all_other_equality"
    ]

    print()
    print("=" * 88)
    print("COMPOSITION TEST CONTROL")
    print("=" * 88)
    print(
        f"Workbooks processed       : "
        f"{len(workbook_metadata)}"
    )
    print(
        f"WA PHA-month observations : "
        f"{analysis['observation_count']:,}"
    )
    print(
        f"Distinct WA PHAs          : "
        f"{analysis['distinct_pha_count']:,}"
    )
    print(
        f"Distinct months           : "
        f"{analysis['distinct_month_count']:,}"
    )
    print(
        f"Duplicate PHA-month keys  : "
        f"{analysis['duplicate_key_count']:,}"
    )
    print(
        f"Distinct HAP fields       : "
        f"{len(field_metadata):,}"
    )
    print()
    print(
        f"Primary exact matches     : "
        f"{primary['exact_count']:,}/"
        f"{primary['tested_count']:,}"
    )
    print(
        f"All-HAP exact matches     : "
        f"{all_other['exact_count']:,}/"
        f"{all_other['tested_count']:,}"
    )
    print()

    for year, value in analysis[
        "annual_hap_total"
    ].items():
        print(
            f"CY{year} HAP Total        : "
            f"{money(value)}"
        )

    print()
    print(
        "Report          : "
        f"{report_path.relative_to(REPO_ROOT)}"
    )
    print(
        "Field inventory : "
        f"{field_csv_path.relative_to(REPO_ROOT)}"
    )
    print(
        "Observations    : "
        f"{observation_csv_path.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON            : "
        f"{json_path.relative_to(REPO_ROOT)}"
    )
    print()

    structural_pass = (
        len(workbook_metadata)
        == EXPECTED_WORKBOOK_COUNT
        and analysis[
            "observation_count"
        ]
        == EXPECTED_OBSERVATIONS
        and analysis[
            "distinct_month_count"
        ]
        == EXPECTED_MONTHS
        and analysis[
            "duplicate_key_count"
        ]
        == 0
    )

    if structural_pass:
        print(
            "RESULT: PASS — COMPOSITION TEST COMPLETED"
        )
    else:
        print(
            "RESULT: REVIEW REQUIRED — "
            "POPULATION CONTROL FAILED"
        )

    print()
    print(
        "No HAP formula was automatically certified."
    )
    print(
        "No accounting-source certification was made."
    )
    print(
        "No USAspending reconciliation was performed."
    )
    print(
        "No county allocation or ledger decision was made."
    )

    return 0 if structural_pass else 2


if __name__ == "__main__":
    sys.exit(main())