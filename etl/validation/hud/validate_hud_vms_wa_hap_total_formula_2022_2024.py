"""
ETL-174B-5Y-2C.1
HUD VMS Washington HAP Total official-formula replication,
CY2022-CY2024.

PURPOSE
-------
Test whether the official HUD VMS definition of "HAP Total" reproduces
the HAP Total contained in the preserved official quarterly workbooks.

SOURCE AUTHORITY
----------------
HUD Voucher Management Systems (VMS) User Manual.

HUD states that HAP Total is a calculated field and lists the HAP
fields included in that calculation.

METHOD
------
1. Read Vms1 from all 12 preserved CY2022-CY2024 workbooks.
2. Restrict observations to Washington PHA codes.
3. Discover actual workbook headers.
4. Map HUD's documented HAP Total components to workbook headers using
   explicit aliases only.
5. STOP if any required component cannot be mapped for a workbook.
6. Sum the documented components for every PHA-month.
7. Compare the reconstructed amount with HUD's reported HAP Total.
8. Preserve every residual without adjustment.

READ ONLY.

NO:
    database writes,
    USAspending reconciliation,
    Vms2/EHV addition,
    county allocation,
    accounting-source certification,
    ledger decision.

Checkpoint:
    ETL-174B-5Y-2C.1
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# =============================================================================
# PATHS / CONTROLS
# =============================================================================

REPO_ROOT = Path(__file__).resolve().parents[3]

RAW_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "quarterly"
)

OUTPUT_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "hud"
    / "vms"
    / "validation"
    / "validate_wa_hap_total_formula_2022_2024"
)

EXPECTED_YEARS = (2022, 2023, 2024)
EXPECTED_WORKBOOK_COUNT = 12
EXPECTED_OBSERVATIONS = 1008
EXPECTED_MONTHS = 36
EXPECTED_PHAS = 28

TARGET_SHEET = "Vms1"
WA_PREFIX = "WA"

EXACT_TOLERANCE = 0.01
NEAR_TOLERANCE = 1.00


# =============================================================================
# HUD OFFICIAL FORMULA
# =============================================================================
#
# Canonical labels below follow the current HUD VMS User Manual.
#
# Each canonical component has one or more EXPLICIT normalized-header aliases.
#
# IMPORTANT:
# We do NOT use fuzzy matching.
#
# If none of the aliases exists in a workbook, the script stops for review.
# This prevents a schema change from being silently interpreted.
#

HUD_HAP_TOTAL_COMPONENTS = {

    "RAD 1 HAP": (
        "rental_assistance_demonstration_1_public_housing_conversion_hap",
    ),

    "RAD 2 HAP": (
        "rental_assistance_demonstration_2_mod_rehab_rent_supp_rap_conversion_hap",
    ),

    "RAD 3 HAP": (
        "rental_assistance_demonstration_3_prac_conversion_hap",
    ),

    "Litigation HAP": (
        "litigation_hap",
    ),

    "Homeownership HAP": (
        "homeownership_hap",
    ),

    "Moving to Work HAP": (
        "moving_to_work_hap",
    ),

    "One Year Mainstream MTW HAP": (
        "one_year_mainstream_mtw_hap",
        "one_year_mainstream_hap",
    ),

    "Family Unification Non MTW HAP": (
        "family_unification_non_mtw_hap",
        "family_unification_hap",
    ),

    "Family Unification Pre-2008 MTW HAP": (
        "family_unification_pre_2008_mtw_hap",
        "family_unification_pre2008_mtw_hap",
    ),

    "Family Unification 2008 Forward MTW HAP": (
        "family_unification_2008_forward_mtw_hap",
        "family_unification_2008_forward_hap",
    ),

    "Non-Elderly Disabled Non-MTW HAP": (
        "non_elderly_disabled_non_mtw_hap",
        "non_elderly_disabled_hap",
    ),

    "Non-Elderly Disabled 2008 Forward MTW HAP": (
        "non_elderly_disabled_2008_forward_mtw_hap",
        "non_elderly_2008_forward_mtw_hap",
        "non_elderly_disabled_2008_forward_hap",
    ),

    "Portable Voucher HAP": (
        "portable_voucher_paid_hap",
        "portable_voucher_hap",
    ),

    "HOPE VI HAP": (
        "hope_vi_hap",
        "hope_6_hap",
    ),

    "Tenant Protection HAP": (
        "tenant_protection_hap",
    ),

    "VASH HAP": (
        "veterans_affairs_supportive_housing_vash_hap",
        "veterans_affair_supported_housing_vash_hap",
        "veterans_affairs_supported_housing_vash_hap",
        "vash_hap",
    ),

    "Housing Stability Vouchers HAP": (
        "housing_stability_vouchers_hap",
        "housing_stability_vouchers_hap_expenses",
    ),

    "Foster Youth to Independence HAP": (
        "foster_youth_to_independence_hap",
        "foster_youth_to_independence_fyi_hap",
    ),

    "Mobility Demonstration HAP": (
        "mobility_demonstration_hap",
    ),

    "Emergency Transfer Vouchers for VAWA HAP": (
        "emergency_transfer_vouchers_for_vawa_hap",
        "emergency_transfer_vouchers_vawa_hap",
    ),

    "DHAP to HCV Vouchers HAP": (
        "dhap_to_hcv_voucher_hap",
    ),

    "All Other Vouchers HAP": (
        "all_other_vouchers_hap",
    ),

    "MTW FUP 2008 Forward After First HAP": (
        "mtw_family_unification_2008_forward_hap_expenses_after_the_first_of_the_month",
        "mtw_family_unification_2008_forward_hap_after_the_first_of_the_month",
    ),

    "MTW FUP Pre-2008 After First HAP": (
        "mtw_family_unification_pre_2008_hap_expenses_after_the_first_of_the_month",
        "mtw_family_unification_pre_2008_hap_after_the_first_of_the_month",
        "mtw_family_unification_pre2008_hap_expenses_after_the_first_of_the_month",
    ),

    "MTW NED 2008 Forward After First HAP": (
        "mtw_non_elderly_disabled_2008_forward_hap_expenses_after_the_first_of_the_month",
        "mtw_non_elderly_disabled_2008_forward_hap_after_the_first_of_the_month",
    ),

    "MTW VASH After First HAP": (
        "mtw_vash_hap_expenses_after_the_first_of_the_month",
        "mtw_vash_hap_after_the_first_of_the_month",
    ),

    "MTW One Year Mainstream After First HAP": (
        "mtw_one_year_mainstream_hap_expenses_after_the_first_of_the_month",
        "mtw_one_year_mainstream_hap_after_the_first_of_the_month",
    ),

    "MTW FYI After First HAP": (
        "mtw_foster_youth_to_independence_hap_expenses_after_the_first_of_the_month",
        "mtw_foster_youth_to_independence_hap_after_the_first_of_the_month",
    ),

    "MTW Mobility Demonstration After First HAP": (
        "mtw_moblity_demonstration_hap_expenses_after_the_first_of_the_month",
    ),

    "MTW Housing Stability After First HAP": (
        "mtw_housing_stability_voucher_hap_expenses_after_the_first_of_the_month",
        "mtw_housing_stability_vouchers_hap_expenses_after_the_first_of_the_month",
        "mtw_housing_stability_voucher_hap_after_the_first_of_the_month",
    ),

    "MTW VAWA After First HAP": (
        "mtw_emergency_transfer_vouchers_for_vawa_hap_expenses_after_the_first_of_the_month",
        "mtw_emergency_transfer_vouchers_for_vawa_hap_after_the_first_of_the_month",
    ),

    "HCV FSS Escrow Deposits": (
        "hcv_fss_escrow_deposits",
        "fss_escrow_deposits",
    ),

    "All Voucher HAP After First of Month": (
        "hcv_all_voucher_hap_expenses_after_the_first_of_the_month",
        "all_voucher_hap_expenses_after_the_first_of_the_month",
        "all_voucher_hap_expenses_for_contracts_effective_after_the_first_of_month",
    ),
}


# =============================================================================
# BASIC HELPERS
# =============================================================================

def clean_text(value: Any) -> str:
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


def normalize_header(value: Any) -> str:
    text = clean_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def numeric(value: Any) -> float | None:
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)

    text = clean_text(value)

    if not text:
        return None

    text = (
        text
        .replace("$", "")
        .replace(",", "")
        .strip()
    )

    try:
        return float(text)
    except ValueError:
        return None


def parse_month(value: Any) -> date | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = clean_text(value)

    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m/%d/%y",
    )

    for fmt in formats:
        try:
            return datetime.strptime(
                text,
                fmt,
            ).date()
        except ValueError:
            pass

    return None


def month_key(value: date) -> str:
    return value.strftime("%Y-%m")


def money(value: float | None) -> str:
    if value is None:
        return "NULL"

    return f"${value:,.2f}"


# =============================================================================
# WORKBOOK DISCOVERY
# =============================================================================

def discover_workbooks() -> list[Path]:
    paths = []

    for year in EXPECTED_YEARS:
        year_dir = RAW_ROOT / f"CY{year}"

        if not year_dir.exists():
            raise RuntimeError(
                f"Missing expected directory: {year_dir}"
            )

        paths.extend(
            sorted(year_dir.glob("*.xlsx"))
        )

    if len(paths) != EXPECTED_WORKBOOK_COUNT:
        raise RuntimeError(
            f"Expected {EXPECTED_WORKBOOK_COUNT} workbooks; "
            f"found {len(paths)}."
        )

    return paths


# =============================================================================
# FORMULA MAPPING
# =============================================================================

def map_formula_fields(
    headers: dict[str, int],
) -> tuple[
    dict[str, str],
    dict[str, tuple[str, ...]],
]:

    mapping = {}
    missing = {}

    for canonical, aliases in (
        HUD_HAP_TOTAL_COMPONENTS.items()
    ):
        matches = [
            alias
            for alias in aliases
            if alias in headers
        ]

        if len(matches) == 1:
            mapping[canonical] = matches[0]

        elif len(matches) > 1:
            raise RuntimeError(
                f"Ambiguous mapping for {canonical}: "
                + ", ".join(matches)
            )

        else:
            missing[canonical] = aliases

    return mapping, missing


# =============================================================================
# EXTRACTION
# =============================================================================

def extract_workbook(
    path: Path,
) -> tuple[
    list[dict[str, Any]],
    dict[str, Any],
]:

    wb = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:
        if TARGET_SHEET not in wb.sheetnames:
            raise RuntimeError(
                f"{path.name}: missing {TARGET_SHEET}"
            )

        ws = wb[TARGET_SHEET]
        rows = ws.iter_rows(values_only=True)

        try:
            raw_headers = next(rows)
        except StopIteration:
            raise RuntimeError(
                f"{path.name}: empty {TARGET_SHEET}"
            )

        headers = {}

        for position, raw in enumerate(
            raw_headers
        ):
            normalized = normalize_header(raw)

            if (
                normalized
                and normalized not in headers
            ):
                headers[normalized] = position

        required_identity = (
            "pha_code",
            "pha_name",
            "submission_month",
            "hap_total",
        )

        missing_identity = [
            field
            for field in required_identity
            if field not in headers
        ]

        if missing_identity:
            raise RuntimeError(
                f"{path.name}: missing identity fields: "
                + ", ".join(missing_identity)
            )

        formula_mapping, missing_formula = (
            map_formula_fields(headers)
        )

        metadata = {
            "source_workbook": path.name,
            "header_count": len(headers),
            "formula_component_count":
                len(HUD_HAP_TOTAL_COMPONENTS),
            "mapped_component_count":
                len(formula_mapping),
            "formula_mapping":
                formula_mapping,
            "missing_formula_components":
                {
                    key: list(value)
                    for key, value
                    in missing_formula.items()
                },
        }

        # STRICT STOP.
        #
        # We do not compute a partial formula.
        if missing_formula:
            return [], metadata

        extracted = []

        for worksheet_row, row in enumerate(
            rows,
            start=2,
        ):
            pha_position = headers["pha_code"]

            if pha_position >= len(row):
                continue

            pha_code = clean_text(
                row[pha_position]
            )

            if not pha_code.upper().startswith(
                WA_PREFIX
            ):
                continue

            submission_position = headers[
                "submission_month"
            ]

            submission_value = (
                row[submission_position]
                if submission_position < len(row)
                else None
            )

            submission_month = parse_month(
                submission_value
            )

            if submission_month is None:
                raise RuntimeError(
                    f"{path.name} row {worksheet_row}: "
                    "invalid submission month."
                )

            hap_total_position = headers[
                "hap_total"
            ]

            hap_total = numeric(
                row[hap_total_position]
                if hap_total_position < len(row)
                else None
            )

            if hap_total is None:
                raise RuntimeError(
                    f"{path.name} row {worksheet_row}: "
                    "HAP Total is not numeric."
                )

            components = {}

            for canonical, actual_header in (
                formula_mapping.items()
            ):
                position = headers[
                    actual_header
                ]

                value = numeric(
                    row[position]
                    if position < len(row)
                    else None
                )

                # HUD workbooks may represent a blank component as
                # no activity. Preserve both raw-null status and use
                # zero only for the arithmetic replication.
                components[canonical] = {
                    "header":
                        actual_header,
                    "value":
                        value,
                    "arithmetic_value":
                        0.0 if value is None else value,
                }

            reconstructed = sum(
                item["arithmetic_value"]
                for item in components.values()
            )

            residual = (
                hap_total - reconstructed
            )

            extracted.append(
                {
                    "source_workbook":
                        path.name,
                    "source_row_number":
                        worksheet_row,
                    "pha_code":
                        pha_code,
                    "pha_name":
                        clean_text(
                            row[
                                headers["pha_name"]
                            ]
                        ),
                    "submission_month":
                        submission_month.isoformat(),
                    "submission_month_key":
                        month_key(
                            submission_month
                        ),
                    "hap_total":
                        hap_total,
                    "reconstructed_hap_total":
                        reconstructed,
                    "residual":
                        residual,
                    "absolute_residual":
                        abs(residual),
                    "components":
                        components,
                }
            )

        metadata["wa_rows"] = len(
            extracted
        )

        return extracted, metadata

    finally:
        wb.close()


# =============================================================================
# ANALYSIS
# =============================================================================

def analyze(
    records: list[dict[str, Any]],
) -> dict[str, Any]:

    keys = Counter(
        (
            record["pha_code"],
            record["submission_month_key"],
        )
        for record in records
    )

    months = sorted(
        {
            record["submission_month_key"]
            for record in records
        }
    )

    phas = sorted(
        {
            record["pha_code"]
            for record in records
        }
    )

    residuals = [
        record["residual"]
        for record in records
    ]

    exact = [
        value
        for value in residuals
        if abs(value) <= EXACT_TOLERANCE
    ]

    near = [
        value
        for value in residuals
        if abs(value) <= NEAR_TOLERANCE
    ]

    non_exact = [
        record
        for record in records
        if abs(
            record["residual"]
        ) > EXACT_TOLERANCE
    ]

    annual_reported = defaultdict(float)
    annual_reconstructed = defaultdict(float)

    component_totals = defaultdict(float)
    component_non_null = Counter()
    component_positive = Counter()

    for record in records:
        year = int(
            record[
                "submission_month_key"
            ][:4]
        )

        annual_reported[year] += (
            record["hap_total"]
        )

        annual_reconstructed[year] += (
            record[
                "reconstructed_hap_total"
            ]
        )

        for canonical, item in (
            record["components"].items()
        ):
            if item["value"] is not None:
                component_non_null[
                    canonical
                ] += 1

                component_totals[
                    canonical
                ] += item["value"]

                if item["value"] > 0:
                    component_positive[
                        canonical
                    ] += 1

    return {
        "observation_count":
            len(records),

        "distinct_month_count":
            len(months),

        "distinct_pha_count":
            len(phas),

        "duplicate_key_count":
            sum(
                count > 1
                for count in keys.values()
            ),

        "exact_count":
            len(exact),

        "near_count":
            len(near),

        "non_exact_count":
            len(non_exact),

        "exact_rate":
            (
                len(exact) / len(residuals)
                if residuals
                else None
            ),

        "residual_sum":
            (
                sum(residuals)
                if residuals
                else None
            ),

        "absolute_residual_sum":
            (
                sum(
                    abs(value)
                    for value in residuals
                )
                if residuals
                else None
            ),

        "residual_min":
            (
                min(residuals)
                if residuals
                else None
            ),

        "residual_max":
            (
                max(residuals)
                if residuals
                else None
            ),

        "annual_reported":
            dict(
                sorted(
                    annual_reported.items()
                )
            ),

        "annual_reconstructed":
            dict(
                sorted(
                    annual_reconstructed.items()
                )
            ),

        "component_totals":
            dict(
                sorted(
                    component_totals.items()
                )
            ),

        "component_non_null":
            dict(
                sorted(
                    component_non_null.items()
                )
            ),

        "component_positive":
            dict(
                sorted(
                    component_positive.items()
                )
            ),

        "non_exact_records":
            non_exact,
    }


# =============================================================================
# OUTPUT
# =============================================================================

def write_observation_csv(
    path: Path,
    records: list[dict[str, Any]],
) -> None:

    fields = [
        "source_workbook",
        "source_row_number",
        "pha_code",
        "pha_name",
        "submission_month",
        "submission_month_key",
        "hap_total",
        "reconstructed_hap_total",
        "residual",
        "absolute_residual",
    ]

    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=fields,
        )

        writer.writeheader()

        for record in records:
            writer.writerow(
                {
                    field: record.get(field)
                    for field in fields
                }
            )


def render_report(
    workbook_metadata: list[dict[str, Any]],
    analysis: dict[str, Any] | None,
) -> str:

    lines = []

    def emit(text: str = "") -> None:
        lines.append(text)

    emit("=" * 100)
    emit("HUD VMS WASHINGTON HAP TOTAL OFFICIAL-FORMULA REPLICATION")
    emit("ETL-174B-5Y-2C.1")
    emit("=" * 100)
    emit()
    emit(
        "Official-definition replication test. "
        "No heuristic or fuzzy field matching."
    )
    emit()

    emit("HUD FORMULA CONTROL")
    emit("-" * 100)
    emit(
        f"Documented formula components : "
        f"{len(HUD_HAP_TOTAL_COMPONENTS)}"
    )
    emit()

    for number, canonical in enumerate(
        HUD_HAP_TOTAL_COMPONENTS,
        start=1,
    ):
        emit(
            f"{number:>2}. {canonical}"
        )

    emit()
    emit("=" * 100)
    emit("WORKBOOK FIELD MAPPING")
    emit("=" * 100)

    mapping_failure = False

    for metadata in workbook_metadata:
        emit()
        emit(metadata["source_workbook"])
        emit(
            f"    mapped components : "
            f"{metadata['mapped_component_count']}/"
            f"{metadata['formula_component_count']}"
        )

        missing = metadata[
            "missing_formula_components"
        ]

        if missing:
            mapping_failure = True

            emit(
                "    RESULT            : "
                "HOLD — UNMAPPED OFFICIAL COMPONENTS"
            )

            for canonical, aliases in (
                missing.items()
            ):
                emit(
                    f"        {canonical}"
                )
                emit(
                    "            aliases tried: "
                    + " | ".join(aliases)
                )

        else:
            emit(
                "    RESULT            : "
                "MAPPING COMPLETE"
            )

            for canonical, header in (
                metadata[
                    "formula_mapping"
                ].items()
            ):
                emit(
                    f"        {canonical}"
                    f" -> {header}"
                )

    if mapping_failure:
        emit()
        emit("=" * 100)
        emit("RESULT")
        emit("=" * 100)
        emit()
        emit(
            "HOLD — OFFICIAL HUD FORMULA COULD NOT "
            "BE MAPPED COMPLETELY."
        )
        emit()
        emit(
            "No partial formula was calculated."
        )
        emit(
            "Review actual workbook headers before adding aliases."
        )

        return "\n".join(lines) + "\n"

    assert analysis is not None

    emit()
    emit("=" * 100)
    emit("POPULATION CONTROL")
    emit("=" * 100)
    emit()
    emit(
        f"PHA-month observations : "
        f"{analysis['observation_count']:,}"
    )
    emit(
        f"Distinct WA PHAs        : "
        f"{analysis['distinct_pha_count']:,}"
    )
    emit(
        f"Distinct months         : "
        f"{analysis['distinct_month_count']:,}"
    )
    emit(
        f"Duplicate keys          : "
        f"{analysis['duplicate_key_count']:,}"
    )

    emit()
    emit("=" * 100)
    emit("OFFICIAL FORMULA REPLICATION")
    emit("=" * 100)
    emit()
    emit(
        f"Exact within $0.01      : "
        f"{analysis['exact_count']:,}/"
        f"{analysis['observation_count']:,}"
    )
    emit(
        f"Near within $1.00       : "
        f"{analysis['near_count']:,}/"
        f"{analysis['observation_count']:,}"
    )

    exact_rate = analysis[
        "exact_rate"
    ]

    emit(
        "Exact rate              : "
        + (
            "NULL"
            if exact_rate is None
            else f"{exact_rate:.6%}"
        )
    )

    emit(
        f"Signed residual sum     : "
        f"{money(analysis['residual_sum'])}"
    )
    emit(
        f"Absolute residual sum   : "
        f"{money(analysis['absolute_residual_sum'])}"
    )
    emit(
        f"Minimum residual        : "
        f"{money(analysis['residual_min'])}"
    )
    emit(
        f"Maximum residual        : "
        f"{money(analysis['residual_max'])}"
    )

    emit()
    emit("=" * 100)
    emit("ANNUAL REPLICATION CONTROL")
    emit("=" * 100)
    emit()

    for year in EXPECTED_YEARS:
        reported = analysis[
            "annual_reported"
        ].get(year)

        reconstructed = analysis[
            "annual_reconstructed"
        ].get(year)

        difference = (
            None
            if (
                reported is None
                or reconstructed is None
            )
            else reported - reconstructed
        )

        emit(
            f"CY{year}"
        )
        emit(
            f"    HUD HAP Total        : "
            f"{money(reported)}"
        )
        emit(
            f"    Reconstructed        : "
            f"{money(reconstructed)}"
        )
        emit(
            f"    Difference           : "
            f"{money(difference)}"
        )

    emit()
    emit("=" * 100)
    emit("OFFICIAL COMPONENT TOTALS")
    emit("=" * 100)

    for canonical in (
        HUD_HAP_TOTAL_COMPONENTS
    ):
        emit()
        emit(canonical)
        emit(
            f"    non-null observations : "
            f"{analysis['component_non_null'].get(canonical, 0):,}"
        )
        emit(
            f"    positive observations : "
            f"{analysis['component_positive'].get(canonical, 0):,}"
        )
        emit(
            f"    three-year sum         : "
            f"{money(analysis['component_totals'].get(canonical, 0.0))}"
        )

    if analysis["non_exact_count"]:
        emit()
        emit("=" * 100)
        emit("NON-EXACT OBSERVATIONS")
        emit("=" * 100)
        emit()

        for record in analysis[
            "non_exact_records"
        ]:
            emit(
                f"{record['pha_code']} "
                f"{record['submission_month_key']} "
                f"HUD={money(record['hap_total'])} "
                f"reconstructed="
                f"{money(record['reconstructed_hap_total'])} "
                f"residual="
                f"{money(record['residual'])}"
            )

    emit()
    emit("=" * 100)
    emit("METHODOLOGICAL STOP")
    emit("=" * 100)
    emit()
    emit(
        "This test evaluates replication of HUD's documented "
        "HAP Total calculation."
    )
    emit(
        "It does not by itself certify HAP Total as the "
        "WCFFBP accounting source."
    )
    emit(
        "No USAspending reconciliation was performed."
    )
    emit(
        "No Vms2/EHV values were added."
    )
    emit(
        "No county allocation was performed."
    )
    emit(
        "No database writes occurred."
    )
    emit(
        "No ledger decision was made."
    )

    return "\n".join(lines) + "\n"


# =============================================================================
# MAIN
# =============================================================================

def main() -> int:

    print("=" * 88)
    print("HUD VMS WASHINGTON HAP TOTAL OFFICIAL-FORMULA REPLICATION")
    print("ETL-174B-5Y-2C.1")
    print("=" * 88)
    print()
    print(f"Input : {RAW_ROOT}")
    print(f"Output: {OUTPUT_ROOT}")
    print()
    print(
        f"HUD documented components : "
        f"{len(HUD_HAP_TOTAL_COMPONENTS)}"
    )
    print("Mode                      : READ ONLY")
    print("DB                        : NO DATABASE WRITES")
    print()

    try:
        paths = discover_workbooks()
    except RuntimeError as exc:
        print(f"ERROR: {exc}")
        return 1

    OUTPUT_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    all_records = []
    workbook_metadata = []

    mapping_failure = False

    for index, path in enumerate(
        paths,
        start=1,
    ):
        print(
            f"[{index:02}/{len(paths):02}] "
            f"{path.name}",
            flush=True,
        )

        try:
            records, metadata = (
                extract_workbook(path)
            )

        except Exception as exc:
            print(
                f"         ERROR: {exc}",
                flush=True,
            )
            return 1

        workbook_metadata.append(
            metadata
        )

        if metadata[
            "missing_formula_components"
        ]:
            mapping_failure = True

            print(
                f"         mapped="
                f"{metadata['mapped_component_count']}/"
                f"{metadata['formula_component_count']} "
                f"— HOLD",
                flush=True,
            )

        else:
            all_records.extend(records)

            print(
                f"         mapped="
                f"{metadata['mapped_component_count']}/"
                f"{metadata['formula_component_count']}; "
                f"WA rows={len(records):,}",
                flush=True,
            )

    analysis = None

    if not mapping_failure:
        analysis = analyze(
            all_records
        )

    timestamp = datetime.now(
        timezone.utc
    ).strftime("%Y%m%dT%H%M%SZ")

    report_path = (
        OUTPUT_ROOT
        / (
            "validate_hud_vms_wa_hap_total_formula_"
            f"2022_2024_{timestamp}.txt"
        )
    )

    json_path = (
        OUTPUT_ROOT
        / (
            "validate_hud_vms_wa_hap_total_formula_"
            f"2022_2024_{timestamp}.json"
        )
    )

    report_path.write_text(
        render_report(
            workbook_metadata,
            analysis,
        ),
        encoding="utf-8",
    )

    json_path.write_text(
        json.dumps(
            {
                "checkpoint":
                    "ETL-174B-5Y-2C.1",
                "generated_utc":
                    datetime.now(
                        timezone.utc
                    ).isoformat(),
                "mode":
                    "READ_ONLY_OFFICIAL_FORMULA_REPLICATION",
                "hud_documented_components":
                    list(
                        HUD_HAP_TOTAL_COMPONENTS
                    ),
                "workbook_metadata":
                    workbook_metadata,
                "analysis":
                    analysis,
            },
            indent=2,
            ensure_ascii=False,
            default=str,
        )
        + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 88)
    print("OFFICIAL FORMULA CONTROL")
    print("=" * 88)

    if mapping_failure:
        print(
            "Formula mapping          : HOLD"
        )
        print(
            "Partial calculation      : NOT PERFORMED"
        )
        print()
        print(
            "One or more HUD-documented components "
            "could not be mapped."
        )
        print(
            "This is expected to be resolvable from "
            "the actual workbook headers."
        )

    else:
        assert analysis is not None

        print(
            f"Formula mapping          : "
            f"{len(HUD_HAP_TOTAL_COMPONENTS)}/"
            f"{len(HUD_HAP_TOTAL_COMPONENTS)}"
        )
        print(
            f"WA observations          : "
            f"{analysis['observation_count']:,}"
        )
        print(
            f"Distinct WA PHAs         : "
            f"{analysis['distinct_pha_count']:,}"
        )
        print(
            f"Distinct months          : "
            f"{analysis['distinct_month_count']:,}"
        )
        print(
            f"Duplicate keys           : "
            f"{analysis['duplicate_key_count']:,}"
        )
        print()
        print(
            f"Exact formula matches    : "
            f"{analysis['exact_count']:,}/"
            f"{analysis['observation_count']:,}"
        )
        print(
            f"Near formula matches     : "
            f"{analysis['near_count']:,}/"
            f"{analysis['observation_count']:,}"
        )
        print(
            f"Signed residual          : "
            f"{money(analysis['residual_sum'])}"
        )
        print(
            f"Absolute residual        : "
            f"{money(analysis['absolute_residual_sum'])}"
        )

        print()

        for year in EXPECTED_YEARS:
            reported = analysis[
                "annual_reported"
            ].get(year)

            reconstructed = analysis[
                "annual_reconstructed"
            ].get(year)

            print(
                f"CY{year} HUD / reconstructed : "
                f"{money(reported)} / "
                f"{money(reconstructed)}"
            )

    print()
    print(
        "Report : "
        f"{report_path.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON   : "
        f"{json_path.relative_to(REPO_ROOT)}"
    )
    print()

    if mapping_failure:
        print(
            "RESULT: HOLD — OFFICIAL FORMULA "
            "MAPPING INCOMPLETE"
        )
        return 2

    population_pass = (
        analysis["observation_count"]
        == EXPECTED_OBSERVATIONS
        and analysis["distinct_month_count"]
        == EXPECTED_MONTHS
        and analysis["distinct_pha_count"]
        == EXPECTED_PHAS
        and analysis["duplicate_key_count"]
        == 0
    )

    formula_pass = (
        analysis["exact_count"]
        == EXPECTED_OBSERVATIONS
    )

    if population_pass and formula_pass:
        print(
            "RESULT: PASS — HUD OFFICIAL HAP TOTAL "
            "FORMULA EXACTLY REPLICATED"
        )
        return 0

    if population_pass:
        print(
            "RESULT: REVIEW REQUIRED — POPULATION PASSES "
            "BUT FORMULA HAS RESIDUALS"
        )
        return 2

    print(
        "RESULT: REVIEW REQUIRED — "
        "POPULATION CONTROL FAILED"
    )
    return 2


if __name__ == "__main__":
    sys.exit(main())