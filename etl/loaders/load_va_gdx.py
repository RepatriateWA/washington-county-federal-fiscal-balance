"""
ETL-VA-GDX-001 — Load VA Geographic Distribution of Expenditures (GDX)
Washington county observations for FY2015-FY2024.

Source regimes
--------------
FY2015-FY2023:
    LEGACY_STATE_SHEETS
    Worksheet: WA
    Source expenditure unit: $000s

FY2024:
    REVISED_DIRECT_INDIRECT
    Worksheets:
        County Direct Exp
        County Indirect Exp
    Source expenditure unit: dollars

Design principles
-----------------
- Raw workbooks are read-only and never modified.
- Source row/column provenance is preserved.
- Legacy $000s are normalized to dollars while retaining source units.
- FY2024 DIRECT and INDIRECT classes remain explicit.
- Veteran Population and Unique Patients are loaded as metrics, never
  as expenditures.
- County names must resolve deterministically to reference.county_registry.
- Rerunning an already completed identical fiscal-year/source-hash batch
  skips it unless --force is supplied.
- Unexpected headers, county populations, or category mappings stop the load.

Examples
--------
Load all FY2015-FY2024 workbooks:

    python -m etl.loaders.load_va_gdx

Load selected years:

    python -m etl.loaders.load_va_gdx --years 2015 2024

Force reload an already completed identical batch:

    python -m etl.loaders.load_va_gdx --years 2024 --force
"""

from __future__ import annotations

import argparse
import hashlib
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from psycopg2.extras import execute_values

from etl.common.db import get_connection


REPO_ROOT = Path(__file__).resolve().parents[2]
RAW_ROOT = REPO_ROOT / "data" / "raw" / "va" / "gdx"

DEFAULT_YEARS = tuple(range(2015, 2025))
EXPECTED_COUNTY_COUNT = 39

LEGACY_SCHEMA = "LEGACY_STATE_SHEETS"
REVISED_SCHEMA = "REVISED_DIRECT_INDIRECT"

LEGACY_METHODOLOGY = "LEGACY_PRE_FY2024"
REVISED_METHODOLOGY = "REVISED_FY2024"

USD = "USD"
USD_THOUSANDS = "USD_THOUSANDS"

LEGACY_EXPECTED_HEADERS = [
    "County/ Congressional District",
    "Veteran Population*",
    "Total Expenditure",
    "Compensation & Pension",
    "Construction",
    "Education & Vocational Rehabilitation/ Employment",
    "Loan Guaranty#",
    "General Operating Expenses",
    "Insurance & Indemnities",
    "Medical Care",
    "Unique Patients**",
]

FY2024_DIRECT_EXPECTED_HEADERS = [
    "State",
    "County",
    "Total Direct Expenditure",
    "Compensation & Pension",
    "Insurance",
    "Dependency & Indemnity Compensation and Survivors Pension",
]

FY2024_INDIRECT_EXPECTED_HEADERS = [
    "State",
    "County",
    "Total Indirect Expenditure**",
    "Memorial Benefits",
    "Construction",
    "Education & Vocational Rehabilitation/ Employment",
    "Loan Guaranty",
    "Medical Care**",
    "Specially Adapted Housing and Automobiles and Adaptive Equipment",
]


LEGACY_EXPENDITURE_MAP = {
    "Total Expenditure": "TOTAL_EXPENDITURE",
    "Compensation & Pension": "COMPENSATION_PENSION",
    "Construction": "CONSTRUCTION",
    "Education & Vocational Rehabilitation/ Employment": "EDUCATION_VRE",
    "Loan Guaranty#": "LOAN_GUARANTY",
    "General Operating Expenses": "GENERAL_OPERATING_EXPENSES",
    "Insurance & Indemnities": "INSURANCE_INDEMNITIES",
    "Medical Care": "MEDICAL_CARE",
}

LEGACY_METRIC_MAP = {
    "Veteran Population*": (
        "VETERAN_POPULATION",
        "Veteran Population",
        "PERSON_ESTIMATE",
    ),
    "Unique Patients**": (
        "UNIQUE_PATIENTS",
        "Unique Patients",
        "PERSON_COUNT",
    ),
}

FY2024_DIRECT_MAP = {
    "Total Direct Expenditure": "TOTAL_DIRECT_EXPENDITURE",
    "Compensation & Pension": "COMPENSATION_PENSION",
    "Insurance": "INSURANCE",
    "Dependency & Indemnity Compensation and Survivors Pension":
        "DIC_SURVIVORS_PENSION",
}

FY2024_INDIRECT_MAP = {
    "Total Indirect Expenditure**": "TOTAL_INDIRECT_EXPENDITURE",
    "Memorial Benefits": "MEMORIAL_BENEFITS",
    "Construction": "CONSTRUCTION",
    "Education & Vocational Rehabilitation/ Employment": "EDUCATION_VRE",
    "Loan Guaranty": "LOAN_GUARANTY",
    "Medical Care**": "MEDICAL_CARE",
    "Specially Adapted Housing and Automobiles and Adaptive Equipment":
        "SPECIALLY_ADAPTED_HOUSING_AUTOMOBILES",
}


@dataclass(frozen=True)
class CountyReference:
    county_id: int
    county_fips: str
    full_fips: str
    county_name: str


@dataclass(frozen=True)
class ExpenditureObservation:
    fiscal_year: int
    county_id: int
    county_name_source: str
    schema_regime: str
    expenditure_class: str
    category_code: str
    source_category_name: str
    source_amount: Decimal | None
    source_unit: str
    expenditure_amount: Decimal | None
    source_sheet: str
    source_row_number: int
    source_column_number: int


@dataclass(frozen=True)
class MetricObservation:
    fiscal_year: int
    county_id: int
    county_name_source: str
    metric_code: str
    metric_name: str
    source_value: Decimal | None
    source_unit: str
    source_sheet: str
    source_row_number: int
    source_column_number: int


@dataclass(frozen=True)
class ParsedYear:
    fiscal_year: int
    source_file: Path
    source_sha256: str
    schema_regime: str
    methodology_regime: str
    methodology_break: bool
    source_county_rows: int
    expenditures: list[ExpenditureObservation]
    metrics: list[MetricObservation]


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Load Washington county VA GDX expenditures "
            "for FY2015-FY2024."
        )
    )
    parser.add_argument(
        "--years",
        type=int,
        nargs="*",
        default=list(DEFAULT_YEARS),
        help=(
            "Fiscal years to load. Default: 2015 through 2024."
        ),
    )
    parser.add_argument(
        "--source-root",
        type=Path,
        default=RAW_ROOT,
        help=f"VA GDX raw root. Default: {RAW_ROOT}",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help=(
            "Reload an identical already-completed fiscal-year/source-hash "
            "batch instead of skipping it."
        ),
    )
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def discover_workbook(source_root: Path, fiscal_year: int) -> Path:
    year_dir = source_root / f"FY{fiscal_year}"

    if not year_dir.exists():
        raise FileNotFoundError(
            f"FY{fiscal_year} source directory not found: {year_dir}"
        )

    candidates = sorted(
        p
        for p in year_dir.glob("*.xlsx")
        if not p.name.startswith("~$")
    )

    if len(candidates) != 1:
        names = ", ".join(p.name for p in candidates) or "NONE"
        raise RuntimeError(
            f"FY{fiscal_year}: expected exactly one XLSX source workbook "
            f"in {year_dir}; found {len(candidates)}: {names}"
        )

    return candidates[0]


def clean_header(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def nonblank_headers(values: Iterable[Any]) -> list[str]:
    return [
        header
        for header in (clean_header(v) for v in values)
        if header is not None
    ]


def to_decimal(value: Any) -> Decimal | None:
    if value is None:
        return None

    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        text = text.replace(",", "").replace("$", "")
    else:
        text = str(value)

    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(
            f"Invalid numeric source value: {value!r}"
        ) from exc


def normalize_dollars(
    source_amount: Decimal | None,
    source_unit: str,
) -> Decimal | None:
    if source_amount is None:
        return None

    if source_unit == USD_THOUSANDS:
        amount = source_amount * Decimal("1000")
    elif source_unit == USD:
        amount = source_amount
    else:
        raise ValueError(
            f"Unsupported source unit: {source_unit}"
        )

    return amount.quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )


def normalize_county_key(value: str) -> str:
    text = value.upper().strip()
    text = re.sub(r"\bCOUNTY\b", "", text)
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def load_washington_county_reference(cur) -> dict[str, CountyReference]:
    cur.execute(
        """
        SELECT
            c.county_id,
            c.county_fips,
            c.full_fips,
            c.county_name
        FROM reference.county_registry AS c
        JOIN reference.state_registry AS s
          ON s.state_id = c.state_id
        WHERE s.state_abbreviation = 'WA'
        ORDER BY c.county_fips;
        """
    )

    rows = cur.fetchall()

    if len(rows) != EXPECTED_COUNTY_COUNT:
        raise RuntimeError(
            "reference.county_registry does not contain exactly "
            f"{EXPECTED_COUNTY_COUNT} Washington counties; "
            f"found {len(rows)}."
        )

    mapping: dict[str, CountyReference] = {}

    for county_id, county_fips, full_fips, county_name in rows:
        ref = CountyReference(
            county_id=int(county_id),
            county_fips=str(county_fips),
            full_fips=str(full_fips),
            county_name=str(county_name),
        )
        key = normalize_county_key(ref.county_name)

        if key in mapping:
            raise RuntimeError(
                f"Duplicate normalized Washington county key: {key}"
            )

        mapping[key] = ref

    return mapping


def resolve_county(
    source_name: str,
    county_map: dict[str, CountyReference],
) -> CountyReference:
    key = normalize_county_key(source_name)
    ref = county_map.get(key)

    if ref is None:
        raise ValueError(
            f"VA GDX county name did not resolve to "
            f"reference.county_registry: {source_name!r} "
            f"(normalized={key!r})"
        )

    return ref


def assert_exact_headers(
    actual: list[str],
    expected: list[str],
    *,
    fiscal_year: int,
    sheet_name: str,
) -> None:
    if actual != expected:
        raise RuntimeError(
            f"FY{fiscal_year} worksheet {sheet_name!r} headers differ "
            "from the certified schema profile.\n"
            f"Expected: {expected}\n"
            f"Actual:   {actual}"
        )


def parse_legacy_year(
    fiscal_year: int,
    workbook_path: Path,
    county_map: dict[str, CountyReference],
) -> ParsedYear:
    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        if "WA" not in workbook.sheetnames:
            raise RuntimeError(
                f"FY{fiscal_year} legacy workbook has no WA worksheet."
            )

        ws = workbook["WA"]

        headers = nonblank_headers(
            next(
                ws.iter_rows(
                    min_row=3,
                    max_row=3,
                    values_only=True,
                )
            )
        )

        assert_exact_headers(
            headers,
            LEGACY_EXPECTED_HEADERS,
            fiscal_year=fiscal_year,
            sheet_name="WA",
        )

        header_to_column = {
            header: index
            for index, header in enumerate(
                [
                    clean_header(v)
                    for v in next(
                        ws.iter_rows(
                            min_row=3,
                            max_row=3,
                            values_only=True,
                        )
                    )
                ],
                start=1,
            )
            if header is not None
        }

        expenditures: list[ExpenditureObservation] = []
        metrics: list[MetricObservation] = []

        county_names_seen: list[str] = []

        for source_row_number, row in enumerate(
            ws.iter_rows(
                min_row=4,
                values_only=True,
            ),
            start=4,
        ):
            raw_county = row[0]

            if raw_county is None:
                continue

            county_name_source = str(raw_county).strip()

            if not county_name_source:
                continue

            # Legacy WA worksheets contain the 39 county records first,
            # followed by non-county summary/CD material.
            if len(county_names_seen) >= EXPECTED_COUNTY_COUNT:
                break

            county = resolve_county(
                county_name_source,
                county_map,
            )

            county_names_seen.append(
                county.county_name
            )

            for source_header, category_code in LEGACY_EXPENDITURE_MAP.items():
                source_column_number = header_to_column[source_header]
                source_value = row[source_column_number - 1]
                source_amount = to_decimal(source_value)

                expenditures.append(
                    ExpenditureObservation(
                        fiscal_year=fiscal_year,
                        county_id=county.county_id,
                        county_name_source=county_name_source,
                        schema_regime=LEGACY_SCHEMA,
                        expenditure_class="LEGACY",
                        category_code=category_code,
                        source_category_name=source_header,
                        source_amount=source_amount,
                        source_unit=USD_THOUSANDS,
                        expenditure_amount=normalize_dollars(
                            source_amount,
                            USD_THOUSANDS,
                        ),
                        source_sheet="WA",
                        source_row_number=source_row_number,
                        source_column_number=source_column_number,
                    )
                )

            for source_header, (
                metric_code,
                metric_name,
                metric_unit,
            ) in LEGACY_METRIC_MAP.items():
                source_column_number = header_to_column[source_header]
                source_value = row[source_column_number - 1]

                metrics.append(
                    MetricObservation(
                        fiscal_year=fiscal_year,
                        county_id=county.county_id,
                        county_name_source=county_name_source,
                        metric_code=metric_code,
                        metric_name=metric_name,
                        source_value=to_decimal(source_value),
                        source_unit=metric_unit,
                        source_sheet="WA",
                        source_row_number=source_row_number,
                        source_column_number=source_column_number,
                    )
                )

        if len(county_names_seen) != EXPECTED_COUNTY_COUNT:
            raise RuntimeError(
                f"FY{fiscal_year}: parsed {len(county_names_seen)} "
                "Washington county rows; expected 39."
            )

        if len(set(county_names_seen)) != EXPECTED_COUNTY_COUNT:
            raise RuntimeError(
                f"FY{fiscal_year}: duplicate Washington county rows "
                "were detected."
            )

        expected_expenditure_count = (
            EXPECTED_COUNTY_COUNT
            * len(LEGACY_EXPENDITURE_MAP)
        )
        expected_metric_count = (
            EXPECTED_COUNTY_COUNT
            * len(LEGACY_METRIC_MAP)
        )

        if len(expenditures) != expected_expenditure_count:
            raise RuntimeError(
                f"FY{fiscal_year}: expected "
                f"{expected_expenditure_count} expenditure observations; "
                f"parsed {len(expenditures)}."
            )

        if len(metrics) != expected_metric_count:
            raise RuntimeError(
                f"FY{fiscal_year}: expected "
                f"{expected_metric_count} metric observations; "
                f"parsed {len(metrics)}."
            )

        return ParsedYear(
            fiscal_year=fiscal_year,
            source_file=workbook_path,
            source_sha256=sha256_file(workbook_path),
            schema_regime=LEGACY_SCHEMA,
            methodology_regime=LEGACY_METHODOLOGY,
            methodology_break=False,
            source_county_rows=EXPECTED_COUNTY_COUNT,
            expenditures=expenditures,
            metrics=metrics,
        )

    finally:
        workbook.close()


def parse_fy2024_sheet(
    *,
    fiscal_year: int,
    ws,
    sheet_name: str,
    expected_headers: list[str],
    category_map: dict[str, str],
    expenditure_class: str,
    county_map: dict[str, CountyReference],
) -> tuple[list[ExpenditureObservation], list[str]]:
    raw_header_row = list(
        next(
            ws.iter_rows(
                min_row=3,
                max_row=3,
                values_only=True,
            )
        )
    )

    headers = nonblank_headers(raw_header_row)

    assert_exact_headers(
        headers,
        expected_headers,
        fiscal_year=fiscal_year,
        sheet_name=sheet_name,
    )

    header_to_column = {
        header: index
        for index, header in enumerate(
            [clean_header(v) for v in raw_header_row],
            start=1,
        )
        if header is not None
    }

    observations: list[ExpenditureObservation] = []
    county_names_seen: list[str] = []

    for source_row_number, row in enumerate(
        ws.iter_rows(
            min_row=6,
            values_only=True,
        ),
        start=6,
    ):
        state_name = row[1]
        county_value = row[2]

        if state_name != "Washington":
            continue

        if county_value is None:
            raise RuntimeError(
                f"FY2024 {sheet_name}: Washington row "
                f"{source_row_number} has no county name."
            )

        county_name_source = str(county_value).strip()
        county = resolve_county(
            county_name_source,
            county_map,
        )

        county_names_seen.append(
            county.county_name
        )

        for source_header, category_code in category_map.items():
            source_column_number = header_to_column[source_header]
            source_value = row[source_column_number - 1]
            source_amount = to_decimal(source_value)

            observations.append(
                ExpenditureObservation(
                    fiscal_year=fiscal_year,
                    county_id=county.county_id,
                    county_name_source=county_name_source,
                    schema_regime=REVISED_SCHEMA,
                    expenditure_class=expenditure_class,
                    category_code=category_code,
                    source_category_name=source_header,
                    source_amount=source_amount,
                    source_unit=USD,
                    expenditure_amount=normalize_dollars(
                        source_amount,
                        USD,
                    ),
                    source_sheet=sheet_name,
                    source_row_number=source_row_number,
                    source_column_number=source_column_number,
                )
            )

    if len(county_names_seen) != EXPECTED_COUNTY_COUNT:
        raise RuntimeError(
            f"FY2024 {sheet_name}: parsed "
            f"{len(county_names_seen)} Washington counties; expected 39."
        )

    if len(set(county_names_seen)) != EXPECTED_COUNTY_COUNT:
        raise RuntimeError(
            f"FY2024 {sheet_name}: duplicate Washington county "
            "rows detected."
        )

    return observations, county_names_seen


def parse_fy2024(
    workbook_path: Path,
    county_map: dict[str, CountyReference],
) -> ParsedYear:
    fiscal_year = 2024

    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        required = {
            "County Direct Exp",
            "County Indirect Exp",
            "Data Guide",
        }

        missing = sorted(
            required - set(workbook.sheetnames)
        )

        if missing:
            raise RuntimeError(
                "FY2024 revised workbook is missing required sheet(s): "
                + ", ".join(missing)
            )

        direct_observations, direct_counties = parse_fy2024_sheet(
            fiscal_year=fiscal_year,
            ws=workbook["County Direct Exp"],
            sheet_name="County Direct Exp",
            expected_headers=FY2024_DIRECT_EXPECTED_HEADERS,
            category_map=FY2024_DIRECT_MAP,
            expenditure_class="DIRECT",
            county_map=county_map,
        )

        indirect_observations, indirect_counties = parse_fy2024_sheet(
            fiscal_year=fiscal_year,
            ws=workbook["County Indirect Exp"],
            sheet_name="County Indirect Exp",
            expected_headers=FY2024_INDIRECT_EXPECTED_HEADERS,
            category_map=FY2024_INDIRECT_MAP,
            expenditure_class="INDIRECT",
            county_map=county_map,
        )

        if set(direct_counties) != set(indirect_counties):
            raise RuntimeError(
                "FY2024 direct and indirect Washington county "
                "populations do not match."
            )

        expenditures = (
            direct_observations
            + indirect_observations
        )

        expected_count = (
            EXPECTED_COUNTY_COUNT
            * (
                len(FY2024_DIRECT_MAP)
                + len(FY2024_INDIRECT_MAP)
            )
        )

        if len(expenditures) != expected_count:
            raise RuntimeError(
                f"FY2024: expected {expected_count} expenditure "
                f"observations; parsed {len(expenditures)}."
            )

        return ParsedYear(
            fiscal_year=fiscal_year,
            source_file=workbook_path,
            source_sha256=sha256_file(workbook_path),
            schema_regime=REVISED_SCHEMA,
            methodology_regime=REVISED_METHODOLOGY,
            methodology_break=True,
            source_county_rows=EXPECTED_COUNTY_COUNT,
            expenditures=expenditures,
            metrics=[],
        )

    finally:
        workbook.close()


def parse_year(
    fiscal_year: int,
    workbook_path: Path,
    county_map: dict[str, CountyReference],
) -> ParsedYear:
    if fiscal_year == 2024:
        return parse_fy2024(
            workbook_path,
            county_map,
        )

    if 2015 <= fiscal_year <= 2023:
        return parse_legacy_year(
            fiscal_year,
            workbook_path,
            county_map,
        )

    raise ValueError(
        f"Unsupported VA GDX fiscal year: {fiscal_year}"
    )


def get_category_ids(cur) -> dict[str, int]:
    cur.execute(
        """
        SELECT category_code, gdx_category_id
        FROM va.gdx_expenditure_category;
        """
    )

    mapping = {
        str(code): int(category_id)
        for code, category_id in cur.fetchall()
    }

    required = (
        set(LEGACY_EXPENDITURE_MAP.values())
        | set(FY2024_DIRECT_MAP.values())
        | set(FY2024_INDIRECT_MAP.values())
    )

    missing = sorted(
        required - set(mapping)
    )

    if missing:
        raise RuntimeError(
            "va.gdx_expenditure_category is missing required "
            "category codes: "
            + ", ".join(missing)
        )

    return mapping


def find_existing_batch(
    cur,
    fiscal_year: int,
    source_sha256: str,
) -> tuple[int, str] | None:
    cur.execute(
        """
        SELECT gdx_load_batch_id, load_status
        FROM va.gdx_load_batch
        WHERE fiscal_year = %s
          AND source_sha256 = %s;
        """,
        (fiscal_year, source_sha256),
    )

    row = cur.fetchone()

    if row is None:
        return None

    return int(row[0]), str(row[1])


def create_or_reset_batch(
    cur,
    parsed: ParsedYear,
    *,
    force: bool,
) -> tuple[int, bool]:
    existing = find_existing_batch(
        cur,
        parsed.fiscal_year,
        parsed.source_sha256,
    )

    if existing is not None:
        batch_id, status = existing

        if status == "COMPLETED" and not force:
            return batch_id, True

        cur.execute(
            """
            DELETE FROM va.gdx_county_metric
            WHERE gdx_load_batch_id = %s;
            """,
            (batch_id,),
        )

        cur.execute(
            """
            DELETE FROM va.gdx_county_expenditure
            WHERE gdx_load_batch_id = %s;
            """,
            (batch_id,),
        )

        cur.execute(
            """
            UPDATE va.gdx_load_batch
            SET
                source_file = %s,
                schema_regime = %s,
                methodology_regime = %s,
                methodology_break = %s,
                load_started_at = CURRENT_TIMESTAMP,
                load_completed_at = NULL,
                source_county_rows = %s,
                loaded_expenditure_rows = 0,
                loaded_metric_rows = 0,
                rejected_row_count = 0,
                load_status = 'STARTED',
                notes = %s
            WHERE gdx_load_batch_id = %s;
            """,
            (
                str(parsed.source_file),
                parsed.schema_regime,
                parsed.methodology_regime,
                parsed.methodology_break,
                parsed.source_county_rows,
                "ETL-VA-GDX-001 reload of existing source hash.",
                batch_id,
            ),
        )

        return batch_id, False

    cur.execute(
        """
        INSERT INTO va.gdx_load_batch (
            fiscal_year,
            source_file,
            source_sha256,
            schema_regime,
            methodology_regime,
            methodology_break,
            source_county_rows,
            load_status,
            notes
        )
        VALUES (
            %s, %s, %s, %s, %s, %s, %s, 'STARTED', %s
        )
        RETURNING gdx_load_batch_id;
        """,
        (
            parsed.fiscal_year,
            str(parsed.source_file),
            parsed.source_sha256,
            parsed.schema_regime,
            parsed.methodology_regime,
            parsed.methodology_break,
            parsed.source_county_rows,
            "ETL-VA-GDX-001 initial load.",
        ),
    )

    return int(cur.fetchone()[0]), False


def insert_expenditures(
    cur,
    batch_id: int,
    observations: list[ExpenditureObservation],
    category_ids: dict[str, int],
) -> None:
    rows = [
        (
            batch_id,
            item.fiscal_year,
            item.county_id,
            item.county_name_source,
            item.schema_regime,
            item.expenditure_class,
            category_ids[item.category_code],
            item.source_category_name,
            item.source_amount,
            item.source_unit,
            item.expenditure_amount,
            item.source_sheet,
            item.source_row_number,
            item.source_column_number,
        )
        for item in observations
    ]

    execute_values(
        cur,
        """
        INSERT INTO va.gdx_county_expenditure (
            gdx_load_batch_id,
            fiscal_year,
            county_id,
            county_name_source,
            schema_regime,
            expenditure_class,
            gdx_category_id,
            source_category_name,
            source_amount,
            source_unit,
            expenditure_amount,
            source_sheet,
            source_row_number,
            source_column_number
        )
        VALUES %s
        """,
        rows,
        page_size=1000,
    )


def insert_metrics(
    cur,
    batch_id: int,
    observations: list[MetricObservation],
) -> None:
    if not observations:
        return

    rows = [
        (
            batch_id,
            item.fiscal_year,
            item.county_id,
            item.county_name_source,
            item.metric_code,
            item.metric_name,
            item.source_value,
            item.source_unit,
            item.source_sheet,
            item.source_row_number,
            item.source_column_number,
        )
        for item in observations
    ]

    execute_values(
        cur,
        """
        INSERT INTO va.gdx_county_metric (
            gdx_load_batch_id,
            fiscal_year,
            county_id,
            county_name_source,
            metric_code,
            metric_name,
            source_value,
            source_unit,
            source_sheet,
            source_row_number,
            source_column_number
        )
        VALUES %s
        """,
        rows,
        page_size=1000,
    )


def validate_loaded_batch(
    cur,
    batch_id: int,
    parsed: ParsedYear,
) -> None:
    cur.execute(
        """
        SELECT COUNT(*), COUNT(DISTINCT county_id)
        FROM va.gdx_county_expenditure
        WHERE gdx_load_batch_id = %s;
        """,
        (batch_id,),
    )

    expenditure_count, county_count = cur.fetchone()

    if int(expenditure_count) != len(parsed.expenditures):
        raise RuntimeError(
            f"Batch {batch_id}: expenditure count mismatch. "
            f"expected={len(parsed.expenditures)}, "
            f"actual={expenditure_count}"
        )

    if int(county_count) != EXPECTED_COUNTY_COUNT:
        raise RuntimeError(
            f"Batch {batch_id}: expenditure county count mismatch. "
            f"expected=39, actual={county_count}"
        )

    cur.execute(
        """
        SELECT COUNT(*)
        FROM va.gdx_county_metric
        WHERE gdx_load_batch_id = %s;
        """,
        (batch_id,),
    )

    metric_count = int(cur.fetchone()[0])

    if metric_count != len(parsed.metrics):
        raise RuntimeError(
            f"Batch {batch_id}: metric count mismatch. "
            f"expected={len(parsed.metrics)}, actual={metric_count}"
        )

    cur.execute(
        """
        SELECT COUNT(*)
        FROM (
            SELECT
                source_sheet,
                source_row_number,
                source_column_number,
                COUNT(*) AS n
            FROM va.gdx_county_expenditure
            WHERE gdx_load_batch_id = %s
            GROUP BY
                source_sheet,
                source_row_number,
                source_column_number
            HAVING COUNT(*) > 1
        ) AS duplicates;
        """,
        (batch_id,),
    )

    duplicate_source_locations = int(
        cur.fetchone()[0]
    )

    if duplicate_source_locations:
        raise RuntimeError(
            f"Batch {batch_id}: duplicate expenditure source "
            f"locations detected: {duplicate_source_locations}"
        )


def mark_batch_completed(
    cur,
    batch_id: int,
    parsed: ParsedYear,
) -> None:
    cur.execute(
        """
        UPDATE va.gdx_load_batch
        SET
            load_completed_at = CURRENT_TIMESTAMP,
            loaded_expenditure_rows = %s,
            loaded_metric_rows = %s,
            rejected_row_count = 0,
            load_status = 'COMPLETED',
            notes = %s
        WHERE gdx_load_batch_id = %s;
        """,
        (
            len(parsed.expenditures),
            len(parsed.metrics),
            (
                "ETL-VA-GDX-001 completed successfully. "
                "County observations resolved to reference.county_registry. "
                "Source units preserved and expenditures normalized to USD."
            ),
            batch_id,
        ),
    )


def load_one_year(
    conn,
    *,
    fiscal_year: int,
    source_root: Path,
    force: bool,
) -> tuple[int, bool, int, int]:
    workbook_path = discover_workbook(
        source_root,
        fiscal_year,
    )

    with conn.cursor() as cur:
        county_map = load_washington_county_reference(
            cur
        )

    parsed = parse_year(
        fiscal_year,
        workbook_path,
        county_map,
    )

    with conn.cursor() as cur:
        category_ids = get_category_ids(
            cur
        )

        batch_id, skipped = create_or_reset_batch(
            cur,
            parsed,
            force=force,
        )

        if skipped:
            conn.rollback()
            return (
                batch_id,
                True,
                len(parsed.expenditures),
                len(parsed.metrics),
            )

        insert_expenditures(
            cur,
            batch_id,
            parsed.expenditures,
            category_ids,
        )

        insert_metrics(
            cur,
            batch_id,
            parsed.metrics,
        )

        validate_loaded_batch(
            cur,
            batch_id,
            parsed,
        )

        mark_batch_completed(
            cur,
            batch_id,
            parsed,
        )

    conn.commit()

    return (
        batch_id,
        False,
        len(parsed.expenditures),
        len(parsed.metrics),
    )


def main() -> None:
    args = parse_arguments()

    requested_years = sorted(
        set(args.years)
    )

    unsupported = [
        year
        for year in requested_years
        if year not in DEFAULT_YEARS
    ]

    if unsupported:
        raise SystemExit(
            "Unsupported fiscal year(s): "
            + ", ".join(str(year) for year in unsupported)
        )

    print("=" * 88)
    print("ETL-VA-GDX-001 — LOAD WASHINGTON COUNTY VA EXPENDITURES")
    print("=" * 88)
    print()
    print(
        "Fiscal years ....................... "
        + ", ".join(str(year) for year in requested_years)
    )
    print(f"Source root ........................ {args.source_root}")
    print(f"Force reload ....................... {args.force}")
    print()

    conn = get_connection()

    results: list[
        tuple[int, int, bool, int, int]
    ] = []

    try:
        for fiscal_year in requested_years:
            print("-" * 88)
            print(f"FY{fiscal_year}")
            print("-" * 88)

            try:
                (
                    batch_id,
                    skipped,
                    expenditure_count,
                    metric_count,
                ) = load_one_year(
                    conn,
                    fiscal_year=fiscal_year,
                    source_root=args.source_root,
                    force=args.force,
                )

            except Exception:
                conn.rollback()
                print(
                    f"FY{fiscal_year} ......................... FAILED"
                )
                raise

            status = (
                "SKIPPED — IDENTICAL COMPLETED BATCH"
                if skipped
                else "LOADED"
            )

            print(f"Load batch ........................ {batch_id}")
            print(f"Status ............................ {status}")
            print(
                f"Expenditure observations .......... "
                f"{expenditure_count:,}"
            )
            print(
                f"Metric observations ............... "
                f"{metric_count:,}"
            )
            print()

            results.append(
                (
                    fiscal_year,
                    batch_id,
                    skipped,
                    expenditure_count,
                    metric_count,
                )
            )

    finally:
        conn.close()

    print("=" * 88)
    print("ETL-VA-GDX-001 SUMMARY")
    print("=" * 88)

    total_expenditures = 0
    total_metrics = 0

    for (
        fiscal_year,
        batch_id,
        skipped,
        expenditure_count,
        metric_count,
    ) in results:
        status = "SKIPPED" if skipped else "LOADED"
        print(
            f"FY{fiscal_year}: "
            f"batch={batch_id} "
            f"{status} "
            f"expenditures={expenditure_count:,} "
            f"metrics={metric_count:,}"
        )

        total_expenditures += expenditure_count
        total_metrics += metric_count

    print()
    print(
        f"Expenditure observations represented ... "
        f"{total_expenditures:,}"
    )
    print(
        f"Metric observations represented ......... "
        f"{total_metrics:,}"
    )
    print()
    print("ETL-VA-GDX-001 RESULT: COMPLETE")
    print("=" * 88)


if __name__ == "__main__":
    main()