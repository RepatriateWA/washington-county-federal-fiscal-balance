"""
WCFFBP
CMS MSSP Historical Linkage and Washington Applicability Profiler v0.2

READ-ONLY.

Changes from v0.1
-----------------
v0.1 established that all 1,296 direct organization-link nonmatches are
confined to 2015-2017:

    2015: 392
    2016: 432
    2017: 472

v0.2 adds only two controlled historical identity bridges:

1. CY2016-CY2017:
   Financial ACO_Num -> official CMS ACO ID Crosswalk -> ACO_ID

2. CY2015:
   Financial ACO name -> deterministic normalized exact name ->
   historical Organization / Participant ACO name

There is NO fuzzy name matching.

CY2018-CY2024 and 2019A continue to use direct ACO_ID linkage.

This profiler DOES NOT:
    - infer beneficiary residence;
    - infer payment-dollar incidence;
    - allocate MSSP payments to Washington;
    - aggregate MSSP dollars as Washington inflow;
    - determine MSSP-vs-FFS economic additivity;
    - net advance/AIM/AIP payments;
    - write to PostgreSQL;
    - authorize ledger inclusion.

Fail-closed identity rules
--------------------------
Every financial ACO-period must resolve to at most one organization
identity.

2015 exact-name normalization may be used only when the normalized name
is unique in the historical organization source.

No fuzzy, substring, edit-distance, manual, or semantic name repair is
permitted.

Crosswalk linkage may be used only when ACO_NUM maps uniquely to one
ACO_ID.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]

FINANCIAL_ZIP = (
    REPO_ROOT
    / "data"
    / "raw"
    / "cms"
    / "medicare_aco"
    / "mssp_performance_results"
    / "Performance Year Financial and Quality Results.zip"
)

ORGANIZATION_ZIP = (
    REPO_ROOT
    / "data"
    / "raw"
    / "cms"
    / "medicare_aco"
    / "mssp_organizations"
    / "Accountable Care Organizations.zip"
)

PARTICIPANT_ZIP = (
    REPO_ROOT
    / "data"
    / "raw"
    / "cms"
    / "medicare_aco"
    / "mssp_participants"
    / "Accountable Care Organization Participants.zip"
)

CROSSWALK_XLSX = (
    REPO_ROOT
    / "data"
    / "raw"
    / "cms"
    / "medicare_aco"
    / "mssp_aco_id_crosswalk"
    / "ssp-aco-id-crosswalk.xlsx"
)

OUTPUT_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "cms"
    / "medicare_aco"
    / "validation"
    / "historical_linkage_v0_2"
)

SOURCE_INVENTORY_CSV = (
    OUTPUT_DIR / "cms_mssp_historical_source_inventory_v0_2.csv"
)

CROSSWALK_PROFILE_CSV = (
    OUTPUT_DIR / "cms_mssp_aco_id_crosswalk_profile_v0_2.csv"
)

PERIOD_PROFILE_CSV = (
    OUTPUT_DIR / "cms_mssp_historical_period_profile_v0_2.csv"
)

LINKAGE_CSV = (
    OUTPUT_DIR / "cms_mssp_historical_linkage_v0_2.csv"
)

WA_CANDIDATES_CSV = (
    OUTPUT_DIR / "cms_mssp_washington_connection_candidates_v0_2.csv"
)

UNRESOLVED_CSV = (
    OUTPUT_DIR / "cms_mssp_historical_unresolved_identity_v0_2.csv"
)

CHECKS_CSV = (
    OUTPUT_DIR / "cms_mssp_historical_linkage_checks_v0_2.csv"
)

REPORT_TXT = (
    OUTPUT_DIR / "cms_mssp_historical_linkage_report_v0_2.txt"
)


PERIODS = [
    "2015",
    "2016",
    "2017",
    "2018",
    "2019",
    "2019A",
    "2020",
    "2021",
    "2022",
    "2023",
    "2024",
]


FINANCIAL_FIELDS = [
    "EarnSaveLoss",
    "GenSaveLoss",
    "ABtotExp",
    "ABtotBnchmk",
    "Adv_Pay",
    "AIM",
    "Adv_Pay_Amt",
    "Adv_Pay_Recoup",
    "AIP",
    "AIPBalance",
    "AIPRecoup",
    "AIPOwe",
]


ID_CANDIDATES = [
    "ACO_ID",
    "ACO_Num",
    "ACO_NUM",
    "ACO ID",
    "ACO Number",
    "ACO_Number",
]


FINANCIAL_NAME_CANDIDATES = [
    "ACO_Name",
    "ACO Name",
    "ACO_NAME",
]


ORGANIZATION_NAME_CANDIDATES = [
    "aco_name",
    "ACO_Name",
    "ACO Name",
    "ACO_NAME",
]


PARTICIPANT_ACO_NAME_CANDIDATES = [
    "ACO Legal or Name/Doing Business As",
    "ACO_Name",
    "ACO Name",
    "ACO_NAME",
]


ORG_GEO_PATTERNS = [
    r"state",
    r"address",
    r"service",
    r"zip",
    r"county",
]


PARTICIPANT_GEO_PATTERNS = [
    r"state",
    r"address",
    r"service",
    r"zip",
    r"county",
]


WA_VALUES = {
    "WA",
    "WASHINGTON",
    "WASHINGTON STATE",
}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as handle:
        while True:
            block = handle.read(1024 * 1024)

            if not block:
                break

            h.update(block)

    return h.hexdigest().upper()


def normalize_header(value: str) -> str:
    return re.sub(
        r"[^a-z0-9]+",
        "",
        (value or "").lower(),
    )


def normalize_text(value: str | None) -> str:
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


def normalize_name(value: str | None) -> str:
    """
    Deterministic exact-name normalization.

    This is NOT fuzzy matching.

    Normalization:
        uppercase
        '&' -> 'AND'
        punctuation -> space
        whitespace collapse

    Corporate suffixes are deliberately retained.
    """
    text = normalize_text(value).upper()
    text = text.replace("&", " AND ")
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()

    return text


def normalized_upper(value: str | None) -> str:
    return normalize_text(value).upper()


def decode_csv(data: bytes) -> str:
    for encoding in [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ]:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise RuntimeError("Unable to decode CSV.")


def read_csv(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = decode_csv(data)

    reader = csv.DictReader(
        io.StringIO(text, newline="")
    )

    headers = reader.fieldnames or []
    rows = []

    for row in reader:
        rows.append(
            {
                str(key): normalize_text(value)
                for key, value in row.items()
                if key is not None
            }
        )

    return headers, rows


def period_from_member(member_name: str) -> str:
    text = member_name.replace("\\", "/").lower()

    if (
        "2019_off-cycle_1" in text
        or "2019_off_cycle_1" in text
        or "2019a" in text
    ):
        return "2019A"

    match = re.search(
        r"/(20(?:15|16|17|18|19|20|21|22|23|24))/",
        text,
    )

    if match:
        return match.group(1)

    for year in range(2015, 2025):
        if re.search(
            rf"(?<!\d){year}(?!\d)",
            text,
        ):
            return str(year)

    return "UNRESOLVED"


def find_header(
    headers: list[str],
    candidates: list[str],
) -> str | None:
    normalized = {
        normalize_header(header): header
        for header in headers
    }

    for candidate in candidates:
        key = normalize_header(candidate)

        if key in normalized:
            return normalized[key]

    return None


def geo_headers(
    headers: list[str],
    patterns: list[str],
) -> list[str]:
    result = []

    for header in headers:
        lowered = header.lower()

        if any(
            re.search(pattern, lowered)
            for pattern in patterns
        ):
            result.append(header)

    return result


def is_explicit_wa_value(value: str | None) -> bool:
    text = normalized_upper(value)

    if text in WA_VALUES:
        return True

    tokens = {
        token.strip()
        for token in re.split(
            r"[,;|/]+",
            text,
        )
        if token.strip()
    }

    return (
        "WA" in tokens
        or "WASHINGTON" in tokens
        or "WASHINGTON STATE" in tokens
    )


def washington_evidence(
    row: dict[str, str],
    geography_headers: list[str],
) -> tuple[bool, str]:
    evidence = []

    for header in geography_headers:
        value = row.get(header, "")

        if is_explicit_wa_value(value):
            evidence.append(
                f"{header}={value}"
            )

    return (
        bool(evidence),
        "; ".join(evidence),
    )


def write_csv(
    path: Path,
    fieldnames: list[str],
    rows: list[dict],
) -> None:
    with path.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=fieldnames,
            extrasaction="ignore",
        )

        writer.writeheader()
        writer.writerows(rows)


def load_source_zip(
    source_name: str,
    zip_path: Path,
) -> tuple[
    dict[str, dict],
    list[dict],
    list[str],
]:
    period_data = {}
    inventory_rows = []
    errors = []

    if not zip_path.exists():
        errors.append(
            f"{source_name}: ZIP not found: {zip_path}"
        )

        return period_data, inventory_rows, errors

    with zipfile.ZipFile(zip_path, "r") as archive:
        members = [
            info
            for info in archive.infolist()
            if (
                not info.is_dir()
                and info.filename.lower().endswith(".csv")
            )
        ]

        for info in members:
            period = period_from_member(
                info.filename
            )

            data = archive.read(info)
            digest = sha256_bytes(data)

            try:
                headers, rows = read_csv(data)
                parse_status = "PASS"
            except Exception as exc:
                headers = []
                rows = []
                parse_status = "FAIL"

                errors.append(
                    f"{source_name} {info.filename}: "
                    f"CSV parse failure: {exc}"
                )

            inventory_rows.append(
                {
                    "source": source_name,
                    "performance_period": period,
                    "member_name": info.filename,
                    "member_bytes": len(data),
                    "sha256": digest,
                    "row_count": len(rows),
                    "column_count": len(headers),
                    "parse_status": parse_status,
                }
            )

            if period == "UNRESOLVED":
                errors.append(
                    f"{source_name}: unresolved period: "
                    f"{info.filename}"
                )
                continue

            if period in period_data:
                errors.append(
                    f"{source_name}: duplicate period {period}"
                )
                continue

            period_data[period] = {
                "member_name": info.filename,
                "sha256": digest,
                "headers": headers,
                "rows": rows,
            }

    return period_data, inventory_rows, errors


def load_crosswalk() -> tuple[
    dict[str, str],
    list[dict],
    list[str],
]:
    errors = []
    profile_rows = []

    if not CROSSWALK_XLSX.exists():
        return {}, [], [
            f"Crosswalk not found: {CROSSWALK_XLSX}"
        ]

    workbook = load_workbook(
        CROSSWALK_XLSX,
        read_only=True,
        data_only=True,
    )

    if "ACO ID Crosswalk" not in workbook.sheetnames:
        workbook.close()

        return {}, [], [
            "ACO ID Crosswalk worksheet not found."
        ]

    sheet = workbook["ACO ID Crosswalk"]

    rows = sheet.iter_rows(
        values_only=True
    )

    try:
        header_row = next(rows)
    except StopIteration:
        workbook.close()

        return {}, [], [
            "ACO ID Crosswalk worksheet is empty."
        ]

    headers = [
        normalize_text(value)
        for value in header_row
    ]

    normalized_headers = {
        normalize_header(header): index
        for index, header in enumerate(headers)
    }

    aco_id_index = normalized_headers.get(
        normalize_header("ACO_ID")
    )

    aco_num_index = normalized_headers.get(
        normalize_header("ACO_NUM")
    )

    if aco_id_index is None:
        errors.append(
            "Crosswalk missing ACO_ID."
        )

    if aco_num_index is None:
        errors.append(
            "Crosswalk missing ACO_NUM."
        )

    if errors:
        workbook.close()
        return {}, [], errors

    num_to_ids: dict[
        str,
        set[str],
    ] = defaultdict(set)

    for excel_row_number, values in enumerate(
        rows,
        start=2,
    ):
        aco_id = normalize_text(
            values[aco_id_index]
            if aco_id_index < len(values)
            else ""
        )

        aco_num = normalize_text(
            values[aco_num_index]
            if aco_num_index < len(values)
            else ""
        )

        if aco_num and aco_id:
            num_to_ids[aco_num].add(
                aco_id
            )

            profile_rows.append(
                {
                    "excel_row_number": excel_row_number,
                    "aco_num": aco_num,
                    "aco_id": aco_id,
                    "mapping_status": "OBSERVED",
                }
            )

    workbook.close()

    ambiguous = {
        key: values
        for key, values in num_to_ids.items()
        if len(values) != 1
    }

    if ambiguous:
        errors.append(
            f"Crosswalk has {len(ambiguous)} "
            "ACO_NUM values mapping to !=1 ACO_ID."
        )

    mapping = {
        key: next(iter(values))
        for key, values in num_to_ids.items()
        if len(values) == 1
    }

    return mapping, profile_rows, errors


def main() -> int:
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    errors = []

    financial, fin_inventory, e = load_source_zip(
        "FINANCIAL",
        FINANCIAL_ZIP,
    )
    errors.extend(e)

    organizations, org_inventory, e = load_source_zip(
        "ORGANIZATION",
        ORGANIZATION_ZIP,
    )
    errors.extend(e)

    participants, part_inventory, e = load_source_zip(
        "PARTICIPANT",
        PARTICIPANT_ZIP,
    )
    errors.extend(e)

    crosswalk, crosswalk_rows, e = load_crosswalk()
    errors.extend(e)

    source_inventory_rows = (
        fin_inventory
        + org_inventory
        + part_inventory
    )

    period_profile_rows = []
    linkage_rows = []
    wa_candidate_rows = []
    unresolved_rows = []

    total_financial_rows = 0
    total_resolved = 0
    total_unresolved = 0
    total_ambiguous = 0

    linkage_method_counts = Counter()

    # ------------------------------------------------------------
    # Period linkage
    # ------------------------------------------------------------
    for period in PERIODS:
        fin = financial.get(period)
        org = organizations.get(period)
        part = participants.get(period)

        if not fin or not org or not part:
            continue

        fin_headers = fin["headers"]
        org_headers = org["headers"]
        part_headers = part["headers"]

        fin_id_header = find_header(
            fin_headers,
            ID_CANDIDATES,
        )

        org_id_header = find_header(
            org_headers,
            ID_CANDIDATES,
        )

        part_id_header = find_header(
            part_headers,
            ID_CANDIDATES,
        )

        fin_name_header = find_header(
            fin_headers,
            FINANCIAL_NAME_CANDIDATES,
        )

        org_name_header = find_header(
            org_headers,
            ORGANIZATION_NAME_CANDIDATES,
        )

        part_name_header = find_header(
            part_headers,
            PARTICIPANT_ACO_NAME_CANDIDATES,
        )

        org_geo = geo_headers(
            org_headers,
            ORG_GEO_PATTERNS,
        )

        part_geo = geo_headers(
            part_headers,
            PARTICIPANT_GEO_PATTERNS,
        )

        # Direct organization ID index for years where available.
        org_id_index = defaultdict(list)

        if org_id_header:
            for row in org["rows"]:
                value = normalize_text(
                    row.get(
                        org_id_header,
                        ""
                    )
                )

                if value:
                    org_id_index[value].append(
                        row
                    )

        # 2015 deterministic organization-name index.
        org_name_index = defaultdict(list)

        if org_name_header:
            for row in org["rows"]:
                value = normalize_name(
                    row.get(
                        org_name_header,
                        ""
                    )
                )

                if value:
                    org_name_index[value].append(
                        row
                    )

        # Participant indexes by ID and by ACO name.
        part_id_index = defaultdict(list)

        if part_id_header:
            for row in part["rows"]:
                value = normalize_text(
                    row.get(
                        part_id_header,
                        ""
                    )
                )

                if value:
                    part_id_index[value].append(
                        row
                    )

        part_name_index = defaultdict(list)

        if part_name_header:
            for row in part["rows"]:
                value = normalize_name(
                    row.get(
                        part_name_header,
                        ""
                    )
                )

                if value:
                    part_name_index[value].append(
                        row
                    )

        period_resolved = 0
        period_unresolved = 0
        period_ambiguous = 0
        period_wa_org = 0
        period_wa_part = 0
        period_wa_either = 0
        period_wa_both = 0

        method_counts = Counter()

        for fin_row_number, fin_row in enumerate(
            fin["rows"],
            start=2,
        ):
            total_financial_rows += 1

            source_aco_id = (
                normalize_text(
                    fin_row.get(
                        fin_id_header,
                        ""
                    )
                )
                if fin_id_header
                else ""
            )

            financial_name = (
                normalize_text(
                    fin_row.get(
                        fin_name_header,
                        ""
                    )
                )
                if fin_name_header
                else ""
            )

            normalized_fin_name = normalize_name(
                financial_name
            )

            canonical_aco_id = ""
            org_matches = []
            participant_rows = []
            linkage_method = ""
            linkage_status = ""

            # ----------------------------------------------------
            # 2015: exact normalized name only.
            # ----------------------------------------------------
            if period == "2015":
                org_matches = org_name_index.get(
                    normalized_fin_name,
                    [],
                )

                participant_rows = part_name_index.get(
                    normalized_fin_name,
                    [],
                )

                linkage_method = (
                    "EXACT_NORMALIZED_NAME"
                )

                # Financial ACO_Num is retained as source identifier,
                # but no 2015 organization ID exists to promote.
                canonical_aco_id = source_aco_id

            # ----------------------------------------------------
            # 2016-2017: official ACO_NUM -> ACO_ID crosswalk.
            # ----------------------------------------------------
            elif period in {"2016", "2017"}:
                canonical_aco_id = crosswalk.get(
                    source_aco_id,
                    ""
                )

                linkage_method = (
                    "CMS_ACO_ID_CROSSWALK"
                )

                if canonical_aco_id:
                    org_matches = org_id_index.get(
                        canonical_aco_id,
                        [],
                    )

                # Participants still use ACO_Num in these years.
                participant_rows = part_id_index.get(
                    source_aco_id,
                    [],
                )

            # ----------------------------------------------------
            # 2018+: direct ACO_ID.
            # ----------------------------------------------------
            else:
                canonical_aco_id = source_aco_id

                linkage_method = (
                    "DIRECT_ACO_ID"
                )

                org_matches = org_id_index.get(
                    canonical_aco_id,
                    [],
                )

                # 2018 participant source may still use ACO_Num.
                # If participant ID vocabulary differs, first try the
                # source financial ID, then crosswalk-derived legacy
                # ACO_NUM values are not guessed here.
                participant_rows = part_id_index.get(
                    source_aco_id,
                    [],
                )

            if len(org_matches) == 1:
                linkage_status = "RESOLVED_UNIQUE"
                period_resolved += 1
                total_resolved += 1

            elif len(org_matches) == 0:
                linkage_status = "UNRESOLVED_NO_ORG_MATCH"
                period_unresolved += 1
                total_unresolved += 1

            else:
                linkage_status = "AMBIGUOUS_MULTI_ORG_MATCH"
                period_ambiguous += 1
                total_ambiguous += 1

            method_counts[
                linkage_method
            ] += 1

            linkage_method_counts[
                linkage_method
            ] += 1

            org_row = (
                org_matches[0]
                if len(org_matches) == 1
                else {}
            )

            organization_name = (
                normalize_text(
                    org_row.get(
                        org_name_header,
                        "",
                    )
                )
                if org_name_header
                else ""
            )

            wa_org, wa_org_detail = washington_evidence(
                org_row,
                org_geo,
            )

            wa_participant_count = 0
            participant_evidence = []

            for participant_row in participant_rows:
                is_wa, detail = washington_evidence(
                    participant_row,
                    part_geo,
                )

                if is_wa:
                    wa_participant_count += 1

                    if detail:
                        participant_evidence.append(
                            detail
                        )

            wa_part = (
                wa_participant_count > 0
            )

            if wa_org:
                period_wa_org += 1

            if wa_part:
                period_wa_part += 1

            if wa_org or wa_part:
                period_wa_either += 1

            if wa_org and wa_part:
                period_wa_both += 1

            if wa_org and wa_part:
                wa_status = (
                    "WA_ORGANIZATION_AND_PARTICIPANT"
                )
            elif wa_org:
                wa_status = (
                    "WA_ORGANIZATION_ONLY"
                )
            elif wa_part:
                wa_status = (
                    "WA_PARTICIPANT_ONLY"
                )
            else:
                wa_status = (
                    "NO_DIRECT_WA_CONNECTION_OBSERVED"
                )

            output = {
                "performance_period": period,
                "financial_row_number": fin_row_number,
                "source_aco_identifier": source_aco_id,
                "canonical_aco_id": canonical_aco_id,
                "financial_aco_name": financial_name,
                "normalized_financial_aco_name": normalized_fin_name,
                "identity_linkage_method": linkage_method,
                "identity_linkage_status": linkage_status,
                "organization_match_count": len(org_matches),
                "organization_aco_name": organization_name,
                "participant_row_count": len(participant_rows),
                "wa_organization_evidence": (
                    "YES" if wa_org else "NO"
                ),
                "wa_organization_evidence_detail": wa_org_detail,
                "wa_participant_evidence": (
                    "YES" if wa_part else "NO"
                ),
                "wa_participant_row_count": wa_participant_count,
                "wa_participant_evidence_detail": (
                    " || ".join(
                        sorted(
                            set(
                                participant_evidence
                            )
                        )
                    )
                ),
                "wa_connection_status": wa_status,
                "geographic_interpretation": (
                    "ORGANIZATIONAL_OR_PARTICIPANT_CONNECTION_ONLY"
                ),
                "payment_dollar_incidence": "NOT_ESTABLISHED",
            }

            for field_name in FINANCIAL_FIELDS:
                output[field_name] = fin_row.get(
                    field_name,
                    ""
                )

            linkage_rows.append(output)

            if linkage_status != "RESOLVED_UNIQUE":
                unresolved_rows.append(
                    dict(output)
                )

            if (
                linkage_status == "RESOLVED_UNIQUE"
                and (wa_org or wa_part)
            ):
                wa_candidate_rows.append(
                    dict(output)
                )

        period_profile_rows.append(
            {
                "performance_period": period,
                "financial_rows": len(fin["rows"]),
                "organization_rows": len(org["rows"]),
                "participant_rows": len(part["rows"]),
                "financial_id_field": fin_id_header or "",
                "organization_id_field": org_id_header or "",
                "participant_id_field": part_id_header or "",
                "financial_name_field": fin_name_header or "",
                "organization_name_field": org_name_header or "",
                "participant_aco_name_field": part_name_header or "",
                "resolved_unique": period_resolved,
                "unresolved_no_org_match": period_unresolved,
                "ambiguous_multi_org_match": period_ambiguous,
                "identity_methods": ";".join(
                    f"{key}={value}"
                    for key, value in sorted(
                        method_counts.items()
                    )
                ),
                "wa_organization_aco_periods": period_wa_org,
                "wa_participant_aco_periods": period_wa_part,
                "wa_either_aco_periods": period_wa_either,
                "wa_both_aco_periods": period_wa_both,
            }
        )

    # ------------------------------------------------------------
    # Checks
    # ------------------------------------------------------------
    checks = []

    def add_check(
        name: str,
        passed: bool,
        detail: str,
    ) -> None:
        checks.append(
            {
                "check": name,
                "status": "PASS" if passed else "FAIL",
                "detail": detail,
            }
        )

    required = set(PERIODS)

    add_check(
        "Financial source exact 11-period universe",
        set(financial) == required,
        f"{len(financial)}/11",
    )

    add_check(
        "Organization source exact 11-period universe",
        set(organizations) == required,
        f"{len(organizations)}/11",
    )

    add_check(
        "Participant source exact 11-period universe",
        set(participants) == required,
        f"{len(participants)}/11",
    )

    add_check(
        "All source CSVs parse",
        all(
            row["parse_status"] == "PASS"
            for row in source_inventory_rows
        ),
        (
            f"{sum(row['parse_status'] != 'PASS' for row in source_inventory_rows)} "
            "parse failures"
        ),
    )

    add_check(
        "Official ACO ID crosswalk loaded",
        bool(crosswalk),
        (
            f"{len(crosswalk):,} unique ACO_NUM->ACO_ID mappings | "
            f"SHA-256={sha256_file(CROSSWALK_XLSX) if CROSSWALK_XLSX.exists() else 'MISSING'}"
        ),
    )

    add_check(
        "No ambiguous organization identity linkage",
        total_ambiguous == 0,
        f"{total_ambiguous:,} ambiguous financial ACO-periods",
    )

    add_check(
        "Every financial ACO-period resolves uniquely",
        total_resolved == total_financial_rows,
        (
            f"{total_resolved:,}/{total_financial_rows:,} resolved | "
            f"{total_unresolved:,} unresolved | "
            f"{total_ambiguous:,} ambiguous"
        ),
    )

    # Explicitly certify 2015 deterministic method.
    p2015 = next(
        (
            row
            for row in period_profile_rows
            if row["performance_period"] == "2015"
        ),
        None,
    )

    add_check(
        "2015 exact normalized-name linkage resolves all financial observations",
        bool(
            p2015
            and p2015["resolved_unique"]
            == p2015["financial_rows"]
        ),
        (
            (
                f"{p2015['resolved_unique']}/"
                f"{p2015['financial_rows']} resolved; "
                f"{p2015['unresolved_no_org_match']} unresolved; "
                f"{p2015['ambiguous_multi_org_match']} ambiguous"
            )
            if p2015
            else "2015 profile missing"
        ),
    )

    # Explicitly certify 2016-17 crosswalk method.
    for period in ["2016", "2017"]:
        profile = next(
            (
                row
                for row in period_profile_rows
                if row["performance_period"] == period
            ),
            None,
        )

        add_check(
            f"{period} CMS crosswalk linkage resolves all financial observations",
            bool(
                profile
                and profile["resolved_unique"]
                == profile["financial_rows"]
            ),
            (
                (
                    f"{profile['resolved_unique']}/"
                    f"{profile['financial_rows']} resolved; "
                    f"{profile['unresolved_no_org_match']} unresolved; "
                    f"{profile['ambiguous_multi_org_match']} ambiguous"
                )
                if profile
                else f"{period} profile missing"
            ),
        )

    add_check(
        "Washington candidates retain payment-dollar-incidence boundary",
        all(
            row["payment_dollar_incidence"]
            == "NOT_ESTABLISHED"
            for row in wa_candidate_rows
        ),
        f"{len(wa_candidate_rows):,} resolved WA-connected observations",
    )

    add_check(
        "No Washington candidate promoted beyond connection evidence",
        all(
            row["geographic_interpretation"]
            == "ORGANIZATIONAL_OR_PARTICIPANT_CONNECTION_ONLY"
            for row in wa_candidate_rows
        ),
        "No beneficiary/payment-dollar geography inferred",
    )

    # ------------------------------------------------------------
    # Write outputs
    # ------------------------------------------------------------
    write_csv(
        SOURCE_INVENTORY_CSV,
        [
            "source",
            "performance_period",
            "member_name",
            "member_bytes",
            "sha256",
            "row_count",
            "column_count",
            "parse_status",
        ],
        source_inventory_rows,
    )

    write_csv(
        CROSSWALK_PROFILE_CSV,
        [
            "excel_row_number",
            "aco_num",
            "aco_id",
            "mapping_status",
        ],
        crosswalk_rows,
    )

    write_csv(
        PERIOD_PROFILE_CSV,
        [
            "performance_period",
            "financial_rows",
            "organization_rows",
            "participant_rows",
            "financial_id_field",
            "organization_id_field",
            "participant_id_field",
            "financial_name_field",
            "organization_name_field",
            "participant_aco_name_field",
            "resolved_unique",
            "unresolved_no_org_match",
            "ambiguous_multi_org_match",
            "identity_methods",
            "wa_organization_aco_periods",
            "wa_participant_aco_periods",
            "wa_either_aco_periods",
            "wa_both_aco_periods",
        ],
        period_profile_rows,
    )

    linkage_fields = [
        "performance_period",
        "financial_row_number",
        "source_aco_identifier",
        "canonical_aco_id",
        "financial_aco_name",
        "normalized_financial_aco_name",
        "identity_linkage_method",
        "identity_linkage_status",
        "organization_match_count",
        "organization_aco_name",
        "participant_row_count",
        "wa_organization_evidence",
        "wa_organization_evidence_detail",
        "wa_participant_evidence",
        "wa_participant_row_count",
        "wa_participant_evidence_detail",
        "wa_connection_status",
        "geographic_interpretation",
        "payment_dollar_incidence",
    ] + FINANCIAL_FIELDS

    write_csv(
        LINKAGE_CSV,
        linkage_fields,
        linkage_rows,
    )

    write_csv(
        WA_CANDIDATES_CSV,
        linkage_fields,
        wa_candidate_rows,
    )

    write_csv(
        UNRESOLVED_CSV,
        linkage_fields,
        unresolved_rows,
    )

    write_csv(
        CHECKS_CSV,
        [
            "check",
            "status",
            "detail",
        ],
        checks,
    )

    overall_pass = (
        not errors
        and all(
            row["status"] == "PASS"
            for row in checks
        )
    )

    # ------------------------------------------------------------
    # Report
    # ------------------------------------------------------------
    lines = []

    lines.append("=" * 100)
    lines.append(
        "CMS MSSP HISTORICAL LINKAGE AND WASHINGTON APPLICABILITY PROFILE v0.2"
    )
    lines.append("=" * 100)
    lines.append("")

    lines.append("BOUNDARY")
    lines.append("-" * 100)
    lines.append("READ-ONLY historical identity/linkage profiling.")
    lines.append("No fuzzy name matching.")
    lines.append("No beneficiary residence inferred.")
    lines.append("No Washington payment-dollar incidence inferred.")
    lines.append("No MSSP dollar aggregation as Washington inflow.")
    lines.append("No FFS overlap/additivity conclusion.")
    lines.append("No advance/AIM/AIP netting.")
    lines.append("No database writes.")
    lines.append("No ledger inclusion.")
    lines.append("")

    lines.append("IDENTITY METHODS")
    lines.append("-" * 100)
    lines.append(
        "2015      EXACT_NORMALIZED_NAME "
        "(punctuation/case/whitespace normalization only; no fuzzy matching)"
    )
    lines.append(
        "2016-2017 CMS_ACO_ID_CROSSWALK "
        "(financial ACO_Num -> official ACO_ID)"
    )
    lines.append(
        "2018-2024 DIRECT_ACO_ID "
        "(2019A preserved independently)"
    )
    lines.append("")

    lines.append("PERIOD PROFILE")
    lines.append("-" * 100)

    for row in period_profile_rows:
        lines.append(
            f"{row['performance_period']:<6} | "
            f"financial={row['financial_rows']:,} | "
            f"resolved={row['resolved_unique']:,} | "
            f"unresolved={row['unresolved_no_org_match']:,} | "
            f"ambiguous={row['ambiguous_multi_org_match']:,} | "
            f"{row['identity_methods']}"
        )

        lines.append(
            f"       WA org={row['wa_organization_aco_periods']:,} | "
            f"WA participant={row['wa_participant_aco_periods']:,} | "
            f"either={row['wa_either_aco_periods']:,} | "
            f"both={row['wa_both_aco_periods']:,}"
        )

    lines.append("")

    lines.append("IDENTITY ACCOUNTING")
    lines.append("-" * 100)
    lines.append(
        f"Financial ACO-period observations: {total_financial_rows:,}"
    )
    lines.append(
        f"Resolved uniquely:                {total_resolved:,}"
    )
    lines.append(
        f"Unresolved:                       {total_unresolved:,}"
    )
    lines.append(
        f"Ambiguous:                        {total_ambiguous:,}"
    )
    lines.append("")

    for method, count in sorted(
        linkage_method_counts.items()
    ):
        lines.append(
            f"{method:<30} {count:>6,}"
        )

    lines.append("")

    lines.append("WASHINGTON CONNECTION PROFILE")
    lines.append("-" * 100)
    lines.append(
        f"Resolved WA-connected ACO-period observations: "
        f"{len(wa_candidate_rows):,}"
    )

    status_counts = Counter(
        row["wa_connection_status"]
        for row in wa_candidate_rows
    )

    for status, count in sorted(
        status_counts.items()
    ):
        lines.append(
            f"{status:<40} {count:>6,}"
        )

    lines.append("")
    lines.append(
        "Financial values remain attached observations only; "
        "not certified Washington inflow."
    )

    for row in wa_candidate_rows:
        lines.append(
            f"{row['performance_period']} | "
            f"{row['canonical_aco_id']} | "
            f"{row['financial_aco_name']} | "
            f"{row['identity_linkage_method']} | "
            f"{row['wa_connection_status']} | "
            f"EarnSaveLoss={row.get('EarnSaveLoss', '')} | "
            f"Adv_Pay_Recoup={row.get('Adv_Pay_Recoup', '')} | "
            f"AIPRecoup={row.get('AIPRecoup', '')}"
        )

    lines.append("")
    lines.append("CHECKS")
    lines.append("-" * 100)

    for check in checks:
        lines.append(
            f"{check['status']} | "
            f"{check['check']} | "
            f"{check['detail']}"
        )

    if errors:
        lines.append("")
        lines.append("ERRORS")
        lines.append("-" * 100)

        for error in errors:
            lines.append(
                f"FAIL | {error}"
            )

    lines.append("")
    lines.append("=" * 100)
    lines.append(
        "OVERALL CMS MSSP HISTORICAL LINKAGE PROFILE v0.2: "
        + ("PASS" if overall_pass else "FAIL")
    )
    lines.append("=" * 100)
    lines.append("")

    lines.append("Interpretive boundary:")
    lines.append(
        "  Identity linkage certification is distinct from geographic "
        "payment-dollar attribution."
    )
    lines.append(
        "  2015 exact-name linkage is deterministic and contains no "
        "fuzzy or manual matching."
    )
    lines.append(
        "  2016-2017 identity bridging uses the preserved official CMS "
        "ACO ID crosswalk."
    )
    lines.append(
        "  Washington organization/participant evidence establishes "
        "connection only, not beneficiary residence or payment incidence."
    )
    lines.append(
        "  EarnSaveLoss and related values remain source observations "
        "pending separate economic-overlap and geography certification."
    )
    lines.append(
        "  No source is additive by default."
    )

    report = "\n".join(lines)

    REPORT_TXT.write_text(
        report,
        encoding="utf-8",
    )

    print(report)
    print("")
    print("Evidence written:")

    for path in [
        SOURCE_INVENTORY_CSV,
        CROSSWALK_PROFILE_CSV,
        PERIOD_PROFILE_CSV,
        LINKAGE_CSV,
        WA_CANDIDATES_CSV,
        UNRESOLVED_CSV,
        CHECKS_CSV,
        REPORT_TXT,
    ]:
        print(
            f"  {path.relative_to(REPO_ROOT)}"
        )

    return 0 if overall_pass else 1


if __name__ == "__main__":
    raise SystemExit(main())
