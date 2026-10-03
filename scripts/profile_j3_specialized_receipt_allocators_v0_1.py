"""
ETL-172J-3B
Profile and Certify Specialized Remaining-Receipts Allocator Sources

Sources
-------
FHWA Highway Statistics 2024, Table FE-221B
    Payments Into the Highway Trust Fund, Highway Account
    Fiscal Years 2015-2024.

NIAAA pcyr1970-2023.txt
    Apparent alcohol consumption by geography and beverage type.

Purpose
-------
Certify source structure, semantics, study-period coverage, and the
Washington / United States observations required for ETL-172J-3.

This script does NOT:
    - allocate OMB federal receipts;
    - carry the NIAAA 2023 share into 2024;
    - decide final Transportation Trust Fund scope;
    - write to PostgreSQL;
    - authorize final-ledger inclusion.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook


# =============================================================================
# PATHS
# =============================================================================

ROOT = Path(__file__).resolve().parents[1]

FHWA_FILE = (
    ROOT
    / "data"
    / "raw"
    / "fhwa"
    / "highway_trust_fund"
    / "source"
    / "fhwa_highway_statistics_2024_fe221b.xlsx"
)

NIAAA_FILE = (
    ROOT
    / "data"
    / "raw"
    / "niaaa"
    / "alcohol_consumption"
    / "source"
    / "pcyr1970-2023.txt"
)

NIAAA_PDF = (
    ROOT
    / "data"
    / "raw"
    / "niaaa"
    / "alcohol_consumption"
    / "source"
    / "niaaa_surveillance_report_122.pdf"
)

OUTPUT_DIR = (
    ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "specialized_receipt_allocator_profile_v0_1"
)

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

FHWA_EXTRACT = OUTPUT_DIR / "fhwa_allocator_extract.csv"
NIAAA_EXTRACT = OUTPUT_DIR / "niaaa_allocator_extract.csv"
SOURCE_PROFILE = OUTPUT_DIR / "source_profile.json"
CERTIFICATION_JSON = OUTPUT_DIR / "certification.json"
CERTIFICATION_REPORT = OUTPUT_DIR / "certification_report.txt"

YEARS = list(range(2015, 2025))
NIAAA_YEARS = list(range(2015, 2024))


# =============================================================================
# HELPERS
# =============================================================================

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)

    return h.hexdigest().upper()


def require_file(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)


def pass_fail(condition: bool) -> str:
    return "PASS" if condition else "FAIL"


def safe_float(value) -> float:
    if value is None:
        raise ValueError("Unexpected blank numeric value.")

    result = float(value)

    if not math.isfinite(result):
        raise ValueError(
            f"Non-finite numeric value: {value!r}"
        )

    return result


def write_csv(
    path: Path,
    rows: list[dict],
    fieldnames: list[str],
) -> None:
    with path.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
        )
        writer.writeheader()
        writer.writerows(rows)


# =============================================================================
# FHWA
# =============================================================================

def profile_fhwa() -> tuple[list[dict], dict, dict]:
    wb = load_workbook(
        FHWA_FILE,
        data_only=True,
        read_only=False,
    )

    checks = {}

    checks["fhwa_expected_sheet"] = pass_fail(
        "Payments 2024" in wb.sheetnames
    )

    if checks["fhwa_expected_sheet"] != "PASS":
        raise ValueError(
            "FHWA workbook does not contain expected "
            "'Payments 2024' sheet."
        )

    ws = wb["Payments 2024"]

    title_1 = str(ws["A5"].value or "")
    title_2 = str(ws["A6"].value or "")
    period = str(ws["A7"].value or "")
    units = str(ws["B10"].value or "")
    heading = str(ws["B11"].value or "")

    checks["fhwa_title_semantics"] = pass_fail(
        "FEDERAL HIGHWAY TRUST FUND" in title_1.upper()
        and "HIGHWAY ACCOUNT" in title_1.upper()
        and "HIGHWAY ACCOUNT" in title_2.upper()
    )

    checks["fhwa_period_semantics"] = pass_fail(
        "2015-2024" in period
    )

    checks["fhwa_units_semantics"] = pass_fail(
        "THOUSANDS OF DOLLARS" in units.upper()
    )

    checks["fhwa_payment_heading"] = pass_fail(
        "PAYMENTS INTO THE HIGHWAY TRUST FUND"
        in heading.upper()
    )

    year_columns: dict[int, int] = {}

    for column in range(1, ws.max_column + 1):
        value = ws.cell(15, column).value

        if isinstance(value, (int, float)):
            year = int(value)

            if year in YEARS:
                year_columns[year] = column

    checks["fhwa_year_coverage"] = pass_fail(
        set(year_columns) == set(YEARS)
    )

    row_lookup: dict[str, int] = {}

    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value

        if value is None:
            continue

        label = str(value).strip()

        if label:
            row_lookup[label] = row

    washington_row = row_lookup.get("Washington")
    total_row = row_lookup.get("Total")
    grand_total_row = row_lookup.get("Grand Total")

    checks["fhwa_washington_row"] = pass_fail(
        washington_row is not None
    )

    checks["fhwa_total_row"] = pass_fail(
        total_row is not None
    )

    checks["fhwa_grand_total_row"] = pass_fail(
        grand_total_row is not None
    )

    if (
        washington_row is None
        or total_row is None
        or grand_total_row is None
    ):
        raise ValueError(
            "Could not identify Washington, Total, and Grand Total "
            "rows in FHWA Payments sheet."
        )

    rows = []

    total_grand_equal = True

    for year in YEARS:
        column = year_columns[year]

        wa_thousands = safe_float(
            ws.cell(washington_row, column).value
        )

        total_thousands = safe_float(
            ws.cell(total_row, column).value
        )

        grand_total_thousands = safe_float(
            ws.cell(grand_total_row, column).value
        )

        if abs(total_thousands - grand_total_thousands) > 0.000001:
            total_grand_equal = False

        if grand_total_thousands == 0:
            raise ValueError(
                f"FHWA Grand Total is zero for {year}."
            )

        share = wa_thousands / grand_total_thousands

        rows.append(
            {
                "fiscal_year": year,
                "washington_payments_thousands": wa_thousands,
                "us_total_payments_thousands": total_thousands,
                "us_grand_total_payments_thousands":
                    grand_total_thousands,
                "washington_share": share,
                "washington_share_pct": share * 100.0,
                "source_sheet": ws.title,
                "source_units": "THOUSANDS_OF_DOLLARS",
                "source_scope":
                    "HIGHWAY_ACCOUNT_PAYMENTS_ATTRIBUTABLE_TO_HIGHWAY_USERS",
            }
        )

    checks["fhwa_total_grand_total_identity"] = pass_fail(
        total_grand_equal
    )

    checks["fhwa_share_calculation"] = pass_fail(
        len(rows) == 10
        and all(
            0.0
            < float(row["washington_share"])
            < 1.0
            for row in rows
        )
    )

    note_1 = " ".join(
        str(ws.cell(r, 1).value or "")
        for r in [75, 76, 77]
    )

    note_2 = " ".join(
        str(ws.cell(r, 1).value or "")
        for r in [78, 79]
    )

    checks["fhwa_highway_account_scope_note"] = pass_fail(
        "HIGHWAY ACCOUNT" in note_1.upper()
        and "EXCLUDED" in note_1.upper()
    )

    checks["fhwa_state_estimation_note"] = pass_fail(
        "ESTIMATED BY THE FEDERAL HIGHWAY ADMINISTRATION"
        in note_2.upper()
    )

    profile = {
        "file": str(FHWA_FILE.relative_to(ROOT)),
        "sha256": sha256_file(FHWA_FILE),
        "sheet_names": wb.sheetnames,
        "profiled_sheet": ws.title,
        "rows": ws.max_row,
        "columns": ws.max_column,
        "title": title_1,
        "period": period,
        "units": units,
        "heading": heading,
        "washington_row": washington_row,
        "total_row": total_row,
        "grand_total_row": grand_total_row,
        "scope_interpretation":
            "Highway Account payments attributable to highway users; "
            "not a generic measure of every Transportation Trust Fund receipt.",
    }

    return rows, profile, checks


# =============================================================================
# NIAAA
# =============================================================================

DATA_RE = re.compile(r"^\d{4}\s+\d{1,2}\s+[1-4]\s+")


def parse_fixed_numeric(
    text: str,
    start: int,
    end: int,
):
    raw = text[start - 1:end].strip()

    if raw in {"", "."}:
        return None

    return float(raw)


def profile_niaaa() -> tuple[list[dict], dict, dict]:
    text = NIAAA_FILE.read_text(
        encoding="utf-8",
        errors="replace",
    )

    lines = text.splitlines()

    checks = {}

    checks["niaaa_header_year_semantics"] = pass_fail(
        "years from 1970 through 2023" in text
    )

    checks["niaaa_washington_geography_semantics"] = pass_fail(
        re.search(
            r"^\s*53\s+Washington\s*$",
            text,
            flags=re.MULTILINE,
        )
        is not None
    )

    checks["niaaa_us_geography_semantics"] = pass_fail(
        re.search(
            r"^\s*99\s+United States\s*$",
            text,
            flags=re.MULTILINE,
        )
        is not None
    )

    checks["niaaa_all_beverages_semantics"] = pass_fail(
        re.search(
            r"^\s*4\s+All beverages\s*$",
            text,
            flags=re.MULTILINE,
        )
        is not None
    )

    checks["niaaa_ethanol_semantics"] = pass_fail(
        "Gallons of ethanol (absolute alcohol)" in text
    )

    records = []

    for line in lines:
        if not DATA_RE.match(line):
            continue

        year = int(line[0:4])
        geo = int(line[5:7].strip())
        beverage = int(line[8:9])

        gallons_beverage = parse_fixed_numeric(
            line,
            11,
            20,
        )

        gallons_ethanol = parse_fixed_numeric(
            line,
            22,
            30,
        )

        population_14_plus = parse_fixed_numeric(
            line,
            32,
            40,
        )

        ethanol_per_capita_14_raw = parse_fixed_numeric(
            line,
            43,
            47,
        )

        population_21_plus = parse_fixed_numeric(
            line,
            52,
            60,
        )

        ethanol_per_capita_21_raw = parse_fixed_numeric(
            line,
            63,
            67,
        )

        abv_raw = parse_fixed_numeric(
            line,
            74,
            76,
        )

        time_varying_ethanol = parse_fixed_numeric(
            line,
            78,
            86,
        )

        records.append(
            {
                "year": year,
                "geography_id": geo,
                "beverage_type": beverage,
                "gallons_beverage": gallons_beverage,
                "gallons_ethanol": gallons_ethanol,
                "population_14_plus": population_14_plus,
                "ethanol_per_capita_14":
                    None
                    if ethanol_per_capita_14_raw is None
                    else ethanol_per_capita_14_raw / 10000.0,
                "population_21_plus": population_21_plus,
                "ethanol_per_capita_21":
                    None
                    if ethanol_per_capita_21_raw is None
                    else ethanol_per_capita_21_raw / 10000.0,
                "abv":
                    None
                    if abv_raw is None
                    else abv_raw / 1000.0,
                "time_varying_ethanol_gallons":
                    time_varying_ethanol,
            }
        )

    checks["niaaa_records_parsed"] = pass_fail(
        len(records) > 0
    )

    lookup = {
        (
            int(row["year"]),
            int(row["geography_id"]),
            int(row["beverage_type"]),
        ): row
        for row in records
    }

    rows = []

    coverage_ok = True
    uniqueness_ok = (
        len(lookup) == len(records)
    )

    for year in NIAAA_YEARS:
        wa = lookup.get((year, 53, 4))
        us = lookup.get((year, 99, 4))

        if wa is None or us is None:
            coverage_ok = False
            continue

        wa_ethanol = wa["gallons_ethanol"]
        us_ethanol = us["gallons_ethanol"]

        if (
            wa_ethanol is None
            or us_ethanol is None
            or us_ethanol <= 0
        ):
            coverage_ok = False
            continue

        share = wa_ethanol / us_ethanol

        rows.append(
            {
                "calendar_year": year,
                "washington_geography_id": 53,
                "us_geography_id": 99,
                "beverage_type": 4,
                "beverage_label": "All beverages",
                "washington_ethanol_gallons": wa_ethanol,
                "us_ethanol_gallons": us_ethanol,
                "washington_share": share,
                "washington_share_pct": share * 100.0,
                "allocator_measure":
                    "GALLONS_OF_ETHANOL_ABSOLUTE_ALCOHOL",
                "time_varying_abv_measure_used": "NO",
            }
        )

    checks["niaaa_record_uniqueness"] = pass_fail(
        uniqueness_ok
    )

    checks["niaaa_2015_2023_coverage"] = pass_fail(
        coverage_ok
        and len(rows) == 9
    )

    checks["niaaa_share_calculation"] = pass_fail(
        len(rows) == 9
        and all(
            0.0
            < float(row["washington_share"])
            < 1.0
            for row in rows
        )
    )

    has_2024 = (
        (2024, 53, 4) in lookup
        or (2024, 99, 4) in lookup
    )

    checks["niaaa_2024_source_gap"] = pass_fail(
        not has_2024
    )

    # Confirm that code-4 historical ethanol is the sum of
    # beverage codes 1, 2, and 3 for WA and U.S. over 2015-2023.
    all_beverage_identity_ok = True

    for year in NIAAA_YEARS:
        for geo in [53, 99]:
            total = lookup.get((year, geo, 4))

            parts = [
                lookup.get((year, geo, beverage))
                for beverage in [1, 2, 3]
            ]

            if total is None or any(p is None for p in parts):
                all_beverage_identity_ok = False
                continue

            total_ethanol = total["gallons_ethanol"]

            part_ethanol = sum(
                float(p["gallons_ethanol"])
                for p in parts
                if p["gallons_ethanol"] is not None
            )

            if total_ethanol is None:
                all_beverage_identity_ok = False
                continue

            if abs(
                float(total_ethanol)
                - part_ethanol
            ) > 1.0:
                all_beverage_identity_ok = False

    checks["niaaa_all_beverage_identity"] = pass_fail(
        all_beverage_identity_ok
    )

    profile = {
        "file": str(NIAAA_FILE.relative_to(ROOT)),
        "sha256": sha256_file(NIAAA_FILE),
        "pdf_file": str(NIAAA_PDF.relative_to(ROOT)),
        "pdf_sha256": sha256_file(NIAAA_PDF),
        "record_count": len(records),
        "minimum_year": min(
            row["year"]
            for row in records
        ),
        "maximum_year": max(
            row["year"]
            for row in records
        ),
        "washington_geography_id": 53,
        "united_states_geography_id": 99,
        "all_beverages_code": 4,
        "principal_allocator_measure":
            "Gallons of ethanol (absolute alcohol)",
        "principal_allocator_measure_columns": "22-30",
        "time_varying_abv_measure":
            "Preserved as separate source field; not used by J-3B.",
        "ffy2024_status":
            "SOURCE_YEAR_UNAVAILABLE_NO_CARRY_FORWARD_DECISION_IN_J3B",
    }

    return rows, profile, checks


# =============================================================================
# MAIN
# =============================================================================

def main() -> None:
    require_file(FHWA_FILE)
    require_file(NIAAA_FILE)
    require_file(NIAAA_PDF)

    fhwa_rows, fhwa_profile, fhwa_checks = profile_fhwa()
    niaaa_rows, niaaa_profile, niaaa_checks = profile_niaaa()

    write_csv(
        FHWA_EXTRACT,
        fhwa_rows,
        [
            "fiscal_year",
            "washington_payments_thousands",
            "us_total_payments_thousands",
            "us_grand_total_payments_thousands",
            "washington_share",
            "washington_share_pct",
            "source_sheet",
            "source_units",
            "source_scope",
        ],
    )

    write_csv(
        NIAAA_EXTRACT,
        niaaa_rows,
        [
            "calendar_year",
            "washington_geography_id",
            "us_geography_id",
            "beverage_type",
            "beverage_label",
            "washington_ethanol_gallons",
            "us_ethanol_gallons",
            "washington_share",
            "washington_share_pct",
            "allocator_measure",
            "time_varying_abv_measure_used",
        ],
    )

    source_profile = {
        "etl": "ETL-172J-3B",
        "purpose":
            "Profile and certify FHWA and NIAAA allocator sources.",
        "fhwa": fhwa_profile,
        "niaaa": niaaa_profile,
        "boundaries": {
            "omb_receipts_allocated": False,
            "niaaa_2023_carried_to_2024": False,
            "transportation_scope_decision": False,
            "database_writes": False,
            "final_ledger_authorization": False,
        },
    }

    SOURCE_PROFILE.write_text(
        json.dumps(
            source_profile,
            indent=2,
        ),
        encoding="utf-8",
    )

    checks = {
        **fhwa_checks,
        **niaaa_checks,
    }

    checks["database_write_boundary"] = "PASS"
    checks["final_ledger_boundary"] = "PASS"

    overall = (
        "PASS"
        if all(
            value == "PASS"
            for value in checks.values()
        )
        else "FAIL"
    )

    certification = {
        "etl": "ETL-172J-3B",
        "result": overall,
        "checks": checks,
        "source_profile": str(
            SOURCE_PROFILE.relative_to(ROOT)
        ),
        "fhwa_extract": str(
            FHWA_EXTRACT.relative_to(ROOT)
        ),
        "niaaa_extract": str(
            NIAAA_EXTRACT.relative_to(ROOT)
        ),
    }

    CERTIFICATION_JSON.write_text(
        json.dumps(
            certification,
            indent=2,
        ),
        encoding="utf-8",
    )

    print("=" * 120)
    print(
        "SPECIALIZED RECEIPT ALLOCATOR SOURCE PROFILE - ETL-172J-3B"
    )
    print("=" * 120)
    print()

    print("FHWA HIGHWAY ACCOUNT PAYMENTS")
    print("-" * 120)
    print(
        f"{'FFY':<6}"
        f"{'WA $000':>14}"
        f"{'US $000':>16}"
        f"{'WA %':>12}"
    )

    for row in fhwa_rows:
        print(
            f"{int(row['fiscal_year']):<6}"
            f"{float(row['washington_payments_thousands']):>14,.0f}"
            f"{float(row['us_grand_total_payments_thousands']):>16,.0f}"
            f"{float(row['washington_share_pct']):>11.4f}%"
        )

    print()
    print("NIAAA ALL-BEVERAGE ETHANOL")
    print("-" * 120)
    print(
        f"{'YEAR':<6}"
        f"{'WA GALLONS':>16}"
        f"{'US GALLONS':>18}"
        f"{'WA %':>12}"
    )

    for row in niaaa_rows:
        print(
            f"{int(row['calendar_year']):<6}"
            f"{float(row['washington_ethanol_gallons']):>16,.0f}"
            f"{float(row['us_ethanol_gallons']):>18,.0f}"
            f"{float(row['washington_share_pct']):>11.4f}%"
        )

    print()
    print(
        "NIAAA 2024 STATUS: "
        "SOURCE_YEAR_UNAVAILABLE_NO_CARRY_FORWARD_DECISION_IN_J3B"
    )

    print()
    print("CERTIFICATION")
    print("-" * 120)

    for name, status in checks.items():
        print(
            f"{name:<70}: {status}"
        )

    print()
    print(
        "FHWA extract :",
        FHWA_EXTRACT.relative_to(ROOT),
    )
    print(
        "NIAAA extract:",
        NIAAA_EXTRACT.relative_to(ROOT),
    )
    print(
        "Profile      :",
        SOURCE_PROFILE.relative_to(ROOT),
    )
    print(
        "JSON         :",
        CERTIFICATION_JSON.relative_to(ROOT),
    )
    print(
        "Report       :",
        CERTIFICATION_REPORT.relative_to(ROOT),
    )

    print()
    print(f"ETL-172J-3B RESULT: {overall}")

    report_lines = [
        "=" * 120,
        "SPECIALIZED RECEIPT ALLOCATOR SOURCE PROFILE - ETL-172J-3B",
        "=" * 120,
        "",
        "FHWA source semantics:",
        (
            "FE-221B Payments 2024 contains estimated state-attributable "
            "payments into the Highway Account of the Federal Highway "
            "Trust Fund for FFY2015-FFY2024."
        ),
        (
            "The source explicitly excludes specified transfers to other "
            "trust funds/accounts and the General Fund."
        ),
        "",
        "NIAAA source semantics:",
        (
            "Geography 53 is Washington; geography 99 is United States; "
            "beverage code 4 is All beverages."
        ),
        (
            "Principal J-3B allocator measure is gallons of ethanol "
            "(absolute alcohol), not the separate time-varying-ABV-derived "
            "ethanol field."
        ),
        (
            "Calendar-year observations are available through 2023. "
            "No 2024 carry-forward decision is made in J-3B."
        ),
        "",
        "CERTIFICATION",
        "-" * 120,
    ]

    for name, status in checks.items():
        report_lines.append(
            f"{name:<70}: {status}"
        )

    report_lines.extend(
        [
            "",
            "BOUNDARIES",
            "-" * 120,
            "No OMB receipt controls allocated.",
            "No NIAAA 2023-to-2024 carry-forward performed.",
            "No final Transportation Trust Fund scope decision made.",
            "No PostgreSQL writes.",
            "No final-ledger authorization.",
            "",
            f"ETL-172J-3B RESULT: {overall}",
            "",
        ]
    )

    CERTIFICATION_REPORT.write_text(
        "\n".join(report_lines),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()