#!/usr/bin/env python3
"""
validate_cms_partd_payment_missingness_wa_exposure_v0_1.py

CMS Medicare Part D — Washington Exposure to Missing PartDPlans Inputs v0.1

PURPOSE
-------
Read-only gate measuring whether CMS PartDPlans missing payment inputs actually
intersect Washington numeric CPSC enrollment for CY2015-CY2024.

Questions:
  1. Can all 117 required Washington CPSC months (Jan 2015-Sep 2024) be resolved?
  2. Among exact same-year Contract+PBP PartDPlans matches with numeric WA
     enrollment, how many member-months encounter BLANK/NONNUMERIC inputs in:
       Direct Subsidy, Risk Score, Reinsurance, or LICS?
  3. What are the exact years, plans, organization types, fields, and
     missingness patterns involved?

BOUNDARY
--------
* No blank -> zero conversion.
* No imputation.
* No payment reconstruction.
* No sibling-PBP or other-year propagation.
* No database writes.
* No ledger inclusion.
* PartDReconciliation remains CONTROL_ONLY.

IMPORTANT
---------
A nonzero Washington exposure is NOT a profiler failure. PASS means the source
resolution and accounting mechanics are valid. The exposure is measured evidence.

Run from repository root:
    python scripts/validate_cms_partd_payment_missingness_wa_exposure_v0_1.py
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
MA_ROOT = REPO_ROOT / "data" / "raw" / "cms" / "medicare_advantage"
PLAN_PAYMENT_ROOT = MA_ROOT / "plan_payment"
PARTD_ROOT = REPO_ROOT / "data" / "raw" / "cms" / "medicare_partd"
OUTPUT_DIR = (
    PARTD_ROOT / "validation" / "partd_payment_missingness_wa_exposure_v0_1"
)

REQUIRED_MONTHS = [(y, m) for y in range(2015, 2024) for m in range(1, 13)]
REQUIRED_MONTHS += [(2024, m) for m in range(1, 10)]
EXPECTED_MONTHS = 117

CONTRACT_ALIASES = {
    "contract", "contract number", "contract_number", "contract id",
    "contract_id", "contractid", "hnumber"
}
PBP_ALIASES = {
    "plan benefit package", "plan benefit package id",
    "plan benefit package number", "plan id", "plan_id",
    "plan identifier", "plan_identifier", "pbp", "pbp id", "pbp_id"
}
ENROLLMENT_ALIASES = {
    "enrollment", "enrollment count", "enrollment_count",
    "total enrollment", "total_enrollment"
}

def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip())

def low(v) -> str:
    return norm(v).lower()

def norm_contract(v) -> str:
    return norm(v).upper()

def norm_pbp(v) -> str:
    s = norm(v)
    if re.fullmatch(r"\d+(?:\.0+)?", s):
        return str(int(float(s))).zfill(3)
    return s.upper()

def parse_int(v) -> Optional[int]:
    if v is None:
        return None
    s = norm(v).replace(",", "")
    if not s:
        return None
    try:
        x = float(s)
    except Exception:
        return None
    if abs(x - round(x)) > 1e-9:
        return None
    return int(round(x))

def classify(v) -> str:
    if v is None or norm(v) == "":
        return "BLANK"
    s = norm(v).replace(",", "").replace("$", "")
    try:
        x = float(s)
    except Exception:
        return "NONNUMERIC_OTHER"
    if x > 0:
        return "NUMERIC_POSITIVE"
    if x < 0:
        return "NUMERIC_NEGATIVE"
    return "NUMERIC_ZERO"

def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1", errors="replace")

def detect_delimiter(text: str) -> str:
    lines = [x for x in text.splitlines() if x.strip()]
    first = lines[0] if lines else text
    candidates = {"\t": first.count("\t"), ",": first.count(","), "|": first.count("|")}
    delim, count = max(candidates.items(), key=lambda x: x[1])
    if count <= 0:
        raise RuntimeError("Unable to detect delimiter")
    return delim

def find_col(headers, aliases):
    by_low = {low(h): h for h in headers if h is not None}
    for alias in aliases:
        if alias in by_low:
            return by_low[alias]
    return None

def find_contains(headers, tokens):
    for h in headers:
        lh = low(h)
        if all(t in lh for t in tokens):
            return h
    return None

def infer_year(text: str) -> Optional[int]:
    vals = {int(x) for x in re.findall(r"(?<!\d)(20(?:1[4-9]|2[0-5]))(?!\d)", text)}
    return next(iter(vals)) if len(vals) == 1 else None

def infer_year_month(text: str) -> Optional[tuple[int, int]]:
    s = text.lower()
    found = set()

    # Numeric YYYYMM / YYYY-MM / YYYY_MM / YYYY M patterns.
    for y, m in re.findall(
        r"(?<!\d)(20(?:1[5-9]|2[0-4]))(?:[-_ ]?)(0?[1-9]|1[0-2])(?!\d)", s
    ):
        found.add((int(y), int(m)))

    month_names = {
        "jan": 1, "january": 1, "feb": 2, "february": 2,
        "mar": 3, "march": 3, "apr": 4, "april": 4,
        "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
        "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
        "oct": 10, "october": 10, "nov": 11, "november": 11,
        "dec": 12, "december": 12,
    }
    years = [int(x) for x in re.findall(
        r"(?<!\d)(20(?:1[5-9]|2[0-4]))(?!\d)", s
    )]
    for name, month in month_names.items():
        if re.search(rf"(?<![a-z]){re.escape(name)}(?![a-z])", s):
            for y in years:
                found.add((y, month))

    return next(iter(found)) if len(found) == 1 else None

def sha256_file(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

# ---------------- PartDPlans ----------------

def discover_plan_payment_zip(year):
    candidates = []
    for p in PLAN_PAYMENT_ROOT.rglob("*.zip"):
        if "validation" in {x.lower() for x in p.parts}:
            continue
        y = infer_year(p.name) or infer_year(str(p.parent))
        if y == year:
            candidates.append(p.resolve())
    candidates = sorted(set(candidates))
    if not candidates:
        raise FileNotFoundError(f"CY{year}: no Plan Payment ZIP found")

    if len(candidates) > 1:
        by_hash = defaultdict(list)
        for p in candidates:
            by_hash[sha256_file(p)].append(p)
        if len(by_hash) == 1:
            return sorted(candidates, key=lambda p: (len(str(p)), str(p)))[0]
        scored = []
        for p in candidates:
            score = (4 if str(year) in p.name else 0) + (3 if "payment" in p.name.lower() else 0)
            scored.append((score, p))
        best_score = max(x[0] for x in scored)
        best = [p for score, p in scored if score == best_score]
        if len(best) != 1:
            raise RuntimeError(f"CY{year}: ambiguous Plan Payment ZIPs: {candidates}")
        return best[0]
    return candidates[0]

def find_partdplans_member(zf):
    hits = []
    for n in zf.namelist():
        if n.endswith("/"):
            continue
        b = Path(n).name.lower().replace(" ", "").replace("_", "").replace("-", "")
        if b.endswith("partdplans.xlsx"):
            hits.append(n)
    if len(hits) != 1:
        raise RuntimeError(f"Expected one PartDPlans workbook; found {len(hits)}: {hits}")
    return hits[0]

def load_partdplans(year):
    zpath = discover_plan_payment_zip(year)
    with zipfile.ZipFile(zpath) as zf:
        member = find_partdplans_member(zf)
        raw = zf.read(member)

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out = {}
    found = False

    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = None
        cc = pc = None
        for rn, candidate in enumerate(rows, start=1):
            if rn > 25:
                break
            hs = [norm(x) for x in candidate]
            cc = find_col(hs, CONTRACT_ALIASES)
            pc = find_col(hs, PBP_ALIASES)
            if cc and pc:
                header = candidate
                break
        if header is None:
            continue

        headers = [norm(x) for x in header]
        idx = {h: i for i, h in enumerate(headers)}
        org = find_contains(headers, ("organization", "type"))
        direct = find_contains(headers, ("direct", "subsidy"))
        risk = find_contains(headers, ("risk", "score"))
        reins = find_contains(headers, ("reinsurance",))
        lics = find_contains(headers, ("low", "income", "cost", "sharing")) or find_contains(headers, ("lics",))
        missing = [k for k, v in {
            "organization": org, "direct": direct, "risk": risk,
            "reinsurance": reins, "lics": lics
        }.items() if not v]
        if missing:
            raise RuntimeError(f"CY{year}: missing fields {missing}")

        found = True
        for row in rows:
            c = norm_contract(row[idx[cc]] if idx[cc] < len(row) else "")
            p = norm_pbp(row[idx[pc]] if idx[pc] < len(row) else "")
            if not c or not p:
                continue
            key = (c, p)
            if key in out:
                raise RuntimeError(f"CY{year}: duplicate exact PartDPlans key {key}")
            vals = {
                "DIRECT_SUBSIDY": row[idx[direct]] if idx[direct] < len(row) else None,
                "RISK_SCORE": row[idx[risk]] if idx[risk] < len(row) else None,
                "REINSURANCE": row[idx[reins]] if idx[reins] < len(row) else None,
                "LICS": row[idx[lics]] if idx[lics] < len(row) else None,
            }
            classes = {k: classify(v) for k, v in vals.items()}
            out[key] = {
                "organization_type": norm(row[idx[org]] if idx[org] < len(row) else ""),
                "classes": classes,
                "has_missing_input": any(
                    x in {"BLANK", "NONNUMERIC_OTHER"} for x in classes.values()
                ),
                "missing_fields": ";".join(
                    k for k, x in classes.items()
                    if x in {"BLANK", "NONNUMERIC_OTHER"}
                ),
                "pattern": "|".join(f"{k}={classes[k]}" for k in (
                    "DIRECT_SUBSIDY", "RISK_SCORE", "REINSURANCE", "LICS"
                )),
            }

    wb.close()
    if not found:
        raise RuntimeError(f"CY{year}: no PartDPlans sheet resolved")
    return out

# ---------------- CPSC ----------------

def candidate_cpsc_archives():
    out = []
    for p in MA_ROOT.rglob("*.zip"):
        parts = {x.lower() for x in p.parts}
        if "validation" in parts or "plan_payment" in parts:
            continue
        out.append(p.resolve())
    return sorted(set(out))

def inspect_cpsc_header(raw):
    text = decode_bytes(raw[:256 * 1024])
    delim = detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    for _ in range(10):
        try:
            row = next(reader)
        except StopIteration:
            break
        hs = [norm(x) for x in row]
        cc = find_col(hs, CONTRACT_ALIASES)
        pc = find_col(hs, PBP_ALIASES)
        ec = find_col(hs, ENROLLMENT_ALIASES) or find_contains(hs, ("enrollment",))
        if cc and pc and ec:
            return delim
    raise RuntimeError("No Contract+PBP+Enrollment header")

def build_cpsc_index():
    required = set(REQUIRED_MONTHS)
    admitted = defaultdict(list)
    scan_errors = []

    archives = candidate_cpsc_archives()
    for zpath in archives:
        try:
            with zipfile.ZipFile(zpath) as zf:
                for member in zf.namelist():
                    if member.endswith("/") or not member.lower().endswith((".csv", ".txt")):
                        continue
                    ym = infer_year_month(f"{zpath} {member}")
                    if ym not in required:
                        continue
                    try:
                        raw = zf.read(member)
                        inspect_cpsc_header(raw)
                    except Exception:
                        continue
                    admitted[ym].append((zpath, member, hashlib.sha256(raw).hexdigest()))
        except Exception as exc:
            scan_errors.append({"scope": str(zpath), "message": str(exc)})

    index = {}
    resolution = []
    for ym in REQUIRED_MONTHS:
        hits = admitted.get(ym, [])
        # collapse byte-identical packaging duplicates
        by_hash = defaultdict(list)
        for z, m, h in hits:
            by_hash[h].append((z, m))
        if len(by_hash) == 1 and hits:
            chosen = sorted(next(iter(by_hash.values())), key=lambda x: (str(x[0]), x[1]))[0]
            index[ym] = chosen
            resolution.append({
                "year": ym[0], "month": ym[1], "status": "RESOLVED",
                "candidate_count": len(hits), "distinct_hashes": 1,
                "zip_path": str(chosen[0]), "member": chosen[1],
            })
        elif len(hits) == 1:
            index[ym] = (hits[0][0], hits[0][1])
            resolution.append({
                "year": ym[0], "month": ym[1], "status": "RESOLVED",
                "candidate_count": 1, "distinct_hashes": 1,
                "zip_path": str(hits[0][0]), "member": hits[0][1],
            })
        else:
            resolution.append({
                "year": ym[0], "month": ym[1],
                "status": "UNRESOLVED",
                "candidate_count": len(hits),
                "distinct_hashes": len(by_hash),
                "zip_path": "", "member": "",
            })
    return index, resolution, scan_errors, len(archives)

def read_cpsc(zpath, member):
    with zipfile.ZipFile(zpath) as zf:
        raw = zf.read(member)
    text = decode_bytes(raw)
    delim = detect_delimiter(text[:256 * 1024])
    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delim)
    hs = reader.fieldnames or []
    cc = find_col(hs, CONTRACT_ALIASES)
    pc = find_col(hs, PBP_ALIASES)
    ec = find_col(hs, ENROLLMENT_ALIASES) or find_contains(hs, ("enrollment",))
    state = find_contains(hs, ("state",))
    if not cc or not pc or not ec:
        raise RuntimeError("Required CPSC fields not found")
    for r in reader:
        if state:
            sv = norm(r.get(state)).upper()
            if sv and sv not in {"WA", "WASHINGTON"}:
                continue
        yield norm_contract(r.get(cc)), norm_pbp(r.get(pc)), r.get(ec)

def write_csv(path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 78)
    print("CMS PART D — WA EXPOSURE TO MISSING PAYMENT INPUTS v0.1")
    print("=" * 78)
    print()
    print("BOUNDARY:")
    print("  Measurement only. No blank-to-zero conversion, imputation, payment")
    print("  reconstruction, database writes, or ledger inclusion.")
    print()

    errors = []

    payment = {}
    for y in range(2015, 2025):
        try:
            payment[y] = load_partdplans(y)
        except Exception as exc:
            errors.append({"scope": f"PartDPlans CY{y}", "message": str(exc)})

    cpsc_index, resolution, scan_errors, archive_count = build_cpsc_index()
    errors.extend(scan_errors)

    detail = []
    annual = defaultdict(lambda: {
        "numeric_wa_enrollment": 0,
        "exact_payment_match_enrollment": 0,
        "exact_match_complete_inputs_enrollment": 0,
        "exact_match_missing_inputs_enrollment": 0,
        "no_exact_payment_match_enrollment": 0,
        "suppressed_or_nonnumeric_rows": 0,
    })
    missing_field_mm = Counter()
    missing_pattern_mm = Counter()

    for y, m in REQUIRED_MONTHS:
        source = cpsc_index.get((y, m))
        if source is None or y not in payment:
            continue
        zpath, member = source
        try:
            for c, p, raw_enr in read_cpsc(zpath, member):
                enr = parse_int(raw_enr)
                if enr is None:
                    annual[y]["suppressed_or_nonnumeric_rows"] += 1
                    continue
                if enr < 0:
                    errors.append({
                        "scope": f"CPSC {y}-{m:02d} {c}-{p}",
                        "message": f"Negative enrollment {enr}",
                    })
                    continue

                annual[y]["numeric_wa_enrollment"] += enr
                rec = payment[y].get((c, p))
                if rec is None:
                    annual[y]["no_exact_payment_match_enrollment"] += enr
                    continue

                annual[y]["exact_payment_match_enrollment"] += enr
                if rec["has_missing_input"]:
                    annual[y]["exact_match_missing_inputs_enrollment"] += enr
                    for fld in rec["missing_fields"].split(";"):
                        if fld:
                            missing_field_mm[(y, fld)] += enr
                    missing_pattern_mm[(y, rec["pattern"])] += enr
                    detail.append({
                        "year": y, "month": m, "contract": c, "pbp": p,
                        "organization_type": rec["organization_type"],
                        "numeric_enrollment": enr,
                        "missing_fields": rec["missing_fields"],
                        "missingness_pattern": rec["pattern"],
                        "source_zip": str(zpath),
                        "source_member": member,
                    })
                else:
                    annual[y]["exact_match_complete_inputs_enrollment"] += enr
        except Exception as exc:
            errors.append({
                "scope": f"CPSC {y}-{m:02d}",
                "message": f"{zpath}::{member}: {exc}",
            })

    annual_rows = []
    for y in range(2015, 2025):
        a = annual[y]
        exact_identity = (
            a["exact_payment_match_enrollment"]
            - a["exact_match_complete_inputs_enrollment"]
            - a["exact_match_missing_inputs_enrollment"]
        )
        total_identity = (
            a["numeric_wa_enrollment"]
            - a["exact_payment_match_enrollment"]
            - a["no_exact_payment_match_enrollment"]
        )
        annual_rows.append({
            "year": y,
            "coverage": "JAN-SEP" if y == 2024 else "JAN-DEC",
            **a,
            "exact_match_accounting_difference": exact_identity,
            "numeric_accounting_difference": total_identity,
        })

    total = {
        k: sum(r[k] for r in annual_rows)
        for k in (
            "numeric_wa_enrollment",
            "exact_payment_match_enrollment",
            "exact_match_complete_inputs_enrollment",
            "exact_match_missing_inputs_enrollment",
            "no_exact_payment_match_enrollment",
            "suppressed_or_nonnumeric_rows",
        )
    }
    exact_diff = (
        total["exact_payment_match_enrollment"]
        - total["exact_match_complete_inputs_enrollment"]
        - total["exact_match_missing_inputs_enrollment"]
    )
    numeric_diff = (
        total["numeric_wa_enrollment"]
        - total["exact_payment_match_enrollment"]
        - total["no_exact_payment_match_enrollment"]
    )

    resolved_months = sum(r["status"] == "RESOLVED" for r in resolution)

    checks = [
        ("All 10 PartDPlans years resolved",
         len(payment) == 10, f"{len(payment)}/10"),
        ("All 117 required CPSC months resolved",
         resolved_months == EXPECTED_MONTHS, f"{resolved_months}/{EXPECTED_MONTHS}"),
        ("Exact-match enrollment accounting identity",
         exact_diff == 0, f"difference={exact_diff:,}"),
        ("Total numeric enrollment accounting identity",
         numeric_diff == 0, f"difference={numeric_diff:,}"),
        ("No processing/source errors",
         len(errors) == 0, f"errors={len(errors)}"),
    ]
    overall = all(ok for _, ok, _ in checks)

    field_rows = [
        {"year": y, "field": fld, "member_months": mm}
        for (y, fld), mm in sorted(missing_field_mm.items())
    ]
    pattern_rows = [
        {"year": y, "missingness_pattern": pat, "member_months": mm}
        for (y, pat), mm in sorted(missing_pattern_mm.items())
    ]

    write_csv(
        OUTPUT_DIR / "cms_partd_wa_missing_input_exposure_detail_v0_1.csv",
        detail,
        [
            "year", "month", "contract", "pbp", "organization_type",
            "numeric_enrollment", "missing_fields", "missingness_pattern",
            "source_zip", "source_member",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_wa_missing_input_exposure_annual_v0_1.csv",
        annual_rows,
        [
            "year", "coverage", "numeric_wa_enrollment",
            "exact_payment_match_enrollment",
            "exact_match_complete_inputs_enrollment",
            "exact_match_missing_inputs_enrollment",
            "no_exact_payment_match_enrollment",
            "suppressed_or_nonnumeric_rows",
            "exact_match_accounting_difference",
            "numeric_accounting_difference",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_wa_missing_input_exposure_by_field_v0_1.csv",
        field_rows, ["year", "field", "member_months"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_wa_missing_input_exposure_patterns_v0_1.csv",
        pattern_rows, ["year", "missingness_pattern", "member_months"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_cpsc_source_resolution_v0_1.csv",
        resolution,
        [
            "year", "month", "status", "candidate_count", "distinct_hashes",
            "zip_path", "member",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_wa_missing_input_exposure_checks_v0_1.csv",
        [
            {"check": n, "status": "PASS" if ok else "FAIL", "detail": d}
            for n, ok, d in checks
        ],
        ["check", "status", "detail"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_wa_missing_input_exposure_errors_v0_1.csv",
        errors, ["scope", "message"],
    )

    print("SOURCE RESOLUTION:")
    print(f"  Candidate MA ZIP archives scanned: {archive_count:,}")
    print(f"  PartDPlans years resolved: {len(payment)}/10")
    print(f"  Required CPSC months resolved: {resolved_months}/{EXPECTED_MONTHS}")
    print()
    print("WASHINGTON NUMERIC ENROLLMENT:")
    print(f"  Total numeric enrollment:             {total['numeric_wa_enrollment']:,}")
    print(f"  Exact PartDPlans match:               {total['exact_payment_match_enrollment']:,}")
    print(f"    Complete payment inputs:            {total['exact_match_complete_inputs_enrollment']:,}")
    print(f"    >=1 missing payment input:          {total['exact_match_missing_inputs_enrollment']:,}")
    print(f"  No exact PartDPlans match:            {total['no_exact_payment_match_enrollment']:,}")
    print(f"  Suppressed/non-numeric rows:          {total['suppressed_or_nonnumeric_rows']:,}")
    print(f"  Exact-match accounting difference:    {exact_diff:,}")
    print(f"  Numeric accounting difference:        {numeric_diff:,}")
    print()
    print("ANNUAL WA EXPOSURE:")
    for r in annual_rows:
        print(
            f"  CY{r['year']} {r['coverage']}: "
            f"numeric={r['numeric_wa_enrollment']:,} | "
            f"exact={r['exact_payment_match_enrollment']:,} | "
            f"complete={r['exact_match_complete_inputs_enrollment']:,} | "
            f"missing-input={r['exact_match_missing_inputs_enrollment']:,} | "
            f"no-exact={r['no_exact_payment_match_enrollment']:,}"
        )
    print()
    print("MISSING-INPUT MEMBER-MONTH EXPOSURE BY FIELD:")
    if field_rows:
        for r in field_rows:
            print(f"  CY{r['year']} {r['field']:<16} {r['member_months']:,}")
    else:
        print("  NONE")
    print()
    print("CHECKS:")
    for n, ok, d in checks:
        print(f"  [{'PASS' if ok else 'FAIL'}] {n} — {d}")
    if errors:
        print()
        print(f"ERRORS/DIAGNOSTICS: {len(errors)}")
        for e in errors[:30]:
            print(f"  {e['scope']}: {e['message']}")
    print()
    print(f"Output directory: {OUTPUT_DIR}")
    print(f"OVERALL: {'PASS' if overall else 'FAIL'}")
    print()
    print("NOTE:")
    print("  Nonzero missing-input exposure does not cause FAIL.")
    print("  PASS certifies source resolution and exposure accounting only.")
    print("  No missing-value semantics or substitution rule is authorized.")
    return 0 if overall else 1

if __name__ == "__main__":
    sys.exit(main())
