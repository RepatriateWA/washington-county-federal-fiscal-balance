#!/usr/bin/env python3
"""
profile_cms_partd_cpsc_enrollment_linkage_v0_2.py

READ-ONLY diagnostic profiler for Washington Medicare Part D linkage to the
already-preserved CMS Monthly Enrollment by Contract/Plan/State/County (CPSC)
source universe.

Scope
-----
Jan 2015 through Sep 2024 = 117 months.

v0.2 changes from v0.1
----------------------
1. Indexes the preserved source tree ONCE instead of recursively searching it
   117 times.
2. Discovers CPSC files from their CONTENT/headers, not merely filename tokens.
3. Infers source month from path/name text where possible and reports every
   selected file/member explicitly.
4. Produces a year x month source-resolution matrix.
5. Reclassifies CPSC rows conservatively:
      PARTD_PAYMENT_MATCH
      NO_PARTD_PAYMENT_MATCH
   rather than treating every non-match as a failed Part D linkage.
6. Preserves suppressed/non-numeric enrollment without imputation.
7. Does not reconstruct dollars, write to the database, or authorize ledger
   inclusion.

Important methodological boundary
---------------------------------
CPSC is broader than the Part D payment universe. Therefore a CPSC
Contract+PBP with no PartDPlans match is evidence to profile, not automatically
a linkage failure.

RDS remains out of scope.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Optional

import openpyxl


START = (2015, 1)
END = (2024, 9)

PAYMENT_FIELDS = {
    "CONTRACT_NUMBER",
    "PLAN_BENEFIT_PACKAGE",
    "CONTRACT_NAME",
    "ORGANIZATION_TYPE",
    "AVERAGE_PART_D_DIRECT_SUBSIDY_PM_PM_PAYMENT",
    "AVERAGE_PART_D_RISK_SCORE",
    "AVERAGE_REINSURANCE_PM_PM_PAYMENT",
    "AVERAGE_LOW_INCOME_COST_SHARING_PM_PM_PAYMENT",
}

CPSC_REQUIRED = {
    "contract": {"CONTRACT_NUMBER", "CONTRACT_ID", "CONTRACT"},
    "pbp": {"PLAN_ID", "PLAN_BENEFIT_PACKAGE", "PBP", "PLAN_BENEFIT_PACKAGE_ID"},
    "state": {"STATE", "STATE_ABBREVIATION", "STATE_ABBR"},
    "county": {"COUNTY", "COUNTY_NAME"},
    "enrollment": {"ENROLLMENT", "ENROLLMENT_COUNT", "ENROLLEES"},
}

CPSC_OPTIONAL = {
    "plan_type": {"PLAN_TYPE", "PLAN_TYPE_DESCRIPTION"},
    "organization_type": {"ORGANIZATION_TYPE", "ORG_TYPE"},
}

MONTH_NAMES = {
    1: ("JAN", "JANUARY"),
    2: ("FEB", "FEBRUARY"),
    3: ("MAR", "MARCH"),
    4: ("APR", "APRIL"),
    5: ("MAY",),
    6: ("JUN", "JUNE"),
    7: ("JUL", "JULY"),
    8: ("AUG", "AUGUST"),
    9: ("SEP", "SEPT", "SEPTEMBER"),
    10: ("OCT", "OCTOBER"),
    11: ("NOV", "NOVEMBER"),
    12: ("DEC", "DECEMBER"),
}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def norm(value: Any) -> str:
    s = clean(value).upper().replace("/", " ")
    s = re.sub(r"[^A-Z0-9]+", "_", s)
    return s.strip("_")


def canonical_pbp(value: Any) -> str:
    s = clean(value)
    if not s:
        return ""
    try:
        return f"{int(float(s)):03d}"
    except (ValueError, TypeError):
        return s.upper()


def parse_enrollment(value: Any) -> tuple[Optional[int], str]:
    if value is None:
        return None, "BLANK"
    if isinstance(value, bool):
        return None, "NONNUMERIC"
    if isinstance(value, (int, float)):
        x = float(value)
        if x >= 0 and x.is_integer():
            return int(x), "NUMERIC"
        return None, "NONNUMERIC"
    s = str(value).strip().replace(",", "")
    if not s:
        return None, "BLANK"
    if re.fullmatch(r"\d+", s):
        return int(s), "NUMERIC"
    return None, "SUPPRESSED_OR_NONNUMERIC"


def month_iter():
    y, m = START
    while (y, m) <= END:
        yield y, m
        m += 1
        if m == 13:
            y, m = y + 1, 1


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def decode_delimited(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:8192]
    candidates = {
        "\t": sample.count("\t"),
        ",": sample.count(","),
        "|": sample.count("|"),
    }
    delimiter = max(candidates, key=candidates.get)
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def detect_cpsc_header(rows: list[list[str]], max_rows: int = 50):
    for idx, row in enumerate(rows[:max_rows]):
        header = [norm(x) for x in row]
        mapping = {}
        ok = True
        for logical, aliases in CPSC_REQUIRED.items():
            pos = next((header.index(a) for a in aliases if a in header), None)
            if pos is None:
                ok = False
                break
            mapping[logical] = pos
        if not ok:
            continue
        for logical, aliases in CPSC_OPTIONAL.items():
            mapping[logical] = next(
                (header.index(a) for a in aliases if a in header), None
            )
        return idx, mapping
    return None, {}


def iter_container_members(path: Path) -> Iterable[tuple[str, bytes]]:
    suffix = path.suffix.lower()
    if suffix == ".zip":
        try:
            with zipfile.ZipFile(path) as zf:
                for name in zf.namelist():
                    if Path(name).suffix.lower() in {".csv", ".txt"}:
                        yield name, zf.read(name)
        except (zipfile.BadZipFile, OSError):
            return
    elif suffix in {".csv", ".txt"}:
        try:
            yield path.name, path.read_bytes()
        except OSError:
            return


def inspect_cpsc_member(data: bytes):
    try:
        rows = decode_delimited(data)
    except Exception:
        return None
    header_idx, mapping = detect_cpsc_header(rows)
    if header_idx is None:
        return None
    return rows, header_idx, mapping


def infer_year_month(text: str) -> tuple[Optional[int], Optional[int], str]:
    """
    Conservative date inference from path/member text.

    Returns (year, month, method). If ambiguous, returns None values.
    """
    upper = text.upper()
    compact = re.sub(r"[^A-Z0-9]", "", upper)

    # YYYYMM
    matches = re.findall(r"(20(?:1[4-9]|2[0-4]))(0[1-9]|1[0-2])", compact)
    unique = {(int(y), int(m)) for y, m in matches}
    if len(unique) == 1:
        y, m = next(iter(unique))
        return y, m, "YYYYMM"

    # MMYYYY
    matches = re.findall(r"(0[1-9]|1[0-2])(20(?:1[4-9]|2[0-4]))", compact)
    unique = {(int(y), int(m)) for m, y in matches}
    if len(unique) == 1:
        y, m = next(iter(unique))
        return y, m, "MMYYYY"

    years = {int(x) for x in re.findall(r"20(?:1[4-9]|2[0-4])", upper)}
    found_months = set()
    for month, names in MONTH_NAMES.items():
        for name in names:
            if re.search(rf"(?<![A-Z]){name}(?![A-Z])", upper):
                found_months.add(month)
                break

    if len(years) == 1 and len(found_months) == 1:
        return next(iter(years)), next(iter(found_months)), "YEAR_MONTH_NAME"

    return None, None, "UNRESOLVED"


def index_cpsc_sources(root: Path):
    """
    Scan the tree once. Only files whose content contains the required CPSC
    header universe are admitted to the index.
    """
    indexed = []
    scanned_files = 0
    admitted_members = 0

    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".zip", ".csv", ".txt"}:
            continue

        # Exclude our own validation outputs and Plan Payment products.
        lowparts = [p.lower() for p in path.parts]
        if "validation" in lowparts:
            continue
        if "plan_payment" in lowparts:
            continue

        scanned_files += 1
        for member_name, data in iter_container_members(path):
            inspected = inspect_cpsc_member(data)
            if inspected is None:
                continue

            admitted_members += 1
            combined = f"{path} {member_name}"
            year, month, method = infer_year_month(combined)

            indexed.append(
                {
                    "path": path,
                    "member_name": member_name,
                    "data": data,
                    "member_sha256": sha256_bytes(data),
                    "year": year,
                    "month": month,
                    "date_method": method,
                    "rows_header": inspected,
                }
            )

    return indexed, scanned_files, admitted_members


def choose_month_source(indexed, year: int, month: int):
    candidates = [
        x for x in indexed if x["year"] == year and x["month"] == month
    ]
    if not candidates:
        return None, []

    def score(item):
        text = f"{item['path']} {item['member_name']}".lower()
        compact = re.sub(r"[^a-z0-9]", "", text)
        s = 0
        if "cpsc" in compact:
            s += 1000
        if "contractplanstatecounty" in compact:
            s += 1000
        if "enrollment" in compact:
            s += 100
        if "snp" in compact:
            s -= 500
        if "scc" in compact and "cpsc" not in compact:
            s -= 500
        # Prefer ZIP-contained official source member over loose derivative copy.
        if item["path"].suffix.lower() == ".zip":
            s += 20
        return (-s, str(item["path"]).lower(), item["member_name"].lower())

    candidates = sorted(candidates, key=score)

    # If several candidates have byte-identical members, they are harmless
    # duplicate containers. If not identical, select deterministically but
    # report the competing hashes for review.
    return candidates[0], candidates


def parse_selected_rows(item):
    rows, header_idx, mapping = item["rows_header"]
    parsed = []

    for source_row, row in enumerate(rows[header_idx + 1 :], header_idx + 2):
        def get(logical):
            pos = mapping.get(logical)
            return row[pos] if pos is not None and pos < len(row) else ""

        parsed.append(
            {
                "source_row": source_row,
                "contract": clean(get("contract")).upper(),
                "pbp": canonical_pbp(get("pbp")),
                "state": clean(get("state")).upper(),
                "county": clean(get("county")),
                "enrollment_raw": clean(get("enrollment")),
                "plan_type": clean(get("plan_type")),
                "organization_type": clean(get("organization_type")),
            }
        )
    return parsed


def find_partd_workbook(plan_root: Path, year: int):
    year_dir = plan_root / str(year)
    found = []

    for zp in sorted(year_dir.glob("*.zip")):
        try:
            with zipfile.ZipFile(zp) as zf:
                for member in zf.namelist():
                    base = Path(member).name.lower().replace(" ", "")
                    if (
                        base.endswith(".xlsx")
                        and "partdplans" in base
                        and "recon" not in base
                    ):
                        data = zf.read(member)
                        found.append((zp, member, data, sha256_bytes(data)))
        except Exception:
            continue

    if not found:
        raise RuntimeError(f"{year}: PartDPlans workbook not found")

    hashes = {x[3] for x in found}
    if len(hashes) != 1:
        raise RuntimeError(
            f"{year}: conflicting PartDPlans workbook bytes across annual ZIPs"
        )

    return sorted(found, key=lambda x: (str(x[0]), x[1]))[0]


def load_partd_keys(plan_root: Path, year: int):
    zp, member, data, member_hash = find_partd_workbook(plan_root, year)
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)

    for ws in wb.worksheets:
        for header_row, row in enumerate(
            ws.iter_rows(min_row=1, max_row=30, values_only=True), 1
        ):
            header = [norm(x) for x in row]
            if not PAYMENT_FIELDS.issubset(set(header)):
                continue

            pos = {field: header.index(field) for field in PAYMENT_FIELDS}
            keys = {}
            duplicates = []

            for excel_row, values in enumerate(
                ws.iter_rows(min_row=header_row + 1, values_only=True),
                header_row + 1,
            ):
                def val(field):
                    idx = pos[field]
                    return values[idx] if idx < len(values) else None

                contract = clean(val("CONTRACT_NUMBER")).upper()
                pbp = canonical_pbp(val("PLAN_BENEFIT_PACKAGE"))
                if not contract or not pbp:
                    continue

                key = (contract, pbp)
                record = {
                    "contract": contract,
                    "pbp": pbp,
                    "contract_name": clean(val("CONTRACT_NAME")),
                    "payment_organization_type": clean(val("ORGANIZATION_TYPE")),
                    "sheet_name": ws.title,
                    "excel_row": excel_row,
                }

                if key in keys:
                    duplicates.append(key)
                else:
                    keys[key] = record

            return keys, duplicates, {
                "zip_path": str(zp),
                "zip_sha256": sha256_file(zp),
                "member_name": member,
                "member_sha256": member_hash,
                "sheet_name": ws.title,
                "header_row": header_row,
            }

    raise RuntimeError(f"{year}: PartDPlans semantic table not found")


def write_csv(path: Path, rows, fields):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--cpsc-root",
        type=Path,
        default=Path("data/raw/cms/medicare_advantage"),
    )
    ap.add_argument(
        "--plan-payment-root",
        type=Path,
        default=Path("data/raw/cms/medicare_advantage/plan_payment"),
    )
    ap.add_argument(
        "--output-dir",
        type=Path,
        default=Path(
            "data/raw/cms/medicare_advantage/validation/"
            "partd_cpsc_enrollment_linkage_profile_v0_2"
        ),
    )
    args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    errors = []
    warnings = []

    print("Indexing preserved CMS enrollment source tree once...", flush=True)
    indexed, scanned_files, admitted_members = index_cpsc_sources(args.cpsc_root)
    print(
        f"Indexed {admitted_members:,} CPSC-like source members "
        f"from {scanned_files:,} candidate files.",
        flush=True,
    )

    annual_keys = {}
    payment_manifest = []
    payment_duplicate_rows = []

    for year in range(2015, 2025):
        try:
            keys, duplicates, meta = load_partd_keys(args.plan_payment_root, year)
            annual_keys[year] = keys
            payment_manifest.append(
                {"payment_year": year, "payment_key_count": len(keys), **meta}
            )
            for contract, pbp in duplicates:
                payment_duplicate_rows.append(
                    {"payment_year": year, "contract": contract, "pbp": pbp}
                )
        except Exception as exc:
            annual_keys[year] = {}
            errors.append(str(exc))

    monthly = []
    source_manifest = []
    vocabulary = Counter()
    no_payment_match_rows = []
    payment_without_wa_match = []

    for year, month in month_iter():
        selected, candidates = choose_month_source(indexed, year, month)

        if selected is None:
            errors.append(f"{year}-{month:02d}: no indexed CPSC source resolved")
            monthly.append(
                {
                    "year": year,
                    "month": month,
                    "source_status": "MISSING",
                    "wa_rows": 0,
                    "numeric_enrollment": 0,
                    "partd_match_numeric_enrollment": 0,
                    "no_partd_payment_match_numeric_enrollment": 0,
                    "partd_match_nonnumeric_rows": 0,
                    "no_partd_payment_match_nonnumeric_rows": 0,
                    "numeric_accounting_difference": 0,
                }
            )
            print(f"{year}-{month:02d}: MISSING", flush=True)
            continue

        hashes = {x["member_sha256"] for x in candidates}
        candidate_hash_status = (
            "BYTE_IDENTICAL_DUPLICATES" if len(hashes) == 1
            else "MULTIPLE_DISTINCT_CANDIDATES"
        )
        if len(hashes) > 1:
            warnings.append(
                f"{year}-{month:02d}: {len(candidates)} CPSC candidates with "
                f"{len(hashes)} distinct member hashes; deterministic source selected"
            )

        source_manifest.append(
            {
                "year": year,
                "month": month,
                "selected_path": str(selected["path"]),
                "selected_member": selected["member_name"],
                "selected_member_sha256": selected["member_sha256"],
                "date_inference_method": selected["date_method"],
                "candidate_count": len(candidates),
                "distinct_candidate_hashes": len(hashes),
                "candidate_hash_status": candidate_hash_status,
            }
        )

        rows = parse_selected_rows(selected)
        wa = [r for r in rows if r["state"] in {"WA", "WASHINGTON"}]
        payment_keys = annual_keys.get(year, {})

        numeric_total = 0
        matched_numeric = 0
        nonmatch_numeric = 0
        matched_nonnum = 0
        nonmatch_nonnum = 0
        observed_matched_keys = set()

        for r in wa:
            enrollment, status = parse_enrollment(r["enrollment_raw"])
            key = (r["contract"], r["pbp"])
            has_partd_payment = key in payment_keys

            if has_partd_payment:
                observed_matched_keys.add(key)
                p = payment_keys[key]
                vocabulary[
                    (
                        year,
                        r["plan_type"],
                        r["organization_type"],
                        p["payment_organization_type"],
                    )
                ] += enrollment or 0

            if status == "NUMERIC":
                numeric_total += enrollment or 0
                if has_partd_payment:
                    matched_numeric += enrollment or 0
                else:
                    nonmatch_numeric += enrollment or 0
                    no_payment_match_rows.append(
                        {
                            "year": year,
                            "month": month,
                            "contract": r["contract"],
                            "pbp": r["pbp"],
                            "county": r["county"],
                            "numeric_enrollment": enrollment,
                            "cpsc_plan_type": r["plan_type"],
                            "cpsc_organization_type": r["organization_type"],
                            "classification": "NO_PARTD_PAYMENT_MATCH",
                        }
                    )
            else:
                if has_partd_payment:
                    matched_nonnum += 1
                else:
                    nonmatch_nonnum += 1

        for key, p in payment_keys.items():
            if key not in observed_matched_keys:
                payment_without_wa_match.append(
                    {
                        "year": year,
                        "month": month,
                        "contract": key[0],
                        "pbp": key[1],
                        "contract_name": p["contract_name"],
                        "payment_organization_type": p[
                            "payment_organization_type"
                        ],
                        "classification": "NO_OBSERVED_WA_CPSC_MATCH_THIS_MONTH",
                    }
                )

        diff = numeric_total - matched_numeric - nonmatch_numeric
        monthly.append(
            {
                "year": year,
                "month": month,
                "source_status": "PARSED",
                "wa_rows": len(wa),
                "numeric_enrollment": numeric_total,
                "partd_match_numeric_enrollment": matched_numeric,
                "no_partd_payment_match_numeric_enrollment": nonmatch_numeric,
                "partd_match_nonnumeric_rows": matched_nonnum,
                "no_partd_payment_match_nonnumeric_rows": nonmatch_nonnum,
                "numeric_accounting_difference": diff,
            }
        )

        print(
            f"{year}-{month:02d}: PARSED | WA numeric={numeric_total:,} | "
            f"PartD match={matched_numeric:,} | no PartD payment match={nonmatch_numeric:,}",
            flush=True,
        )

    vocab_rows = [
        {
            "year": y,
            "cpsc_plan_type": pt,
            "cpsc_organization_type": co,
            "payment_organization_type": po,
            "partd_match_numeric_enrollment": n,
        }
        for (y, pt, co, po), n in sorted(vocabulary.items())
    ]

    # Year x month matrix.
    monthly_lookup = {(r["year"], r["month"]): r for r in monthly}
    matrix_rows = []
    for year in range(2015, 2025):
        row = {"year": year}
        for month in range(1, 13):
            if (year, month) < START or (year, month) > END:
                row[f"m{month:02d}"] = "OUT_OF_SCOPE"
            else:
                r = monthly_lookup.get((year, month))
                row[f"m{month:02d}"] = r["source_status"] if r else "NOT_RECORDED"
        matrix_rows.append(row)

    write_csv(
        args.output_dir / "partd_cpsc_monthly_profile.csv",
        monthly,
        [
            "year", "month", "source_status", "wa_rows", "numeric_enrollment",
            "partd_match_numeric_enrollment",
            "no_partd_payment_match_numeric_enrollment",
            "partd_match_nonnumeric_rows",
            "no_partd_payment_match_nonnumeric_rows",
            "numeric_accounting_difference",
        ],
    )
    write_csv(
        args.output_dir / "partd_cpsc_source_resolution_matrix.csv",
        matrix_rows,
        ["year"] + [f"m{x:02d}" for x in range(1, 13)],
    )
    write_csv(
        args.output_dir / "partd_cpsc_source_manifest.csv",
        source_manifest,
        [
            "year", "month", "selected_path", "selected_member",
            "selected_member_sha256", "date_inference_method",
            "candidate_count", "distinct_candidate_hashes",
            "candidate_hash_status",
        ],
    )
    write_csv(
        args.output_dir / "partd_cpsc_partd_match_vocabulary.csv",
        vocab_rows,
        [
            "year", "cpsc_plan_type", "cpsc_organization_type",
            "payment_organization_type", "partd_match_numeric_enrollment",
        ],
    )
    write_csv(
        args.output_dir / "partd_cpsc_no_partd_payment_match_numeric_rows.csv",
        no_payment_match_rows,
        [
            "year", "month", "contract", "pbp", "county",
            "numeric_enrollment", "cpsc_plan_type",
            "cpsc_organization_type", "classification",
        ],
    )
    write_csv(
        args.output_dir / "partd_payment_rows_without_wa_cpsc_match.csv",
        payment_without_wa_match,
        [
            "year", "month", "contract", "pbp", "contract_name",
            "payment_organization_type", "classification",
        ],
    )
    write_csv(
        args.output_dir / "partd_payment_source_manifest.csv",
        payment_manifest,
        [
            "payment_year", "payment_key_count", "zip_path", "zip_sha256",
            "member_name", "member_sha256", "sheet_name", "header_row",
        ],
    )
    write_csv(
        args.output_dir / "partd_payment_duplicate_keys.csv",
        payment_duplicate_rows,
        ["payment_year", "contract", "pbp"],
    )

    required_months = 117
    parsed_months = sum(r["source_status"] == "PARSED" for r in monthly)
    total_numeric = sum(int(r["numeric_enrollment"]) for r in monthly)
    matched_numeric = sum(
        int(r["partd_match_numeric_enrollment"]) for r in monthly
    )
    nonmatch_numeric = sum(
        int(r["no_partd_payment_match_numeric_enrollment"]) for r in monthly
    )
    accounting_diff = sum(
        int(r["numeric_accounting_difference"]) for r in monthly
    )
    matched_nonnum = sum(int(r["partd_match_nonnumeric_rows"]) for r in monthly)
    nonmatch_nonnum = sum(
        int(r["no_partd_payment_match_nonnumeric_rows"]) for r in monthly
    )

    report = [
        "CMS MEDICARE PART D CPSC ENROLLMENT LINKAGE PROFILE v0.2",
        "=" * 72,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        "Scope: Washington, Jan 2015-Sep 2024 (117 months)",
        "",
        "BOUNDARY",
        "-" * 72,
        "Read-only diagnostic profile.",
        "CPSC is broader than the Part D payment universe.",
        "A CPSC row with no PartDPlans match is NOT automatically a linkage failure.",
        "No payment reconstruction, imputation, database writes, or ledger inclusion.",
        "RDS remains out of scope.",
        "",
        "SOURCE INDEX",
        "-" * 72,
        f"Candidate files scanned once: {scanned_files:,}",
        f"CPSC-like members admitted by content/header: {admitted_members:,}",
        f"Required months: {required_months}",
        f"Parsed months: {parsed_months}",
        "",
        "ENROLLMENT ACCOUNTING",
        "-" * 72,
        f"WA numeric enrollment across monthly county rows: {total_numeric:,}",
        f"PARTD_PAYMENT_MATCH numeric enrollment: {matched_numeric:,}",
        f"NO_PARTD_PAYMENT_MATCH numeric enrollment: {nonmatch_numeric:,}",
        f"PARTD_PAYMENT_MATCH suppressed/nonnumeric rows: {matched_nonnum:,}",
        f"NO_PARTD_PAYMENT_MATCH suppressed/nonnumeric rows: {nonmatch_nonnum:,}",
        f"Numeric accounting difference: {accounting_diff:,}",
        "",
        "OTHER DIAGNOSTICS",
        "-" * 72,
        f"Payment-plan/month rows without observed WA CPSC match: "
        f"{len(payment_without_wa_match):,}",
        f"Duplicate annual Part D payment keys: {len(payment_duplicate_rows):,}",
        f"Warnings: {len(warnings)}",
        f"Errors: {len(errors)}",
        "",
        "CHECKS",
        "-" * 72,
        f"All 117 months resolved and parsed: "
        f"{'PASS' if parsed_months == required_months else 'FAIL'}",
        f"Numeric enrollment accounting identity: "
        f"{'PASS' if accounting_diff == 0 else 'FAIL'}",
        f"No duplicate annual Part D payment keys: "
        f"{'PASS' if not payment_duplicate_rows else 'FAIL'}",
        "",
        "INTERPRETATION",
        "-" * 72,
        "This remains a PROFILE, not a payment reconstruction certification.",
        "NO_PARTD_PAYMENT_MATCH is a classification for investigation, not an error.",
        "Payment rows without WA enrollment are expected in a national payment file",
        "and are retained for transparency rather than treated as automatic failures.",
        "",
        "WARNINGS",
        "-" * 72,
    ]
    report += warnings if warnings else ["None"]
    report += ["", "ERRORS", "-" * 72]
    report += errors if errors else ["None"]

    report_text = "\n".join(report) + "\n"
    (args.output_dir / "partd_cpsc_enrollment_linkage_profile_report.txt").write_text(
        report_text, encoding="utf-8"
    )

    print("\n" + report_text)
    print(f"Outputs: {args.output_dir}")

    # Exit success only for the mechanical/source-integrity gates. Warnings about
    # distinct candidates remain visible but do not silently alter source choice.
    mechanical_pass = (
        parsed_months == required_months
        and accounting_diff == 0
        and not payment_duplicate_rows
        and not errors
    )
    return 0 if mechanical_pass else 1


if __name__ == "__main__":
    sys.exit(main())
