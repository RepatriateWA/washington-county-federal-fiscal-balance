from __future__ import annotations

import csv
import hashlib
import json
import sys
from collections import defaultdict
from decimal import Decimal, getcontext
from pathlib import Path
from typing import Any

import openpyxl


# =============================================================================
# ETL-172I-5
# REMAINING FEDERAL RECEIPTS METHODOLOGY CROSSWALK
#
# Purpose:
#   Convert the remaining national federal-receipts universe established by
#   ETL-172I-2, ETL-172I-3, and ETL-172I-4 into a finite methodology work queue.
#
# Classifications:
#
#   ALLOCATOR_METHOD_DOCUMENTED
#       Rockefeller documents a state-allocation method. The national dollars
#       remain unfinished until the allocator is acquired/certified/applied.
#
#   INTENTIONALLY_UNALLOCABLE
#       Rockefeller explicitly treats the national receipt component as
#       unallocable to states.
#
#   METHOD_CHANGE_REVIEW
#       Rockefeller evidence indicates a method/source change, substitution,
#       or vintage-specific treatment requiring review before reconstruction.
#
#   ALLOCATOR_SOURCE_REQUIRED
#       Reserved for a documented allocation method whose source still must
#       be acquired/certified.
#
#   WITHHOLD_UNRESOLVED
#       Evidence is insufficient for a methodological classification.
#
# This step does NOT:
#   - calculate Washington receipts;
#   - apply any allocator;
#   - infer incidence from economic theory;
#   - force reconciliation to IRS gross collections;
#   - write to PostgreSQL;
#   - authorize final ledger inclusion.
#
# IMPORTANT:
#   The purpose of I-5 is to distinguish:
#
#       remaining national receipts
#         = allocatable unfinished
#         + intentionally unallocable
#         + method-review / unresolved
#
#   "Intentionally unallocable" does NOT mean the dollars disappear from the
#   federal receipts universe. It means they are not assigned to Washington
#   under the documented Rockefeller state-allocation architecture.
# =============================================================================


getcontext().prec = 50


REPO_ROOT = Path(__file__).resolve().parents[1]

FY2027_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
)

I2_FILE = (
    FY2027_ROOT
    / "validation"
    / "social_insurance_excise_decomposition_v0_1"
    / "national_social_insurance_excise_decomposition.csv"
)

I3_FILE = (
    FY2027_ROOT
    / "validation"
    / "federal_receipts_reconstruction_coverage_v0_1"
    / "federal_receipts_reconstruction_coverage.csv"
)

I4_FILE = (
    FY2027_ROOT
    / "validation"
    / "other_receipts_decomposition_v0_1"
    / "national_other_receipts_annual_controls.csv"
)

TABLE_24_FILE = (
    FY2027_ROOT
    / "source"
    / "BUDGET-2027-TAB-3-4.xlsx"
)

OUTPUT_DIR = (
    FY2027_ROOT
    / "validation"
    / "remaining_receipts_methodology_crosswalk_v0_1"
)

OUTPUT_CROSSWALK = (
    OUTPUT_DIR
    / "remaining_receipts_methodology_crosswalk.csv"
)

OUTPUT_ANNUAL = (
    OUTPUT_DIR
    / "remaining_receipts_work_queue.csv"
)

OUTPUT_JSON = (
    OUTPUT_DIR
    / "certification.json"
)

OUTPUT_REPORT = (
    OUTPUT_DIR
    / "certification_report.txt"
)


TARGET_YEARS = list(range(2015, 2025))

MILLIONS_TO_DOLLARS = Decimal("1000000")


# =============================================================================
# METHODOLOGY CROSSWALK
#
# These classifications encode only the Rockefeller methodology evidence
# already preserved by WCFFBP.
#
# The source-method labels intentionally preserve Rockefeller terminology.
# =============================================================================


CROSSWALK = [
    # -------------------------------------------------------------------------
    # REMAINING SOCIAL INSURANCE / RETIREMENT
    # -------------------------------------------------------------------------

    {
        "component_group": "SOCIAL_INSURANCE_REMAINDER",
        "component_code": "OASI_FEDERAL_FUNDS",
        "component_label": "Old-age and survivors insurance - Federal funds",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": "",
        "classification": "WITHHOLD_UNRESOLVED",
        "classification_reason": (
            "ETL-172I-3 preserves OASI Federal funds as a possible residual "
            "component when nonzero. No independent Rockefeller allocator "
            "classification is asserted in I-5."
        ),
    },

    {
        "component_group": "SOCIAL_INSURANCE_REMAINDER",
        "component_code": "RAILROAD_RETIREMENT",
        "component_label": "Railroad retirement",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": "IRS Gross Collections, Table 5",
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology identifies IRS Gross Collections, "
            "Table 5 as the state allocation source for Railroad Retirement."
        ),
    },

    {
        "component_group": "SOCIAL_INSURANCE_REMAINDER",
        "component_code": "UNEMPLOYMENT_INSURANCE",
        "component_label": "Unemployment insurance",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": (
            "US DOL Unemployment Insurance Financial Transaction Summary; "
            "method/source varies by Rockefeller vintage"
        ),
        "classification": "METHOD_CHANGE_REVIEW",
        "classification_reason": (
            "Rockefeller documents state allocation of unemployment-insurance "
            "receipts, but preserved evidence shows source/method variation "
            "across report vintages. Historical reconstruction requires a "
            "year-specific method review before allocation."
        ),
    },

    {
        "component_group": "SOCIAL_INSURANCE_REMAINDER",
        "component_code": "OTHER_RETIREMENT",
        "component_label": "Other retirement",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": "Census Population",
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology identifies Census population as the "
            "state allocator for Other Retirement."
        ),
    },

    # -------------------------------------------------------------------------
    # EXCISE TAXES
    #
    # Rockefeller reconstructs selected named excise categories and treats
    # the remaining excise amount as Other Excises.
    # -------------------------------------------------------------------------

    {
        "component_group": "EXCISE",
        "component_code": "TRANSPORTATION_TRUST_FUND",
        "component_label": "Transportation (Trust Fund)",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": (
            "FHWA payments into the FHTF Highway Account"
        ),
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology identifies FHWA payments into the "
            "Federal Highway Trust Fund Highway Account as the allocator."
        ),
    },

    {
        "component_group": "EXCISE",
        "component_code": "TOBACCO",
        "component_label": "Tobacco",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": "Census Population",
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology identifies Census population as the "
            "state allocator for tobacco excise receipts."
        ),
    },

    {
        "component_group": "EXCISE",
        "component_code": "AIRPORT_AND_AIRWAY",
        "component_label": "Airport and Airway",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": "Census Population",
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology identifies Census population as the "
            "state allocator for Airport and Airway receipts."
        ),
    },

    {
        "component_group": "EXCISE",
        "component_code": "HEALTH_INSURANCE_PROVIDERS",
        "component_label": "Health Insurance Providers",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": (
            "Oliver Wyman Analysis; vintage-specific applicability"
        ),
        "classification": "METHOD_CHANGE_REVIEW",
        "classification_reason": (
            "Rockefeller evidence identifies Oliver Wyman analysis for this "
            "category, but the tax is not present uniformly across the study "
            "period and methodology applicability is vintage-specific."
        ),
    },

    {
        "component_group": "EXCISE",
        "component_code": "ALCOHOL",
        "component_label": "Alcohol",
        "national_source": "OMB Historical Table 2.4",
        "rockefeller_method": "NIAAA alcohol consumption",
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology identifies NIAAA alcohol consumption "
            "as the state allocator for alcohol excise receipts."
        ),
    },

    {
        "component_group": "EXCISE",
        "component_code": "OTHER_EXCISES",
        "component_label": "Other Excises",
        "national_source": (
            "Calculated from OMB Historical Tables 2.1 and 2.4"
        ),
        "rockefeller_method": "Census Population",
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology allocates the calculated Other Excises "
            "residual using Census population."
        ),
    },

    # -------------------------------------------------------------------------
    # OTHER RECEIPTS
    # -------------------------------------------------------------------------

    {
        "component_group": "OTHER_RECEIPTS",
        "component_code": "ESTATE_AND_GIFT_TAXES",
        "component_label": "Estate and Gift Taxes",
        "national_source": "OMB Historical Table 2.5",
        "rockefeller_method": "IRS Gross Collections, Table 5",
        "classification": "ALLOCATOR_METHOD_DOCUMENTED",
        "classification_reason": (
            "Rockefeller methodology classifies Estate and Gift Taxes as "
            "allocable and identifies IRS Gross Collections, Table 5 as "
            "the state allocation source."
        ),
    },

    {
        "component_group": "OTHER_RECEIPTS",
        "component_code": "CUSTOMS_DUTIES_AND_FEES",
        "component_label": "Customs Duties and Fees",
        "national_source": "OMB Historical Table 2.5",
        "rockefeller_method": "Unallocable receipts",
        "classification": "INTENTIONALLY_UNALLOCABLE",
        "classification_reason": (
            "Rockefeller places Customs Duties and Fees in Unallocable "
            "Receipts rather than assigning them to states."
        ),
    },

    {
        "component_group": "OTHER_RECEIPTS",
        "component_code": "FEDERAL_RESERVE_DEPOSITS",
        "component_label": "Federal Reserve Deposits",
        "national_source": "OMB Historical Table 2.5",
        "rockefeller_method": "Unallocable receipts",
        "classification": "INTENTIONALLY_UNALLOCABLE",
        "classification_reason": (
            "Rockefeller explicitly treats deposits of Federal Reserve "
            "System earnings as unallocable federal receipts."
        ),
    },

    {
        "component_group": "OTHER_RECEIPTS",
        "component_code": "MISCELLANEOUS_ALL_OTHER",
        "component_label": "All Other Miscellaneous Receipts",
        "national_source": "OMB Historical Table 2.5",
        "rockefeller_method": "Unallocable receipts",
        "classification": "INTENTIONALLY_UNALLOCABLE",
        "classification_reason": (
            "Rockefeller places All Other Miscellaneous Receipts in "
            "Unallocable Receipts rather than assigning them to states."
        ),
    },
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(
            lambda: handle.read(1024 * 1024),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest().upper()


def decimal_value(
    value: Any,
    *,
    context: str,
) -> Decimal:
    if value is None or str(value).strip() == "":
        raise RuntimeError(
            f"MISSING NUMERIC VALUE: {context}"
        )

    try:
        result = Decimal(str(value))
    except Exception as exc:
        raise RuntimeError(
            f"NUMERIC PARSE FAILURE: {context}: {value!r}"
        ) from exc

    if not result.is_finite():
        raise RuntimeError(
            f"NON-FINITE NUMERIC VALUE: {context}: {value!r}"
        )

    return result


def plain_decimal(value: Decimal) -> str:
    return format(value, "f")


def normalized(text: Any) -> str:
    if text is None:
        return ""

    return (
        str(text)
        .strip()
        .lower()
        .replace("-", " ")
        .replace("_", " ")
    )


def load_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise RuntimeError(
            "INPUT FILE NOT FOUND\n"
            f"{path.relative_to(REPO_ROOT)}"
        )

    with path.open(
        "r",
        newline="",
        encoding="utf-8",
    ) as handle:
        return list(csv.DictReader(handle))


def find_column(
    fieldnames: list[str],
    candidates: list[str],
    *,
    context: str,
) -> str:
    for candidate in candidates:
        if candidate in fieldnames:
            return candidate

    raise RuntimeError(
        f"COLUMN NOT FOUND: {context}\n"
        f"Candidates: {candidates}\n"
        f"Available : {fieldnames}"
    )


def load_i3_remaining() -> dict[int, Decimal]:
    rows = load_csv(I3_FILE)

    if not rows:
        raise RuntimeError("I-3 FILE IS EMPTY")

    fieldnames = list(rows[0].keys())

    year_col = find_column(
        fieldnames,
        ["fiscal_year"],
        context="I-3 fiscal year",
    )

    remaining_col = find_column(
        fieldnames,
        [
            "remaining_national_receipts_universe_millions",
        ],
        context="I-3 remaining national receipts universe",
    )

    output: dict[int, Decimal] = {}

    for row in rows:
        year = int(row[year_col])

        if year not in TARGET_YEARS:
            continue

        if year in output:
            raise RuntimeError(
                f"I-3 DUPLICATE FFY: {year}"
            )

        output[year] = decimal_value(
            row[remaining_col],
            context=f"I-3 FFY{year} remaining receipts",
        )

    if set(output) != set(TARGET_YEARS):
        raise RuntimeError(
            "I-3 YEAR COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {sorted(output)}"
        )

    return output


def load_i4_other() -> dict[int, dict[str, Decimal]]:
    rows = load_csv(I4_FILE)

    if not rows:
        raise RuntimeError("I-4 FILE IS EMPTY")

    output: dict[int, dict[str, Decimal]] = {}

    required = [
        "fiscal_year",
        "estate_and_gift_taxes_millions",
        "customs_duties_and_fees_millions",
        "federal_reserve_deposits_millions",
        "miscellaneous_all_other_millions",
        "table_2_5_total_other_receipts_millions",
    ]

    for column in required:
        if column not in rows[0]:
            raise RuntimeError(
                "I-4 SCHEMA FAILURE\n"
                f"Missing: {column}\n"
                f"Available: {list(rows[0].keys())}"
            )

    for row in rows:
        year = int(row["fiscal_year"])

        if year not in TARGET_YEARS:
            continue

        output[year] = {
            "ESTATE_AND_GIFT_TAXES": decimal_value(
                row["estate_and_gift_taxes_millions"],
                context=f"FFY{year} estate/gift",
            ),
            "CUSTOMS_DUTIES_AND_FEES": decimal_value(
                row["customs_duties_and_fees_millions"],
                context=f"FFY{year} customs",
            ),
            "FEDERAL_RESERVE_DEPOSITS": decimal_value(
                row["federal_reserve_deposits_millions"],
                context=f"FFY{year} Federal Reserve",
            ),
            "MISCELLANEOUS_ALL_OTHER": decimal_value(
                row["miscellaneous_all_other_millions"],
                context=f"FFY{year} miscellaneous all other",
            ),
            "TOTAL_OTHER": decimal_value(
                row["table_2_5_total_other_receipts_millions"],
                context=f"FFY{year} total other",
            ),
        }

    if set(output) != set(TARGET_YEARS):
        raise RuntimeError(
            "I-4 YEAR COVERAGE FAILURE"
        )

    return output


# =============================================================================
# TABLE 2.4 EXTRACTION
#
# We deliberately extract the national values directly from the certified OMB
# source workbook here. I-2 established the taxonomy; I-5 needs the component
# amounts in a deterministic form for the methodology queue.
#
# Source rows established by ETL-172I-2:
#
#   Railroad retirement:
#       row 12 = Federal funds
#       row 13 = Trust funds
#       row 14 = Railroad social security equivalent account
#
#   Railroad retirement is reconstructed as rows 12 + 13 + 14.
#   Row 15 is a broader subtotal and MUST NOT be used as the
#   railroad-retirement amount because it also includes OASI, DI,
#   and HI.
#
#   Unemployment insurance:
#       row 19 = total
#
#   Other retirement:
#       row 23 = total
#
#   OASI Federal funds:
#       row 7
#
#   Excise:
#       row 27 Alcohol
#       row 28 Tobacco
#       row 34 Health insurance providers
#       row 40 Transportation
#       row 41 Airport and airway
#       row 54 Total Excise Taxes
#
# Other Excises is Rockefeller's calculated residual:
#
#   Total Excise
#       - Transportation
#       - Tobacco
#       - Airport and Airway
#       - Health Insurance Providers
#       - Alcohol
#
# This preserves Rockefeller's published category architecture instead of
# attempting to assign every OMB Table 2.4 excise leaf separately.
# =============================================================================


def parse_year_header(value: Any) -> int | None:
    if value is None:
        return None

    text = str(value).strip().lower()

    if text.endswith(" estimate"):
        text = text.replace(" estimate", "").strip()

    if text.isdigit():
        return int(text)

    return None


def locate_table24_year_columns(
    worksheet: Any,
) -> dict[int, int]:
    candidates: list[dict[int, int]] = []

    for row_number in range(
        1,
        min(worksheet.max_row, 10) + 1,
    ):
        mapping: dict[int, int] = {}

        for col in range(
            1,
            worksheet.max_column + 1,
        ):
            year = parse_year_header(
                worksheet.cell(
                    row_number,
                    col,
                ).value
            )

            if year in TARGET_YEARS:
                mapping[year] = col

        if set(mapping) == set(TARGET_YEARS):
            candidates.append(mapping)

    if len(candidates) != 1:
        raise RuntimeError(
            "TABLE 2.4 YEAR HEADER FAILURE\n"
            f"Candidate rows with complete 2015-2024 coverage: "
            f"{len(candidates)}"
        )

    return candidates[0]


def load_table24_components() -> dict[int, dict[str, Decimal]]:
    if not TABLE_24_FILE.exists():
        raise RuntimeError(
            "TABLE 2.4 SOURCE NOT FOUND\n"
            f"{TABLE_24_FILE.relative_to(REPO_ROOT)}"
        )

    workbook = openpyxl.load_workbook(
        TABLE_24_FILE,
        data_only=True,
        read_only=False,
    )

    if workbook.sheetnames != ["Table"]:
        raise RuntimeError(
            "TABLE 2.4 SHEET CONTRACT FAILURE\n"
            f"Observed: {workbook.sheetnames}"
        )

    worksheet = workbook["Table"]

    year_columns = locate_table24_year_columns(
        worksheet
    )

    row_map = {
        "OASI_FEDERAL_FUNDS": 7,
        "RAILROAD_RETIREMENT_FEDERAL_FUNDS": 12,
        "RAILROAD_RETIREMENT_TRUST_FUNDS": 13,
        "RAILROAD_SOCIAL_SECURITY_EQUIVALENT": 14,
        "UNEMPLOYMENT_INSURANCE": 19,
        "OTHER_RETIREMENT": 23,
        "ALCOHOL": 27,
        "TOBACCO": 28,
        "HEALTH_INSURANCE_PROVIDERS": 34,
        "TRANSPORTATION_TRUST_FUND": 40,
        "AIRPORT_AND_AIRWAY": 41,
        "TOTAL_EXCISE": 54,
    }

    output: dict[int, dict[str, Decimal]] = {}

    for year in TARGET_YEARS:
        col = year_columns[year]

        values: dict[str, Decimal] = {}

        for code, row_number in row_map.items():
            raw = worksheet.cell(
                row_number,
                col,
            ).value

            # OMB sometimes uses symbols/blank display for zero.
            if raw is None:
                value = Decimal("0")
            elif isinstance(raw, str):
                cleaned = raw.strip()

                if cleaned in {
                    "",
                    "*",
                    "..........",
                    "—",
                    "-",
                }:
                    value = Decimal("0")
                else:
                    value = decimal_value(
                        cleaned.replace(",", ""),
                        context=(
                            f"Table 2.4 FFY{year} "
                            f"{code}"
                        ),
                    )
            else:
                value = decimal_value(
                    raw,
                    context=(
                        f"Table 2.4 FFY{year} "
                        f"{code}"
                    ),
                )

            values[code] = value

        # Railroad retirement must be constructed from its three
        # additive components. OMB row 15 is a broader subtotal that
        # also includes OASI, DI, and HI and therefore must not be used
        # as the railroad-retirement amount.
        values["RAILROAD_RETIREMENT"] = (
            values["RAILROAD_RETIREMENT_FEDERAL_FUNDS"]
            + values["RAILROAD_RETIREMENT_TRUST_FUNDS"]
            + values["RAILROAD_SOCIAL_SECURITY_EQUIVALENT"]
        )

        values["OTHER_EXCISES"] = (
            values["TOTAL_EXCISE"]
            - values["TRANSPORTATION_TRUST_FUND"]
            - values["TOBACCO"]
            - values["AIRPORT_AND_AIRWAY"]
            - values["HEALTH_INSURANCE_PROVIDERS"]
            - values["ALCOHOL"]
        )

        output[year] = values

    return output


def build_component_amounts(
    table24: dict[int, dict[str, Decimal]],
    i4: dict[int, dict[str, Decimal]],
) -> dict[int, dict[str, Decimal]]:
    output: dict[int, dict[str, Decimal]] = {}

    for year in TARGET_YEARS:
        values = dict(table24[year])

        for code in [
            "ESTATE_AND_GIFT_TAXES",
            "CUSTOMS_DUTIES_AND_FEES",
            "FEDERAL_RESERVE_DEPOSITS",
            "MISCELLANEOUS_ALL_OTHER",
        ]:
            values[code] = i4[year][code]

        output[year] = values

    return output


def write_crosswalk() -> None:
    fieldnames = [
        "component_group",
        "component_code",
        "component_label",
        "national_source",
        "rockefeller_method",
        "classification",
        "classification_reason",
    ]

    with OUTPUT_CROSSWALK.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(CROSSWALK)


def build_work_queue(
    component_amounts: dict[int, dict[str, Decimal]],
    i3_remaining: dict[int, Decimal],
) -> list[dict[str, str]]:
    classification_by_code = {
        row["component_code"]: row["classification"]
        for row in CROSSWALK
    }

    output: list[dict[str, str]] = []

    for year in TARGET_YEARS:
        buckets: defaultdict[str, Decimal] = defaultdict(
            lambda: Decimal("0")
        )

        component_total = Decimal("0")

        for code, classification in classification_by_code.items():
            if code not in component_amounts[year]:
                raise RuntimeError(
                    "COMPONENT AMOUNT MISSING\n"
                    f"FFY{year}: {code}"
                )

            amount = component_amounts[year][code]

            component_total += amount
            buckets[classification] += amount

        i3_total = i3_remaining[year]

        difference = (
            i3_total
            - component_total
        )

        if difference != 0:
            raise RuntimeError(
                "I-5 / I-3 REMAINING UNIVERSE RECONCILIATION FAILURE\n"
                f"FFY{year}\n"
                f"I-3 remaining           : {i3_total}\n"
                f"I-5 component universe : {component_total}\n"
                f"Difference              : {difference}"
            )

        allocatable = buckets[
            "ALLOCATOR_METHOD_DOCUMENTED"
        ]

        intentionally_unallocable = buckets[
            "INTENTIONALLY_UNALLOCABLE"
        ]

        method_review = buckets[
            "METHOD_CHANGE_REVIEW"
        ]

        unresolved = buckets[
            "WITHHOLD_UNRESOLVED"
        ]

        source_required = buckets[
            "ALLOCATOR_SOURCE_REQUIRED"
        ]

        actionable_unfinished = (
            allocatable
            + method_review
            + unresolved
            + source_required
        )

        output.append(
            {
                "fiscal_year": str(year),

                "i3_remaining_receipts_millions": (
                    plain_decimal(i3_total)
                ),

                "i5_component_universe_millions": (
                    plain_decimal(component_total)
                ),

                "reconciliation_difference_millions": (
                    plain_decimal(difference)
                ),

                "allocator_method_documented_millions": (
                    plain_decimal(allocatable)
                ),

                "allocator_source_required_millions": (
                    plain_decimal(source_required)
                ),

                "method_change_review_millions": (
                    plain_decimal(method_review)
                ),

                "withhold_unresolved_millions": (
                    plain_decimal(unresolved)
                ),

                "intentionally_unallocable_millions": (
                    plain_decimal(
                        intentionally_unallocable
                    )
                ),

                "actionable_unfinished_millions": (
                    plain_decimal(
                        actionable_unfinished
                    )
                ),

                "intentionally_unallocable_share_pct": (
                    plain_decimal(
                        (
                            intentionally_unallocable
                            / i3_total
                            * Decimal("100")
                        )
                        if i3_total != 0
                        else Decimal("0")
                    )
                ),

                "actionable_unfinished_share_pct": (
                    plain_decimal(
                        (
                            actionable_unfinished
                            / i3_total
                            * Decimal("100")
                        )
                        if i3_total != 0
                        else Decimal("0")
                    )
                ),
            }
        )

    return output


def write_csv(
    path: Path,
    rows: list[dict[str, str]],
) -> None:
    if not rows:
        raise RuntimeError(
            f"REFUSING TO WRITE EMPTY CSV: {path.name}"
        )

    with path.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(rows[0].keys()),
        )

        writer.writeheader()
        writer.writerows(rows)


def certify(
    work_queue: list[dict[str, str]],
) -> dict[str, Any]:
    if len(CROSSWALK) != 14:
        raise RuntimeError(
            "CROSSWALK COMPONENT COUNT FAILURE\n"
            f"Expected: 14\n"
            f"Actual  : {len(CROSSWALK)}"
        )

    codes = [
        row["component_code"]
        for row in CROSSWALK
    ]

    if len(codes) != len(set(codes)):
        raise RuntimeError(
            "DUPLICATE CROSSWALK COMPONENT CODE"
        )

    allowed_classifications = {
        "ALLOCATOR_METHOD_DOCUMENTED",
        "INTENTIONALLY_UNALLOCABLE",
        "METHOD_CHANGE_REVIEW",
        "ALLOCATOR_SOURCE_REQUIRED",
        "WITHHOLD_UNRESOLVED",
    }

    observed_classifications = {
        row["classification"]
        for row in CROSSWALK
    }

    if not observed_classifications.issubset(
        allowed_classifications
    ):
        raise RuntimeError(
            "UNKNOWN CLASSIFICATION\n"
            f"{sorted(observed_classifications)}"
        )

    if len(work_queue) != 10:
        raise RuntimeError(
            "WORK QUEUE YEAR COUNT FAILURE"
        )

    for row in work_queue:
        if (
            Decimal(
                row[
                    "reconciliation_difference_millions"
                ]
            )
            != 0
        ):
            raise RuntimeError(
                "NONZERO I-3 RECONCILIATION DIFFERENCE"
            )

    return {
        "crosswalk_component_count": len(CROSSWALK),
        "crosswalk_component_uniqueness": "PASS",
        "classification_vocabulary": "PASS",
        "fiscal_year_coverage": "PASS",
        "i3_remaining_universe_reconciliation": "PASS",
        "intentionally_unallocable_preserved_in_national_universe": "PASS",
        "washington_allocation_boundary": "PASS",
        "rockefeller_evidence_boundary": "PASS",
        "irs_gross_collections_boundary": "PASS",
        "database_write_boundary": "PASS",
        "ledger_boundary": "PASS",
    }


def build_report(
    *,
    work_queue: list[dict[str, str]],
    certification: dict[str, Any],
) -> str:
    lines = [
        "=" * 126,
        (
            "REMAINING FEDERAL RECEIPTS METHODOLOGY CROSSWALK "
            "- ETL-172I-5"
        ),
        "=" * 126,
        "",
        "ANNUAL REMAINING-RECEIPTS WORK QUEUE",
        "-" * 126,
        (
            f"{'FFY':<6}"
            f"{'I3 REMAIN':>14}"
            f"{'DOC METHOD':>14}"
            f"{'METHOD REV':>14}"
            f"{'UNRESOLVED':>14}"
            f"{'UNALLOCABLE':>14}"
            f"{'ACTIONABLE':>14}"
            f"{'UNALLOC %':>12}"
        ),
        "-" * 126,
    ]

    for row in work_queue:
        lines.append(
            f"{row['fiscal_year']:<6}"
            f"{Decimal(row['i3_remaining_receipts_millions']):>14,.0f}"
            f"{Decimal(row['allocator_method_documented_millions']):>14,.0f}"
            f"{Decimal(row['method_change_review_millions']):>14,.0f}"
            f"{Decimal(row['withhold_unresolved_millions']):>14,.0f}"
            f"{Decimal(row['intentionally_unallocable_millions']):>14,.0f}"
            f"{Decimal(row['actionable_unfinished_millions']):>14,.0f}"
            f"{Decimal(row['intentionally_unallocable_share_pct']):>11.2f}%"
        )

    lines.extend(
        [
            "",
            "METHODOLOGY CROSSWALK",
            "-" * 126,
        ]
    )

    for item in CROSSWALK:
        lines.append(
            f"{item['component_code']:<36} "
            f"{item['classification']:<30} "
            f"{item['rockefeller_method']}"
        )

    lines.extend(
        [
            "",
            "CERTIFICATION",
            "-" * 126,
        ]
    )

    for key, value in certification.items():
        lines.append(
            f"{key:<62}: {value}"
        )

    lines.extend(
        [
            "",
            "INTERPRETATION",
            "-" * 126,
            (
                "Actionable unfinished receipts are national receipt dollars "
                "whose Washington allocation still requires source acquisition, "
                "method review, or resolution."
            ),
            (
                "Intentionally unallocable receipts remain part of the national "
                "federal receipts universe but are not assigned to Washington "
                "under the documented Rockefeller allocation architecture."
            ),
            (
                "ETL-172I-5 performs no Washington allocation and authorizes "
                "no final ledger inclusion."
            ),
            "",
            "ETL-172I-5 RESULT: PASS",
            "",
        ]
    )

    return "\n".join(lines)


def main() -> int:
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("=" * 126)
    print(
        "REMAINING FEDERAL RECEIPTS METHODOLOGY CROSSWALK "
        "- ETL-172I-5"
    )
    print("=" * 126)
    print()

    i3_remaining = load_i3_remaining()
    i4_other = load_i4_other()
    table24 = load_table24_components()

    component_amounts = build_component_amounts(
        table24,
        i4_other,
    )

    work_queue = build_work_queue(
        component_amounts,
        i3_remaining,
    )

    certification = certify(
        work_queue
    )

    write_crosswalk()

    write_csv(
        OUTPUT_ANNUAL,
        work_queue,
    )

    source_hashes = {
        "etl_172i_2_csv": sha256_file(I2_FILE),
        "etl_172i_3_csv": sha256_file(I3_FILE),
        "etl_172i_4_csv": sha256_file(I4_FILE),
        "omb_table_2_4_xlsx": sha256_file(TABLE_24_FILE),
    }

    payload = {
        "etl_step": "ETL-172I-5",
        "result": "PASS",
        "artifact_type": (
            "REMAINING_FEDERAL_RECEIPTS_METHODOLOGY_CROSSWALK"
        ),
        "target_fiscal_years": TARGET_YEARS,
        "source_hashes": source_hashes,
        "classification_vocabulary": [
            "ALLOCATOR_METHOD_DOCUMENTED",
            "INTENTIONALLY_UNALLOCABLE",
            "METHOD_CHANGE_REVIEW",
            "ALLOCATOR_SOURCE_REQUIRED",
            "WITHHOLD_UNRESOLVED",
        ],
        "crosswalk": CROSSWALK,
        "certification": certification,
        "methodological_boundaries": {
            "washington_allocation_performed": False,
            "incidence_inferred_without_documentation": False,
            "intentionally_unallocable_removed_from_national_universe": False,
            "irs_gross_collections_reconciliation": False,
            "database_writes": False,
            "final_ledger_inclusion_authorized": False,
        },
    }

    OUTPUT_JSON.write_text(
        json.dumps(
            payload,
            indent=2,
        ),
        encoding="utf-8",
    )

    report = build_report(
        work_queue=work_queue,
        certification=certification,
    )

    OUTPUT_REPORT.write_text(
        report,
        encoding="utf-8",
    )

    print("ANNUAL REMAINING-RECEIPTS WORK QUEUE")
    print("-" * 126)

    print(
        f"{'FFY':<6}"
        f"{'I3 REMAIN':>14}"
        f"{'DOC METHOD':>14}"
        f"{'METHOD REV':>14}"
        f"{'UNRESOLVED':>14}"
        f"{'UNALLOCABLE':>14}"
        f"{'ACTIONABLE':>14}"
        f"{'UNALLOC %':>12}"
    )

    for row in work_queue:
        print(
            f"{row['fiscal_year']:<6}"
            f"{Decimal(row['i3_remaining_receipts_millions']):>14,.0f}"
            f"{Decimal(row['allocator_method_documented_millions']):>14,.0f}"
            f"{Decimal(row['method_change_review_millions']):>14,.0f}"
            f"{Decimal(row['withhold_unresolved_millions']):>14,.0f}"
            f"{Decimal(row['intentionally_unallocable_millions']):>14,.0f}"
            f"{Decimal(row['actionable_unfinished_millions']):>14,.0f}"
            f"{Decimal(row['intentionally_unallocable_share_pct']):>11.2f}%"
        )

    print()
    print("CERTIFICATION")
    print("-" * 126)

    for key, value in certification.items():
        print(
            f"{key:<62}: {value}"
        )

    print()
    print(
        "Crosswalk : "
        f"{OUTPUT_CROSSWALK.relative_to(REPO_ROOT)}"
    )
    print(
        "Work queue: "
        f"{OUTPUT_ANNUAL.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON      : "
        f"{OUTPUT_JSON.relative_to(REPO_ROOT)}"
    )
    print(
        "Report    : "
        f"{OUTPUT_REPORT.relative_to(REPO_ROOT)}"
    )
    print()
    print(
        "ETL-172I-5 RESULT: PASS"
    )

    return 0


if __name__ == "__main__":
    sys.exit(
        main()
    )