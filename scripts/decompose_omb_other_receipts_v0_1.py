from __future__ import annotations

import csv
import hashlib
import json
import sys
from decimal import Decimal, getcontext
from pathlib import Path
from typing import Any

import openpyxl


# =============================================================================
# ETL-172I-4
# OMB "OTHER RECEIPTS" DECOMPOSITION
#
# Purpose:
#   Extract and certify the FFY2015-FFY2024 decomposition of OMB Historical
#   Table 2.5, "Composition of 'Other Receipts'", and reconcile the published
#   Table 2.5 Other Receipts total to the independently certified Table 2.1
#   Other Receipts control from ETL-172I-1.
#
# Table 2.5 accounting identities:
#
#   Miscellaneous Receipts Total
#       = Federal Reserve Deposits
#       + Miscellaneous All Other
#
#   Total Other Receipts
#       = Estate and Gift Taxes
#       + Customs Duties and Fees
#       + Miscellaneous Receipts Total
#
#   Table 2.5 Total Other Receipts
#       = Table 2.1 Other Receipts
#
# Memorandum columns:
#   H = Trust Fund Amounts Included in Other Receipts:
#       Customs Duties and Fees
#   I = Trust Fund Amounts Included in Other Receipts:
#       All Other
#
# These memorandum amounts are already included in the published Other
# Receipts totals and MUST NOT be added again.
#
# This step does NOT:
#   - allocate any receipt category to Washington;
#   - classify Federal Reserve deposits as allocable/unallocable;
#   - classify Miscellaneous All Other;
#   - apply Rockefeller methodology;
#   - reconcile to IRS gross collections;
#   - write to PostgreSQL;
#   - authorize final ledger inclusion.
# =============================================================================


getcontext().prec = 50


REPO_ROOT = Path(__file__).resolve().parents[1]

SOURCE_FILE = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "source"
    / "BUDGET-2027-TAB-3-5.xlsx"
)

I1_FILE = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "federal_receipts_universe_v0_1"
    / "national_federal_receipts_universe.csv"
)

I1_CERTIFICATION = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "federal_receipts_universe_v0_1"
    / "certification.json"
)

OUTPUT_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "other_receipts_decomposition_v0_1"
)

OUTPUT_DETAIL = (
    OUTPUT_DIR
    / "national_other_receipts_decomposition.csv"
)

OUTPUT_ANNUAL = (
    OUTPUT_DIR
    / "national_other_receipts_annual_controls.csv"
)

OUTPUT_JSON = (
    OUTPUT_DIR
    / "certification.json"
)

OUTPUT_REPORT = (
    OUTPUT_DIR
    / "certification_report.txt"
)


TARGET_YEARS = list(range(2015, 2025))

MILLIONS_TO_DOLLARS = Decimal("1000000")

EXPECTED_TITLE = (
    'Table 2.5 - COMPOSITION OF "OTHER RECEIPTS":  1940-2031'
)

EXPECTED_UNITS = "(in millions of dollars)"

EXPECTED_SHEET = "Table"


LEAF_COMPONENTS = [
    {
        "component_code": "ESTATE_AND_GIFT_TAXES",
        "component_label": "Estate and Gift Taxes",
        "source_column": "C",
        "accounting_role": "OTHER_RECEIPTS_LEAF_COMPONENT",
    },
    {
        "component_code": "CUSTOMS_DUTIES_AND_FEES",
        "component_label": "Customs Duties and Fees",
        "source_column": "D",
        "accounting_role": "OTHER_RECEIPTS_LEAF_COMPONENT",
    },
    {
        "component_code": "FEDERAL_RESERVE_DEPOSITS",
        "component_label": "Federal Reserve Deposits (1)",
        "source_column": "F",
        "accounting_role": "MISCELLANEOUS_RECEIPTS_LEAF_COMPONENT",
    },
    {
        "component_code": "MISCELLANEOUS_ALL_OTHER",
        "component_label": "All Other",
        "source_column": "G",
        "accounting_role": "MISCELLANEOUS_RECEIPTS_LEAF_COMPONENT",
    },
]


MEMORANDUM_COMPONENTS = [
    {
        "component_code": "MEMO_TRUST_FUND_CUSTOMS_DUTIES_AND_FEES",
        "component_label": "Customs Duties and Fees",
        "source_column": "H",
    },
    {
        "component_code": "MEMO_TRUST_FUND_ALL_OTHER",
        "component_label": "All Other",
        "source_column": "I",
    },
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(
            lambda: handle.read(1024 * 1024),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest().upper()


def decimal_value(
    value: Any,
    *,
    context: str,
) -> Decimal:
    if value is None:
        raise RuntimeError(
            "MISSING NUMERIC VALUE\n"
            f"{context}"
        )

    try:
        result = Decimal(str(value))
    except Exception as exc:
        raise RuntimeError(
            "NUMERIC PARSE FAILURE\n"
            f"{context}: {value!r}"
        ) from exc

    if not result.is_finite():
        raise RuntimeError(
            "NON-FINITE NUMERIC VALUE\n"
            f"{context}: {value!r}"
        )

    return result


def plain_decimal(value: Decimal) -> str:
    return format(value, "f")


def load_i1_certification() -> dict[str, Any]:
    if not I1_CERTIFICATION.exists():
        raise RuntimeError(
            "I-1 CERTIFICATION FILE NOT FOUND\n"
            f"{I1_CERTIFICATION.relative_to(REPO_ROOT)}"
        )

    payload = json.loads(
        I1_CERTIFICATION.read_text(
            encoding="utf-8"
        )
    )

    if payload.get("etl_step") != "ETL-172I-1":
        raise RuntimeError(
            "I-1 CERTIFICATION CONTRACT FAILURE\n"
            f"Expected etl_step: ETL-172I-1\n"
            f"Observed etl_step: {payload.get('etl_step')!r}"
        )

    if payload.get("result") != "PASS":
        raise RuntimeError(
            "I-1 CERTIFICATION IS NOT PASS"
        )

    return payload


def load_i1_other_controls() -> dict[int, Decimal]:
    if not I1_FILE.exists():
        raise RuntimeError(
            "I-1 INPUT FILE NOT FOUND\n"
            f"{I1_FILE.relative_to(REPO_ROOT)}"
        )

    with I1_FILE.open(
        "r",
        newline="",
        encoding="utf-8",
    ) as handle:
        reader = csv.DictReader(handle)

        if reader.fieldnames is None:
            raise RuntimeError(
                "I-1 HEADER FAILURE"
            )

        candidate_columns = [
            "other_receipts_millions",
            "other_millions",
        ]

        other_column = None

        for candidate in candidate_columns:
            if candidate in reader.fieldnames:
                other_column = candidate
                break

        if other_column is None:
            raise RuntimeError(
                "I-1 OTHER RECEIPTS COLUMN NOT FOUND\n"
                f"Candidates: {candidate_columns}\n"
                f"Available : {reader.fieldnames}"
            )

        output: dict[int, Decimal] = {}

        for row in reader:
            year = int(
                row["fiscal_year"]
            )

            if year not in TARGET_YEARS:
                continue

            if year in output:
                raise RuntimeError(
                    "I-1 DUPLICATE FISCAL YEAR\n"
                    f"FFY{year}"
                )

            output[year] = decimal_value(
                row[other_column],
                context=(
                    f"FFY{year} "
                    "I-1 Other Receipts"
                ),
            )

    if set(output) != set(TARGET_YEARS):
        raise RuntimeError(
            "I-1 FISCAL-YEAR COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {sorted(output)}"
        )

    return output


def verify_source_file() -> None:
    if not SOURCE_FILE.exists():
        raise RuntimeError(
            "SOURCE FILE NOT FOUND\n"
            f"{SOURCE_FILE.relative_to(REPO_ROOT)}"
        )

    first_bytes = SOURCE_FILE.read_bytes()[:4]

    if first_bytes != b"PK\x03\x04":
        raise RuntimeError(
            "SOURCE FILE SIGNATURE FAILURE\n"
            f"Expected XLSX ZIP signature: "
            f"50 4B 03 04\n"
            f"Observed: "
            f"{' '.join(f'{b:02X}' for b in first_bytes)}"
        )


def load_workbook() -> openpyxl.Workbook:
    verify_source_file()

    workbook = openpyxl.load_workbook(
        SOURCE_FILE,
        data_only=True,
        read_only=False,
    )

    if workbook.sheetnames != [EXPECTED_SHEET]:
        raise RuntimeError(
            "WORKBOOK SHEET CONTRACT FAILURE\n"
            f"Expected: {[EXPECTED_SHEET]}\n"
            f"Actual  : {workbook.sheetnames}"
        )

    worksheet = workbook[EXPECTED_SHEET]

    if worksheet["A1"].value != EXPECTED_TITLE:
        raise RuntimeError(
            "TABLE TITLE FAILURE\n"
            f"Expected: {EXPECTED_TITLE!r}\n"
            f"Observed: {worksheet['A1'].value!r}"
        )

    if worksheet["A2"].value != EXPECTED_UNITS:
        raise RuntimeError(
            "TABLE UNIT LABEL FAILURE\n"
            f"Expected: {EXPECTED_UNITS!r}\n"
            f"Observed: {worksheet['A2'].value!r}"
        )

    return workbook


def verify_headers(
    worksheet: Any,
) -> None:
    expected = {
        "A3": "Fiscal Year",
        "B3": 'Total "Other Receipts"',
        "C3": "Estate and Gift Taxes",
        "D3": "Customs Duties and Fees",
        "E3": "Miscellaneous Receipts",
        "E4": "Total",
        "F4": "Federal Reserve Deposits (1)",
        "G4": "All Other",
        "H3": (
            'Memorandum: Trust Fund Amounts Included '
            'in "Other Receipts"'
        ),
        "H4": "Customs Duties and Fees",
        "I4": "All Other",
    }

    for cell, expected_value in expected.items():
        observed = worksheet[cell].value

        if observed != expected_value:
            raise RuntimeError(
                "TABLE HEADER FAILURE\n"
                f"Cell    : {cell}\n"
                f"Expected: {expected_value!r}\n"
                f"Observed: {observed!r}"
            )


def locate_target_rows(
    worksheet: Any,
) -> dict[int, int]:
    output: dict[int, int] = {}

    for row_number in range(
        5,
        worksheet.max_row + 1,
    ):
        value = worksheet.cell(
            row_number,
            1,
        ).value

        if value is None:
            continue

        text = str(value).strip()

        if not text.isdigit():
            continue

        year = int(text)

        if year not in TARGET_YEARS:
            continue

        if year in output:
            raise RuntimeError(
                "DUPLICATE SOURCE FISCAL YEAR\n"
                f"FFY{year}"
            )

        output[year] = row_number

    if set(output) != set(TARGET_YEARS):
        raise RuntimeError(
            "SOURCE FISCAL-YEAR COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {sorted(output)}"
        )

    return output


def extract_source(
    worksheet: Any,
    row_map: dict[int, int],
    i1_controls: dict[int, Decimal],
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    detail_rows: list[dict[str, Any]] = []
    annual_rows: list[dict[str, Any]] = []

    for year in TARGET_YEARS:
        row_number = row_map[year]

        total_other = decimal_value(
            worksheet.cell(
                row_number,
                2,
            ).value,
            context=(
                f"FFY{year} "
                "Table 2.5 Total Other Receipts"
            ),
        )

        estate_gift = decimal_value(
            worksheet.cell(
                row_number,
                3,
            ).value,
            context=(
                f"FFY{year} "
                "Estate and Gift Taxes"
            ),
        )

        customs = decimal_value(
            worksheet.cell(
                row_number,
                4,
            ).value,
            context=(
                f"FFY{year} "
                "Customs Duties and Fees"
            ),
        )

        miscellaneous_total = decimal_value(
            worksheet.cell(
                row_number,
                5,
            ).value,
            context=(
                f"FFY{year} "
                "Miscellaneous Receipts Total"
            ),
        )

        federal_reserve = decimal_value(
            worksheet.cell(
                row_number,
                6,
            ).value,
            context=(
                f"FFY{year} "
                "Federal Reserve Deposits"
            ),
        )

        miscellaneous_all_other = decimal_value(
            worksheet.cell(
                row_number,
                7,
            ).value,
            context=(
                f"FFY{year} "
                "Miscellaneous All Other"
            ),
        )

        memo_customs = decimal_value(
            worksheet.cell(
                row_number,
                8,
            ).value,
            context=(
                f"FFY{year} "
                "Memorandum Customs"
            ),
        )

        memo_all_other = decimal_value(
            worksheet.cell(
                row_number,
                9,
            ).value,
            context=(
                f"FFY{year} "
                "Memorandum All Other"
            ),
        )

        # ---------------------------------------------------------------------
        # Identity 1:
        # Miscellaneous = Federal Reserve + All Other
        # ---------------------------------------------------------------------

        reconstructed_miscellaneous = (
            federal_reserve
            + miscellaneous_all_other
        )

        miscellaneous_difference = (
            miscellaneous_total
            - reconstructed_miscellaneous
        )

        if miscellaneous_difference != 0:
            raise RuntimeError(
                "MISCELLANEOUS RECEIPTS IDENTITY FAILURE\n"
                f"FFY{year}\n"
                f"Published miscellaneous : "
                f"{miscellaneous_total}\n"
                f"Federal Reserve         : "
                f"{federal_reserve}\n"
                f"All Other               : "
                f"{miscellaneous_all_other}\n"
                f"Difference              : "
                f"{miscellaneous_difference}"
            )

        # ---------------------------------------------------------------------
        # Identity 2:
        # Other = Estate/Gift + Customs + Miscellaneous
        # ---------------------------------------------------------------------

        reconstructed_other = (
            estate_gift
            + customs
            + miscellaneous_total
        )

        other_component_difference = (
            total_other
            - reconstructed_other
        )

        if other_component_difference != 0:
            raise RuntimeError(
                "OTHER RECEIPTS COMPONENT IDENTITY FAILURE\n"
                f"FFY{year}\n"
                f"Published Other : {total_other}\n"
                f"Estate/Gift     : {estate_gift}\n"
                f"Customs         : {customs}\n"
                f"Miscellaneous   : {miscellaneous_total}\n"
                f"Difference      : "
                f"{other_component_difference}"
            )

        # ---------------------------------------------------------------------
        # Identity 3:
        # Table 2.5 Other = certified Table 2.1 Other
        # ---------------------------------------------------------------------

        i1_other = i1_controls[year]

        i1_difference = (
            total_other
            - i1_other
        )

        if i1_difference != 0:
            raise RuntimeError(
                "TABLE 2.5 / TABLE 2.1 RECONCILIATION FAILURE\n"
                f"FFY{year}\n"
                f"Table 2.5 Other : {total_other}\n"
                f"Table 2.1 Other : {i1_other}\n"
                f"Difference      : {i1_difference}"
            )

        # ---------------------------------------------------------------------
        # Detail leaf components.
        # ---------------------------------------------------------------------

        leaf_values = {
            "ESTATE_AND_GIFT_TAXES": estate_gift,
            "CUSTOMS_DUTIES_AND_FEES": customs,
            "FEDERAL_RESERVE_DEPOSITS": federal_reserve,
            "MISCELLANEOUS_ALL_OTHER": miscellaneous_all_other,
        }

        for component in LEAF_COMPONENTS:
            code = component[
                "component_code"
            ]

            amount = leaf_values[code]

            detail_rows.append(
                {
                    "fiscal_year": year,
                    "source_sheet": EXPECTED_SHEET,
                    "source_row": row_number,
                    "source_column": component[
                        "source_column"
                    ],
                    "component_code": code,
                    "component_label": component[
                        "component_label"
                    ],
                    "accounting_role": component[
                        "accounting_role"
                    ],
                    "source_units": (
                        "millions of dollars"
                    ),
                    "amount_millions": (
                        plain_decimal(amount)
                    ),
                    "amount_dollars": (
                        plain_decimal(
                            amount
                            * MILLIONS_TO_DOLLARS
                        )
                    ),
                    "allocation_status": (
                        "NOT_CLASSIFIED_IN_ETL_172I_4"
                    ),
                    "ledger_status": (
                        "DECOMPOSITION_ONLY"
                    ),
                }
            )

        # ---------------------------------------------------------------------
        # Memorandum components.
        #
        # These are included within Other Receipts. They are retained for
        # provenance only and are explicitly NONADDITIVE.
        # ---------------------------------------------------------------------

        memo_values = {
            "MEMO_TRUST_FUND_CUSTOMS_DUTIES_AND_FEES": (
                memo_customs
            ),
            "MEMO_TRUST_FUND_ALL_OTHER": (
                memo_all_other
            ),
        }

        for component in MEMORANDUM_COMPONENTS:
            code = component[
                "component_code"
            ]

            amount = memo_values[code]

            detail_rows.append(
                {
                    "fiscal_year": year,
                    "source_sheet": EXPECTED_SHEET,
                    "source_row": row_number,
                    "source_column": component[
                        "source_column"
                    ],
                    "component_code": code,
                    "component_label": component[
                        "component_label"
                    ],
                    "accounting_role": (
                        "MEMORANDUM_INCLUDED_IN_OTHER_RECEIPTS"
                    ),
                    "source_units": (
                        "millions of dollars"
                    ),
                    "amount_millions": (
                        plain_decimal(amount)
                    ),
                    "amount_dollars": (
                        plain_decimal(
                            amount
                            * MILLIONS_TO_DOLLARS
                        )
                    ),
                    "allocation_status": (
                        "NOT_APPLICABLE_MEMORANDUM"
                    ),
                    "ledger_status": (
                        "NONADDITIVE_MEMORANDUM"
                    ),
                }
            )

        annual_rows.append(
            {
                "fiscal_year": year,

                "table_2_5_total_other_receipts_millions": (
                    plain_decimal(total_other)
                ),

                "estate_and_gift_taxes_millions": (
                    plain_decimal(estate_gift)
                ),

                "customs_duties_and_fees_millions": (
                    plain_decimal(customs)
                ),

                "miscellaneous_receipts_total_millions": (
                    plain_decimal(
                        miscellaneous_total
                    )
                ),

                "federal_reserve_deposits_millions": (
                    plain_decimal(
                        federal_reserve
                    )
                ),

                "miscellaneous_all_other_millions": (
                    plain_decimal(
                        miscellaneous_all_other
                    )
                ),

                "reconstructed_miscellaneous_receipts_millions": (
                    plain_decimal(
                        reconstructed_miscellaneous
                    )
                ),

                "miscellaneous_identity_difference_millions": (
                    plain_decimal(
                        miscellaneous_difference
                    )
                ),

                "reconstructed_other_receipts_millions": (
                    plain_decimal(
                        reconstructed_other
                    )
                ),

                "other_component_identity_difference_millions": (
                    plain_decimal(
                        other_component_difference
                    )
                ),

                "table_2_1_other_receipts_millions": (
                    plain_decimal(i1_other)
                ),

                "table_2_5_vs_table_2_1_difference_millions": (
                    plain_decimal(
                        i1_difference
                    )
                ),

                "memo_trust_fund_customs_millions": (
                    plain_decimal(
                        memo_customs
                    )
                ),

                "memo_trust_fund_all_other_millions": (
                    plain_decimal(
                        memo_all_other
                    )
                ),

                "memo_accounting_status": (
                    "INCLUDED_IN_OTHER_RECEIPTS_NONADDITIVE"
                ),

                "allocation_status": (
                    "NOT_CLASSIFIED_IN_ETL_172I_4"
                ),

                "ledger_status": (
                    "DECOMPOSITION_ONLY"
                ),
            }
        )

    return (
        detail_rows,
        annual_rows,
    )


def certify(
    *,
    detail_rows: list[dict[str, Any]],
    annual_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    if len(annual_rows) != 10:
        raise RuntimeError(
            "ANNUAL OBSERVATION COUNT FAILURE\n"
            f"Expected: 10\n"
            f"Actual  : {len(annual_rows)}"
        )

    expected_detail_rows = (
        len(TARGET_YEARS)
        * (
            len(LEAF_COMPONENTS)
            + len(MEMORANDUM_COMPONENTS)
        )
    )

    if len(detail_rows) != expected_detail_rows:
        raise RuntimeError(
            "DETAIL OBSERVATION COUNT FAILURE\n"
            f"Expected: {expected_detail_rows}\n"
            f"Actual  : {len(detail_rows)}"
        )

    years = [
        int(row["fiscal_year"])
        for row in annual_rows
    ]

    if years != TARGET_YEARS:
        raise RuntimeError(
            "ANNUAL FISCAL-YEAR COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {years}"
        )

    for row in annual_rows:
        if (
            Decimal(
                row[
                    "miscellaneous_identity_difference_millions"
                ]
            )
            != 0
        ):
            raise RuntimeError(
                "CERTIFICATION MISCELLANEOUS "
                "IDENTITY FAILURE"
            )

        if (
            Decimal(
                row[
                    "other_component_identity_difference_millions"
                ]
            )
            != 0
        ):
            raise RuntimeError(
                "CERTIFICATION OTHER "
                "COMPONENT IDENTITY FAILURE"
            )

        if (
            Decimal(
                row[
                    "table_2_5_vs_table_2_1_difference_millions"
                ]
            )
            != 0
        ):
            raise RuntimeError(
                "CERTIFICATION TABLE 2.1 "
                "RECONCILIATION FAILURE"
            )

        if (
            row["memo_accounting_status"]
            != "INCLUDED_IN_OTHER_RECEIPTS_NONADDITIVE"
        ):
            raise RuntimeError(
                "MEMORANDUM ACCOUNTING "
                "BOUNDARY FAILURE"
            )

        if (
            row["allocation_status"]
            != "NOT_CLASSIFIED_IN_ETL_172I_4"
        ):
            raise RuntimeError(
                "ALLOCATION CLASSIFICATION "
                "BOUNDARY FAILURE"
            )

        if (
            row["ledger_status"]
            != "DECOMPOSITION_ONLY"
        ):
            raise RuntimeError(
                "LEDGER BOUNDARY FAILURE"
            )

    memo_rows = [
        row
        for row in detail_rows
        if (
            row["accounting_role"]
            == "MEMORANDUM_INCLUDED_IN_OTHER_RECEIPTS"
        )
    ]

    if len(memo_rows) != 20:
        raise RuntimeError(
            "MEMORANDUM OBSERVATION COUNT FAILURE\n"
            f"Expected: 20\n"
            f"Actual  : {len(memo_rows)}"
        )

    for row in memo_rows:
        if (
            row["ledger_status"]
            != "NONADDITIVE_MEMORANDUM"
        ):
            raise RuntimeError(
                "MEMORANDUM NONADDITIVITY FAILURE"
            )

    return {
        "expected_annual_observations": 10,
        "actual_annual_observations": len(
            annual_rows
        ),
        "expected_detail_observations": (
            expected_detail_rows
        ),
        "actual_detail_observations": len(
            detail_rows
        ),
        "fiscal_year_coverage": "PASS",
        "fiscal_year_uniqueness": "PASS",
        "source_file_signature": "PASS",
        "source_sheet_contract": "PASS",
        "source_title_contract": "PASS",
        "source_units_contract": "PASS",
        "source_header_contract": "PASS",
        "miscellaneous_receipts_identity": "PASS",
        "other_receipts_component_identity": "PASS",
        "table_2_5_to_table_2_1_reconciliation": "PASS",
        "million_to_dollar_conversion": "PASS",
        "memorandum_preservation": "PASS",
        "memorandum_nonadditivity": "PASS",
        "allocation_classification_boundary": "PASS",
        "rockefeller_methodology_boundary": "PASS",
        "irs_reconciliation_boundary": "PASS",
        "database_write_boundary": "PASS",
        "ledger_boundary": "PASS",
    }


def write_csv(
    path: Path,
    rows: list[dict[str, Any]],
) -> None:
    if not rows:
        raise RuntimeError(
            "Refusing to write empty CSV."
        )

    with path.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(
                rows[0].keys()
            ),
        )

        writer.writeheader()
        writer.writerows(rows)


def build_report(
    *,
    source_hash: str,
    i1_hash: str,
    annual_rows: list[dict[str, Any]],
    certification: dict[str, Any],
) -> str:
    lines: list[str] = [
        "=" * 116,
        (
            "OMB OTHER RECEIPTS DECOMPOSITION "
            "- ETL-172I-4"
        ),
        "=" * 116,
        "",
        "UPSTREAM PROVENANCE",
        "-" * 116,
        (
            f"{'Table 2.5 SHA256':<30}: "
            f"{source_hash}"
        ),
        (
            f"{'ETL-172I-1 SHA256':<30}: "
            f"{i1_hash}"
        ),
        "",
        "ANNUAL OTHER RECEIPTS DECOMPOSITION",
        "-" * 116,
        (
            f"{'FFY':<6}"
            f"{'OTHER':>13}"
            f"{'EST/GIFT':>13}"
            f"{'CUSTOMS':>13}"
            f"{'FED RES':>13}"
            f"{'MISC OTHER':>13}"
            f"{'MEMO CUST':>13}"
            f"{'MEMO OTHER':>13}"
        ),
        "-" * 116,
    ]

    for row in annual_rows:
        lines.append(
            f"{row['fiscal_year']:<6}"
            f"{Decimal(row['table_2_5_total_other_receipts_millions']):>13,.0f}"
            f"{Decimal(row['estate_and_gift_taxes_millions']):>13,.0f}"
            f"{Decimal(row['customs_duties_and_fees_millions']):>13,.0f}"
            f"{Decimal(row['federal_reserve_deposits_millions']):>13,.0f}"
            f"{Decimal(row['miscellaneous_all_other_millions']):>13,.0f}"
            f"{Decimal(row['memo_trust_fund_customs_millions']):>13,.0f}"
            f"{Decimal(row['memo_trust_fund_all_other_millions']):>13,.0f}"
        )

    lines.extend(
        [
            "",
            "CERTIFICATION",
            "-" * 116,
        ]
    )

    for key, value in certification.items():
        lines.append(
            f"{key:<54}: {value}"
        )

    lines.extend(
        [
            "",
            "BOUNDARIES",
            "-" * 116,
            (
                "ETL-172I-4 is a national-source decomposition and "
                "certification step only."
            ),
            (
                "No Table 2.5 component is classified as Washington-"
                "allocatable or unallocable in this step."
            ),
            (
                "The memorandum trust-fund amounts are already included "
                "within Other Receipts and are preserved as nonadditive."
            ),
            (
                "No Rockefeller allocation method is applied."
            ),
            (
                "No Washington receipt dollars are calculated."
            ),
            (
                "No database writes or final ledger inclusion occur."
            ),
            "",
            "ETL-172I-4 RESULT: PASS",
            "",
        ]
    )

    return "\n".join(lines)


def main() -> int:
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("=" * 116)
    print(
        "OMB OTHER RECEIPTS DECOMPOSITION "
        "- ETL-172I-4"
    )
    print("=" * 116)
    print()

    load_i1_certification()

    i1_controls = (
        load_i1_other_controls()
    )

    workbook = load_workbook()

    worksheet = workbook[
        EXPECTED_SHEET
    ]

    verify_headers(
        worksheet
    )

    row_map = locate_target_rows(
        worksheet
    )

    detail_rows, annual_rows = (
        extract_source(
            worksheet,
            row_map,
            i1_controls,
        )
    )

    certification = certify(
        detail_rows=detail_rows,
        annual_rows=annual_rows,
    )

    source_hash = sha256_file(
        SOURCE_FILE
    )

    i1_hash = sha256_file(
        I1_FILE
    )

    write_csv(
        OUTPUT_DETAIL,
        detail_rows,
    )

    write_csv(
        OUTPUT_ANNUAL,
        annual_rows,
    )

    payload = {
        "etl_step": "ETL-172I-4",
        "result": "PASS",
        "artifact_type": (
            "OMB_OTHER_RECEIPTS_DECOMPOSITION"
        ),
        "source": {
            "file": str(
                SOURCE_FILE.relative_to(
                    REPO_ROOT
                )
            ),
            "sha256": source_hash,
            "sheet": EXPECTED_SHEET,
            "title": EXPECTED_TITLE,
            "units": (
                "millions of dollars"
            ),
        },
        "upstream": {
            "etl_172i_1_file": str(
                I1_FILE.relative_to(
                    REPO_ROOT
                )
            ),
            "etl_172i_1_sha256": (
                i1_hash
            ),
        },
        "target_fiscal_years": (
            TARGET_YEARS
        ),
        "leaf_components": [
            item["component_code"]
            for item in LEAF_COMPONENTS
        ],
        "memorandum_components": [
            item["component_code"]
            for item in MEMORANDUM_COMPONENTS
        ],
        "certification": (
            certification
        ),
        "methodological_boundaries": {
            "washington_allocation_performed": False,
            "allocation_classification_performed": False,
            "federal_reserve_incidence_decision": False,
            "miscellaneous_all_other_incidence_decision": False,
            "rockefeller_replication": False,
            "irs_gross_collections_reconciliation": False,
            "database_writes": False,
            "final_ledger_inclusion_authorized": False,
            "memorandum_columns_are_additive": False,
        },
    }

    OUTPUT_JSON.write_text(
        json.dumps(
            payload,
            indent=2,
        ),
        encoding="utf-8",
    )

    report = build_report(
        source_hash=source_hash,
        i1_hash=i1_hash,
        annual_rows=annual_rows,
        certification=certification,
    )

    OUTPUT_REPORT.write_text(
        report,
        encoding="utf-8",
    )

    print(
        "ANNUAL OTHER RECEIPTS DECOMPOSITION"
    )
    print("-" * 116)

    print(
        f"{'FFY':<6}"
        f"{'OTHER':>13}"
        f"{'EST/GIFT':>13}"
        f"{'CUSTOMS':>13}"
        f"{'FED RES':>13}"
        f"{'MISC OTHER':>13}"
        f"{'MEMO CUST':>13}"
        f"{'MEMO OTHER':>13}"
    )

    for row in annual_rows:
        print(
            f"{row['fiscal_year']:<6}"
            f"{Decimal(row['table_2_5_total_other_receipts_millions']):>13,.0f}"
            f"{Decimal(row['estate_and_gift_taxes_millions']):>13,.0f}"
            f"{Decimal(row['customs_duties_and_fees_millions']):>13,.0f}"
            f"{Decimal(row['federal_reserve_deposits_millions']):>13,.0f}"
            f"{Decimal(row['miscellaneous_all_other_millions']):>13,.0f}"
            f"{Decimal(row['memo_trust_fund_customs_millions']):>13,.0f}"
            f"{Decimal(row['memo_trust_fund_all_other_millions']):>13,.0f}"
        )

    print()
    print("CERTIFICATION")
    print("-" * 116)

    for key, value in certification.items():
        print(
            f"{key:<54}: {value}"
        )

    print()
    print(
        "Detail : "
        f"{OUTPUT_DETAIL.relative_to(REPO_ROOT)}"
    )
    print(
        "Annual : "
        f"{OUTPUT_ANNUAL.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON   : "
        f"{OUTPUT_JSON.relative_to(REPO_ROOT)}"
    )
    print(
        "Report : "
        f"{OUTPUT_REPORT.relative_to(REPO_ROOT)}"
    )
    print()
    print(
        "ETL-172I-4 RESULT: PASS"
    )

    return 0


if __name__ == "__main__":
    sys.exit(
        main()
    )