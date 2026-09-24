#!/usr/bin/env python3
"""
profile_cms_partd_contract_pbp_relationship_v0_1.py

Read-only relationship profiler for the CMS Medicare Part D CPSC enrollment
classified by nonmatch validator v0.1 as SAME_CONTRACT_OTHER_PBP.

Purpose
-------
Describe exactly how unmatched Washington CPSC PBPs relate to the PBPs
published in the corresponding annual CMS PartDPlans payment workbook.

This script:
- starts from the certified nonmatch-classification output;
- does NOT reread the monthly CPSC corpus;
- inventories same-contract payment PBPs and their published payment fields;
- tests exact Contract+PBP presence in ANY CY2015-CY2024 PartDPlans year;
- describes coexistence of matched/unmatched PBPs using v0.3 matched vocabulary
  when that file exposes Contract/PBP/month fields;
- does NOT copy payment values onto unmatched PBPs;
- does NOT reconstruct dollars, impute enrollment, write to PostgreSQL, or
  authorize ledger inclusion.
"""
from __future__ import annotations
import argparse, csv, hashlib, io, re, sys, zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
import openpyxl

EXPECTED_CONTROL = 5_203_895
PAYMENT_FIELD_ALIASES = {
    "contract": ("CONTRACT_NUMBER", "CONTRACT"),
    "pbp": ("PLAN_BENEFIT_PACKAGE", "PLAN_BENEFIT_PACKAGE_NUMBER", "PBP"),
    "contract_name": ("CONTRACT_NAME",),
    "organization_type": ("ORGANIZATION_TYPE",),
    "direct_subsidy": ("AVERAGE_PART_D_DIRECT_SUBSIDY_PM_PM_PAYMENT",),
    "risk_score": ("AVERAGE_PART_D_RISK_SCORE",),
    "reinsurance": ("AVERAGE_REINSURANCE_PM_PM_PAYMENT",),
    "lics": ("AVERAGE_LOW_INCOME_COST_SHARING_PM_PM_PAYMENT",),
}

def clean(v: Any) -> str:
    return "" if v is None else str(v).strip()

def norm(v: Any) -> str:
    s = clean(v).upper().replace("/", " ")
    return re.sub(r"[^A-Z0-9]+", "_", s).strip("_")

def canon_pbp(v: Any) -> str:
    s = clean(v)
    if not s: return ""
    try: return f"{int(float(s)):03d}"
    except (ValueError, TypeError): return s.upper()

def sha_file(p: Path) -> str:
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1024*1024), b""): h.update(b)
    return h.hexdigest()

def read_csv(p: Path):
    with p.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def write_csv(p: Path, rows, fields):
    with p.open("w", encoding="utf-8-sig", newline="") as f:
        w=csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)

def find_col(header, aliases):
    pos={norm(x):i for i,x in enumerate(header)}
    for a in aliases:
        if norm(a) in pos: return pos[norm(a)]
    return None

def find_partdplans(root: Path, year: int):
    hits=[]
    yd=root/str(year)
    for zp in sorted(yd.glob("*.zip")):
        try:
            with zipfile.ZipFile(zp) as z:
                for m in z.namelist():
                    b=Path(m).name.lower().replace(" ","")
                    if b.endswith(".xlsx") and "partdplans" in b and "recon" not in b:
                        data=z.read(m)
                        hits.append((zp,m,data,hashlib.sha256(data).hexdigest()))
        except zipfile.BadZipFile: pass
    if not hits: raise RuntimeError(f"CY{year}: PartDPlans not found")
    if len({x[3] for x in hits}) != 1:
        raise RuntimeError(f"CY{year}: conflicting PartDPlans bytes")
    return sorted(hits,key=lambda x:(str(x[0]),x[1]))[0]

def load_payment_year(root: Path, year: int):
    zp,member,data,mhash=find_partdplans(root,year)
    wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    try:
        for ws in wb.worksheets:
            for hr,row in enumerate(ws.iter_rows(min_row=1,max_row=30,values_only=True),1):
                idx={k:find_col(row,v) for k,v in PAYMENT_FIELD_ALIASES.items()}
                if idx["contract"] is None or idx["pbp"] is None: continue
                records={}; dups=[]
                for vals in ws.iter_rows(min_row=hr+1,values_only=True):
                    def g(k):
                        i=idx[k]
                        return "" if i is None or i>=len(vals) else clean(vals[i])
                    c=g("contract").upper(); p=canon_pbp(g("pbp"))
                    if not c or not p: continue
                    key=(c,p)
                    rec={
                        "contract":c,"pbp":p,"contract_name":g("contract_name"),
                        "payment_organization_type":g("organization_type"),
                        "direct_subsidy_std_pmpm":g("direct_subsidy"),
                        "partd_risk_score":g("risk_score"),
                        "reinsurance_pmpm":g("reinsurance"),
                        "lics_pmpm":g("lics"),
                    }
                    if key in records: dups.append(key)
                    else: records[key]=rec
                return records,dups,{
                    "year":year,"zip_path":str(zp),"zip_sha256":sha_file(zp),
                    "member_name":member,"member_sha256":mhash,
                    "sheet_name":ws.title,"header_row":hr,
                    "payment_key_count":len(records)}
    finally: wb.close()
    raise RuntimeError(f"CY{year}: payment table not found")

def normalized_map(r):
    return {norm(k):v for k,v in r.items()}

def pick(r,*names):
    m=normalized_map(r)
    for n in names:
        if norm(n) in m: return clean(m[norm(n)])
    return ""

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--classification-dir",type=Path,default=Path(
        "data/raw/cms/medicare_advantage/validation/partd_cpsc_nonmatch_classification_v0_1"))
    ap.add_argument("--v03-dir",type=Path,default=Path(
        "data/raw/cms/medicare_advantage/validation/partd_cpsc_enrollment_linkage_profile_v0_3"))
    ap.add_argument("--plan-payment-root",type=Path,default=Path(
        "data/raw/cms/medicare_advantage/plan_payment"))
    ap.add_argument("--output-dir",type=Path,default=Path(
        "data/raw/cms/medicare_advantage/validation/partd_contract_pbp_relationship_v0_1"))
    a=ap.parse_args(); a.output_dir.mkdir(parents=True,exist_ok=True)
    errors=[]; warnings=[]

    inp=a.classification_dir/"partd_cpsc_nonmatch_classified_rows.csv"
    if not inp.exists():
        print(f"ERROR missing {inp}",file=sys.stderr); return 2
    rows=read_csv(inp)
    target=[]
    for i,r in enumerate(rows,2):
        if pick(r,"classification_v0_1")!="SAME_CONTRACT_OTHER_PBP": continue
        try:
            target.append({
                "year":int(pick(r,"year")),"month":int(pick(r,"month")),
                "contract":pick(r,"contract","contract_number").upper(),
                "pbp":canon_pbp(pick(r,"pbp","plan_benefit_package")),
                "cpsc_plan_type":pick(r,"cpsc_plan_type","plan_type"),
                "cpsc_organization_type":pick(r,"cpsc_organization_type","organization_type"),
                "numeric_enrollment":int(float(pick(r,"numeric_enrollment","enrollment"))),
            })
        except Exception as e: errors.append(f"input row {i}: {e}")

    pay={}; manifests=[]; duprows=[]
    for y in range(2015,2025):
        try:
            recs,dups,mf=load_payment_year(a.plan_payment_root,y)
            pay[y]=recs; manifests.append(mf)
            duprows += [{"year":y,"contract":c,"pbp":p} for c,p in dups]
        except Exception as e:
            pay[y]={}; errors.append(str(e))

    # All-year exact appearance map.
    years_for_key=defaultdict(list)
    for y,recs in pay.items():
        for k in recs: years_for_key[k].append(y)

    # Aggregate target to contract-year/PBP and month sets.
    enroll=Counter(); months=defaultdict(set); vocab=defaultdict(set)
    for r in target:
        k=(r["year"],r["contract"],r["pbp"])
        enroll[k]+=r["numeric_enrollment"]; months[k].add(r["month"])
        vocab[k].add((r["cpsc_plan_type"],r["cpsc_organization_type"]))

    # Optional matched-vocabulary evidence from v0.3.
    matched_path=a.v03_dir/"partd_cpsc_partd_match_vocabulary.csv"
    matched_by_cym=defaultdict(set); matched_by_cy=defaultdict(set)
    matched_vocabulary_available=False
    if matched_path.exists():
        mr=read_csv(matched_path)
        for r in mr:
            try:
                c=pick(r,"contract","contract_number").upper()
                p=canon_pbp(pick(r,"pbp","plan_benefit_package","plan_id"))
                ys=pick(r,"year","calendar_year"); ms=pick(r,"month","calendar_month")
                if c and p and ys:
                    y=int(ys); matched_by_cy[(y,c)].add(p)
                    if ms:
                        matched_by_cym[(y,int(ms),c)].add(p)
                    matched_vocabulary_available=True
            except Exception: pass
    else:
        warnings.append("v0.3 matched vocabulary file not found; coexistence test unavailable.")

    relationship=[]
    contract_years=defaultdict(lambda:{"unmatched":set(),"enrollment":0})
    for (y,c,p),n in sorted(enroll.items()):
        same_year_pbps=sorted(pp for (cc,pp) in pay.get(y,{}) if cc==c)
        other_years=sorted(yy for yy in years_for_key[(c,p)] if yy!=y)
        month_list=sorted(months[(y,c,p)])
        coexist_months=0
        if matched_vocabulary_available:
            coexist_months=sum(1 for m in month_list if matched_by_cym.get((y,m,c)))
        vt=sorted(vocab[(y,c,p)])
        relationship.append({
            "year":y,"contract":c,"unmatched_cpsc_pbp":p,
            "numeric_enrollment":n,"distinct_months":len(month_list),
            "months":",".join(f"{m:02d}" for m in month_list),
            "cpsc_plan_type_vocabulary":" | ".join(sorted({x[0] for x in vt if x[0]})),
            "cpsc_organization_type_vocabulary":" | ".join(sorted({x[1] for x in vt if x[1]})),
            "same_year_partdplans_pbp_count":len(same_year_pbps),
            "same_year_partdplans_pbps":",".join(same_year_pbps),
            "exact_pbp_other_partdplans_years":",".join(map(str,other_years)),
            "exact_pbp_appears_any_other_year":"YES" if other_years else "NO",
            "months_with_same_contract_matched_pbp_observed":coexist_months if matched_vocabulary_available else "",
        })
        contract_years[(y,c)]["unmatched"].add(p)
        contract_years[(y,c)]["enrollment"]+=n

    # Payment PBP detail for affected contract-years. No propagation.
    payment_detail=[]
    for (y,c),d in sorted(contract_years.items()):
        recs=[v for (cc,p),v in pay.get(y,{}).items() if cc==c]
        # fingerprint published four-field payment vector, useful only descriptively.
        vectors=defaultdict(list)
        for r in recs:
            vec=(r["direct_subsidy_std_pmpm"],r["partd_risk_score"],
                 r["reinsurance_pmpm"],r["lics_pmpm"])
            vectors[vec].append(r["pbp"])
        for r in sorted(recs,key=lambda x:x["pbp"]):
            vec=(r["direct_subsidy_std_pmpm"],r["partd_risk_score"],
                 r["reinsurance_pmpm"],r["lics_pmpm"])
            payment_detail.append({
                "year":y,"contract":c,"payment_pbp":r["pbp"],
                "contract_name":r["contract_name"],
                "payment_organization_type":r["payment_organization_type"],
                "direct_subsidy_std_pmpm":r["direct_subsidy_std_pmpm"],
                "partd_risk_score":r["partd_risk_score"],
                "reinsurance_pmpm":r["reinsurance_pmpm"],
                "lics_pmpm":r["lics_pmpm"],
                "payment_vector_shared_by_pbp_count":len(vectors[vec]),
                "payment_vector_shared_by_pbps":",".join(sorted(vectors[vec])),
            })

    # Contract-year summary.
    cyrows=[]
    for (y,c),d in sorted(contract_years.items()):
        represented=sorted(p for (cc,p) in pay.get(y,{}) if cc==c)
        unmatched=sorted(d["unmatched"])
        vectors=set()
        for p in represented:
            r=pay[y][(c,p)]
            vectors.add((r["direct_subsidy_std_pmpm"],r["partd_risk_score"],
                         r["reinsurance_pmpm"],r["lics_pmpm"]))
        cyrows.append({
            "year":y,"contract":c,"unmatched_numeric_enrollment":d["enrollment"],
            "unmatched_pbp_count":len(unmatched),"unmatched_pbps":",".join(unmatched),
            "partdplans_pbp_count":len(represented),"partdplans_pbps":",".join(represented),
            "distinct_published_payment_vectors":len(vectors),
            "all_represented_pbps_share_one_payment_vector":"YES" if len(vectors)==1 and represented else "NO",
            "matched_vocabulary_pbps":",".join(sorted(matched_by_cy.get((y,c),set()))) if matched_vocabulary_available else "",
        })

    control=sum(r["numeric_enrollment"] for r in target)
    reltotal=sum(r["numeric_enrollment"] for r in relationship)
    cytotal=sum(r["unmatched_numeric_enrollment"] for r in cyrows)
    exact_other=sum(r["numeric_enrollment"] for r in relationship if r["exact_pbp_appears_any_other_year"]=="YES")
    no_other=control-exact_other

    write_csv(a.output_dir/"partd_contract_pbp_relationship.csv",relationship,
              ["year","contract","unmatched_cpsc_pbp","numeric_enrollment","distinct_months","months",
               "cpsc_plan_type_vocabulary","cpsc_organization_type_vocabulary",
               "same_year_partdplans_pbp_count","same_year_partdplans_pbps",
               "exact_pbp_other_partdplans_years","exact_pbp_appears_any_other_year",
               "months_with_same_contract_matched_pbp_observed"])
    write_csv(a.output_dir/"partd_affected_contract_year_summary.csv",cyrows,
              ["year","contract","unmatched_numeric_enrollment","unmatched_pbp_count","unmatched_pbps",
               "partdplans_pbp_count","partdplans_pbps","distinct_published_payment_vectors",
               "all_represented_pbps_share_one_payment_vector","matched_vocabulary_pbps"])
    write_csv(a.output_dir/"partd_affected_contract_payment_pbp_detail.csv",payment_detail,
              ["year","contract","payment_pbp","contract_name","payment_organization_type",
               "direct_subsidy_std_pmpm","partd_risk_score","reinsurance_pmpm","lics_pmpm",
               "payment_vector_shared_by_pbp_count","payment_vector_shared_by_pbps"])
    write_csv(a.output_dir/"partd_payment_source_manifest.csv",manifests,
              ["year","zip_path","zip_sha256","member_name","member_sha256","sheet_name",
               "header_row","payment_key_count"])
    write_csv(a.output_dir/"partd_payment_duplicate_keys.csv",duprows,["year","contract","pbp"])

    checks=[
        ("SAME_CONTRACT_OTHER_PBP control equals 5,203,895",control==EXPECTED_CONTROL),
        ("relationship accounting identity",reltotal==control),
        ("contract-year accounting identity",cytotal==control),
        ("every target contract has >=1 same-year PartDPlans PBP",
         all(r["same_year_partdplans_pbp_count"]>=1 for r in relationship)),
        ("no duplicate annual Part D payment keys",not duprows),
        ("no source/profile errors",not errors),
    ]
    overall=all(x[1] for x in checks)
    lines=[
        "CMS MEDICARE PART D CONTRACT/PBP RELATIONSHIP PROFILE v0.1","="*76,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        "Scope: SAME_CONTRACT_OTHER_PBP population from nonmatch classifier v0.1","",
        "BOUNDARY","-"*76,
        "Read-only descriptive relationship profiling.",
        "No payment propagation to unmatched PBPs.",
        "No reconstruction, imputation, database writes, or ledger inclusion.","",
        "CONTROL","-"*76,
        f"Expected member-month control: {EXPECTED_CONTROL:,}",
        f"Input member-months: {control:,}",
        f"Relationship member-months: {reltotal:,}",
        f"Accounting difference: {control-reltotal:,}","",
        "RELATIONSHIP COUNTS","-"*76,
        f"Distinct unmatched Contract+PBP+year relationships: {len(relationship):,}",
        f"Affected contract-years: {len(cyrows):,}",
        f"Member-months whose exact Contract+PBP appears in another CY2015-CY2024 PartDPlans year: {exact_other:,}",
        f"Member-months with no exact Contract+PBP in any other target year: {no_other:,}",
        f"Matched-vocabulary coexistence evidence available: {'YES' if matched_vocabulary_available else 'NO'}","",
        "CHECKS","-"*76,
    ]
    lines += [f"{k}: {'PASS' if v else 'FAIL'}" for k,v in checks]
    lines += ["","WARNINGS","-"*76] + (warnings or ["None"])
    lines += ["","ERRORS","-"*76] + (errors or ["None"])
    lines += ["","OVERALL","-"*76,f"OVERALL: {'PASS' if overall else 'FAIL'}","",
              "INTERPRETATION BOUNDARY","-"*76,
              "PASS certifies source resolution and relationship accounting only.",
              "It does not authorize applying one PBP's payment values to another PBP.",
              "Identical payment vectors are descriptive evidence, not a propagation rule."]
    report="\n".join(lines)+"\n"
    (a.output_dir/"partd_contract_pbp_relationship_report.txt").write_text(report,encoding="utf-8")
    print(report); print(f"Outputs: {a.output_dir}")
    return 0 if overall else 1

if __name__=="__main__":
    sys.exit(main())
