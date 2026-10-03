from __future__ import annotations

import csv
import hashlib
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]

SOURCE_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "ssa"
    / "earnings_employment"
    / "source"
)

VALIDATION_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "ssa"
    / "earnings_employment"
    / "validation"
)

OUTPUT_DIR = VALIDATION_DIR / "source_certification_v0_1"

CERTIFICATION_CSV = OUTPUT_DIR / "manual_acquisition_certification_v0_1.csv"
SHEET_INVENTORY_CSV = OUTPUT_DIR / "sheet_inventory.csv"
SEMANTIC_PROFILE_TXT = OUTPUT_DIR / "semantic_profile.txt"

START_YEAR = 2015
END_YEAR = 2023

BASE_URL = (
    "https://www.ssa.gov/policy/docs/statcomps/"
    "eedata_sc/{year}/eedata_sc{yy}.xlsx"
)

TARGET_SHEETS = (
    "Table 1",
    "Table 4",
)

MAX_PROFILE_ROWS = 20
MAX_PROFILE_COLS = 16


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest().upper()


def validate_xlsx(path: Path) -> tuple[bool, str]:
    try:
        with path.open("rb") as handle:
            signature = handle.read(4)

        if signature != b"PK\x03\x04":
            return False, f"Bad ZIP/XLSX signature: {signature!r}"

        if not zipfile.is_zipfile(path):
            return False, "Not a valid ZIP/XLSX archive"

        with zipfile.ZipFile(path, "r") as archive:
            names = set(archive.namelist())

            required = {
                "[Content_Types].xml",
                "xl/workbook.xml",
            }

            missing = required - names

            if missing:
                return (
                    False,
                    f"Missing XLSX members: {sorted(missing)}",
                )

        return True, ""

    except Exception as exc:
        return False, str(exc)


def clean(value) -> str:
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\r", " ")
    text = text.replace("\n", " ")

    return " ".join(text.split())


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    certification_rows = []
    sheet_rows = []
    profile_lines = []

    failures = []

    profile_lines.append("=" * 110)
    profile_lines.append(
        "SSA EARNINGS AND EMPLOYMENT SOURCE CERTIFICATION "
        "AND SEMANTIC PROFILE v0.1 - ETL-172E-4"
    )
    profile_lines.append("=" * 110)
    profile_lines.append("")
    profile_lines.append(
        "Manual browser acquisition from official SSA publication."
    )
    profile_lines.append(
        "No receipt allocation, state-share calculation, "
        "Rockefeller replication, or DB writes."
    )
    profile_lines.append("")

    print("=" * 100)
    print(
        "SSA EARNINGS AND EMPLOYMENT SOURCE CERTIFICATION "
        "- ETL-172E-4"
    )
    print("=" * 100)
    print()

    for year in range(START_YEAR, END_YEAR + 1):
        yy = str(year)[-2:]
        filename = f"eedata_sc{yy}.xlsx"
        path = SOURCE_DIR / filename

        source_url = BASE_URL.format(
            year=year,
            yy=yy,
        )

        if not path.exists():
            failures.append(f"{year}: source file missing")

            certification_rows.append(
                {
                    "data_year": year,
                    "source_url": source_url,
                    "local_path": str(
                        path.relative_to(REPO_ROOT)
                    ),
                    "acquisition_method": "MANUAL_BROWSER_DOWNLOAD",
                    "bytes": "",
                    "sha256": "",
                    "xlsx_integrity": "FAIL",
                    "sheet_count": "",
                    "table_1_present": "",
                    "table_4_present": "",
                    "known_source_limitation": (
                        "2015 HI state-level methodological flaw "
                        "documented by SSA"
                        if year == 2015
                        else ""
                    ),
                    "certified_utc": (
                        datetime.now(timezone.utc)
                        .isoformat(timespec="seconds")
                    ),
                    "status": "FAIL",
                }
            )

            continue

        valid, integrity_error = validate_xlsx(path)

        if not valid:
            failures.append(
                f"{year}: XLSX integrity failure: "
                f"{integrity_error}"
            )
            continue

        size_bytes = path.stat().st_size
        sha256 = sha256_file(path)

        wb = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        normalized = {
            name.strip(): name
            for name in wb.sheetnames
        }

        table_1_present = "Table 1" in normalized
        table_4_present = "Table 4" in normalized

        if not table_1_present:
            failures.append(
                f"{year}: Table 1 not found"
            )

        if not table_4_present:
            failures.append(
                f"{year}: Table 4 not found"
            )

        for position, sheet_name in enumerate(
            wb.sheetnames,
            start=1,
        ):
            sheet_rows.append(
                {
                    "data_year": year,
                    "sheet_position": position,
                    "sheet_name": sheet_name,
                }
            )

        profile_lines.append("#" * 110)
        profile_lines.append(f"YEAR {year}")
        profile_lines.append(
            f"SOURCE: {path.relative_to(REPO_ROOT)}"
        )
        profile_lines.append(
            f"BYTES: {size_bytes}"
        )
        profile_lines.append(
            f"SHA256: {sha256}"
        )
        profile_lines.append(
            f"SHEETS: {len(wb.sheetnames)}"
        )
        profile_lines.append("#" * 110)

        for requested_sheet in TARGET_SHEETS:
            profile_lines.append("")
            profile_lines.append("-" * 110)
            profile_lines.append(
                f"SHEET: {requested_sheet}"
            )

            actual_sheet = normalized.get(
                requested_sheet
            )

            if actual_sheet is None:
                profile_lines.append("STATUS: NOT FOUND")
                continue

            profile_lines.append(
                f"ACTUAL SHEET NAME: {actual_sheet}"
            )
            profile_lines.append("-" * 110)

            ws = wb[actual_sheet]

            for row_number, row in enumerate(
                ws.iter_rows(
                    min_row=1,
                    max_row=MAX_PROFILE_ROWS,
                    max_col=MAX_PROFILE_COLS,
                    values_only=True,
                ),
                start=1,
            ):
                values = [
                    clean(value)
                    for value in row
                ]

                if not any(values):
                    continue

                rendered = " | ".join(
                    f"C{column_number}={value}"
                    for column_number, value in enumerate(
                        values,
                        start=1,
                    )
                    if value
                )

                profile_lines.append(
                    f"R{row_number:02d} | {rendered}"
                )

        status = (
            "PASS"
            if table_1_present and table_4_present
            else "FAIL"
        )

        certification_rows.append(
            {
                "data_year": year,
                "source_url": source_url,
                "local_path": str(
                    path.relative_to(REPO_ROOT)
                ),
                "acquisition_method": "MANUAL_BROWSER_DOWNLOAD",
                "bytes": size_bytes,
                "sha256": sha256,
                "xlsx_integrity": "PASS",
                "sheet_count": len(wb.sheetnames),
                "table_1_present": table_1_present,
                "table_4_present": table_4_present,
                "known_source_limitation": (
                    "2015 HI state-level methodological flaw "
                    "documented by SSA"
                    if year == 2015
                    else ""
                ),
                "certified_utc": (
                    datetime.now(timezone.utc)
                    .isoformat(timespec="seconds")
                ),
                "status": status,
            }
        )

        wb.close()

        print(
            f"{year}: PASS | "
            f"{size_bytes:,} bytes | "
            f"{len(sheet_rows)} cumulative sheet records"
        )

    certification_fields = [
        "data_year",
        "source_url",
        "local_path",
        "acquisition_method",
        "bytes",
        "sha256",
        "xlsx_integrity",
        "sheet_count",
        "table_1_present",
        "table_4_present",
        "known_source_limitation",
        "certified_utc",
        "status",
    ]

    with CERTIFICATION_CSV.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=certification_fields,
        )
        writer.writeheader()
        writer.writerows(certification_rows)

    with SHEET_INVENTORY_CSV.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "data_year",
                "sheet_position",
                "sheet_name",
            ],
        )
        writer.writeheader()
        writer.writerows(sheet_rows)

    profile_lines.append("")
    profile_lines.append("=" * 110)
    profile_lines.append("PROFILE COMPLETE")
    profile_lines.append("=" * 110)
    profile_lines.append("")
    profile_lines.append(
        "This profile does not certify receipt allocator semantics."
    )

    SEMANTIC_PROFILE_TXT.write_text(
        "\n".join(profile_lines) + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 100)
    print("CERTIFICATION SUMMARY")
    print("=" * 100)
    print(
        f"Expected years       : "
        f"{END_YEAR - START_YEAR + 1}"
    )
    print(
        f"Certified workbooks  : "
        f"{sum(row['status'] == 'PASS' for row in certification_rows)}"
    )
    print(
        f"Failures             : {len(failures)}"
    )
    print(
        f"Certification CSV    : "
        f"{CERTIFICATION_CSV.relative_to(REPO_ROOT)}"
    )
    print(
        f"Sheet inventory      : "
        f"{SHEET_INVENTORY_CSV.relative_to(REPO_ROOT)}"
    )
    print(
        f"Semantic profile     : "
        f"{SEMANTIC_PROFILE_TXT.relative_to(REPO_ROOT)}"
    )

    if failures:
        print()
        print("FAILURES:")
        for failure in failures:
            print(f"  - {failure}")

        print()
        print("ETL-172E-4 RESULT: FAIL")
        return 1

    print()
    print("ETL-172E-4 RESULT: PASS")
    print()
    print(
        "PASS certifies manual source preservation, XLSX "
        "integrity, hashes, and target-sheet presence only."
    )
    print(
        "PASS does NOT certify allocator semantics, "
        "Washington shares, fiscal-year mapping, "
        "Rockefeller replication, or ledger inclusion."
    )

    return 0


if __name__ == "__main__":
    sys.exit(main())