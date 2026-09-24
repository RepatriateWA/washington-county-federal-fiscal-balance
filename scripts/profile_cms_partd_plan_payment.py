#!/usr/bin/env python
"""Read-only structural profiler for CMS Part D Plan Payment archives, CY2015-CY2024."""
from __future__ import annotations
import argparse,csv,hashlib,io,re,sys,zipfile
from collections import Counter,defaultdict
from datetime import datetime,timezone
from pathlib import Path
import openpyxl

PLANS="PART_D_PLANS"; RECON="PART_D_RECONCILIATION"; RDS="RDS"
ROOT=Path("data/raw/cms/medicare_advantage/plan_payment")
OUT=Path("data/raw/cms/medicare_advantage/validation/partd_payment_profile_v0_1")

def hbytes(b): return hashlib.sha256(b).hexdigest()
def hfile(p):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for x in iter(lambda:f.read(1048576),b""): h.update(x)
    return h.hexdigest()
def clean(x): return re.sub(r"\s+"," ",str(x).strip()) if x is not None else ""
def norm(x): return re.sub(r"[^A-Z0-9]+","_",clean(x).upper()).strip("_")
def yearof(p):
    for x in reversed(p.parts):
        if re.fullmatch(r"20\d{2}",x): return int(x)
    m=re.search(r"(20\d{2})",p.name); return int(m.group(1)) if m else None
def cls(n):
    s=re.sub(r"[^a-z0-9]+","",Path(n).name.lower())
    if "rds" in s:return RDS
    if "partd" in s and "reconcil" in s:return RECON
    if "partdplans" in s:return PLANS
    return None
def writecsv(p,rows,fields=None):
    p.parent.mkdir(parents=True,exist_ok=True)
    fields=fields or (list(rows[0]) if rows else ["status"])
    with p.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore");w.writeheader();w.writerows(rows)
def header(ws):
    keys=("contract","plan","organization","name","type","state","county","risk","payment","pmpm",
          "subsid","reinsurance","reconcil","premium","enrollment","member","beneficiar","amount","cost","direct")
    best=None
    for r in range(1,min(ws.max_row or 0,30)+1):
        vals=[clean(ws.cell(r,c).value) for c in range(1,min(ws.max_column or 0,100)+1)]
        non=[v for v in vals if v]
        score=(sum(any(k in v.lower() for k in keys) for v in non),len(non))
        if best is None or score>best[0]:best=(score,r,vals)
    if not best or best[0][1]<2:return None,[]
    vals=best[2]
    while vals and not vals[-1]:vals.pop()
    return best[1],vals
def kind(v):
    if v is None or clean(v)=="":return "BLANK"
    if isinstance(v,bool):return "BOOLEAN"
    if isinstance(v,(int,float)):return "NUMERIC"
    return "TEXT"

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--root",type=Path,default=ROOT);ap.add_argument("--output-dir",type=Path,default=OUT)
    ap.add_argument("--start-year",type=int,default=2015);ap.add_argument("--end-year",type=int,default=2024)
    a=ap.parse_args();a.output_dir.mkdir(parents=True,exist_ok=True)
    errs=[];audit=[];members=[];sheets=[];cols=[]
    by=defaultdict(list)
    for p in sorted(a.root.rglob("*.zip")):
        y=yearof(p)
        if y and a.start_year<=y<=a.end_year:by[y].append((p.resolve(),hfile(p),p.stat().st_size))
    chosen={}
    for y in range(a.start_year,a.end_year+1):
        items=by.get(y,[]); hashes={x[1] for x in items}
        if not items: errs.append(f"CY{y}: no ZIP");continue
        if len(hashes)>1:
            errs.append(f"CY{y}: multiple non-identical ZIPs; manual selection required")
            for x in items:audit.append({"payment_year":y,"zip_path":str(x[0]),"zip_sha256":x[1],"selection":"CONFLICT"})
            continue
        x=sorted(items,key=lambda z:str(z[0]).lower())[0];chosen[y]=x
        for q in items:audit.append({"payment_year":y,"zip_path":str(q[0]),"zip_sha256":q[1],
                                     "selection":"SELECTED" if q[0]==x[0] else "BYTE_IDENTICAL_DUPLICATE"})
    for y,x in chosen.items():
        zp,zsha,_=x
        try:
            with zipfile.ZipFile(zp) as z:
                rel=[(i,cls(i.filename)) for i in z.infolist() if cls(i.filename)]
                cc=Counter(c for _,c in rel)
                for fam in (PLANS,RECON):
                    if cc[fam]!=1:errs.append(f"CY{y}: expected 1 {fam}, found {cc[fam]}")
                for i,fam in rel:
                    b=z.read(i);msha=hbytes(b)
                    mr={"payment_year":y,"zip_path":str(zp),"zip_sha256":zsha,"zip_member":i.filename,
                        "member_size_bytes":i.file_size,"member_class":fam,"member_sha256":msha,
                        "disposition":"INVENTORY_ONLY" if fam==RDS else "PROFILE","profile_status":""}
                    if fam==RDS:mr["profile_status"]="INVENTORY_ONLY";members.append(mr);continue
                    try:
                        wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
                        for ws in wb.worksheets:
                            hr,heads=header(ws)
                            sheets.append({"payment_year":y,"workbook_family":fam,"zip_member":i.filename,
                                "member_sha256":msha,"sheet_name":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,
                                "detected_header_row":hr or "","detected_header_columns":len(heads),
                                "profile_status":"PASS" if hr else "NO_HEADER"})
                            if not hr:continue
                            end=min(ws.max_row or hr,hr+250)
                            for ci,h in enumerate(heads,1):
                                if not h:continue
                                c=Counter();ex=[]
                                for ri in range(hr+1,end+1):
                                    v=ws.cell(ri,ci).value;c[kind(v)]+=1;t=clean(v)
                                    if t and t not in ex and len(ex)<5:ex.append(t[:160])
                                cols.append({"payment_year":y,"workbook_family":fam,"zip_member":i.filename,
                                    "sheet_name":ws.title,"header_row":hr,"column_position":ci,"column_name":h,
                                    "column_name_normalized":norm(h),"sample_numeric_count":c["NUMERIC"],
                                    "sample_text_count":c["TEXT"],"sample_blank_count":c["BLANK"],
                                    "sample_examples":" | ".join(ex)})
                        wb.close();mr["profile_status"]="PASS"
                    except Exception as e:
                        mr["profile_status"]="ERROR";errs.append(f"CY{y} {i.filename}: {type(e).__name__}: {e}")
                    members.append(mr)
        except Exception as e:errs.append(f"CY{y} ZIP error: {type(e).__name__}: {e}")
    groups=defaultdict(set);names=defaultdict(Counter)
    for r in cols:
        k=(r["workbook_family"],r["sheet_name"],r["column_name_normalized"]);groups[k].add(r["payment_year"]);names[k][r["column_name"]]+=1
    cross=[]
    for k,ys in sorted(groups.items()):
        cross.append({"workbook_family":k[0],"sheet_name":k[1],"column_name_normalized":k[2],
                      "representative_column_name":names[k].most_common(1)[0][0],
                      "years_present":";".join(map(str,sorted(ys))),"year_count":len(ys)})
    writecsv(a.output_dir/"cms_partd_zip_selection_audit.csv",audit)
    writecsv(a.output_dir/"cms_partd_member_inventory.csv",members)
    writecsv(a.output_dir/"cms_partd_sheet_profile.csv",sheets)
    writecsv(a.output_dir/"cms_partd_column_profile.csv",cols)
    writecsv(a.output_dir/"cms_partd_cross_year_schema.csv",cross)
    writecsv(a.output_dir/"cms_partd_profile_errors.csv",[{"error":e} for e in errs],["error"])
    req=set(range(a.start_year,a.end_year+1))
    py={r["payment_year"] for r in members if r["member_class"]==PLANS and r["profile_status"]=="PASS"}
    ry={r["payment_year"] for r in members if r["member_class"]==RECON and r["profile_status"]=="PASS"}
    ok=(not errs and set(chosen)==req and py==req and ry==req)
    lines=["CMS MEDICARE PART D PLAN PAYMENT STRUCTURAL PROFILE","="*68,
           f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",f"Scope: CY{a.start_year}-CY{a.end_year}",
           "","BOUNDARY","-"*68,"Read-only structural profiling only.",
           "Part D Plans and Part D Reconciliation are profiled separately.",
           "RDS is inventoried only and remains analytically separate.",
           "No reconstruction, imputation, database writes, or ledger inclusion.","","COUNTS","-"*68,
           f"Required years: {len(req)}",f"Selected ZIP years: {len(chosen)}",
           f"Part D Plans years profiled: {len(py)}",f"Part D Reconciliation years profiled: {len(ry)}",
           f"Sheet rows: {len(sheets)}",f"Column rows: {len(cols)}",f"Errors: {len(errs)}","","CHECKS","-"*68,
           f"Unique canonical ZIP every year: {'PASS' if set(chosen)==req else 'FAIL'}",
           f"Part D Plans every year: {'PASS' if py==req else 'FAIL'}",
           f"Part D Reconciliation every year: {'PASS' if ry==req else 'FAIL'}",
           f"No structural/profile errors: {'PASS' if not errs else 'FAIL'}","",f"OVERALL: {'PASS' if ok else 'FAIL'}",
           "","A PASS certifies source presence and structural profiling only; it does not",
           "establish payment semantics, additivity, an accounting equation, or ledger eligibility."]
    (a.output_dir/"cms_partd_profile_summary.txt").write_text("\n".join(lines)+"\n",encoding="utf-8")
    print("\n".join(lines));return 0 if ok else 1
if __name__=="__main__":raise SystemExit(main())
