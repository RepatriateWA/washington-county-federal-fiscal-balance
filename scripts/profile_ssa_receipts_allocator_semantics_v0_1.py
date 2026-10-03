from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
SSA_ROOT = REPO_ROOT / "data" / "raw" / "ssa" / "oasdi"

TARGET_SHEETS = (
    "Table 1",
    "Table 2",
    "Table 3",
    "Table 4 - Washington",
)

MAX_ROWS = 30
MAX_COLS = 14


def clean(value):
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\r", " ")
    text = text.replace("\n", " ")
    return " ".join(text.split())


def main():
    print("=" * 110)
    print("SSA RECEIPTS ALLOCATOR SEMANTIC PROFILE v0.1 - ETL-172E-2")
    print("=" * 110)
    print()
    print("Purpose:")
    print("  Inspect labels/layout only.")
    print("  No receipt allocation.")
    print("  No state-share calculation.")
    print("  No denominator decision.")
    print("  No DB writes.")
    print()

    for year in range(2015, 2024):
        yy = str(year)[-2:]
        path = SSA_ROOT / str(year) / f"oasdi_sc{yy}.xlsx"

        if not path.exists():
            raise FileNotFoundError(path)

        print()
        print("#" * 110)
        print(f"YEAR {year}")
        print(f"SOURCE: {path.relative_to(REPO_ROOT)}")
        print("#" * 110)

        wb = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        normalized_sheet_lookup = {
            name.strip(): name
            for name in wb.sheetnames
        }

        for requested_sheet in TARGET_SHEETS:
            actual_sheet = normalized_sheet_lookup.get(
                requested_sheet
            )

            print()
            print("-" * 110)
            print(f"SHEET: {requested_sheet}")

            if actual_sheet is None:
                print("STATUS: NOT FOUND")
                continue

            print(f"ACTUAL SHEET NAME: {actual_sheet}")
            print("-" * 110)

            ws = wb[actual_sheet]

            for row_number, row in enumerate(
                ws.iter_rows(
                    min_row=1,
                    max_row=MAX_ROWS,
                    max_col=MAX_COLS,
                    values_only=True,
                ),
                start=1,
            ):
                values = [clean(value) for value in row]

                if not any(values):
                    continue

                rendered = " | ".join(
                    f"C{column_number}={value}"
                    for column_number, value in enumerate(
                        values,
                        start=1,
                    )
                    if value
                )

                print(
                    f"R{row_number:02d} | {rendered}"
                )

        wb.close()

    print()
    print("=" * 110)
    print("PROFILE COMPLETE")
    print("=" * 110)
    print()
    print(
        "No semantic certification is implied by this profile."
    )


if __name__ == "__main__":
    main()