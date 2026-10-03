from __future__ import annotations

import csv
import hashlib
import json
import sys
from decimal import Decimal, InvalidOperation, getcontext
from pathlib import Path
from typing import Any

import openpyxl


# =============================================================================
# ETL-172I-2
# OMB SOCIAL INSURANCE AND EXCISE RECEIPTS DECOMPOSITION
#
# Purpose:
#   Decompose and certify the national FFY2015-FFY2024:
#
#       1. Social Insurance and Retirement Receipts
#       2. Excise Taxes
#
#   from OMB Historical Table 2.4, and reconcile both categories to the
#   corresponding certified OMB Historical Table 2.1 controls from ETL-172I-1.
#
# This is a NATIONAL CONTROL / DECOMPOSITION step only.
#
# It does NOT:
#   - allocate any newly decomposed component to Washington;
#   - determine state incidence;
#   - reproduce Rockefeller estimates;
#   - reconcile to IRS gross collections;
#   - write to PostgreSQL;
#   - authorize final ledger inclusion.
#
# Important source convention:
#   Table 2.4 uses textual placeholders such as ".........." for zero / no
#   reported amount. Those placeholders are normalized to Decimal(0) for
#   arithmetic certification while the source-cell text is preserved in
#   provenance fields where relevant.
# =============================================================================


getcontext().prec = 50


REPO_ROOT = Path(__file__).resolve().parents[1]

SOURCE_FILE = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "source"
    / "BUDGET-2027-TAB-3-4.xlsx"
)

I1_FILE = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "federal_receipts_universe_v0_1"
    / "national_federal_receipts_universe.csv"
)

I1_CERTIFICATION = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "federal_receipts_universe_v0_1"
    / "certification.json"
)

OUTPUT_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "social_insurance_excise_decomposition_v0_1"
)

OUTPUT_DETAIL = (
    OUTPUT_DIR
    / "national_social_insurance_excise_decomposition.csv"
)

OUTPUT_ANNUAL = (
    OUTPUT_DIR
    / "national_social_insurance_excise_annual_controls.csv"
)

OUTPUT_JSON = (
    OUTPUT_DIR
    / "certification.json"
)

OUTPUT_REPORT = (
    OUTPUT_DIR
    / "certification_report.txt"
)


SHEET_NAME = "Table"
TARGET_YEARS = list(range(2015, 2025))

EXPECTED_TITLE = (
    "Table 2.4 - COMPOSITION OF SOCIAL INSURANCE AND "
    "RETIREMENT RECEIPTS AND OF EXCISE TAXES:  1940 - 2031"
)

EXPECTED_UNITS = "(in millions of dollars)"

MILLIONS_TO_DOLLARS = Decimal("1000000")


# -----------------------------------------------------------------------------
# Table 2.4 source-row taxonomy.
#
# Parent/header rows are deliberately excluded from arithmetic except where
# they contain an explicit published total.
# -----------------------------------------------------------------------------

SOCIAL_ROWS = {
    7: (
        "OASI_FEDERAL_FUNDS",
        "Old-age and survivors insurance - Federal funds",
        "EMPLOYMENT_GENERAL_RETIREMENT",
    ),
    8: (
        "OASI_TRUST_FUNDS_OFF_BUDGET",
        "Old-age and survivors insurance - Trust funds (Off-Budget)",
        "EMPLOYMENT_GENERAL_RETIREMENT",
    ),
    9: (
        "DISABILITY_INSURANCE_OFF_BUDGET",
        "Disability insurance (Off-Budget)",
        "EMPLOYMENT_GENERAL_RETIREMENT",
    ),
    10: (
        "HOSPITAL_INSURANCE",
        "Hospital insurance",
        "EMPLOYMENT_GENERAL_RETIREMENT",
    ),
    12: (
        "RAILROAD_RETIREMENT_FEDERAL_FUNDS",
        "Railroad retirement/pension fund - Federal funds",
        "EMPLOYMENT_GENERAL_RETIREMENT",
    ),
    13: (
        "RAILROAD_RETIREMENT_TRUST_FUNDS",
        "Railroad retirement/pension fund - Trust funds",
        "EMPLOYMENT_GENERAL_RETIREMENT",
    ),
    14: (
        "RAILROAD_SOCIAL_SECURITY_EQUIVALENT",
        "Railroad social security equivalent account",
        "EMPLOYMENT_GENERAL_RETIREMENT",
    ),
    17: (
        "UNEMPLOYMENT_INSURANCE_FEDERAL_FUNDS",
        "Unemployment insurance - Federal funds",
        "UNEMPLOYMENT_INSURANCE",
    ),
    18: (
        "UNEMPLOYMENT_INSURANCE_TRUST_FUNDS",
        "Unemployment insurance - Trust funds",
        "UNEMPLOYMENT_INSURANCE",
    ),
    21: (
        "FEDERAL_EMPLOYEE_RETIREMENT_EMPLOYEE_SHARE",
        "Federal employees retirement - employee share",
        "OTHER_RETIREMENT",
    ),
    22: (
        "NONFEDERAL_EMPLOYEE_RETIREMENT",
        "Non-Federal employees retirement",
        "OTHER_RETIREMENT",
    ),
}

SOCIAL_PUBLISHED_TOTAL_ROWS = {
    "employment_general_retirement_total": 15,
    "unemployment_insurance_total": 19,
    "other_retirement_total": 23,
    "social_insurance_retirement_total": 24,
}


EXCISE_ROWS = {
    27: ("ALCOHOL", "Alcohol", "FEDERAL_FUNDS"),
    28: ("TOBACCO", "Tobacco", "FEDERAL_FUNDS"),
    29: (
        "CRUDE_OIL_WINDFALL_PROFIT",
        "Crude oil windfall profit",
        "FEDERAL_FUNDS",
    ),
    30: ("TELEPHONE", "Telephone", "FEDERAL_FUNDS"),
    31: (
        "OZONE_DEPLETING_CHEMICALS_PRODUCTS",
        "Ozone depleting chemicals/products",
        "FEDERAL_FUNDS",
    ),
    32: (
        "TRANSPORTATION_FUELS",
        "Transportation fuels",
        "FEDERAL_FUNDS",
    ),
    33: (
        "CORPORATE_STOCK_REPURCHASE",
        "Corporate stock repurchase",
        "FEDERAL_FUNDS",
    ),
    34: (
        "HEALTH_INSURANCE_PROVIDERS",
        "Health insurance providers",
        "FEDERAL_FUNDS",
    ),
    35: (
        "INDOOR_TANNING_SERVICES",
        "Indoor tanning services",
        "FEDERAL_FUNDS",
    ),
    36: (
        "MEDICAL_DEVICES",
        "Medical devices",
        "FEDERAL_FUNDS",
    ),
    37: (
        "OTHER_FEDERAL_FUND_EXCISE",
        "Other",
        "FEDERAL_FUNDS",
    ),
    40: (
        "TRANSPORTATION_TRUST_FUND",
        "Transportation",
        "TRUST_FUNDS",
    ),
    41: (
        "AIRPORT_AIRWAY_TRUST_FUND",
        "Airport and airway",
        "TRUST_FUNDS",
    ),
    42: (
        "BLACK_LUNG_DISABILITY_TRUST_FUND",
        "Black lung disability",
        "TRUST_FUNDS",
    ),
    43: (
        "INLAND_WATERWAY_TRUST_FUND",
        "Inland waterway",
        "TRUST_FUNDS",
    ),
    44: (
        "HAZARDOUS_SUBSTANCE_SUPERFUND",
        "Hazardous substance superfund",
        "TRUST_FUNDS",
    ),
    45: (
        "POST_CLOSURE_LIABILITY_HAZARDOUS_WASTE",
        "Post-closure liability (hazardous waste)",
        "TRUST_FUNDS",
    ),
    46: (
        "OIL_SPILL_LIABILITY",
        "Oil spill liability",
        "TRUST_FUNDS",
    ),
    47: (
        "AQUATIC_RESOURCES",
        "Aquatic resources",
        "TRUST_FUNDS",
    ),
    48: (
        "LEAKING_UNDERGROUND_STORAGE_TANK",
        "Leaking underground storage tank",
        "TRUST_FUNDS",
    ),
    49: (
        "TOBACCO_ASSESSMENTS",
        "Tobacco Assessments",
        "TRUST_FUNDS",
    ),
    50: (
        "VACCINE_INJURY_COMPENSATION",
        "Vaccine injury compensation",
        "TRUST_FUNDS",
    ),
    51: (
        "SUPPLEMENTARY_MEDICAL_INSURANCE",
        "Supplementary medical insurance",
        "TRUST_FUNDS",
    ),
    52: (
        "PATIENT_CENTERED_OUTCOMES_RESEARCH",
        "Patient-centered outcomes research",
        "TRUST_FUNDS",
    ),
}

EXCISE_PUBLISHED_TOTAL_ROWS = {
    "federal_funds_total": 38,
    "trust_funds_total": 53,
    "excise_total": 54,
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(
            lambda: handle.read(1024 * 1024),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest().upper()


def normalize_label(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value).strip().split()
    )


def source_number(
    value: Any,
    *,
    context: str,
) -> Decimal:
    """
    Normalize OMB Table 2.4 source cells to Decimal.

    Dotted placeholders and blank cells are treated as zero.
    The source contains nonnumeric placeholder strings for amounts
    not reported / effectively zero at publication precision.
    """

    if value is None:
        return Decimal("0")

    if isinstance(
        value,
        (int, float, Decimal),
    ):
        return Decimal(str(value))

    text = str(value).strip()

    if not text:
        return Decimal("0")

    # OMB historical tables use strings such as ".........."
    # for zero/no amount.
    if set(text) <= {".", " "}:
        return Decimal("0")

    # Historical cells may use -* for very small negative values.
    # This is outside the target modern period but is handled
    # defensively as zero at millions-of-dollars publication precision.
    if text in {"*", "-*"}:
        return Decimal("0")

    text = text.replace(",", "")

    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise RuntimeError(
            "NUMERIC SOURCE FAILURE\n"
            f"{context}: {value!r}"
        ) from exc

    if not result.is_finite():
        raise RuntimeError(
            "NUMERIC SOURCE FAILURE\n"
            f"{context}: non-finite"
        )

    return result


def plain_decimal(value: Decimal) -> str:
    return format(value, "f")


def dollars(value_millions: Decimal) -> Decimal:
    return (
        value_millions
        * MILLIONS_TO_DOLLARS
    )


def load_i1() -> dict[int, dict[str, str]]:
    if not I1_FILE.exists():
        raise RuntimeError(
            "ETL-172I-1 INPUT NOT FOUND\n"
            f"{I1_FILE.relative_to(REPO_ROOT)}"
        )

    if not I1_CERTIFICATION.exists():
        raise RuntimeError(
            "ETL-172I-1 CERTIFICATION NOT FOUND\n"
            f"{I1_CERTIFICATION.relative_to(REPO_ROOT)}"
        )

    certification = json.loads(
        I1_CERTIFICATION.read_text(
            encoding="utf-8"
        )
    )

    if certification.get("etl_step") != "ETL-172I-1":
        raise RuntimeError(
            "ETL-172I-1 CERTIFICATION CONTRACT FAILURE\n"
            "Unexpected etl_step."
        )

    if certification.get("result") != "PASS":
        raise RuntimeError(
            "ETL-172I-1 IS NOT CERTIFIED PASS."
        )

    rows: dict[int, dict[str, str]] = {}

    with I1_FILE.open(
        "r",
        newline="",
        encoding="utf-8",
    ) as handle:
        reader = csv.DictReader(handle)

        for row in reader:
            year = int(row["fiscal_year"])

            if year in rows:
                raise RuntimeError(
                    "ETL-172I-1 DUPLICATE YEAR\n"
                    f"FFY{year}"
                )

            rows[year] = row

    if set(rows) != set(TARGET_YEARS):
        raise RuntimeError(
            "ETL-172I-1 YEAR COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {sorted(rows)}"
        )

    return rows


def load_source() -> tuple[Any, Any]:
    if not SOURCE_FILE.exists():
        raise RuntimeError(
            "OMB TABLE 2.4 SOURCE NOT FOUND\n"
            f"{SOURCE_FILE.relative_to(REPO_ROOT)}"
        )

    # Deliberately not read_only=True.
    # The GovInfo workbook does not expose reliable calculated
    # worksheet dimensions under openpyxl read-only mode.
    workbook = openpyxl.load_workbook(
        SOURCE_FILE,
        data_only=True,
        read_only=False,
    )

    if SHEET_NAME not in workbook.sheetnames:
        raise RuntimeError(
            "OMB TABLE 2.4 SHEET FAILURE\n"
            f"Expected: {SHEET_NAME}\n"
            f"Observed: {workbook.sheetnames}"
        )

    return workbook, workbook[SHEET_NAME]


def certify_structure(sheet: Any) -> dict[str, Any]:
    if sheet["A1"].value != EXPECTED_TITLE:
        raise RuntimeError(
            "TABLE 2.4 TITLE FAILURE\n"
            f"Expected: {EXPECTED_TITLE!r}\n"
            f"Observed: {sheet['A1'].value!r}"
        )

    if sheet["A2"].value != EXPECTED_UNITS:
        raise RuntimeError(
            "TABLE 2.4 UNIT FAILURE\n"
            f"Expected: {EXPECTED_UNITS!r}\n"
            f"Observed: {sheet['A2'].value!r}"
        )

    expected_labels = {
        4: "Social Insurance and Retirement Receipts",
        5: "Employment and general retirement:",
        6: "Old-age and survivors insurance:",
        7: "Federal funds",
        8: "Trust funds (Off-Budget)",
        9: "Disability insurance (Off-Budget)",
        10: "Hospital insurance",
        11: "Railroad retirement/pension fund:",
        12: "Federal funds",
        13: "Trust funds",
        14: "Railroad social security equivalent account",
        15: "Total (1)",
        16: "Unemployment insurance:",
        17: "Federal funds",
        18: "Trust funds",
        19: "Total",
        20: "Other retirement:",
        21: "Federal employees retirement - employee share",
        22: "Non-Federal employees retirement",
        23: "Total",
        24: "Total, Social Insurance and Retirement Receipts (1)",
        25: "Excise Taxes",
        26: "Federal funds:",
        27: "Alcohol",
        28: "Tobacco",
        29: "Crude oil windfall profit",
        30: "Telephone",
        31: "Ozone depleting chemicals/products",
        32: "Transportation fuels",
        33: "Corporate stock repurchase",
        34: "Health insurance providers",
        35: "Indoor tanning services",
        36: "Medical devices",
        37: "Other",
        38: "Total",
        39: "Trust funds:",
        40: "Transportation",
        41: "Airport and airway",
        42: "Black lung disability",
        43: "Inland waterway",
        44: "Hazardous substance superfund",
        45: "Post-closure liability (hazardous waste)",
        46: "Oil spill liability",
        47: "Aquatic resources",
        48: "Leaking underground storage tank",
        49: "Tobacco Assessments",
        50: "Vaccine injury compensation",
        51: "Supplementary medical insurance",
        52: "Patient-centered outcomes research",
        53: "Total",
        54: "Total, Excise Taxes",
    }

    for row_number, expected in expected_labels.items():
        observed = normalize_label(
            sheet.cell(
                row_number,
                1,
            ).value
        )

        if observed != expected:
            raise RuntimeError(
                "TABLE 2.4 ROW-LABEL FAILURE\n"
                f"Row      : {row_number}\n"
                f"Expected : {expected!r}\n"
                f"Observed : {observed!r}"
            )

    return {
        "title": "PASS",
        "units": "PASS",
        "row_taxonomy": "PASS",
        "worksheet_dimensions": (
            f"{sheet.max_row}x{sheet.max_column}"
        ),
    }


def map_year_columns(
    sheet: Any,
) -> dict[int, int]:
    mapping: dict[int, int] = {}

    for column in range(
        2,
        sheet.max_column + 1,
    ):
        value = sheet.cell(
            3,
            column,
        ).value

        if value is None:
            continue

        try:
            year = int(value)
        except (
            TypeError,
            ValueError,
        ):
            continue

        if year in TARGET_YEARS:
            if year in mapping:
                raise RuntimeError(
                    "DUPLICATE YEAR COLUMN\n"
                    f"FFY{year}"
                )

            mapping[year] = column

    if set(mapping) != set(TARGET_YEARS):
        raise RuntimeError(
            "TABLE 2.4 YEAR-COLUMN COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {sorted(mapping)}"
        )

    return mapping


def extract_component_rows(
    *,
    sheet: Any,
    year_columns: dict[int, int],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []

    for year in TARGET_YEARS:
        column = year_columns[year]

        for source_row, (
            code,
            label,
            group,
        ) in SOCIAL_ROWS.items():
            raw_value = sheet.cell(
                source_row,
                column,
            ).value

            value = source_number(
                raw_value,
                context=(
                    f"FFY{year} row {source_row} "
                    f"{label}"
                ),
            )

            output.append(
                {
                    "fiscal_year": year,
                    "source_column": column,
                    "source_row": source_row,
                    "major_category": (
                        "SOCIAL_INSURANCE_AND_RETIREMENT"
                    ),
                    "component_group": group,
                    "component_code": code,
                    "component_label": label,
                    "source_raw_value": (
                        ""
                        if raw_value is None
                        else str(raw_value)
                    ),
                    "amount_millions": (
                        plain_decimal(value)
                    ),
                    "amount_dollars": (
                        plain_decimal(
                            dollars(value)
                        )
                    ),
                }
            )

        for source_row, (
            code,
            label,
            group,
        ) in EXCISE_ROWS.items():
            raw_value = sheet.cell(
                source_row,
                column,
            ).value

            value = source_number(
                raw_value,
                context=(
                    f"FFY{year} row {source_row} "
                    f"{label}"
                ),
            )

            output.append(
                {
                    "fiscal_year": year,
                    "source_column": column,
                    "source_row": source_row,
                    "major_category": (
                        "EXCISE_TAXES"
                    ),
                    "component_group": group,
                    "component_code": code,
                    "component_label": label,
                    "source_raw_value": (
                        ""
                        if raw_value is None
                        else str(raw_value)
                    ),
                    "amount_millions": (
                        plain_decimal(value)
                    ),
                    "amount_dollars": (
                        plain_decimal(
                            dollars(value)
                        )
                    ),
                }
            )

    return output


def component_sum(
    detail_rows: list[dict[str, Any]],
    *,
    year: int,
    major_category: str,
    component_group: str | None = None,
) -> Decimal:
    total = Decimal("0")

    for row in detail_rows:
        if int(row["fiscal_year"]) != year:
            continue

        if row["major_category"] != major_category:
            continue

        if (
            component_group is not None
            and row["component_group"]
            != component_group
        ):
            continue

        total += Decimal(
            row["amount_millions"]
        )

    return total


def published_value(
    sheet: Any,
    *,
    source_row: int,
    source_column: int,
    context: str,
) -> Decimal:
    return source_number(
        sheet.cell(
            source_row,
            source_column,
        ).value,
        context=context,
    )


def build_annual(
    *,
    sheet: Any,
    year_columns: dict[int, int],
    detail_rows: list[dict[str, Any]],
    i1_rows: dict[int, dict[str, str]],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []

    for year in TARGET_YEARS:
        column = year_columns[year]

        employment_component_sum = component_sum(
            detail_rows,
            year=year,
            major_category=(
                "SOCIAL_INSURANCE_AND_RETIREMENT"
            ),
            component_group=(
                "EMPLOYMENT_GENERAL_RETIREMENT"
            ),
        )

        employment_published_total = published_value(
            sheet,
            source_row=(
                SOCIAL_PUBLISHED_TOTAL_ROWS[
                    "employment_general_retirement_total"
                ]
            ),
            source_column=column,
            context=(
                f"FFY{year} employment/general "
                "retirement published total"
            ),
        )

        unemployment_component_sum = component_sum(
            detail_rows,
            year=year,
            major_category=(
                "SOCIAL_INSURANCE_AND_RETIREMENT"
            ),
            component_group=(
                "UNEMPLOYMENT_INSURANCE"
            ),
        )

        unemployment_published_total = published_value(
            sheet,
            source_row=(
                SOCIAL_PUBLISHED_TOTAL_ROWS[
                    "unemployment_insurance_total"
                ]
            ),
            source_column=column,
            context=(
                f"FFY{year} unemployment insurance "
                "published total"
            ),
        )

        other_retirement_component_sum = component_sum(
            detail_rows,
            year=year,
            major_category=(
                "SOCIAL_INSURANCE_AND_RETIREMENT"
            ),
            component_group=(
                "OTHER_RETIREMENT"
            ),
        )

        other_retirement_published_total = published_value(
            sheet,
            source_row=(
                SOCIAL_PUBLISHED_TOTAL_ROWS[
                    "other_retirement_total"
                ]
            ),
            source_column=column,
            context=(
                f"FFY{year} other retirement "
                "published total"
            ),
        )

        social_component_sum = (
            employment_component_sum
            + unemployment_component_sum
            + other_retirement_component_sum
        )

        social_published_total = published_value(
            sheet,
            source_row=(
                SOCIAL_PUBLISHED_TOTAL_ROWS[
                    "social_insurance_retirement_total"
                ]
            ),
            source_column=column,
            context=(
                f"FFY{year} total social insurance "
                "and retirement receipts"
            ),
        )

        i1_social_total = Decimal(
            i1_rows[year][
                "social_insurance_retirement_total_millions"
            ]
        )

        excise_federal_component_sum = component_sum(
            detail_rows,
            year=year,
            major_category="EXCISE_TAXES",
            component_group="FEDERAL_FUNDS",
        )

        excise_federal_published_total = published_value(
            sheet,
            source_row=(
                EXCISE_PUBLISHED_TOTAL_ROWS[
                    "federal_funds_total"
                ]
            ),
            source_column=column,
            context=(
                f"FFY{year} federal-fund excise total"
            ),
        )

        excise_trust_component_sum = component_sum(
            detail_rows,
            year=year,
            major_category="EXCISE_TAXES",
            component_group="TRUST_FUNDS",
        )

        excise_trust_published_total = published_value(
            sheet,
            source_row=(
                EXCISE_PUBLISHED_TOTAL_ROWS[
                    "trust_funds_total"
                ]
            ),
            source_column=column,
            context=(
                f"FFY{year} trust-fund excise total"
            ),
        )

        excise_component_sum = (
            excise_federal_component_sum
            + excise_trust_component_sum
        )

        excise_published_total = published_value(
            sheet,
            source_row=(
                EXCISE_PUBLISHED_TOTAL_ROWS[
                    "excise_total"
                ]
            ),
            source_column=column,
            context=(
                f"FFY{year} total excise taxes"
            ),
        )

        i1_excise_total = Decimal(
            i1_rows[year][
                "excise_taxes_millions"
            ]
        )

        checks = {
            "employment_general_retirement_identity": (
                employment_published_total
                - employment_component_sum
            ),
            "unemployment_insurance_identity": (
                unemployment_published_total
                - unemployment_component_sum
            ),
            "other_retirement_identity": (
                other_retirement_published_total
                - other_retirement_component_sum
            ),
            "social_total_identity": (
                social_published_total
                - social_component_sum
            ),
            "social_table_2_1_reconciliation": (
                i1_social_total
                - social_published_total
            ),
            "excise_federal_funds_identity": (
                excise_federal_published_total
                - excise_federal_component_sum
            ),
            "excise_trust_funds_identity": (
                excise_trust_published_total
                - excise_trust_component_sum
            ),
            "excise_total_identity": (
                excise_published_total
                - excise_component_sum
            ),
            "excise_table_2_1_reconciliation": (
                i1_excise_total
                - excise_published_total
            ),
        }

        failures = {
            key: value
            for key, value in checks.items()
            if value != 0
        }

        if failures:
            details = "\n".join(
                f"{key}: {value}"
                for key, value in failures.items()
            )

            raise RuntimeError(
                "ETL-172I-2 ACCOUNTING IDENTITY FAILURE\n"
                f"FFY{year}\n"
                f"{details}"
            )

        output.append(
            {
                "fiscal_year": year,
                "source_column": column,

                "employment_general_retirement_component_sum_millions": (
                    plain_decimal(
                        employment_component_sum
                    )
                ),
                "employment_general_retirement_published_total_millions": (
                    plain_decimal(
                        employment_published_total
                    )
                ),
                "employment_general_retirement_status": "PASS",

                "unemployment_insurance_component_sum_millions": (
                    plain_decimal(
                        unemployment_component_sum
                    )
                ),
                "unemployment_insurance_published_total_millions": (
                    plain_decimal(
                        unemployment_published_total
                    )
                ),
                "unemployment_insurance_status": "PASS",

                "other_retirement_component_sum_millions": (
                    plain_decimal(
                        other_retirement_component_sum
                    )
                ),
                "other_retirement_published_total_millions": (
                    plain_decimal(
                        other_retirement_published_total
                    )
                ),
                "other_retirement_status": "PASS",

                "social_component_sum_millions": (
                    plain_decimal(
                        social_component_sum
                    )
                ),
                "social_published_total_millions": (
                    plain_decimal(
                        social_published_total
                    )
                ),
                "social_table_2_1_control_millions": (
                    plain_decimal(
                        i1_social_total
                    )
                ),
                "social_decomposition_status": "PASS",
                "social_table_2_1_reconciliation_status": "PASS",

                "excise_federal_funds_component_sum_millions": (
                    plain_decimal(
                        excise_federal_component_sum
                    )
                ),
                "excise_federal_funds_published_total_millions": (
                    plain_decimal(
                        excise_federal_published_total
                    )
                ),
                "excise_federal_funds_status": "PASS",

                "excise_trust_funds_component_sum_millions": (
                    plain_decimal(
                        excise_trust_component_sum
                    )
                ),
                "excise_trust_funds_published_total_millions": (
                    plain_decimal(
                        excise_trust_published_total
                    )
                ),
                "excise_trust_funds_status": "PASS",

                "excise_component_sum_millions": (
                    plain_decimal(
                        excise_component_sum
                    )
                ),
                "excise_published_total_millions": (
                    plain_decimal(
                        excise_published_total
                    )
                ),
                "excise_table_2_1_control_millions": (
                    plain_decimal(
                        i1_excise_total
                    )
                ),
                "excise_decomposition_status": "PASS",
                "excise_table_2_1_reconciliation_status": "PASS",

                "washington_allocation_status": (
                    "NOT_PERFORMED_IN_ETL_172I_2"
                ),
                "ledger_status": "INVENTORY_ONLY",
            }
        )

    return output


def certify_output(
    *,
    detail_rows: list[dict[str, Any]],
    annual_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    expected_detail = (
        len(TARGET_YEARS)
        * (
            len(SOCIAL_ROWS)
            + len(EXCISE_ROWS)
        )
    )

    if len(detail_rows) != expected_detail:
        raise RuntimeError(
            "DETAIL OBSERVATION COUNT FAILURE\n"
            f"Expected: {expected_detail}\n"
            f"Actual  : {len(detail_rows)}"
        )

    if len(annual_rows) != len(TARGET_YEARS):
        raise RuntimeError(
            "ANNUAL OBSERVATION COUNT FAILURE"
        )

    annual_years = [
        int(row["fiscal_year"])
        for row in annual_rows
    ]

    if annual_years != TARGET_YEARS:
        raise RuntimeError(
            "ANNUAL YEAR COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {annual_years}"
        )

    for row in annual_rows:
        for key, value in row.items():
            if key.endswith("_status"):
                if key in {
                    "washington_allocation_status",
                    "ledger_status",
                }:
                    continue

                if value != "PASS":
                    raise RuntimeError(
                        "CERTIFICATION STATUS FAILURE\n"
                        f"FFY{row['fiscal_year']} "
                        f"{key}: {value}"
                    )

        if (
            row["washington_allocation_status"]
            != "NOT_PERFORMED_IN_ETL_172I_2"
        ):
            raise RuntimeError(
                "WASHINGTON ALLOCATION BOUNDARY FAILURE"
            )

        if row["ledger_status"] != "INVENTORY_ONLY":
            raise RuntimeError(
                "LEDGER BOUNDARY FAILURE"
            )

    return {
        "expected_detail_observations": expected_detail,
        "actual_detail_observations": len(detail_rows),
        "expected_annual_observations": 10,
        "actual_annual_observations": len(annual_rows),
        "fiscal_year_coverage": "PASS",
        "fiscal_year_uniqueness": "PASS",
        "source_year_column_mapping": "PASS",
        "source_units": "PASS",
        "social_employment_general_retirement_identity": "PASS",
        "social_unemployment_insurance_identity": "PASS",
        "social_other_retirement_identity": "PASS",
        "social_total_identity": "PASS",
        "social_table_2_1_reconciliation": "PASS",
        "excise_federal_funds_identity": "PASS",
        "excise_trust_funds_identity": "PASS",
        "excise_total_identity": "PASS",
        "excise_table_2_1_reconciliation": "PASS",
        "washington_allocation_boundary": "PASS",
        "rockefeller_replication_boundary": "PASS",
        "irs_reconciliation_boundary": "PASS",
        "database_write_boundary": "PASS",
        "ledger_boundary": "PASS",
    }


def write_csv(
    path: Path,
    rows: list[dict[str, Any]],
) -> None:
    if not rows:
        raise RuntimeError(
            "Refusing to write empty CSV."
        )

    with path.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(
                rows[0].keys()
            ),
        )

        writer.writeheader()
        writer.writerows(rows)


def build_report(
    *,
    source_sha256: str,
    i1_sha256: str,
    structure: dict[str, Any],
    annual_rows: list[dict[str, Any]],
    certification: dict[str, Any],
) -> str:
    lines: list[str] = [
        "=" * 100,
        (
            "OMB SOCIAL INSURANCE AND EXCISE RECEIPTS "
            "DECOMPOSITION - ETL-172I-2"
        ),
        "=" * 100,
        "",
        "UPSTREAM PROVENANCE",
        "-" * 100,
        (
            "Table 2.4 SHA256 : "
            f"{source_sha256}"
        ),
        (
            "ETL-172I-1 SHA256: "
            f"{i1_sha256}"
        ),
        "",
        "ANNUAL CONTROL RECONCILIATION",
        "-" * 100,
        (
            f"{'FFY':<7}"
            f"{'SOCIAL':>16}"
            f"{'EXCISE':>16}"
            f"{'T2.1 SOCIAL':>16}"
            f"{'T2.1 EXCISE':>16}"
        ),
        "-" * 100,
    ]

    for row in annual_rows:
        social = Decimal(
            row["social_published_total_millions"]
        )

        excise = Decimal(
            row["excise_published_total_millions"]
        )

        i1_social = Decimal(
            row["social_table_2_1_control_millions"]
        )

        i1_excise = Decimal(
            row["excise_table_2_1_control_millions"]
        )

        lines.append(
            f"{row['fiscal_year']:<7}"
            f"{social:>16,.0f}"
            f"{excise:>16,.0f}"
            f"{i1_social:>16,.0f}"
            f"{i1_excise:>16,.0f}"
        )

    lines.extend(
        [
            "",
            "CERTIFICATION",
            "-" * 100,
        ]
    )

    for key, value in structure.items():
        lines.append(
            f"{'structure_' + key:<52}: {value}"
        )

    for key, value in certification.items():
        lines.append(
            f"{key:<52}: {value}"
        )

    lines.extend(
        [
            "",
            "BOUNDARY",
            "-" * 100,
            (
                "ETL-172I-2 certifies national OMB decomposition "
                "and accounting identities only."
            ),
            (
                "No newly decomposed component is allocated to "
                "Washington in this step."
            ),
            (
                "No Rockefeller replication, IRS reconciliation, "
                "database write, or final ledger authorization occurs."
            ),
            "",
            "ETL-172I-2 RESULT: PASS",
            "",
        ]
    )

    return "\n".join(lines)


def main() -> int:
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("=" * 100)
    print(
        "OMB SOCIAL INSURANCE AND EXCISE RECEIPTS "
        "DECOMPOSITION - ETL-172I-2"
    )
    print("=" * 100)
    print()
    print(
        "National decomposition and control reconciliation only. "
        "No Washington allocation."
    )
    print()

    i1_rows = load_i1()

    source_sha256 = sha256_file(
        SOURCE_FILE
    )

    i1_sha256 = sha256_file(
        I1_FILE
    )

    workbook, sheet = load_source()

    structure = certify_structure(
        sheet
    )

    year_columns = map_year_columns(
        sheet
    )

    detail_rows = extract_component_rows(
        sheet=sheet,
        year_columns=year_columns,
    )

    annual_rows = build_annual(
        sheet=sheet,
        year_columns=year_columns,
        detail_rows=detail_rows,
        i1_rows=i1_rows,
    )

    certification = certify_output(
        detail_rows=detail_rows,
        annual_rows=annual_rows,
    )

    write_csv(
        OUTPUT_DETAIL,
        detail_rows,
    )

    write_csv(
        OUTPUT_ANNUAL,
        annual_rows,
    )

    payload = {
        "etl_step": "ETL-172I-2",
        "result": "PASS",
        "artifact_type": (
            "OMB_SOCIAL_INSURANCE_EXCISE_DECOMPOSITION"
        ),
        "source": {
            "file": str(
                SOURCE_FILE.relative_to(
                    REPO_ROOT
                )
            ),
            "sha256": source_sha256,
            "sheet": SHEET_NAME,
            "title": EXPECTED_TITLE,
            "units": EXPECTED_UNITS,
        },
        "upstream": {
            "etl_172i_1_file": str(
                I1_FILE.relative_to(
                    REPO_ROOT
                )
            ),
            "etl_172i_1_sha256": i1_sha256,
        },
        "social_component_codes": [
            value[0]
            for value in SOCIAL_ROWS.values()
        ],
        "excise_component_codes": [
            value[0]
            for value in EXCISE_ROWS.values()
        ],
        "structure_certification": structure,
        "certification": certification,
        "methodological_boundaries": {
            "washington_allocation_performed": False,
            "rockefeller_replication": False,
            "irs_gross_collections_reconciliation": False,
            "database_writes": False,
            "final_ledger_inclusion_authorized": False,
        },
    }

    OUTPUT_JSON.write_text(
        json.dumps(
            payload,
            indent=2,
        ),
        encoding="utf-8",
    )

    report = build_report(
        source_sha256=source_sha256,
        i1_sha256=i1_sha256,
        structure=structure,
        annual_rows=annual_rows,
        certification=certification,
    )

    OUTPUT_REPORT.write_text(
        report,
        encoding="utf-8",
    )

    workbook.close()

    print("UPSTREAM PROVENANCE")
    print("-" * 100)
    print(
        f"Table 2.4 SHA256 : {source_sha256}"
    )
    print(
        f"ETL-172I-1 SHA256: {i1_sha256}"
    )
    print()

    print("ANNUAL CONTROL RECONCILIATION")
    print("-" * 100)
    print(
        f"{'FFY':<7}"
        f"{'SOCIAL':>16}"
        f"{'EXCISE':>16}"
        f"{'T2.1 SOCIAL':>16}"
        f"{'T2.1 EXCISE':>16}"
    )

    for row in annual_rows:
        social = Decimal(
            row["social_published_total_millions"]
        )

        excise = Decimal(
            row["excise_published_total_millions"]
        )

        i1_social = Decimal(
            row["social_table_2_1_control_millions"]
        )

        i1_excise = Decimal(
            row["excise_table_2_1_control_millions"]
        )

        print(
            f"{row['fiscal_year']:<7}"
            f"{social:>16,.0f}"
            f"{excise:>16,.0f}"
            f"{i1_social:>16,.0f}"
            f"{i1_excise:>16,.0f}"
        )

    print()
    print("CERTIFICATION")
    print("-" * 100)

    for key, value in certification.items():
        print(
            f"{key:<52}: {value}"
        )

    print()
    print(
        "Detail : "
        f"{OUTPUT_DETAIL.relative_to(REPO_ROOT)}"
    )
    print(
        "Annual : "
        f"{OUTPUT_ANNUAL.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON   : "
        f"{OUTPUT_JSON.relative_to(REPO_ROOT)}"
    )
    print(
        "Report : "
        f"{OUTPUT_REPORT.relative_to(REPO_ROOT)}"
    )
    print()
    print(
        "ETL-172I-2 RESULT: PASS"
    )

    return 0


if __name__ == "__main__":
    sys.exit(
        main()
    )