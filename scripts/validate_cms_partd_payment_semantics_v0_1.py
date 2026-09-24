#!/usr/bin/env python3
"""
validate_cms_partd_payment_semantics_v0_1.py

Read-only certification artifact for CMS Medicare Part D Plan Payment semantics.

Scope: CY2015-CY2024

This validator encodes and tests the source-supported semantic model:

  risk_adjusted_direct_subsidy_pmpm =
      standardized_direct_subsidy_pmpm * average_part_d_risk_score

Candidate reconstructed Part D plan-payment PMPM:

  part_d_plan_payment_pmpm =
      risk_adjusted_direct_subsidy_pmpm
      + average_reinsurance_pmpm
      + average_lics_pmpm

Boundary:
- The separate Part D Plan Reconciliation workbook is CONTROL_ONLY for
  reconstruction and is NOT added to the Part D Plans PMPM fields.
- RDS remains separate/out of scope.
- No enrollment multiplication.
- No Washington allocation.
- No database writes.
- No ledger inclusion.
- No imputation.

The script independently verifies the required Part D Plans field universe in
the preserved annual CMS Plan Payment ZIPs. It also records documentation
coverage as an explicit evidence boundary:
  * CY2015-CY2022: DIRECT_DOCUMENTATION
  * CY2023-CY2024: SCHEMA_AND_DEFINITION_CONTINUITY

This artifact does not claim that the latter two years have the same direct
documentation coverage as CY2015-CY2022.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl


YEARS = range(2015, 2025)

EXPECTED_FIELDS = {
    "CONTRACT_NUMBER",
    "PLAN_BENEFIT_PACKAGE",
    "CONTRACT_NAME",
    "ORGANIZATION_TYPE",
    "AVERAGE_PART_D_DIRECT_SUBSIDY_PM_PM_PAYMENT",
    "AVERAGE_PART_D_RISK_SCORE",
    "AVERAGE_REINSURANCE_PM_PM_PAYMENT",
    "AVERAGE_LOW_INCOME_COST_SHARING_PM_PM_PAYMENT",
}

SEMANTIC_ROWS = [
    {
        "semantic_id": "PARTD_DIRECT_SUBSIDY_STANDARDIZATION",
        "field": "Average Part D Direct Subsidy PM/PM Payment",
        "classification": "RISK_STANDARDIZED",
        "rule": "Multiply published standardized Direct Subsidy PMPM by Average Part D Risk Score to reverse 1.0-risk-score standardization.",
        "equation": "DirectSubsidy_RA_PMPM = DirectSubsidy_STD_PMPM * PartD_RiskScore",
    },
    {
        "semantic_id": "PARTD_REINSURANCE",
        "field": "Average Reinsurance PM/PM Payment",
        "classification": "ADDITIVE_PLAN_PAYMENT_COMPONENT",
        "rule": "Add published Reinsurance PMPM without risk-score multiplication.",
        "equation": "Reinsurance_PMPM = published Average Reinsurance PM/PM Payment",
    },
    {
        "semantic_id": "PARTD_LICS",
        "field": "Average Low Income Cost Sharing PM/PM Payment",
        "classification": "ADDITIVE_PLAN_PAYMENT_COMPONENT",
        "rule": "Add published LICS PMPM without risk-score multiplication.",
        "equation": "LICS_PMPM = published Average Low Income Cost Sharing PM/PM Payment",
    },
    {
        "semantic_id": "PARTD_RECONCILIATION_WORKBOOK",
        "field": "Part D Plan Reconciliation",
        "classification": "CONTROL_ONLY",
        "rule": "Retain Reconciliation Amount, LICS reconciliation, Reinsurance reconciliation, and Risk Sharing as reconciliation/provenance controls; do not add the separate Reconciliation Amount to reconstructed Part D Plans PMPM.",
        "equation": "PartD_Reconstructed != PartDPlans + ReconciliationAmount",
    },
    {
        "semantic_id": "PARTD_RDS",
        "field": "Retiree Drug Subsidy",
        "classification": "OUT_OF_SCOPE_SEPARATE",
        "rule": "RDS remains analytically separate until an independent additivity/overlap determination is completed.",
        "equation": "",
    },
]


def norm(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip().upper()
    s = s.replace("/", " ")
    s = re.sub(r"[^A-Z0-9]+", "_", s)
    return s.strip("_")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def looks_like_partd_plans(name: str) -> bool:
    base = Path(name).name.lower().replace(" ", "")
    return base.endswith(".xlsx") and "partdplans" in base and "recon" not in base


def discover(root: Path, year: int):
    year_dir = root / str(year)
    if not year_dir.exists():
        return None, [f"{year}: missing directory {year_dir}"]

    found = []
    errors = []
    for zp in sorted(year_dir.glob("*.zip")):
        try:
            with zipfile.ZipFile(zp) as zf:
                for member in zf.namelist():
                    if looks_like_partd_plans(member):
                        data = zf.read(member)
                        found.append((zp, member, data, sha256_bytes(data)))
        except Exception as exc:
            errors.append(f"{year}: cannot inspect {zp}: {exc}")

    if not found:
        return None, errors + [f"{year}: no PartDPlans workbook found"]

    hashes = {x[3] for x in found}
    if len(hashes) != 1:
        return None, errors + [
            f"{year}: conflicting PartDPlans workbook bytes across annual ZIP candidates"
        ]

    chosen = sorted(found, key=lambda x: (str(x[0]), x[1]))[0]
    return chosen, errors


def detect_header(ws, max_rows=30):
    for row_num, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), 1
    ):
        fields = {norm(v) for v in row if v is not None}
        if EXPECTED_FIELDS.issubset(fields):
            return row_num, fields
    return None, set()


def profile_workbook(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    candidates = []
    for ws in wb.worksheets:
        header_row, fields = detect_header(ws)
        if header_row is not None:
            candidates.append((ws.title, header_row, fields))
    return wb.sheetnames, candidates


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--root",
        type=Path,
        default=Path("data/raw/cms/medicare_advantage/plan_payment"),
    )
    ap.add_argument(
        "--output-dir",
        type=Path,
        default=Path(
            "data/raw/cms/medicare_advantage/validation/"
            "partd_payment_semantics_v0_1"
        ),
    )
    args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    manifest = []
    field_evidence = []
    year_checks = []
    errors = []

    for year in YEARS:
        found, errs = discover(args.root, year)
        errors.extend(errs)
        if found is None:
            year_checks.append({
                "payment_year": year,
                "source_resolved": "FAIL",
                "required_field_universe": "NOT_RUN",
                "documentation_coverage": (
                    "DIRECT_DOCUMENTATION" if year <= 2022
                    else "SCHEMA_AND_DEFINITION_CONTINUITY"
                ),
                "year_status": "FAIL",
            })
            continue

        zp, member, data, member_hash = found
        try:
            sheet_names, candidates = profile_workbook(data)
        except Exception as exc:
            errors.append(f"{year}: workbook profile failure: {exc}")
            candidates = []
            sheet_names = []

        manifest.append({
            "payment_year": year,
            "zip_path": str(zp),
            "zip_sha256": sha256_file(zp),
            "member_name": member,
            "member_sha256": member_hash,
            "sheet_names": " | ".join(sheet_names),
        })

        if len(candidates) != 1:
            errors.append(
                f"{year}: expected exactly one Part D Plans semantic table; found {len(candidates)}"
            )
            status = "FAIL"
        else:
            sheet, header_row, fields = candidates[0]
            missing = sorted(EXPECTED_FIELDS - fields)
            status = "PASS" if not missing else "FAIL"
            for field in sorted(EXPECTED_FIELDS):
                field_evidence.append({
                    "payment_year": year,
                    "sheet_name": sheet,
                    "header_row": header_row,
                    "normalized_field": field,
                    "field_present": "YES" if field in fields else "NO",
                })
            if missing:
                errors.append(f"{year}: missing fields: {', '.join(missing)}")

        year_checks.append({
            "payment_year": year,
            "source_resolved": "PASS",
            "required_field_universe": status,
            "documentation_coverage": (
                "DIRECT_DOCUMENTATION" if year <= 2022
                else "SCHEMA_AND_DEFINITION_CONTINUITY"
            ),
            "year_status": status,
        })

    semantic_evidence = []
    for year in YEARS:
        coverage = (
            "DIRECT_DOCUMENTATION"
            if year <= 2022
            else "SCHEMA_AND_DEFINITION_CONTINUITY"
        )
        for row in SEMANTIC_ROWS:
            x = dict(row)
            x["payment_year"] = year
            x["evidence_class"] = coverage
            semantic_evidence.append(x)

    write_csv(
        args.output_dir / "cms_partd_payment_semantics_source_manifest.csv",
        manifest,
        ["payment_year", "zip_path", "zip_sha256", "member_name",
         "member_sha256", "sheet_names"],
    )
    write_csv(
        args.output_dir / "cms_partd_payment_semantics_field_evidence.csv",
        field_evidence,
        ["payment_year", "sheet_name", "header_row",
         "normalized_field", "field_present"],
    )
    write_csv(
        args.output_dir / "cms_partd_payment_semantics_year_checks.csv",
        year_checks,
        ["payment_year", "source_resolved", "required_field_universe",
         "documentation_coverage", "year_status"],
    )
    write_csv(
        args.output_dir / "cms_partd_payment_semantics_policy.csv",
        semantic_evidence,
        ["payment_year", "semantic_id", "field", "classification",
         "rule", "equation", "evidence_class"],
    )

    all_years = len(year_checks) == 10
    structural_pass = all(r["year_status"] == "PASS" for r in year_checks)
    overall = all_years and structural_pass and not errors

    report = [
        "CMS MEDICARE PART D PAYMENT SEMANTICS VALIDATOR",
        "=" * 72,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        "Scope: CY2015-CY2024",
        "",
        "BOUNDARY",
        "-" * 72,
        "Read-only semantic/structural certification.",
        "No enrollment multiplication, Washington allocation, database writes,",
        "imputation, or ledger inclusion.",
        "RDS remains separate/out of scope.",
        "",
        "CERTIFIED SEMANTIC MODEL",
        "-" * 72,
        "DirectSubsidy_RA_PMPM = DirectSubsidy_STD_PMPM * PartD_RiskScore",
        "Candidate PartD_PMPM = DirectSubsidy_RA_PMPM + Reinsurance_PMPM + LICS_PMPM",
        "Separate Part D Plan Reconciliation workbook = CONTROL_ONLY",
        "Reconciliation Amount is NOT added to the Part D Plans PMPM reconstruction.",
        "",
        "DOCUMENTATION EVIDENCE BOUNDARY",
        "-" * 72,
        "CY2015-CY2022: DIRECT_DOCUMENTATION",
        "CY2023-CY2024: SCHEMA_AND_DEFINITION_CONTINUITY",
        "The latter classification is deliberately weaker and is not silently",
        "promoted to direct-documentation status.",
        "",
        "COUNTS",
        "-" * 72,
        f"Required years: 10",
        f"Resolved source years: {len(manifest)}",
        f"Field evidence rows: {len(field_evidence)}",
        f"Semantic policy rows: {len(semantic_evidence)}",
        f"Errors: {len(errors)}",
        "",
        "YEAR CHECKS",
        "-" * 72,
    ]
    for r in year_checks:
        report.append(
            f"{r['payment_year']}: {r['year_status']} | "
            f"source={r['source_resolved']} | fields={r['required_field_universe']} | "
            f"evidence={r['documentation_coverage']}"
        )
    report += ["", "ERRORS", "-" * 72]
    report += errors if errors else ["None"]
    report += [
        "",
        "OVERALL",
        "-" * 72,
        f"OVERALL: {'PASS' if overall else 'FAIL'}",
        "",
        "INTERPRETATION",
        "-" * 72,
        "A PASS certifies the required Part D Plans field universe and records the",
        "source-supported semantic policy and its documentation boundary. It does",
        "not certify Washington enrollment linkage, reconstructed dollars, county",
        "allocation, overlap with other Medicare sources, RDS additivity, or final",
        "ledger eligibility.",
    ]

    report_text = "\n".join(report) + "\n"
    (args.output_dir / "cms_partd_payment_semantics_report.txt").write_text(
        report_text, encoding="utf-8"
    )
    print(report_text)
    print(f"Outputs: {args.output_dir}")
    return 0 if overall else 1


if __name__ == "__main__":
    sys.exit(main())
