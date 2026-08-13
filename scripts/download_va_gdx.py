from __future__ import annotations

import argparse
import hashlib
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

SOURCE_PAGE = "https://www.va.gov/vetdata/Expenditures.asp"
START_FY = 2015
END_FY = 2024
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/151.0 Safari/537.36 "
    "WCFFBP-Research-Acquisition/1.0"
)


def repository_root() -> Path:
    return Path(__file__).resolve().parents[1]


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            block = handle.read(1024 * 1024)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest().upper()


def request_bytes(url: str, timeout: int = 120) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": USER_AGENT, "Accept": "*/*"},
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(
            f"HTTP {exc.code} from {url}\n{body[:2000]}"
        ) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Unable to retrieve {url}: {exc}") from exc


def get_html(url: str) -> str:
    return request_bytes(url).decode("utf-8", errors="replace")


def discover_gdx_links(html: str) -> dict[int, str]:
    links: dict[int, str] = {}
    anchor_pattern = re.compile(
        r'<a\b[^>]*href=["\']([^"\']+)["\'][^>]*>(.*?)</a>',
        flags=re.IGNORECASE | re.DOTALL,
    )
    tag_pattern = re.compile(r"<[^>]+>")

    for href, anchor_html in anchor_pattern.findall(html):
        anchor_text = tag_pattern.sub("", anchor_html).replace("&nbsp;", " ").strip()
        match = re.search(
            r"Fiscal\s+Year\s+(20\d{2})",
            anchor_text,
            flags=re.IGNORECASE,
        )
        if not match:
            continue

        fiscal_year = int(match.group(1))
        if START_FY <= fiscal_year <= END_FY:
            links[fiscal_year] = urllib.parse.urljoin(SOURCE_PAGE, href)

    return links


def filename_from_url(url: str, fiscal_year: int) -> str:
    parsed = urllib.parse.urlparse(url)
    filename = Path(urllib.parse.unquote(parsed.path)).name
    if filename:
        return filename
    return f"GDX_FY{str(fiscal_year)[-2:]}.xlsx"


def write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def methodology_regime(fiscal_year: int) -> str:
    return "REVISED_FY2024" if fiscal_year == 2024 else "LEGACY_PRE_FY2024"


def inspect_legacy_workbook(
    workbook: openpyxl.Workbook,
    fiscal_year: int,
) -> dict[str, Any]:
    sheet_names = list(workbook.sheetnames)

    if "WA" not in sheet_names:
        raise RuntimeError(
            f"FY{fiscal_year} legacy workbook does not contain a WA worksheet."
        )

    sheet = workbook["WA"]

    county_names: list[str] = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        value = row[0]
        if value is None:
            continue

        text = str(value).strip()
        if not text:
            continue

        # In the legacy Washington sheet, the first 39 data rows
        # following the header are the 39 Washington counties.
        county_names.append(text)

        if len(county_names) == 39:
            break

    if len(county_names) != 39:
        raise RuntimeError(
            f"FY{fiscal_year} WA worksheet did not yield 39 county rows. "
            f"Observed: {len(county_names)}"
        )

    return {
        "schema_regime": "LEGACY_STATE_SHEETS",
        "worksheet_count": len(sheet_names),
        "worksheets": sheet_names,
        "washington_worksheet_present": True,
        "washington_worksheet": {
            "name": "WA",
            "max_rows": sheet.max_row,
            "max_columns": sheet.max_column,
        },
        "washington_county_rows": len(county_names),
        "washington_county_names": county_names,
        "methodology_break": False,
    }


def inspect_fy2024_workbook(
    workbook: openpyxl.Workbook,
) -> dict[str, Any]:
    required = {
        "Data Guide",
        "County Direct Exp",
        "County Indirect Exp",
    }

    sheet_names = list(workbook.sheetnames)
    missing = sorted(required - set(sheet_names))

    if missing:
        raise RuntimeError(
            "FY2024 revised workbook is missing required worksheet(s): "
            + ", ".join(missing)
        )

    direct_sheet = workbook["County Direct Exp"]
    indirect_sheet = workbook["County Indirect Exp"]

    direct_rows: list[str] = []
    indirect_rows: list[str] = []

    for row in direct_sheet.iter_rows(min_row=6, values_only=True):
        if row[1] == "Washington":
            direct_rows.append(str(row[2]).strip())

    for row in indirect_sheet.iter_rows(min_row=6, values_only=True):
        if row[1] == "Washington":
            indirect_rows.append(str(row[2]).strip())

    if len(direct_rows) != 39:
        raise RuntimeError(
            "FY2024 County Direct Exp does not contain exactly 39 "
            f"Washington rows. Observed: {len(direct_rows)}"
        )

    if len(indirect_rows) != 39:
        raise RuntimeError(
            "FY2024 County Indirect Exp does not contain exactly 39 "
            f"Washington rows. Observed: {len(indirect_rows)}"
        )

    if set(direct_rows) != set(indirect_rows):
        raise RuntimeError(
            "FY2024 direct and indirect Washington county populations "
            "do not match."
        )

    return {
        "schema_regime": "REVISED_DIRECT_INDIRECT",
        "worksheet_count": len(sheet_names),
        "worksheets": sheet_names,
        "washington_worksheet_present": False,
        "required_revised_sheets_present": True,
        "washington_county_direct_rows": len(direct_rows),
        "washington_county_indirect_rows": len(indirect_rows),
        "washington_county_populations_match": True,
        "washington_county_names": sorted(set(direct_rows)),
        "county_direct_worksheet": {
            "name": "County Direct Exp",
            "max_rows": direct_sheet.max_row,
            "max_columns": direct_sheet.max_column,
        },
        "county_indirect_worksheet": {
            "name": "County Indirect Exp",
            "max_rows": indirect_sheet.max_row,
            "max_columns": indirect_sheet.max_column,
        },
        "methodology_break": True,
        "methodology_note": (
            "FY2024 introduces separate Direct and Indirect expenditure "
            "tables and revised geographic attribution methodology."
        ),
    }


def inspect_workbook(path: Path, fiscal_year: int) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
    )
    try:
        if fiscal_year == 2024:
            return inspect_fy2024_workbook(workbook)
        return inspect_legacy_workbook(workbook, fiscal_year)
    finally:
        workbook.close()


def locate_existing_workbook(
    target_dir: Path,
    expected_filename: str,
) -> Path | None:
    expected_path = target_dir / expected_filename
    if expected_path.exists():
        return expected_path

    candidates = sorted(target_dir.glob("*.xlsx"))
    if len(candidates) == 1:
        return candidates[0]
    if len(candidates) > 1:
        raise RuntimeError(
            f"Multiple XLSX files exist in {target_dir}. "
            "Cannot determine the authoritative raw artifact."
        )
    return None


def acquire_year(
    fiscal_year: int,
    url: str,
    root: Path,
) -> dict[str, Any]:
    target_dir = (
        root
        / "data"
        / "raw"
        / "va"
        / "gdx"
        / f"FY{fiscal_year}"
    )
    target_dir.mkdir(parents=True, exist_ok=True)

    expected_filename = filename_from_url(url, fiscal_year)
    existing_path = locate_existing_workbook(
        target_dir,
        expected_filename,
    )

    print()
    print("-" * 88)
    print(f"FY{fiscal_year}")
    print("-" * 88)
    print(f"Source URL ....................... {url}")

    if existing_path is not None:
        target_path = existing_path
        print(f"Existing source file ............. {target_path}")
        print("Raw-file action .................. PRESERVED — NOT OVERWRITTEN")
        acquisition_method = "PREEXISTING_LOCAL_FILE"
    else:
        target_path = target_dir / expected_filename
        print(f"Target ........................... {target_path}")
        print("Downloading .......................")

        content = request_bytes(url, timeout=300)
        temporary_path = target_path.with_suffix(target_path.suffix + ".part")
        temporary_path.write_bytes(content)
        temporary_path.replace(target_path)

        acquisition_method = "DOWNLOADED_FROM_VA"
        print(
            "Download complete ................. "
            f"{target_path.stat().st_size:,} bytes"
        )

    digest = sha256_file(target_path)
    print(f"SHA-256 .......................... {digest}")

    workbook_info = inspect_workbook(target_path, fiscal_year)
    print(f"Schema regime .................... {workbook_info['schema_regime']}")

    if fiscal_year == 2024:
        print(
            "Washington direct counties ........ "
            f"{workbook_info['washington_county_direct_rows']}"
        )
        print(
            "Washington indirect counties ...... "
            f"{workbook_info['washington_county_indirect_rows']}"
        )
        print("Direct/indirect populations ....... MATCH")
    else:
        print(
            "Washington county rows ............ "
            f"{workbook_info['washington_county_rows']}"
        )

    metadata = {
        "dataset": "VA Geographic Distribution of Expenditures (GDX)",
        "fiscal_year": fiscal_year,
        "source_agency": "U.S. Department of Veterans Affairs",
        "source_organization": (
            "National Center for Veterans Analysis and Statistics (NCVAS)"
        ),
        "source_page": SOURCE_PAGE,
        "download_url": url,
        "filename": target_path.name,
        "local_path": str(target_path),
        "file_size_bytes": target_path.stat().st_size,
        "sha256": digest,
        "acquisition_timestamp_utc": utc_now(),
        "acquisition_method": acquisition_method,
        "schema_regime": workbook_info["schema_regime"],
        "methodology_regime": methodology_regime(fiscal_year),
        "methodology_break": fiscal_year == 2024,
        "raw_file_policy": "Original source artifact preserved unchanged.",
        "analytical_warning": (
            "Do not add GDX expenditures to USAspending-derived VA "
            "expenditures without explicit overlap reconciliation."
        ),
        "workbook": workbook_info,
    }

    metadata_path = target_dir / "acquisition_metadata.json"
    write_json(metadata_path, metadata)

    hash_path = target_dir / "SHA256.txt"
    hash_path.write_text(
        f"{digest}  {target_path.name}\n",
        encoding="utf-8",
    )

    print(f"Metadata .......................... {metadata_path}")
    return metadata


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Acquire VA Geographic Distribution of Expenditures workbooks."
        )
    )
    parser.add_argument("--start-year", type=int, default=START_FY)
    parser.add_argument("--end-year", type=int, default=END_FY)
    parser.add_argument(
        "--discover-only",
        action="store_true",
        help="Discover and display VA source links without downloading files.",
    )
    args = parser.parse_args()

    if args.start_year > args.end_year:
        raise ValueError("start-year cannot exceed end-year")

    print("=" * 88)
    print("WCFFBP — VA GEOGRAPHIC DISTRIBUTION OF EXPENDITURES ACQUISITION")
    print("=" * 88)
    print()
    print(
        f"Requested fiscal years ............ "
        f"FY{args.start_year}-FY{args.end_year}"
    )
    print(f"Official source page .............. {SOURCE_PAGE}")
    print()
    print("Retrieving official VA GDX source page...")

    html = get_html(SOURCE_PAGE)
    links = discover_gdx_links(html)

    print()
    print("Discovered GDX links:")
    print()

    missing: list[int] = []

    for fiscal_year in range(args.start_year, args.end_year + 1):
        url = links.get(fiscal_year)
        if url:
            print(f"  FY{fiscal_year}: {url}")
        else:
            print(f"  FY{fiscal_year}: NOT FOUND")
            missing.append(fiscal_year)

    if missing:
        raise RuntimeError(
            "The official VA page did not expose download links for: "
            + ", ".join(f"FY{year}" for year in missing)
            + ". No downloads were attempted."
        )

    if args.discover_only:
        print()
        print("=" * 88)
        print("DISCOVERY COMPLETE — NO FILES DOWNLOADED")
        print("=" * 88)
        return

    root = repository_root()
    acquisitions: list[dict[str, Any]] = []

    for fiscal_year in range(args.start_year, args.end_year + 1):
        acquisitions.append(
            acquire_year(
                fiscal_year,
                links[fiscal_year],
                root,
            )
        )

    manifest_dir = root / "data" / "raw" / "va" / "gdx"

    manifest = {
        "dataset": "VA Geographic Distribution of Expenditures (GDX)",
        "source_page": SOURCE_PAGE,
        "manifest_generated_utc": utc_now(),
        "fiscal_year_start": args.start_year,
        "fiscal_year_end": args.end_year,
        "file_count": len(acquisitions),
        "county_target": "Washington State — 39 counties",
        "schema_regimes": {
            "FY2015-FY2023": "LEGACY_STATE_SHEETS",
            "FY2024": "REVISED_DIRECT_INDIRECT",
        },
        "methodology_break": {
            "fiscal_year": 2024,
            "description": (
                "VA revised the GDX organization and geographic methodology "
                "for FY2024. Direct and indirect expenditures are reported "
                "separately. FY2024 must not be treated as methodologically "
                "identical to FY2015-FY2023."
            ),
        },
        "double_counting_warning": (
            "GDX includes categories whose source data may overlap "
            "USAspending. Reconcile overlap before combining GDX with "
            "generic USAspending expenditure totals."
        ),
        "files": acquisitions,
    }

    manifest_path = manifest_dir / "acquisition_manifest.json"
    write_json(manifest_path, manifest)

    print()
    print("=" * 88)
    print("VA GDX ACQUISITION COMPLETE")
    print("=" * 88)
    print()
    print(
        f"Fiscal years processed ............ "
        f"{len(acquisitions)}"
    )
    print(f"Master manifest ................... {manifest_path}")


if __name__ == "__main__":
    main()