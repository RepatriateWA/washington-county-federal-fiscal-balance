#!/usr/bin/env python3
"""
validate_cms_partd_cpsc_nonmatch_classification_v0_1.py

Read-only validator for the numeric enrollment classified by the certified
CMS Medicare Part D CPSC linkage profile v0.3 as NO_PARTD_PAYMENT_MATCH.

It does not reread the monthly CPSC corpus, impute enrollment, reconstruct
payments, write to PostgreSQL, or authorize ledger inclusion.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

EXPECTED_CONTROL = 5_203_941

# Narrow evidence vocabulary only. HMO/PPO labels are deliberately NOT treated
# as evidence that a plan is outside Part D.
EXPLICIT_NONPARTD_TERMS = (
    "COST", "PACE", "MMP", "MEDICARE-MEDICAID", "MEDICARE MEDICAID",
    "DEMONSTRATION", "1876",
)

def clean(v: Any) -> str:
    return "" if v is None else str(v).strip()

def norm(v: Any) -> str:
    s = clean(v).upper().replace("/", " ")
    return re.sub(r"[^A-Z0-9]+", "_", s).strip("_")

def pbp(v: Any) -> str:
    s = clean(v)
    if not s:
        return ""
    try:
        return f"{int(float(s)):03d}"
    except (ValueError, TypeError):
        return s.upper()

def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()

def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def read_csv(p: Path) -> list[dict[str, str]]:
    with p.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def write_csv(p: Path, rows, fields) -> None:
    with p.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

def find_partdplans(root: Path, year: int):
    yd = root / str(year)
    hits = []
    for zpath in sorted(yd.glob("*.zip")):
        try:
            with zipfile.ZipFile(zpath) as zf:
                for member in zf.namelist():
                    b = Path(member).name.lower().replace(" ", "")
                    if b.endswith(".xlsx") and "partdplans" in b and "recon" not in b:
                        data = zf.read(member)
                        hits.append((zpath, member, data, sha256_bytes(data)))
        except zipfile.BadZipFile:
            continue
    if not hits:
        raise RuntimeError(f"CY{year}: no PartDPlans workbook found")
    hashes = {h[3] for h in hits}
    if len(hashes) != 1:
        raise RuntimeError(f"CY{year}: conflicting PartDPlans workbook bytes")
    return sorted(hits, key=lambda x: (str(x[0]), x[1]))[0]

def load_payment_universe(root: Path, year: int):
    zpath, member, data, member_hash = find_partdplans(root, year)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for hr, vals in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), 1):
                hdr = [norm(x) for x in vals]
                if "CONTRACT_NUMBER" not in hdr or "PLAN_BENEFIT_PACKAGE" not in hdr:
                    continue
                ci = hdr.index("CONTRACT_NUMBER")
                pi = hdr.index("PLAN_BENEFIT_PACKAGE")
                keys, contracts, dups = set(), set(), []
                for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                    c = clean(row[ci] if ci < len(row) else "").upper()
                    p = pbp(row[pi] if pi < len(row) else "")
                    if not c or not p:
                        continue
                    k = (c, p)
                    if k in keys:
                        dups.append(k)
                    keys.add(k)
                    contracts.add(c)
                return keys, contracts, dups, {
                    "year": year, "zip_path": str(zpath),
                    "zip_sha256": sha256_file(zpath), "member_name": member,
                    "member_sha256": member_hash, "sheet_name": ws.title,
                    "header_row": hr, "payment_key_count": len(keys),
                    "contract_count": len(contracts),
                }
    finally:
        wb.close()
    raise RuntimeError(f"CY{year}: Contract/PBP table not found")

def nonpartd_evidence(plan_type: str, org_type: str) -> str:
    text = f"{plan_type} | {org_type}".upper()
    return ";".join(t for t in EXPLICIT_NONPARTD_TERMS if t in text)

def get_col(row: dict[str, str], candidates: tuple[str, ...]) -> str:
    # Allows harmless v0.3 header naming variations without changing semantics.
    normalized = {norm(k): v for k, v in row.items()}
    for c in candidates:
        if norm(c) in normalized:
            return clean(normalized[norm(c)])
    return ""

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--v03-dir", type=Path, default=Path(
        "data/raw/cms/medicare_advantage/validation/"
        "partd_cpsc_enrollment_linkage_profile_v0_3"))
    ap.add_argument("--plan-payment-root", type=Path, default=Path(
        "data/raw/cms/medicare_advantage/plan_payment"))
    ap.add_argument("--output-dir", type=Path, default=Path(
        "data/raw/cms/medicare_advantage/validation/"
        "partd_cpsc_nonmatch_classification_v0_1"))
    a = ap.parse_args()
    a.output_dir.mkdir(parents=True, exist_ok=True)

    src = a.v03_dir / "partd_cpsc_no_partd_payment_match_aggregates.csv"
    if not src.exists():
        print(f"ERROR: missing input: {src}", file=sys.stderr)
        return 2

    errors = []
    raw = read_csv(src)
    if not raw:
        print("ERROR: v0.3 aggregate is empty", file=sys.stderr)
        return 2

    universes, contracts, manifests, duplicate_rows = {}, {}, [], []
    for y in range(2015, 2025):
        try:
            ks, cs, ds, mf = load_payment_universe(a.plan_payment_root, y)
            universes[y], contracts[y] = ks, cs
            manifests.append(mf)
            duplicate_rows.extend({"year": y, "contract": c, "pbp": p} for c, p in ds)
        except Exception as e:
            universes[y], contracts[y] = set(), set()
            errors.append(str(e))

    parsed = []
    for i, r in enumerate(raw, 2):
        try:
            ys = get_col(r, ("year", "calendar_year"))
            ms = get_col(r, ("month", "calendar_month"))
            cs = get_col(r, ("contract", "contract_number"))
            ps = get_col(r, ("pbp", "plan_benefit_package", "plan_id"))
            pts = get_col(r, ("cpsc_plan_type", "plan_type"))
            ots = get_col(r, ("cpsc_organization_type", "organization_type"))
            ns = get_col(r, ("numeric_enrollment", "enrollment", "numeric_enrollment_sum"))
            y, m, n = int(ys), int(ms), int(float(ns))
            c, p = cs.upper(), pbp(ps)
            if not c or not p:
                raise ValueError("blank Contract/PBP")
            parsed.append({
                "year": y, "month": m, "contract": c, "pbp": p,
                "cpsc_plan_type": pts, "cpsc_organization_type": ots,
                "numeric_enrollment": n,
            })
        except Exception as e:
            errors.append(f"Input row {i}: {e}")

    months_by_key, years_by_key = defaultdict(set), defaultdict(set)
    enrollment_by_key = Counter()
    for r in parsed:
        k = (r["contract"], r["pbp"])
        months_by_key[k].add((r["year"], r["month"]))
        years_by_key[k].add(r["year"])
        enrollment_by_key[k] += r["numeric_enrollment"]

    classified = []
    class_total, year_total = Counter(), Counter()
    plan_total, org_total, key_total = Counter(), Counter(), Counter()

    for r in parsed:
        y, n = r["year"], r["numeric_enrollment"]
        k = (r["contract"], r["pbp"])
        same_contract = r["contract"] in contracts.get(y, set())
        adjacent = [yy for yy in (y - 1, y + 1)
                    if 2015 <= yy <= 2024 and k in universes.get(yy, set())]
        evidence = nonpartd_evidence(r["cpsc_plan_type"], r["cpsc_organization_type"])

        if same_contract:
            cls = "SAME_CONTRACT_OTHER_PBP"
            why = "Contract exists in same-year PartDPlans; exact PBP does not."
        elif adjacent:
            cls = "ADJACENT_YEAR_PARTD_MATCH"
            why = "Exact Contract+PBP appears in adjacent year(s): " + ",".join(map(str, adjacent))
        elif evidence:
            cls = "NON_PARTD_PLAN_EVIDENCE"
            why = "Explicit non-Part-D vocabulary evidence: " + evidence
        elif len(months_by_key[k]) >= 3:
            cls = "PERSISTENT_UNRESOLVED"
            why = f"Unmatched in {len(months_by_key[k])} distinct months."
        else:
            cls = "OTHER_REVIEW_REQUIRED"
            why = "No stronger evidence-based classification established."

        class_total[cls] += n
        year_total[y] += n
        plan_total[(y, r["cpsc_plan_type"])] += n
        org_total[(y, r["cpsc_organization_type"])] += n
        key_total[(y, *k)] += n
        classified.append({
            **r, "classification_v0_1": cls, "classification_rationale": why,
            "same_year_contract_exists": "YES" if same_contract else "NO",
            "adjacent_year_exact_match": ",".join(map(str, adjacent)),
            "explicit_nonpartd_evidence": evidence,
            "distinct_nonmatch_months_for_key": len(months_by_key[k]),
            "distinct_nonmatch_years_for_key": len(years_by_key[k]),
            "all_years_nonmatch_enrollment_for_key": enrollment_by_key[k],
        })

    input_total = sum(r["numeric_enrollment"] for r in parsed)
    classified_total = sum(class_total.values())

    annual = []
    for y in range(2015, 2025):
        ct = Counter()
        for r in classified:
            if r["year"] == y:
                ct[r["classification_v0_1"]] += r["numeric_enrollment"]
        tot = sum(ct.values())
        annual.append({
            "year": y, "source_nonmatch_numeric_enrollment": year_total[y],
            "same_contract_other_pbp": ct["SAME_CONTRACT_OTHER_PBP"],
            "adjacent_year_partd_match": ct["ADJACENT_YEAR_PARTD_MATCH"],
            "non_partd_plan_evidence": ct["NON_PARTD_PLAN_EVIDENCE"],
            "persistent_unresolved": ct["PERSISTENT_UNRESOLVED"],
            "other_review_required": ct["OTHER_REVIEW_REQUIRED"],
            "classified_total": tot,
            "accounting_difference": year_total[y] - tot,
        })

    summary = [{
        "classification": c, "numeric_enrollment": n,
        "share_of_nonmatch": n / input_total if input_total else 0.0
    } for c, n in sorted(class_total.items(), key=lambda x: -x[1])]

    plan_rows = [{"year": y, "cpsc_plan_type": t, "numeric_enrollment": n}
                 for (y, t), n in sorted(plan_total.items())]
    org_rows = [{"year": y, "cpsc_organization_type": t, "numeric_enrollment": n}
                for (y, t), n in sorted(org_total.items())]
    key_rows = [{
        "year": y, "contract": c, "pbp": p, "numeric_enrollment": n,
        "distinct_nonmatch_months_all_years": len(months_by_key[(c, p)]),
        "distinct_nonmatch_years": len(years_by_key[(c, p)]),
    } for (y, c, p), n in sorted(key_total.items(), key=lambda x: (-x[1], x[0]))]

    write_csv(a.output_dir / "partd_cpsc_nonmatch_classified_rows.csv", classified,
              ["year","month","contract","pbp","cpsc_plan_type","cpsc_organization_type",
               "numeric_enrollment","classification_v0_1","classification_rationale",
               "same_year_contract_exists","adjacent_year_exact_match",
               "explicit_nonpartd_evidence","distinct_nonmatch_months_for_key",
               "distinct_nonmatch_years_for_key","all_years_nonmatch_enrollment_for_key"])
    write_csv(a.output_dir / "partd_cpsc_nonmatch_annual_accounting.csv", annual,
              ["year","source_nonmatch_numeric_enrollment","same_contract_other_pbp",
               "adjacent_year_partd_match","non_partd_plan_evidence",
               "persistent_unresolved","other_review_required","classified_total",
               "accounting_difference"])
    write_csv(a.output_dir / "partd_cpsc_nonmatch_classification_summary.csv", summary,
              ["classification","numeric_enrollment","share_of_nonmatch"])
    write_csv(a.output_dir / "partd_cpsc_nonmatch_by_plan_type.csv", plan_rows,
              ["year","cpsc_plan_type","numeric_enrollment"])
    write_csv(a.output_dir / "partd_cpsc_nonmatch_by_organization_type.csv", org_rows,
              ["year","cpsc_organization_type","numeric_enrollment"])
    write_csv(a.output_dir / "partd_cpsc_nonmatch_by_contract_pbp.csv", key_rows,
              ["year","contract","pbp","numeric_enrollment",
               "distinct_nonmatch_months_all_years","distinct_nonmatch_years"])
    write_csv(a.output_dir / "partd_payment_source_manifest.csv", manifests,
              ["year","zip_path","zip_sha256","member_name","member_sha256",
               "sheet_name","header_row","payment_key_count","contract_count"])
    write_csv(a.output_dir / "partd_payment_duplicate_keys.csv", duplicate_rows,
              ["year","contract","pbp"])

    checks = [
        ("v0.3 control total equals 5,203,941", input_total == EXPECTED_CONTROL),
        ("classification accounting identity", classified_total == input_total),
        ("year accounting identity", sum(year_total.values()) == input_total),
        ("plan-type accounting identity", sum(plan_total.values()) == input_total),
        ("organization-type accounting identity", sum(org_total.values()) == input_total),
        ("Contract+PBP accounting identity", sum(key_total.values()) == input_total),
        ("all annual accounting differences zero", all(x["accounting_difference"] == 0 for x in annual)),
        ("no duplicate annual Part D payment keys", not duplicate_rows),
        ("no source/profile errors", not errors),
    ]
    overall = all(v for _, v in checks)

    lines = [
        "CMS MEDICARE PART D CPSC NONMATCH CLASSIFICATION VALIDATOR v0.1",
        "=" * 76,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        "Scope: v0.3 NO_PARTD_PAYMENT_MATCH numeric enrollment",
        "",
        "BOUNDARY", "-" * 76,
        "Read-only diagnostic classification.",
        "No monthly CPSC reread; no imputation; no payment reconstruction.",
        "No database writes or ledger inclusion.",
        "",
        "CONTROL", "-" * 76,
        f"Expected v0.3 control: {EXPECTED_CONTROL:,}",
        f"Input numeric enrollment: {input_total:,}",
        f"Classified numeric enrollment: {classified_total:,}",
        f"Accounting difference: {input_total - classified_total:,}",
        "",
        "CLASSIFICATION TOTALS", "-" * 76,
    ]
    for r in summary:
        lines.append(f"{r['classification']}: {r['numeric_enrollment']:,} ({r['share_of_nonmatch']:.4%})")
    lines += ["", "CHECKS", "-" * 76]
    lines += [f"{label}: {'PASS' if ok else 'FAIL'}" for label, ok in checks]
    lines += ["", "ERRORS", "-" * 76]
    lines += errors if errors else ["None"]
    lines += [
        "", "OVERALL", "-" * 76, f"OVERALL: {'PASS' if overall else 'FAIL'}",
        "", "INTERPRETATION BOUNDARY", "-" * 76,
        "PASS certifies classification/accounting mechanics only.",
        "It does not prove every nonmatch is outside Part D.",
        "PERSISTENT_UNRESOLVED and OTHER_REVIEW_REQUIRED remain review populations.",
        "No observation is force-linked to a payment record.",
    ]
    report = "\n".join(lines) + "\n"
    (a.output_dir / "partd_cpsc_nonmatch_classification_report.txt").write_text(
        report, encoding="utf-8")
    print(report)
    print(f"Outputs: {a.output_dir}")
    return 0 if overall else 1

if __name__ == "__main__":
    sys.exit(main())
