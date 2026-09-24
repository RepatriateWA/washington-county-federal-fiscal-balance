#!/usr/bin/env python3
"""
profile_cms_partd_landscape.py

Read-only structural profiler for the preserved official CMS CY2006-CY2025
Landscape archive. Target analytical years: CY2015-CY2024.

Purpose
-------
Inventory archive members and profile tabular schemas without deciding which
field means "offers Part D". This is deliberately a structural gate only.

No imputation. No classification. No payment reconstruction. No DB writes.
"""

from __future__ import annotations
import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

DEFAULT_ARCHIVE = Path(
    "data/raw/cms/medicare_partd/landscape/cy2006-cy2025-landscape-files.zip"
)
DEFAULT_OUT = Path(
    "data/raw/cms/medicare_partd/landscape/validation/landscape_profile_v0_1"
)
TARGET_YEARS = set(range(2015, 2025))

def clean(v: Any) -> str:
    return "" if v is None else str(v).strip()

def norm(v: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", clean(v).upper()).strip("_")

def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def infer_year(name: str):
    # Prefer explicit 4-digit coverage years.
    vals = [int(x) for x in re.findall(r"(?<!\d)(20(?:0[6-9]|1\d|2[0-5]))(?!\d)", name)]
    return vals[-1] if vals else None

def decode_delimited(data: bytes):
    text = None
    encoding = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            encoding = enc
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        return None
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",|\t;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        header = next(reader)
    except StopIteration:
        return None
    rows = sum(1 for _ in reader)
    return encoding, header, rows

def workbook_profiles(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            best = None
            for rno, vals in enumerate(
                ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 40), values_only=True), 1
            ):
                populated = [clean(x) for x in vals if clean(x)]
                score = len(populated)
                if best is None or score > best[0]:
                    best = (score, rno, [clean(x) for x in vals])
            if best is None:
                out.append((ws.title, "", [], 0))
            else:
                _, hr, header = best
                nonempty = 0
                for vals in ws.iter_rows(min_row=hr + 1, values_only=True):
                    if any(clean(x) for x in vals):
                        nonempty += 1
                out.append((ws.title, hr, header, nonempty))
    finally:
        wb.close()
    return out

def write_csv(path: Path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--archive", type=Path, default=DEFAULT_ARCHIVE)
    p.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    a = p.parse_args()
    a.output_dir.mkdir(parents=True, exist_ok=True)

    if not a.archive.exists():
        print(f"ERROR: archive not found: {a.archive}", file=sys.stderr)
        return 2
    if not zipfile.is_zipfile(a.archive):
        print(f"ERROR: not a ZIP: {a.archive}", file=sys.stderr)
        return 2

    inventory = []
    schemas = []
    errors = []
    year_tabular = Counter()
    candidate_fields = []
    archive_hash = sha256_file(a.archive)

    with zipfile.ZipFile(a.archive) as z:
        bad = z.testzip()
        for info in z.infolist():
            if info.is_dir():
                continue
            name = info.filename
            suffix = Path(name).suffix.lower()
            year = infer_year(name)
            row = {
                "member_name": name,
                "inferred_year": year or "",
                "suffix": suffix,
                "bytes": info.file_size,
                "crc": info.CRC,
                "target_year": "YES" if year in TARGET_YEARS else "NO",
                "profile_status": "INVENTORIED",
            }
            inventory.append(row)

            if year not in TARGET_YEARS or suffix not in (".csv", ".txt", ".xlsx", ".xlsm"):
                continue
            try:
                data = z.read(name)
                mhash = sha256_bytes(data)
                if suffix in (".csv", ".txt"):
                    parsed = decode_delimited(data)
                    if not parsed:
                        raise RuntimeError("unable to decode delimited file")
                    enc, header, nrows = parsed
                    schemas.append({
                        "year": year, "member_name": name, "member_sha256": mhash,
                        "sheet_name": "", "header_row": 1, "encoding": enc,
                        "data_rows": nrows, "column_count": len(header),
                        "columns": " | ".join(header),
                    })
                    for c in header:
                        nc = norm(c)
                        if any(tok in nc for tok in ("CONTRACT","PLAN","PBP","DRUG","PART_D","PDP","MA_PD")):
                            candidate_fields.append({
                                "year": year, "member_name": name, "sheet_name": "",
                                "field_name": c, "normalized_field": nc
                            })
                    year_tabular[year] += 1
                else:
                    for sheet, hr, header, nrows in workbook_profiles(data):
                        schemas.append({
                            "year": year, "member_name": name, "member_sha256": mhash,
                            "sheet_name": sheet, "header_row": hr, "encoding": "",
                            "data_rows": nrows, "column_count": len(header),
                            "columns": " | ".join(header),
                        })
                        for c in header:
                            nc = norm(c)
                            if any(tok in nc for tok in ("CONTRACT","PLAN","PBP","DRUG","PART_D","PDP","MA_PD")):
                                candidate_fields.append({
                                    "year": year, "member_name": name, "sheet_name": sheet,
                                    "field_name": c, "normalized_field": nc
                                })
                        year_tabular[year] += 1
                row["profile_status"] = "PROFILED"
                del data
            except Exception as exc:
                row["profile_status"] = "ERROR"
                errors.append({
                    "year": year, "member_name": name,
                    "error": f"{type(exc).__name__}: {exc}"
                })

    missing_years = [y for y in sorted(TARGET_YEARS) if year_tabular[y] == 0]
    checks = [
        ("archive ZIP integrity", bad is None),
        ("all CY2015-CY2024 have >=1 profiled tabular member", not missing_years),
        ("no target-member profile errors", not errors),
        ("candidate Contract/Plan/PBP/Part-D-like fields observed", bool(candidate_fields)),
    ]
    overall = all(ok for _, ok in checks)

    write_csv(
        a.output_dir / "cms_partd_landscape_member_inventory.csv",
        inventory,
        ["member_name","inferred_year","suffix","bytes","crc","target_year","profile_status"]
    )
    write_csv(
        a.output_dir / "cms_partd_landscape_schema_profile.csv",
        schemas,
        ["year","member_name","member_sha256","sheet_name","header_row","encoding",
         "data_rows","column_count","columns"]
    )
    write_csv(
        a.output_dir / "cms_partd_landscape_candidate_fields.csv",
        candidate_fields,
        ["year","member_name","sheet_name","field_name","normalized_field"]
    )
    write_csv(
        a.output_dir / "cms_partd_landscape_profile_errors.csv",
        errors, ["year","member_name","error"]
    )

    lines = [
        "CMS MEDICARE PART D LANDSCAPE STRUCTURAL PROFILE v0.1",
        "=" * 72,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        f"Archive: {a.archive}",
        f"Archive SHA-256: {archive_hash}",
        "",
        "BOUNDARY",
        "-" * 72,
        "Structural profiling only. No Part D status semantics are assigned here.",
        "No imputation, reconstruction, database writes, or ledger inclusion.",
        "",
        "INVENTORY",
        "-" * 72,
        f"Archive members: {len(inventory):,}",
        f"Profiled target-year tabular schemas: {len(schemas):,}",
        f"Candidate semantic fields: {len(candidate_fields):,}",
        f"Profile errors: {len(errors):,}",
        "",
        "TARGET-YEAR COVERAGE",
        "-" * 72,
    ]
    for y in sorted(TARGET_YEARS):
        lines.append(f"CY{y}: {year_tabular[y]} profiled tabular member/sheet(s)")
    lines += ["", "CHECKS", "-" * 72]
    lines += [f"{label}: {'PASS' if ok else 'FAIL'}" for label, ok in checks]
    lines += ["", "ERRORS", "-" * 72]
    if errors:
        lines += [f"CY{e['year']} {e['member_name']}: {e['error']}" for e in errors]
    else:
        lines.append("None")
    lines += ["", "OVERALL", "-" * 72, f"OVERALL: {'PASS' if overall else 'FAIL'}"]
    report = "\n".join(lines) + "\n"
    (a.output_dir / "cms_partd_landscape_profile_report.txt").write_text(
        report, encoding="utf-8"
    )
    print(report)
    print(f"Outputs: {a.output_dir}")
    return 0 if overall else 1

if __name__ == "__main__":
    sys.exit(main())
