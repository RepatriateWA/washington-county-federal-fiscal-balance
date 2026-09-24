#!/usr/bin/env python3
"""
profile_cms_partd_payment_missing_values_v0_1.py

CMS Medicare Part D — PartDPlans Missing-Value Profiler v0.1

PURPOSE
-------
Read-only profiling of CMS PartDPlans payment fields for CY2015-CY2024.

This script answers a narrow source-semantics question before any missing-value
policy is adopted:

    What is the empirical pattern of blank/non-numeric values in:
      * Average Part D Direct Subsidy PM/PM Payment
      * Average Part D Risk Score
      * Average Reinsurance PM/PM Payment
      * Average Low Income Cost Sharing PM/PM Payment

It reports counts by year, organization type, and field, distinguishing:
    NUMERIC_POSITIVE
    NUMERIC_ZERO
    NUMERIC_NEGATIVE
    BLANK
    NONNUMERIC_OTHER

It also produces row-level exception evidence and co-missingness patterns.

BOUNDARY
--------
* No blank -> zero conversion.
* No imputation.
* No payment reconstruction.
* No CPSC enrollment processing.
* No database writes.
* No ledger inclusion.
* PartDReconciliation remains CONTROL_ONLY and is not read as an additive
  payment source.

Run from repository root:
    python scripts/profile_cms_partd_payment_missing_values_v0_1.py
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
PLAN_PAYMENT_ROOT = (
    REPO_ROOT / "data" / "raw" / "cms" / "medicare_advantage" / "plan_payment"
)
OUTPUT_DIR = (
    REPO_ROOT / "data" / "raw" / "cms" / "medicare_partd"
    / "validation" / "partd_payment_missing_values_v0_1"
)

START_YEAR = 2015
END_YEAR = 2024
EXPECTED_YEARS = 10

CONTRACT_ALIASES = {
    "contract", "contract number", "contract_number", "contract id",
    "contract_id", "contractid", "hnumber"
}
PBP_ALIASES = {
    "plan benefit package", "plan benefit package id",
    "plan benefit package number", "plan id", "plan_id",
    "plan identifier", "plan_identifier", "pbp", "pbp id", "pbp_id"
}

FIELD_SPECS = {
    "DIRECT_SUBSIDY": ("direct", "subsidy"),
    "RISK_SCORE": ("risk", "score"),
    "REINSURANCE": ("reinsurance",),
    "LICS": ("low", "income", "cost", "sharing"),
}

def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip())

def low(v) -> str:
    return norm(v).lower()

def norm_contract(v) -> str:
    return norm(v).upper()

def norm_pbp(v) -> str:
    s = norm(v)
    if re.fullmatch(r"\d+(?:\.0+)?", s):
        return str(int(float(s))).zfill(3)
    return s.upper()

def infer_year(text: str) -> Optional[int]:
    vals = {
        int(x)
        for x in re.findall(r"(?<!\d)(20(?:1[4-9]|2[0-5]))(?!\d)", text)
    }
    return next(iter(vals)) if len(vals) == 1 else None

def find_col(headers: list[str], aliases: set[str]) -> Optional[str]:
    by_low = {low(h): h for h in headers if h is not None}
    for alias in aliases:
        if alias in by_low:
            return by_low[alias]
    return None

def find_contains(headers: list[str], tokens: tuple[str, ...]) -> Optional[str]:
    for h in headers:
        lh = low(h)
        if all(t in lh for t in tokens):
            return h
    return None

def classify_value(v) -> tuple[str, Optional[float], str]:
    """
    Preserve the distinction between an actually blank cell and a textual
    nonnumeric value. Returns (class, numeric_value, raw_text).
    """
    if v is None:
        return "BLANK", None, ""
    raw = norm(v)
    if raw == "":
        return "BLANK", None, ""
    cleaned = raw.replace(",", "").replace("$", "")
    try:
        x = float(cleaned)
    except Exception:
        return "NONNUMERIC_OTHER", None, raw
    if x > 0:
        return "NUMERIC_POSITIVE", x, raw
    if x < 0:
        return "NUMERIC_NEGATIVE", x, raw
    return "NUMERIC_ZERO", x, raw

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def discover_plan_payment_zip(year: int) -> Path:
    candidates = []
    for p in PLAN_PAYMENT_ROOT.rglob("*.zip"):
        if "validation" in {x.lower() for x in p.parts}:
            continue
        y = infer_year(p.name) or infer_year(str(p.parent))
        if y == year:
            candidates.append(p.resolve())

    candidates = sorted(set(candidates))
    if not candidates:
        raise FileNotFoundError(f"CY{year}: no Plan Payment ZIP found")

    if len(candidates) > 1:
        by_hash = defaultdict(list)
        for p in candidates:
            by_hash[sha256_file(p)].append(p)
        if len(by_hash) == 1:
            return sorted(candidates, key=lambda p: (len(str(p)), str(p)))[0]

        scored = []
        for p in candidates:
            s = p.name.lower()
            score = (4 if str(year) in s else 0) + (3 if "payment" in s else 0)
            scored.append((score, p))
        best_score = max(x[0] for x in scored)
        best = [p for score, p in scored if score == best_score]
        if len(best) != 1:
            raise RuntimeError(
                f"CY{year}: ambiguous Plan Payment ZIPs: "
                + " | ".join(map(str, candidates))
            )
        return best[0]

    return candidates[0]

def find_partdplans_member(zf: zipfile.ZipFile) -> str:
    hits = []
    for n in zf.namelist():
        if n.endswith("/"):
            continue
        b = (
            Path(n).name.lower()
            .replace(" ", "")
            .replace("_", "")
            .replace("-", "")
        )
        if b.endswith("partdplans.xlsx"):
            hits.append(n)
    if len(hits) != 1:
        raise RuntimeError(
            f"Expected exactly one PartDPlans workbook; found {len(hits)}: {hits}"
        )
    return hits[0]

def write_csv(path: Path, rows: list[dict], fields: list[str]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

def profile_year(year: int):
    zpath = discover_plan_payment_zip(year)
    with zipfile.ZipFile(zpath) as zf:
        member = find_partdplans_member(zf)
        workbook_bytes = zf.read(member)

    wb = load_workbook(
        io.BytesIO(workbook_bytes), read_only=True, data_only=True
    )

    detail_rows = []
    source_rows = []
    found = False

    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)

        header = None
        header_row_number = None
        contract_col = None
        pbp_col = None

        for row_number, candidate in enumerate(rows, start=1):
            if row_number > 25:
                break
            candidate_headers = [norm(x) for x in candidate]
            cc = find_col(candidate_headers, CONTRACT_ALIASES)
            pc = find_col(candidate_headers, PBP_ALIASES)
            if cc and pc:
                header = candidate
                header_row_number = row_number
                contract_col = cc
                pbp_col = pc
                break

        if header is None:
            continue

        headers = [norm(x) for x in header]
        idx = {h: i for i, h in enumerate(headers)}
        ci, pi = idx[contract_col], idx[pbp_col]

        org_col = find_contains(headers, ("organization", "type"))
        field_cols = {
            field: find_contains(headers, tokens)
            for field, tokens in FIELD_SPECS.items()
        }

        # LICS fallback for a possible abbreviated header.
        if not field_cols["LICS"]:
            field_cols["LICS"] = find_contains(headers, ("lics",))

        missing = []
        if not org_col:
            missing.append("ORGANIZATION_TYPE")
        missing.extend([k for k, v in field_cols.items() if not v])
        if missing:
            raise RuntimeError(
                f"CY{year} {ws.title}: missing required fields: {missing}; "
                f"headers={headers}"
            )

        found = True
        oi = idx[org_col]
        fis = {field: idx[col] for field, col in field_cols.items()}

        row_count = 0
        for source_row_number, row in enumerate(
            rows, start=header_row_number + 1
        ):
            contract = norm_contract(row[ci] if ci < len(row) else "")
            pbp = norm_pbp(row[pi] if pi < len(row) else "")
            if not contract or not pbp:
                continue

            row_count += 1
            org = norm(row[oi] if oi < len(row) else "") or "[BLANK ORGANIZATION TYPE]"

            classes = {}
            for field, i in fis.items():
                value = row[i] if i < len(row) else None
                cls, numeric, raw = classify_value(value)
                classes[field] = cls
                detail_rows.append({
                    "year": year,
                    "sheet": ws.title,
                    "source_row_number": source_row_number,
                    "contract": contract,
                    "pbp": pbp,
                    "organization_type": org,
                    "field": field,
                    "classification": cls,
                    "numeric_value": "" if numeric is None else numeric,
                    "raw_value": raw,
                })

            source_rows.append({
                "year": year,
                "sheet": ws.title,
                "source_row_number": source_row_number,
                "contract": contract,
                "pbp": pbp,
                "organization_type": org,
                "direct_subsidy_class": classes["DIRECT_SUBSIDY"],
                "risk_score_class": classes["RISK_SCORE"],
                "reinsurance_class": classes["REINSURANCE"],
                "lics_class": classes["LICS"],
                "missingness_pattern": "|".join(
                    f"{field}={classes[field]}"
                    for field in (
                        "DIRECT_SUBSIDY", "RISK_SCORE",
                        "REINSURANCE", "LICS"
                    )
                ),
            })

    wb.close()

    if not found:
        raise RuntimeError(
            f"CY{year}: no Contract+PBP PartDPlans worksheet resolved"
        )

    return detail_rows, source_rows, {
        "year": year,
        "zip_path": str(zpath),
        "zip_sha256": sha256_file(zpath),
        "member": member,
        "member_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "detail_field_observations": len(detail_rows),
        "plan_rows": len(source_rows),
    }

def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 78)
    print("CMS MEDICARE PART D PARTDPLANS MISSING-VALUE PROFILER v0.1")
    print("=" * 78)
    print()
    print("BOUNDARY:")
    print("  Source profiling only.")
    print("  No blank-to-zero conversion, imputation, reconstruction, CPSC processing,")
    print("  database writes, or ledger inclusion.")
    print()

    all_detail = []
    all_plan_rows = []
    source_manifest = []
    errors = []

    for year in range(START_YEAR, END_YEAR + 1):
        try:
            detail, plan_rows, source = profile_year(year)
            all_detail.extend(detail)
            all_plan_rows.extend(plan_rows)
            source_manifest.append(source)
        except Exception as exc:
            errors.append({"year": year, "message": str(exc)})

    # Aggregate by year / organization type / field / classification.
    counts = Counter()
    for r in all_detail:
        counts[
            (
                r["year"],
                r["organization_type"],
                r["field"],
                r["classification"],
            )
        ] += 1

    summary_rows = [
        {
            "year": y,
            "organization_type": org,
            "field": field,
            "classification": cls,
            "count": count,
        }
        for (y, org, field, cls), count in sorted(counts.items())
    ]

    # Aggregate by year/field across all organization types.
    yearly_counts = Counter()
    for r in all_detail:
        yearly_counts[(r["year"], r["field"], r["classification"])] += 1
    yearly_rows = [
        {
            "year": y,
            "field": field,
            "classification": cls,
            "count": count,
        }
        for (y, field, cls), count in sorted(yearly_counts.items())
    ]

    # Co-missingness / classification pattern by year and organization type.
    pattern_counts = Counter()
    for r in all_plan_rows:
        pattern_counts[
            (r["year"], r["organization_type"], r["missingness_pattern"])
        ] += 1
    pattern_rows = [
        {
            "year": y,
            "organization_type": org,
            "missingness_pattern": pattern,
            "count": count,
        }
        for (y, org, pattern), count in sorted(pattern_counts.items())
    ]

    # Exception rows: any BLANK or NONNUMERIC_OTHER in any payment field.
    exception_plan_rows = [
        r for r in all_plan_rows
        if any(
            r[k] in {"BLANK", "NONNUMERIC_OTHER"}
            for k in (
                "direct_subsidy_class",
                "risk_score_class",
                "reinsurance_class",
                "lics_class",
            )
        )
    ]

    blank_counts_by_field = Counter()
    nonnumeric_counts_by_field = Counter()
    for r in all_detail:
        if r["classification"] == "BLANK":
            blank_counts_by_field[r["field"]] += 1
        elif r["classification"] == "NONNUMERIC_OTHER":
            nonnumeric_counts_by_field[r["field"]] += 1

    years_resolved = len(source_manifest)
    checks = [
        (
            "All CY2015-CY2024 PartDPlans sources resolved",
            years_resolved == EXPECTED_YEARS,
            f"{years_resolved}/{EXPECTED_YEARS}",
        ),
        (
            "All resolved plan rows yield four field observations",
            len(all_detail) == len(all_plan_rows) * 4,
            f"{len(all_detail):,} == {len(all_plan_rows):,} * 4",
        ),
        (
            "No source/profile errors",
            len(errors) == 0,
            f"errors={len(errors)}",
        ),
    ]
    overall = all(ok for _, ok, _ in checks)

    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missing_values_detail_v0_1.csv",
        all_detail,
        [
            "year", "sheet", "source_row_number", "contract", "pbp",
            "organization_type", "field", "classification",
            "numeric_value", "raw_value",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missing_values_by_org_v0_1.csv",
        summary_rows,
        ["year", "organization_type", "field", "classification", "count"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missing_values_by_year_v0_1.csv",
        yearly_rows,
        ["year", "field", "classification", "count"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missingness_patterns_v0_1.csv",
        pattern_rows,
        ["year", "organization_type", "missingness_pattern", "count"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missing_value_plan_exceptions_v0_1.csv",
        exception_plan_rows,
        [
            "year", "sheet", "source_row_number", "contract", "pbp",
            "organization_type", "direct_subsidy_class", "risk_score_class",
            "reinsurance_class", "lics_class", "missingness_pattern",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missing_values_source_manifest_v0_1.csv",
        source_manifest,
        [
            "year", "zip_path", "zip_sha256", "member", "member_sha256",
            "detail_field_observations", "plan_rows",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missing_values_checks_v0_1.csv",
        [
            {"check": name, "status": "PASS" if ok else "FAIL", "detail": detail}
            for name, ok, detail in checks
        ],
        ["check", "status", "detail"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_missing_values_errors_v0_1.csv",
        errors,
        ["year", "message"],
    )

    print("SOURCE COVERAGE:")
    print(f"  Years required: {EXPECTED_YEARS}")
    print(f"  Years resolved: {years_resolved}")
    print(f"  Plan rows profiled: {len(all_plan_rows):,}")
    print(f"  Field observations: {len(all_detail):,}")
    print()
    print("BLANK COUNTS:")
    for field in FIELD_SPECS:
        print(f"  {field:<16} {blank_counts_by_field[field]:,}")
    print()
    print("OTHER NONNUMERIC COUNTS:")
    for field in FIELD_SPECS:
        print(f"  {field:<16} {nonnumeric_counts_by_field[field]:,}")
    print()
    print(f"PLAN ROWS WITH >=1 BLANK/OTHER NONNUMERIC FIELD: {len(exception_plan_rows):,}")
    print()

    print("YEARLY BLANK PROFILE:")
    for year in range(START_YEAR, END_YEAR + 1):
        vals = {
            field: yearly_counts[(year, field, "BLANK")]
            for field in FIELD_SPECS
        }
        print(
            f"  CY{year}: "
            f"direct={vals['DIRECT_SUBSIDY']:,} | "
            f"risk={vals['RISK_SCORE']:,} | "
            f"reinsurance={vals['REINSURANCE']:,} | "
            f"LICS={vals['LICS']:,}"
        )

    # Show organization types carrying blanks, compactly.
    blank_org = Counter()
    for r in all_detail:
        if r["classification"] == "BLANK":
            blank_org[(r["organization_type"], r["field"])] += 1

    print()
    print("TOP BLANK-BEARING ORGANIZATION TYPE / FIELD COMBINATIONS:")
    for (org, field), count in blank_org.most_common(25):
        print(f"  {count:>6,} | {field:<16} | {org}")

    print()
    print("CHECKS:")
    for name, ok, detail in checks:
        print(f"  [{'PASS' if ok else 'FAIL'}] {name} — {detail}")

    if errors:
        print()
        print(f"ERRORS: {len(errors)}")
        for e in errors:
            print(f"  CY{e['year']}: {e['message']}")

    print()
    print(f"Output directory: {OUTPUT_DIR}")
    print(f"OVERALL: {'PASS' if overall else 'FAIL'}")
    print()
    print("NOTE:")
    print("  PASS certifies the profiling mechanics and source coverage only.")
    print("  It does NOT establish that a blank means zero or authorize any")
    print("  missing-value substitution.")
    return 0 if overall else 1

if __name__ == "__main__":
    sys.exit(main())
