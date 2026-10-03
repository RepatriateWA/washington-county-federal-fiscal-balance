from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

import openpyxl


ROOT = Path(
    "data/raw/cms/medicare_rds/"
    "program_statistics_partd_enrollment"
)

OUTDIR = (
    ROOT
    / "validation"
    / "mdcr_enroll_d3_xlsx_profile_v0_2"
)

SHEETS_OUT = (
    OUTDIR
    / "cms_mdcr_enroll_d3_xlsx_sheets_v0_2.csv"
)

PREVIEW_OUT = (
    OUTDIR
    / "cms_mdcr_enroll_d3_xlsx_preview_v0_2.csv"
)

CHECKS_OUT = (
    OUTDIR
    / "cms_mdcr_enroll_d3_xlsx_checks_v0_2.csv"
)

REPORT_OUT = (
    OUTDIR
    / "cms_mdcr_enroll_d3_xlsx_report_v0_2.txt"
)

YEARS = list(range(2015, 2024))


def clean(value: object) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value).replace("\r", " ").replace("\n", " ").split()
    )


def find_zip(year: int) -> Path:
    matches = sorted(
        (ROOT / str(year)).glob("*.zip")
    )

    if len(matches) != 1:
        raise RuntimeError(
            f"CY{year}: expected exactly one ZIP; "
            f"found {len(matches)}"
        )

    return matches[0]


def is_xlsx(name: str) -> bool:
    return Path(name).suffix.lower() in {
        ".xlsx",
        ".xlsm",
    }


def d3_score(
    member_name: str,
    sheet_name: str,
    preview_text: str,
) -> tuple[int, list[str]]:

    blob = (
        f"{member_name} "
        f"{sheet_name} "
        f"{preview_text}"
    ).lower()

    score = 0
    reasons: list[str] = []

    patterns = [
        (
            r"mdcr\s*enroll\s*d\s*3",
            5,
            "MDCR ENROLL D 3 text",
        ),
        (
            r"enroll[_\s-]*d[_\s-]*3",
            5,
            "ENROLL D 3 text",
        ),
        (
            r"03epd",
            5,
            "03EPD identifier",
        ),
        (
            r"area\s+of\s+residence",
            4,
            "Area of Residence",
        ),
        (
            r"retiree\s+drug\s+subsid",
            3,
            "Retiree Drug Subsidy",
        ),
        (
            r"\brds\b",
            2,
            "RDS text",
        ),
        (
            r"low\s+income\s+subsid",
            1,
            "Low Income Subsidy",
        ),
        (
            r"type\s+of\s+plan",
            1,
            "Type of Plan",
        ),
    ]

    for pattern, points, label in patterns:
        if re.search(
            pattern,
            blob,
            flags=re.IGNORECASE,
        ):
            score += points
            reasons.append(label)

    return score, reasons


def make_check(
    label: str,
    passed: bool,
    detail: str,
) -> dict[str, str]:

    return {
        "check": label,
        "status": (
            "PASS"
            if passed
            else "FAIL"
        ),
        "detail": detail,
    }


def main() -> None:
    print("=" * 100)
    print(
        "CMS MDCR ENROLL D 3 XLSX STRUCTURAL "
        "PROFILER v0.2"
    )
    print("=" * 100)
    print()
    print("READ-ONLY")
    print("No database writes.")
    print("No RDS dollar allocation.")
    print("No geographic interpretation.")
    print()

    OUTDIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    sheet_rows: list[dict[str, str]] = []
    preview_rows: list[dict[str, str]] = []
    errors: list[str] = []

    print("=" * 100)
    print("XLSX WORKBOOK / WORKSHEET INVENTORY")
    print("=" * 100)

    for year in YEARS:
        try:
            zpath = find_zip(year)

            with zipfile.ZipFile(zpath) as zf:
                members = [
                    x
                    for x in zf.infolist()
                    if (
                        not x.is_dir()
                        and is_xlsx(x.filename)
                    )
                ]

                print()
                print(f"CY{year}")
                print(f"ZIP: {zpath}")
                print(
                    f"XLSX members: {len(members)}"
                )

                for info in members:
                    data = zf.read(
                        info.filename
                    )

                    wb = openpyxl.load_workbook(
                        io.BytesIO(data),
                        read_only=True,
                        data_only=True,
                    )

                    try:
                        for ws in wb.worksheets:

                            raw_preview: list[
                                list[str]
                            ] = []

                            # Read enough rows to expose
                            # titles, headers, and first
                            # observations.
                            for row in ws.iter_rows(
                                min_row=1,
                                max_row=min(
                                    ws.max_row,
                                    25,
                                ),
                                values_only=True,
                            ):
                                values = [
                                    clean(v)
                                    for v in row
                                ]

                                raw_preview.append(
                                    values
                                )

                            preview_text = " ".join(
                                " ".join(row)
                                for row in raw_preview
                            )

                            score, reasons = (
                                d3_score(
                                    info.filename,
                                    ws.title,
                                    preview_text,
                                )
                            )

                            sheet_rows.append(
                                {
                                    "year": str(year),
                                    "zip_path": str(
                                        zpath
                                    ),
                                    "member_name": (
                                        info.filename
                                    ),
                                    "sheet_name": (
                                        ws.title
                                    ),
                                    "max_row": str(
                                        ws.max_row
                                    ),
                                    "max_column": str(
                                        ws.max_column
                                    ),
                                    "d3_score": str(
                                        score
                                    ),
                                    "d3_reasons": (
                                        "; ".join(
                                            reasons
                                        )
                                    ),
                                }
                            )

                            print(
                                f"  {info.filename}"
                            )
                            print(
                                f"    SHEET: "
                                f"{ws.title}"
                            )
                            print(
                                f"    rows="
                                f"{ws.max_row:,} | "
                                f"cols="
                                f"{ws.max_column:,} | "
                                f"D3 score={score}"
                            )

                            if reasons:
                                print(
                                    "    evidence: "
                                    + "; ".join(
                                        reasons
                                    )
                                )

                            for row_number, values in (
                                enumerate(
                                    raw_preview,
                                    start=1,
                                )
                            ):
                                if not any(values):
                                    continue

                                preview_rows.append(
                                    {
                                        "year": str(
                                            year
                                        ),
                                        "member_name": (
                                            info.filename
                                        ),
                                        "sheet_name": (
                                            ws.title
                                        ),
                                        "row_number": str(
                                            row_number
                                        ),
                                        "values": (
                                            " | ".join(
                                                values
                                            )
                                        ),
                                        "d3_score": str(
                                            score
                                        ),
                                    }
                                )

                    finally:
                        wb.close()

        except Exception as exc:
            message = (
                f"CY{year}: {exc!r}"
            )
            errors.append(message)
            print(
                f"{year}: ERROR | {exc!r}"
            )

    # --------------------------------------------------
    # Identify strongest candidate sheet(s) per year.
    # --------------------------------------------------

    candidate_by_year: dict[
        int,
        list[dict[str, str]],
    ] = {}

    for year in YEARS:
        rows = [
            row
            for row in sheet_rows
            if int(row["year"]) == year
        ]

        if not rows:
            continue

        max_score = max(
            int(row["d3_score"])
            for row in rows
        )

        # Require positive evidence rather than
        # arbitrarily selecting a sheet.
        if max_score <= 0:
            continue

        candidate_by_year[year] = [
            row
            for row in rows
            if int(row["d3_score"]) == max_score
        ]

    print()
    print("=" * 100)
    print("STRONGEST D3 CANDIDATE BY YEAR")
    print("=" * 100)

    for year in YEARS:
        rows = candidate_by_year.get(
            year,
            [],
        )

        if not rows:
            print(
                f"CY{year}: NO POSITIVE D3 "
                f"CANDIDATE"
            )
            continue

        for row in rows:
            print(
                f"CY{year} | "
                f"{row['member_name']} | "
                f"{row['sheet_name']} | "
                f"score={row['d3_score']} | "
                f"{row['d3_reasons']}"
            )

    # --------------------------------------------------
    # Preview strongest candidates.
    # --------------------------------------------------

    print()
    print("=" * 100)
    print("D3 CANDIDATE CONTENT PREVIEW")
    print("=" * 100)

    for year in YEARS:
        candidates = candidate_by_year.get(
            year,
            [],
        )

        for candidate in candidates:
            print()
            print(
                f"CY{year} | "
                f"{candidate['member_name']} | "
                f"{candidate['sheet_name']}"
            )
            print("-" * 100)

            matching_preview = [
                row
                for row in preview_rows
                if (
                    row["year"]
                    == str(year)
                    and row["member_name"]
                    == candidate[
                        "member_name"
                    ]
                    and row["sheet_name"]
                    == candidate[
                        "sheet_name"
                    ]
                )
            ]

            for row in matching_preview:
                print(
                    f"{int(row['row_number']):3d}"
                    f" | {row['values']}"
                )

    # --------------------------------------------------
    # Checks
    # --------------------------------------------------

    profiled_years = {
        int(row["year"])
        for row in sheet_rows
    }

    candidate_years = set(
        candidate_by_year
    )

    unique_candidate_years = {
        year
        for year, rows
        in candidate_by_year.items()
        if len(rows) == 1
    }

    checks = [
        make_check(
            "CY2015-CY2023 XLSX sources profiled",
            profiled_years == set(YEARS),
            (
                f"{len(profiled_years)}/9 "
                f"years"
            ),
        ),
        make_check(
            "Positive D3 candidate found for every year",
            candidate_years == set(YEARS),
            (
                f"{len(candidate_years)}/9 "
                f"years"
            ),
        ),
        make_check(
            "Unique strongest D3 candidate for every year",
            unique_candidate_years
            == set(YEARS),
            (
                f"{len(unique_candidate_years)}"
                "/9 years"
            ),
        ),
        make_check(
            "No XLSX profiling errors",
            len(errors) == 0,
            (
                "0 errors"
                if not errors
                else "; ".join(errors)
            ),
        ),
    ]

    # --------------------------------------------------
    # Evidence files
    # --------------------------------------------------

    sheet_fields = [
        "year",
        "zip_path",
        "member_name",
        "sheet_name",
        "max_row",
        "max_column",
        "d3_score",
        "d3_reasons",
    ]

    with SHEETS_OUT.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=sheet_fields,
        )
        writer.writeheader()
        writer.writerows(sheet_rows)

    preview_fields = [
        "year",
        "member_name",
        "sheet_name",
        "row_number",
        "values",
        "d3_score",
    ]

    with PREVIEW_OUT.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=preview_fields,
        )
        writer.writeheader()
        writer.writerows(preview_rows)

    with CHECKS_OUT.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "check",
                "status",
                "detail",
            ],
        )
        writer.writeheader()
        writer.writerows(checks)

    print()
    print("=" * 100)
    print("CHECKS")
    print("=" * 100)

    for row in checks:
        print(
            f"{row['status']:4} | "
            f"{row['check']} | "
            f"{row['detail']}"
        )

    overall = all(
        row["status"] == "PASS"
        for row in checks
    )

    report_lines = [
        (
            "CMS MDCR ENROLL D 3 XLSX "
            "STRUCTURAL PROFILER v0.2"
        ),
        "=" * 78,
        "",
        "Scope: CY2015-CY2023",
        (
            f"Worksheet observations: "
            f"{len(sheet_rows)}"
        ),
        (
            f"Preview observations: "
            f"{len(preview_rows)}"
        ),
        (
            f"Years with positive D3 "
            f"candidate: "
            f"{len(candidate_years)}"
        ),
        "",
        "CHECKS",
        "-" * 78,
    ]

    for row in checks:
        report_lines.append(
            f"{row['status']} | "
            f"{row['check']} | "
            f"{row['detail']}"
        )

    report_lines.extend(
        [
            "",
            "BOUNDARY",
            "-" * 78,
            (
                "This profiler identifies the "
                "structurally strongest MDCR ENROLL "
                "D 3 worksheet for each year."
            ),
            (
                "It does not yet certify the "
                "meaning of Area of Residence."
            ),
            (
                "It does not treat RDS enrollment "
                "as a dollar-allocation weight."
            ),
            (
                "No RDS payment has been assigned "
                "to Washington or any county."
            ),
        ]
    )

    REPORT_OUT.write_text(
        "\n".join(report_lines) + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 100)
    print(
        "OVERALL XLSX STRUCTURAL PROFILE: "
        + (
            "PASS"
            if overall
            else "FAIL"
        )
    )
    print("=" * 100)

    print()
    print("Evidence written:")
    print(f"  {SHEETS_OUT}")
    print(f"  {PREVIEW_OUT}")
    print(f"  {CHECKS_OUT}")
    print(f"  {REPORT_OUT}")

    print()
    print(
        "No RDS geographic or dollar "
        "allocation has been made."
    )

    if not overall:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
