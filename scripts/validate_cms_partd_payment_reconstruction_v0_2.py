#!/usr/bin/env python3
"""
validate_cms_partd_payment_reconstruction_v0_2.py

CMS Medicare Part D — Washington Payment Reconstruction v0.2

Scope:
  Washington, Jan 2015-Sep 2024 (117 months).

Certified regression controls:
  WA numeric CPSC enrollment             111,380,906
  Exact same-year PartDPlans match       106,176,965
  No exact same-year PartDPlans match      5,203,941

Payment semantics already established upstream:
  DirectSubsidy_RA_PMPM = DirectSubsidy_STD_PMPM * PartD_RiskScore
  PartD_PMPM = DirectSubsidy_RA_PMPM + Reinsurance_PMPM + LICS_PMPM
  PartD_Dollars = Numeric Enrollment * PartD_PMPM

Boundary:
  * Exact same-year Contract+PBP only.
  * No blank-to-zero conversion.
  * No imputation.
  * No sibling-PBP propagation.
  * No adjacent-year payment carry-forward.
  * No database writes.
  * No ledger inclusion.
  * PartDReconciliation is CONTROL_ONLY and not added.
  * RDS remains separate/out of scope.
  * No-exact-payment-match enrollment receives no reconstructed dollars here.

This validator intentionally reconstructs only the exact-match population. The
separately certified affirmative Part D / source-payment-unavailable exceptions
remain outside reconstructed dollars and must remain NULL/unavailable, not zero.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, getcontext
from pathlib import Path

import openpyxl

getcontext().prec = 40

START = (2015, 1)
END = (2024, 9)

EXPECTED_MONTHS = 117
EXPECTED_WA_NUMERIC = 111_380_906
EXPECTED_EXACT_MATCH = 106_176_965
EXPECTED_NO_EXACT = 5_203_941

PAYMENT_FIELDS = {
    "CONTRACT_NUMBER",
    "PLAN_BENEFIT_PACKAGE",
    "CONTRACT_NAME",
    "ORGANIZATION_TYPE",
    "AVERAGE_PART_D_DIRECT_SUBSIDY_PM_PM_PAYMENT",
    "AVERAGE_PART_D_RISK_SCORE",
    "AVERAGE_REINSURANCE_PM_PM_PAYMENT",
    "AVERAGE_LOW_INCOME_COST_SHARING_PM_PM_PAYMENT",
}

REQ = {
    "contract": {"CONTRACT_NUMBER", "CONTRACT_ID", "CONTRACT"},
    "pbp": {"PLAN_ID", "PLAN_BENEFIT_PACKAGE", "PBP", "PLAN_BENEFIT_PACKAGE_ID"},
    "state": {"STATE", "STATE_ABBREVIATION", "STATE_ABBR"},
    "county": {"COUNTY", "COUNTY_NAME"},
    "enrollment": {"ENROLLMENT", "ENROLLMENT_COUNT", "ENROLLEES"},
}
OPT = {
    "plan_type": {"PLAN_TYPE", "PLAN_TYPE_DESCRIPTION"},
    "organization_type": {"ORGANIZATION_TYPE", "ORG_TYPE"},
}
MONTHS = {
    1: ("JAN", "JANUARY"), 2: ("FEB", "FEBRUARY"), 3: ("MAR", "MARCH"),
    4: ("APR", "APRIL"), 5: ("MAY",), 6: ("JUN", "JUNE"),
    7: ("JUL", "JULY"), 8: ("AUG", "AUGUST"),
    9: ("SEP", "SEPT", "SEPTEMBER"), 10: ("OCT", "OCTOBER"),
    11: ("NOV", "NOVEMBER"), 12: ("DEC", "DECEMBER"),
}


def clean(v):
    return "" if v is None else str(v).strip()


def norm(v):
    return re.sub(r"[^A-Z0-9]+", "_", clean(v).upper().replace("/", " ")).strip("_")


def pbp(v):
    s = clean(v)
    if not s:
        return ""
    try:
        return f"{int(float(s)):03d}"
    except Exception:
        return s.upper()


def enrollment(v):
    s = clean(v).replace(",", "")
    if re.fullmatch(r"\d+", s):
        return int(s), "NUMERIC"
    return None, "NONNUMERIC"


def dec(v, field, year, key):
    s = clean(v).replace(",", "").replace("$", "")
    if not s:
        raise ValueError(f"CY{year} {key[0]}-{key[1]}: blank {field}")
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(
            f"CY{year} {key[0]}-{key[1]}: nonnumeric {field}={v!r}"
        ) from exc


def month_iter():
    y, m = START
    while (y, m) <= END:
        yield y, m
        m += 1
        if m == 13:
            y, m = y + 1, 1


def sha_file(p):
    h = hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1048576), b""):
            h.update(b)
    return h.hexdigest()


def sha_member(p, n):
    h = hashlib.sha256()
    with zipfile.ZipFile(p) as z, z.open(n) as f:
        for b in iter(lambda: f.read(1048576), b""):
            h.update(b)
    return h.hexdigest()


def infer_date(text):
    u = text.upper()
    c = re.sub(r"[^A-Z0-9]", "", u)
    a = {
        (int(y), int(m))
        for y, m in re.findall(r"(20(?:1[4-9]|2[0-4]))(0[1-9]|1[0-2])", c)
    }
    if len(a) == 1:
        return (*next(iter(a)), "YYYYMM")
    a = {
        (int(y), int(m))
        for m, y in re.findall(r"(0[1-9]|1[0-2])(20(?:1[4-9]|2[0-4]))", c)
    }
    if len(a) == 1:
        return (*next(iter(a)), "MMYYYY")
    ys = {int(x) for x in re.findall(r"20(?:1[4-9]|2[0-4])", u)}
    ms = {
        m
        for m, names in MONTHS.items()
        if any(re.search(rf"(?<![A-Z]){n}(?![A-Z])", u) for n in names)
    }
    if len(ys) == len(ms) == 1:
        return next(iter(ys)), next(iter(ms)), "YEAR_MONTH_NAME"
    return None, None, "UNRESOLVED"


@contextmanager
def text_source(path, member=None):
    if path.suffix.lower() == ".zip":
        z = zipfile.ZipFile(path)
        raw = z.open(member)
        t = io.TextIOWrapper(
            raw, encoding="utf-8-sig", errors="replace", newline=""
        )
        try:
            yield t
        finally:
            t.close()
            z.close()
    else:
        f = path.open("r", encoding="utf-8-sig", errors="replace", newline="")
        try:
            yield f
        finally:
            f.close()


def header_info(path, member):
    with text_source(path, member) as f:
        sample = f.read(8192)
        delim = max([",", "\t", "|"], key=lambda x: sample.count(x))
        f.seek(0)
        r = csv.reader(f, delimiter=delim)
        for rn, row in enumerate(r, 1):
            if rn > 50:
                break
            h = [norm(x) for x in row]
            mp = {}
            ok = True
            for k, aliases in REQ.items():
                pos = next((h.index(a) for a in aliases if a in h), None)
                if pos is None:
                    ok = False
                    break
                mp[k] = pos
            if ok:
                for k, aliases in OPT.items():
                    mp[k] = next((h.index(a) for a in aliases if a in h), None)
                return rn, mp, delim
    return None, None, None


def build_index(root):
    out = []
    scanned = 0
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".zip", ".csv", ".txt"}:
            continue
        parts = {x.lower() for x in path.parts}
        if "validation" in parts or "plan_payment" in parts:
            continue
        scanned += 1
        if path.suffix.lower() == ".zip":
            try:
                with zipfile.ZipFile(path) as z:
                    members = [
                        n
                        for n in z.namelist()
                        if Path(n).suffix.lower() in {".csv", ".txt"}
                    ]
            except Exception:
                continue
        else:
            members = [None]
        for member in members:
            try:
                rn, mp, d = header_info(path, member)
            except Exception:
                continue
            if rn is None:
                continue
            y, m, method = infer_date(f"{path} {member or path.name}")
            out.append(
                {
                    "path": path,
                    "member": member,
                    "year": y,
                    "month": m,
                    "method": method,
                    "header": rn,
                    "map": mp,
                    "delim": d,
                }
            )
    return out, scanned


def choose(idx, y, m):
    a = [x for x in idx if x["year"] == y and x["month"] == m]

    def rank(x):
        s = re.sub(
            r"[^a-z0-9]", "", f"{x['path']} {x['member'] or ''}".lower()
        )
        score = (
            (1000 if "cpsc" in s else 0)
            + (1000 if "contractplanstatecounty" in s else 0)
            + (100 if "enrollment" in s else 0)
        )
        if "snp" in s:
            score -= 500
        if "scc" in s and "cpsc" not in s:
            score -= 500
        return -score, str(x["path"]), (x["member"] or "")

    a.sort(key=rank)
    return (a[0] if a else None), a


def wa_rows(src):
    with text_source(src["path"], src["member"]) as f:
        r = csv.reader(f, delimiter=src["delim"])
        mp = src["map"]
        for rn, row in enumerate(r, 1):
            if rn <= src["header"]:
                continue

            def g(k):
                i = mp.get(k)
                return row[i] if i is not None and i < len(row) else ""

            if clean(g("state")).upper() not in {"WA", "WASHINGTON"}:
                continue
            yield {
                "contract": clean(g("contract")).upper(),
                "pbp": pbp(g("pbp")),
                "county": clean(g("county")),
                "enrollment": clean(g("enrollment")),
                "plan_type": clean(g("plan_type")),
                "org": clean(g("organization_type")),
            }


def payment_keys(root, y):
    found = []
    for zpath in sorted((root / str(y)).glob("*.zip")):
        try:
            with zipfile.ZipFile(zpath) as z:
                for n in z.namelist():
                    b = Path(n).name.lower().replace(" ", "")
                    if b.endswith(".xlsx") and "partdplans" in b and "recon" not in b:
                        data = z.read(n)
                        found.append(
                            (zpath, n, data, hashlib.sha256(data).hexdigest())
                        )
        except Exception:
            pass

    if not found:
        raise RuntimeError(f"{y}: PartDPlans workbook not found")
    if len({x[3] for x in found}) != 1:
        raise RuntimeError(f"{y}: conflicting PartDPlans workbook bytes")

    zp, n, data, mhash = sorted(found, key=lambda x: (str(x[0]), x[1]))[0]
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)

    try:
        for ws in wb.worksheets:
            for hr, row in enumerate(
                ws.iter_rows(min_row=1, max_row=30, values_only=True), 1
            ):
                h = [norm(x) for x in row]
                if not PAYMENT_FIELDS.issubset(set(h)):
                    continue
                pos = {k: h.index(k) for k in PAYMENT_FIELDS}
                keys = {}
                dups = []

                for vals in ws.iter_rows(min_row=hr + 1, values_only=True):
                    def v(k):
                        i = pos[k]
                        return vals[i] if i < len(vals) else None

                    key = (
                        clean(v("CONTRACT_NUMBER")).upper(),
                        pbp(v("PLAN_BENEFIT_PACKAGE")),
                    )
                    if not all(key):
                        continue

                    rec = {
                        "name": clean(v("CONTRACT_NAME")),
                        "org": clean(v("ORGANIZATION_TYPE")),
                        "direct_std": v(
                            "AVERAGE_PART_D_DIRECT_SUBSIDY_PM_PM_PAYMENT"
                        ),
                        "risk_score": v("AVERAGE_PART_D_RISK_SCORE"),
                        "reinsurance": v(
                            "AVERAGE_REINSURANCE_PM_PM_PAYMENT"
                        ),
                        "lics": v(
                            "AVERAGE_LOW_INCOME_COST_SHARING_PM_PM_PAYMENT"
                        ),
                    }

                    if key in keys:
                        dups.append(key)
                    else:
                        keys[key] = rec

                return (
                    keys,
                    dups,
                    {
                        "zip_path": str(zp),
                        "zip_sha256": sha_file(zp),
                        "member_name": n,
                        "member_sha256": mhash,
                        "sheet_name": ws.title,
                        "header_row": hr,
                    },
                )
    finally:
        wb.close()

    raise RuntimeError(f"{y}: PartDPlans semantic table not found")


def write_csv(p, rows, fields):
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def money(d):
    return f"{d.quantize(Decimal('0.01')):,.2f}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--cpsc-root",
        type=Path,
        default=Path("data/raw/cms/medicare_advantage"),
    )
    ap.add_argument(
        "--plan-payment-root",
        type=Path,
        default=Path("data/raw/cms/medicare_advantage/plan_payment"),
    )
    ap.add_argument(
        "--output-dir",
        type=Path,
        default=Path(
            "data/raw/cms/medicare_partd/validation/"
            "partd_payment_reconstruction_v0_2"
        ),
    )
    a = ap.parse_args()
    a.output_dir.mkdir(parents=True, exist_ok=True)

    errors = []
    warnings = []

    print("=" * 78)
    print("CMS MEDICARE PART D PAYMENT RECONSTRUCTION v0.2")
    print("=" * 78)
    print()
    print("BOUNDARY:")
    print("  Exact same-year Contract+PBP payment matches only.")
    print("  No blank-to-zero conversion, imputation, sibling propagation,")
    print("  adjacent-year carry-forward, DB writes, or ledger inclusion.")
    print("  No-exact-payment-match population receives no reconstructed dollars.")
    print()

    print("Building metadata-only CPSC source index...", flush=True)
    idx, scanned = build_index(a.cpsc_root)
    print(
        f"Scanned {scanned:,} candidate files; "
        f"admitted {len(idx):,} CPSC-like members.",
        flush=True,
    )

    pkeys = {}
    pmanifest = []
    dups = []
    for y in range(2015, 2025):
        try:
            k, d, meta = payment_keys(a.plan_payment_root, y)
            pkeys[y] = k
            pmanifest.append(
                {"payment_year": y, "payment_key_count": len(k), **meta}
            )
            dups += [
                {"payment_year": y, "contract": x[0], "pbp": x[1]}
                for x in d
            ]
        except Exception as exc:
            errors.append(f"PartDPlans CY{y}: {exc}")
            pkeys[y] = {}

    # Parse and validate all payment inputs once, before touching monthly dollars.
    payment_numeric = {}
    bad_payment_inputs = []
    for y, keys in pkeys.items():
        payment_numeric[y] = {}
        for key, rec in keys.items():
            try:
                direct_std = dec(rec["direct_std"], "DIRECT_SUBSIDY", y, key)
                risk = dec(rec["risk_score"], "RISK_SCORE", y, key)
                reins = dec(rec["reinsurance"], "REINSURANCE", y, key)
                lics = dec(rec["lics"], "LICS", y, key)
                direct_ra = direct_std * risk
                total_pmpm = direct_ra + reins + lics
                payment_numeric[y][key] = {
                    **rec,
                    "direct_std_d": direct_std,
                    "risk_d": risk,
                    "reinsurance_d": reins,
                    "lics_d": lics,
                    "direct_ra_d": direct_ra,
                    "total_pmpm_d": total_pmpm,
                }
            except Exception as exc:
                bad_payment_inputs.append(
                    {
                        "year": y,
                        "contract": key[0],
                        "pbp": key[1],
                        "organization_type": rec["org"],
                        "message": str(exc),
                    }
                )

    # National PartDPlans can contain missing inputs. This is not itself a failure:
    # only WA exact-match exposure matters. Therefore use pkeys for linkage, and
    # payment_numeric only for reconstruction.
    monthly = []
    manifest = []
    detail = []
    wa_bad_exposure = Counter()

    for y, m in month_iter():
        print(f"{y}-{m:02d}: starting...", flush=True)
        src, cands = choose(idx, y, m)
        if not src:
            errors.append(f"{y}-{m:02d}: no indexed CPSC source resolved")
            monthly.append(
                {
                    "year": y,
                    "month": m,
                    "source_status": "MISSING",
                    "wa_rows": 0,
                    "numeric_enrollment": 0,
                    "exact_payment_match_numeric_enrollment": 0,
                    "reconstructed_numeric_enrollment": 0,
                    "exact_match_unusable_payment_input_enrollment": 0,
                    "no_exact_payment_match_numeric_enrollment": 0,
                    "nonnumeric_rows": 0,
                    "direct_subsidy_ra_dollars": "",
                    "reinsurance_dollars": "",
                    "lics_dollars": "",
                    "total_partd_dollars": "",
                    "numeric_accounting_difference": 0,
                    "exact_match_reconstruction_difference": 0,
                    "dollar_identity_difference": "",
                }
            )
            continue

        hashes = []
        for c in cands:
            try:
                hashes.append(
                    sha_member(c["path"], c["member"])
                    if c["path"].suffix.lower() == ".zip"
                    else sha_file(c["path"])
                )
            except Exception as exc:
                warnings.append(
                    f"{y}-{m:02d}: candidate hash error: {exc}"
                )
        if len(set(hashes)) > 1:
            warnings.append(
                f"{y}-{m:02d}: {len(cands)} candidates / "
                f"{len(set(hashes))} distinct hashes"
            )

        sh = (
            sha_member(src["path"], src["member"])
            if src["path"].suffix.lower() == ".zip"
            else sha_file(src["path"])
        )
        manifest.append(
            {
                "year": y,
                "month": m,
                "selected_path": str(src["path"]),
                "selected_member": src["member"] or "",
                "selected_source_sha256": sh,
                "date_inference_method": src["method"],
                "candidate_count": len(cands),
                "distinct_candidate_hashes": len(set(hashes)),
            }
        )

        wr = tot = match = reconstructed = unusable = no = nonnum = 0
        direct_dollars = Decimal("0")
        reins_dollars = Decimal("0")
        lics_dollars = Decimal("0")

        try:
            for r in wa_rows(src):
                wr += 1
                n, status = enrollment(r["enrollment"])
                key = (r["contract"], r["pbp"])

                if status != "NUMERIC":
                    nonnum += 1
                    continue

                tot += n
                raw_payment = pkeys[y].get(key)
                if raw_payment is None:
                    no += n
                    continue

                match += n
                pay = payment_numeric[y].get(key)
                if pay is None:
                    unusable += n
                    wa_bad_exposure[(y, key[0], key[1], raw_payment["org"])] += n
                    continue

                reconstructed += n
                nd = Decimal(n)
                ab = nd * pay["direct_ra_d"]
                rr = nd * pay["reinsurance_d"]
                ll = nd * pay["lics_d"]

                direct_dollars += ab
                reins_dollars += rr
                lics_dollars += ll

                detail.append(
                    {
                        "year": y,
                        "month": m,
                        "contract": key[0],
                        "pbp": key[1],
                        "county": r["county"],
                        "cpsc_plan_type": r["plan_type"],
                        "cpsc_organization_type": r["org"],
                        "payment_organization_type": pay["org"],
                        "numeric_enrollment": n,
                        "direct_subsidy_std_pmpm": str(pay["direct_std_d"]),
                        "risk_score": str(pay["risk_d"]),
                        "direct_subsidy_ra_pmpm": str(pay["direct_ra_d"]),
                        "reinsurance_pmpm": str(pay["reinsurance_d"]),
                        "lics_pmpm": str(pay["lics_d"]),
                        "total_partd_pmpm": str(pay["total_pmpm_d"]),
                        "direct_subsidy_ra_dollars": str(ab),
                        "reinsurance_dollars": str(rr),
                        "lics_dollars": str(ll),
                        "total_partd_dollars": str(ab + rr + ll),
                    }
                )

        except Exception as exc:
            errors.append(f"{y}-{m:02d}: streaming parse failed: {exc}")
            continue

        total_dollars = direct_dollars + reins_dollars + lics_dollars
        component_sum = direct_dollars + reins_dollars + lics_dollars
        dollar_diff = total_dollars - component_sum

        monthly.append(
            {
                "year": y,
                "month": m,
                "source_status": "PARSED",
                "wa_rows": wr,
                "numeric_enrollment": tot,
                "exact_payment_match_numeric_enrollment": match,
                "reconstructed_numeric_enrollment": reconstructed,
                "exact_match_unusable_payment_input_enrollment": unusable,
                "no_exact_payment_match_numeric_enrollment": no,
                "nonnumeric_rows": nonnum,
                "direct_subsidy_ra_dollars": str(direct_dollars),
                "reinsurance_dollars": str(reins_dollars),
                "lics_dollars": str(lics_dollars),
                "total_partd_dollars": str(total_dollars),
                "numeric_accounting_difference": tot - match - no,
                "exact_match_reconstruction_difference": match - reconstructed - unusable,
                "dollar_identity_difference": str(dollar_diff),
            }
        )

        print(
            f"{y}-{m:02d}: complete | WA numeric={tot:,} | "
            f"exact={match:,} | reconstructed={reconstructed:,} | "
            f"unusable={unusable:,} | no-exact={no:,} | "
            f"Part D=${money(total_dollars)}",
            flush=True,
        )

    parsed = sum(r["source_status"] == "PARSED" for r in monthly)
    total = sum(r["numeric_enrollment"] for r in monthly)
    match = sum(r["exact_payment_match_numeric_enrollment"] for r in monthly)
    reconstructed = sum(r["reconstructed_numeric_enrollment"] for r in monthly)
    unusable = sum(
        r["exact_match_unusable_payment_input_enrollment"] for r in monthly
    )
    no = sum(r["no_exact_payment_match_numeric_enrollment"] for r in monthly)
    nonnum = sum(r["nonnumeric_rows"] for r in monthly)
    numeric_diff = sum(r["numeric_accounting_difference"] for r in monthly)
    recon_diff = sum(
        r["exact_match_reconstruction_difference"] for r in monthly
    )

    total_direct = sum(
        (Decimal(r["direct_subsidy_ra_dollars"]) for r in monthly if r["source_status"] == "PARSED"),
        Decimal("0"),
    )
    total_reins = sum(
        (Decimal(r["reinsurance_dollars"]) for r in monthly if r["source_status"] == "PARSED"),
        Decimal("0"),
    )
    total_lics = sum(
        (Decimal(r["lics_dollars"]) for r in monthly if r["source_status"] == "PARSED"),
        Decimal("0"),
    )
    total_partd = sum(
        (Decimal(r["total_partd_dollars"]) for r in monthly if r["source_status"] == "PARSED"),
        Decimal("0"),
    )
    dollar_diff = total_partd - total_direct - total_reins - total_lics

    annual = []
    for y in range(2015, 2025):
        rr = [r for r in monthly if r["year"] == y]
        annual.append(
            {
                "year": y,
                "coverage": "JAN-SEP" if y == 2024 else "JAN-DEC",
                "numeric_enrollment": sum(r["numeric_enrollment"] for r in rr),
                "exact_payment_match_numeric_enrollment": sum(
                    r["exact_payment_match_numeric_enrollment"] for r in rr
                ),
                "reconstructed_numeric_enrollment": sum(
                    r["reconstructed_numeric_enrollment"] for r in rr
                ),
                "exact_match_unusable_payment_input_enrollment": sum(
                    r["exact_match_unusable_payment_input_enrollment"] for r in rr
                ),
                "no_exact_payment_match_numeric_enrollment": sum(
                    r["no_exact_payment_match_numeric_enrollment"] for r in rr
                ),
                "direct_subsidy_ra_dollars": str(
                    sum(
                        (Decimal(r["direct_subsidy_ra_dollars"]) for r in rr if r["source_status"] == "PARSED"),
                        Decimal("0"),
                    )
                ),
                "reinsurance_dollars": str(
                    sum(
                        (Decimal(r["reinsurance_dollars"]) for r in rr if r["source_status"] == "PARSED"),
                        Decimal("0"),
                    )
                ),
                "lics_dollars": str(
                    sum(
                        (Decimal(r["lics_dollars"]) for r in rr if r["source_status"] == "PARSED"),
                        Decimal("0"),
                    )
                ),
                "total_partd_dollars": str(
                    sum(
                        (Decimal(r["total_partd_dollars"]) for r in rr if r["source_status"] == "PARSED"),
                        Decimal("0"),
                    )
                ),
            }
        )

    checks = [
        (
            "All 117 required CPSC months resolved and parsed",
            parsed == EXPECTED_MONTHS,
            f"{parsed}/{EXPECTED_MONTHS}",
        ),
        (
            "All 10 PartDPlans years resolved",
            len(pmanifest) == 10,
            f"{len(pmanifest)}/10",
        ),
        (
            "Certified WA numeric CPSC enrollment regression control",
            total == EXPECTED_WA_NUMERIC,
            f"{total:,} == {EXPECTED_WA_NUMERIC:,}",
        ),
        (
            "Certified exact PartDPlans match regression control",
            match == EXPECTED_EXACT_MATCH,
            f"{match:,} == {EXPECTED_EXACT_MATCH:,}",
        ),
        (
            "Certified no-exact-payment-match regression control",
            no == EXPECTED_NO_EXACT,
            f"{no:,} == {EXPECTED_NO_EXACT:,}",
        ),
        (
            "All WA exact-match numeric enrollment reconstructed",
            reconstructed == EXPECTED_EXACT_MATCH and unusable == 0,
            f"reconstructed={reconstructed:,}; unusable={unusable:,}",
        ),
        (
            "Numeric enrollment accounting identity",
            numeric_diff == 0,
            f"difference={numeric_diff:,}",
        ),
        (
            "Exact-match reconstruction accounting identity",
            recon_diff == 0,
            f"difference={recon_diff:,}",
        ),
        (
            "Dollar component identity",
            dollar_diff == 0,
            f"difference={dollar_diff}",
        ),
        (
            "No duplicate annual Part D payment keys",
            not dups,
            f"duplicates={len(dups)}",
        ),
        (
            "No processing/source errors",
            not errors,
            f"errors={len(errors)}",
        ),
    ]

    overall = all(ok for _, ok, _ in checks)

    write_csv(
        a.output_dir / "cms_partd_payment_reconstruction_monthly_v0_2.csv",
        monthly,
        [
            "year", "month", "source_status", "wa_rows", "numeric_enrollment",
            "exact_payment_match_numeric_enrollment",
            "reconstructed_numeric_enrollment",
            "exact_match_unusable_payment_input_enrollment",
            "no_exact_payment_match_numeric_enrollment", "nonnumeric_rows",
            "direct_subsidy_ra_dollars", "reinsurance_dollars", "lics_dollars",
            "total_partd_dollars", "numeric_accounting_difference",
            "exact_match_reconstruction_difference", "dollar_identity_difference",
        ],
    )
    write_csv(
        a.output_dir / "cms_partd_payment_reconstruction_annual_v0_2.csv",
        annual,
        [
            "year", "coverage", "numeric_enrollment",
            "exact_payment_match_numeric_enrollment",
            "reconstructed_numeric_enrollment",
            "exact_match_unusable_payment_input_enrollment",
            "no_exact_payment_match_numeric_enrollment",
            "direct_subsidy_ra_dollars", "reinsurance_dollars",
            "lics_dollars", "total_partd_dollars",
        ],
    )
    write_csv(
        a.output_dir / "cms_partd_payment_reconstruction_detail_v0_2.csv",
        detail,
        [
            "year", "month", "contract", "pbp", "county",
            "cpsc_plan_type", "cpsc_organization_type",
            "payment_organization_type", "numeric_enrollment",
            "direct_subsidy_std_pmpm", "risk_score",
            "direct_subsidy_ra_pmpm", "reinsurance_pmpm", "lics_pmpm",
            "total_partd_pmpm", "direct_subsidy_ra_dollars",
            "reinsurance_dollars", "lics_dollars", "total_partd_dollars",
        ],
    )
    write_csv(
        a.output_dir / "cms_partd_payment_unusable_wa_exposure_v0_2.csv",
        [
            {
                "year": y,
                "contract": c,
                "pbp": p,
                "organization_type": org,
                "numeric_enrollment": n,
            }
            for (y, c, p, org), n in sorted(wa_bad_exposure.items())
        ],
        ["year", "contract", "pbp", "organization_type", "numeric_enrollment"],
    )
    write_csv(
        a.output_dir / "cms_partd_payment_national_bad_inputs_v0_2.csv",
        bad_payment_inputs,
        ["year", "contract", "pbp", "organization_type", "message"],
    )
    write_csv(
        a.output_dir / "cms_partd_cpsc_source_manifest_v0_2.csv",
        manifest,
        [
            "year", "month", "selected_path", "selected_member",
            "selected_source_sha256", "date_inference_method",
            "candidate_count", "distinct_candidate_hashes",
        ],
    )
    write_csv(
        a.output_dir / "cms_partd_payment_source_manifest_v0_2.csv",
        pmanifest,
        [
            "payment_year", "payment_key_count", "zip_path", "zip_sha256",
            "member_name", "member_sha256", "sheet_name", "header_row",
        ],
    )
    write_csv(
        a.output_dir / "cms_partd_payment_reconstruction_checks_v0_2.csv",
        [
            {"check": n, "status": "PASS" if ok else "FAIL", "detail": d}
            for n, ok, d in checks
        ],
        ["check", "status", "detail"],
    )
    write_csv(
        a.output_dir / "cms_partd_payment_reconstruction_errors_v0_2.csv",
        [{"message": e} for e in errors],
        ["message"],
    )

    report = [
        "CMS MEDICARE PART D PAYMENT RECONSTRUCTION v0.2",
        "=" * 78,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        "Scope: Washington, Jan 2015-Sep 2024 (117 months)",
        "",
        "BOUNDARY",
        "-" * 78,
        "Exact same-year Contract+PBP payment matches only.",
        "No blank-to-zero conversion or imputation.",
        "No sibling-PBP propagation or adjacent-year carry-forward.",
        "No database writes or ledger inclusion.",
        "PartDReconciliation is CONTROL_ONLY; RDS remains separate/out of scope.",
        "",
        "SOURCE / ENROLLMENT ACCOUNTING",
        "-" * 78,
        f"Candidate files scanned once: {scanned:,}",
        f"Required CPSC months parsed: {parsed}/{EXPECTED_MONTHS}",
        f"PartDPlans years resolved: {len(pmanifest)}/10",
        f"WA numeric enrollment: {total:,}",
        f"Exact PartDPlans match: {match:,}",
        f"Reconstructed exact-match enrollment: {reconstructed:,}",
        f"Exact-match unusable payment-input enrollment: {unusable:,}",
        f"No exact PartDPlans match: {no:,}",
        f"Suppressed/non-numeric rows: {nonnum:,}",
        f"Numeric accounting difference: {numeric_diff:,}",
        f"Exact-match reconstruction difference: {recon_diff:,}",
        "",
        "PAYMENT RECONSTRUCTION",
        "-" * 78,
        f"Risk-adjusted Direct Subsidy dollars: ${money(total_direct)}",
        f"Reinsurance dollars:                 ${money(total_reins)}",
        f"LICS dollars:                        ${money(total_lics)}",
        f"Total reconstructed Part D dollars:  ${money(total_partd)}",
        f"Dollar component identity difference: {dollar_diff}",
        "",
        "ANNUAL TOTALS",
        "-" * 78,
    ]

    for r in annual:
        report.append(
            f"CY{r['year']} {r['coverage']}: "
            f"enrollment={r['reconstructed_numeric_enrollment']:,} | "
            f"Direct Subsidy RA=${money(Decimal(r['direct_subsidy_ra_dollars']))} | "
            f"Reinsurance=${money(Decimal(r['reinsurance_dollars']))} | "
            f"LICS=${money(Decimal(r['lics_dollars']))} | "
            f"Total=${money(Decimal(r['total_partd_dollars']))}"
        )

    report += ["", "CHECKS", "-" * 78]
    for n, ok, d in checks:
        report.append(f"[{'PASS' if ok else 'FAIL'}] {n} — {d}")

    report += [
        "",
        "NATIONAL MISSING/NONNUMERIC PAYMENT INPUTS",
        "-" * 78,
        f"PartDPlans records with unusable required inputs nationally: "
        f"{len(bad_payment_inputs):,}",
        "These do not invalidate reconstruction unless they intersect WA exact-match "
        "numeric enrollment.",
        f"WA exact-match numeric enrollment exposed to unusable inputs: {unusable:,}",
        "",
        "WARNINGS",
        "-" * 78,
    ]
    report += warnings or ["None"]
    report += ["", "ERRORS", "-" * 78]
    report += errors or ["None"]
    report += [
        "",
        f"OVERALL: {'PASS' if overall else 'FAIL'}",
        "",
        "NOTE:",
        "This reconstruction does not assign dollars to no-exact-payment-match "
        "enrollment. Separately certified affirmative Part D / source-payment-"
        "unavailable exceptions remain NULL/unavailable, not zero.",
    ]

    txt = "\n".join(report) + "\n"
    (a.output_dir / "cms_partd_payment_reconstruction_report_v0_2.txt").write_text(
        txt, encoding="utf-8"
    )

    print()
    print(txt)
    print(f"Outputs: {a.output_dir}")
    return 0 if overall else 1


if __name__ == "__main__":
    sys.exit(main())
