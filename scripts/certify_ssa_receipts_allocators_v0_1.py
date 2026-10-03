from __future__ import annotations

import csv
import re
import sys
from decimal import Decimal, InvalidOperation, getcontext
from pathlib import Path

from openpyxl import load_workbook


getcontext().prec = 28


REPO_ROOT = Path(__file__).resolve().parents[1]

SOURCE_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "ssa"
    / "earnings_employment"
    / "source"
)

VALIDATION_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "ssa"
    / "earnings_employment"
    / "validation"
    / "allocator_certification_v0_1"
)

EXTRACT_CSV = (
    VALIDATION_DIR
    / "washington_ssa_receipts_allocator_extract.csv"
)

REPORT_TXT = (
    VALIDATION_DIR
    / "certification_report.txt"
)

START_YEAR = 2015
END_YEAR = 2023

TARGETS = {
    "OASDI": {
        "sheet": "Table 1",
        "expected_title_fragment": (
            "Social Security (OASDI) taxable earnings"
        ),
        "expected_contribution_fragment": (
            "OASDI contributions"
        ),
    },
    "MEDICARE_HI": {
        "sheet": "Table 4",
        "expected_title_fragment": (
            "Medicare Part A (HI) taxable earnings"
        ),
        "expected_contribution_fragment": (
            "HI contributions"
        ),
    },
}

# Source structure established by ETL-172E-4:
#
# C8  = total taxable earnings
# C9  = wage-and-salary taxable earnings
# C10 = self-employed taxable earnings
#
# C11 = total contributions
# C12 = wage-and-salary contributions
# C13 = self-employed contributions
#
EARNINGS_TOTAL_COLUMN = 8
EARNINGS_WAGE_COLUMN = 9
EARNINGS_SELF_EMPLOYED_COLUMN = 10

CONTRIBUTIONS_TOTAL_COLUMN = 11
CONTRIBUTIONS_WAGE_COLUMN = 12
CONTRIBUTIONS_SELF_EMPLOYED_COLUMN = 13

MAX_SCAN_ROWS = 250

# SSA publishes component values at finite precision.
# We test arithmetic in the source's published units.
#
# This is intentionally narrow and is NOT a general accounting tolerance.
PUBLICATION_PRECISION_TOLERANCE_SOURCE_UNITS = Decimal("2")


def clean(value) -> str:
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\r", " ")
    text = text.replace("\n", " ")

    return " ".join(text.split())


def as_decimal(value, *, context: str) -> Decimal:
    if value is None:
        raise RuntimeError(
            f"{context}: expected numeric value, found blank"
        )

    if isinstance(value, bool):
        raise RuntimeError(
            f"{context}: Boolean value is not valid numeric data"
        )

    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise RuntimeError(
            f"{context}: cannot convert {value!r} to Decimal"
        ) from exc


def detect_unit_multiplier(header_text: str) -> tuple[str, Decimal]:
    normalized = header_text.lower()

    if "thousands of dollars" in normalized:
        return "THOUSANDS_OF_DOLLARS", Decimal("1000")

    if "millions of dollars" in normalized:
        return "MILLIONS_OF_DOLLARS", Decimal("1000000")

    raise RuntimeError(
        "Unable to identify monetary unit from header: "
        f"{header_text!r}"
    )


def find_row(
    ws,
    target: str,
    *,
    columns: tuple[int, ...] = (1, 2, 3, 4),
) -> int:
    target_normalized = target.strip().lower()

    for row_number in range(1, MAX_SCAN_ROWS + 1):
        for column_number in columns:
            value = clean(
                ws.cell(
                    row=row_number,
                    column=column_number,
                ).value
            )

            if value.lower() == target_normalized:
                return row_number

    raise RuntimeError(
        f"{ws.title}: could not find row label {target!r}"
    )


def get_header_text(ws) -> str:
    parts = []

    for row_number in range(1, 4):
        for column_number in range(1, 14):
            value = clean(
                ws.cell(
                    row=row_number,
                    column=column_number,
                ).value
            )

            if value:
                parts.append(value)

    return " | ".join(parts)


def arithmetic_status(
    total: Decimal,
    wage: Decimal,
    self_employed: Decimal,
) -> tuple[Decimal, str]:
    residual = total - wage - self_employed

    if residual == 0:
        return residual, "EXACT_PASS"

    if (
        abs(residual)
        <= PUBLICATION_PRECISION_TOLERANCE_SOURCE_UNITS
    ):
        return residual, "PUBLICATION_PRECISION_PASS"

    return residual, "FAIL"


def format_decimal(value: Decimal) -> str:
    return format(value, "f")


def main() -> int:
    VALIDATION_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    records: list[dict[str, str | int]] = []
    report_lines: list[str] = []
    failures: list[str] = []

    report_lines.append("=" * 110)
    report_lines.append(
        "SSA RECEIPTS ALLOCATOR CERTIFICATION v0.1 "
        "- ETL-172E-5"
    )
    report_lines.append("=" * 110)
    report_lines.append("")
    report_lines.append(
        "Purpose: certify Washington and All-areas SSA "
        "OASDI and Medicare HI allocator inputs."
    )
    report_lines.append("")
    report_lines.append(
        "No federal Budget receipt dollars are allocated."
    )
    report_lines.append(
        "No Rockefeller replication is performed."
    )
    report_lines.append(
        "No FY2024 extrapolation or carry-forward is performed."
    )
    report_lines.append(
        "No database writes are performed."
    )
    report_lines.append("")

    print("=" * 100)
    print(
        "SSA RECEIPTS ALLOCATOR CERTIFICATION "
        "- ETL-172E-5"
    )
    print("=" * 100)
    print()

    for year in range(START_YEAR, END_YEAR + 1):
        yy = str(year)[-2:]
        path = SOURCE_DIR / f"eedata_sc{yy}.xlsx"

        if not path.exists():
            message = f"{year}: missing source {path}"
            failures.append(message)
            print(f"{year}: FAIL | source missing")
            continue

        wb = load_workbook(
            path,
            read_only=False,
            data_only=True,
        )

        normalized_sheet_lookup = {
            name.strip(): name
            for name in wb.sheetnames
        }

        for program, spec in TARGETS.items():
            sheet_name = spec["sheet"]

            actual_sheet = normalized_sheet_lookup.get(
                sheet_name
            )

            if actual_sheet is None:
                message = (
                    f"{year} {program}: "
                    f"{sheet_name} not found"
                )
                failures.append(message)
                continue

            ws = wb[actual_sheet]

            header_text = get_header_text(ws)

            if (
                spec["expected_title_fragment"].lower()
                not in header_text.lower()
            ):
                message = (
                    f"{year} {program}: expected title "
                    "semantics not found"
                )
                failures.append(message)
                continue

            if (
                spec["expected_contribution_fragment"].lower()
                not in header_text.lower()
            ):
                message = (
                    f"{year} {program}: expected contribution "
                    "semantics not found"
                )
                failures.append(message)
                continue

            try:
                unit_name, unit_multiplier = (
                    detect_unit_multiplier(header_text)
                )

                all_areas_row = find_row(
                    ws,
                    "All areas",
                )

                washington_row = find_row(
                    ws,
                    "Washington",
                )

            except RuntimeError as exc:
                failures.append(
                    f"{year} {program}: {exc}"
                )
                continue

            def value(
                row_number: int,
                column_number: int,
                label: str,
            ) -> Decimal:
                return as_decimal(
                    ws.cell(
                        row=row_number,
                        column=column_number,
                    ).value,
                    context=(
                        f"{year} {program} "
                        f"{label}"
                    ),
                )

            all_earn_total = value(
                all_areas_row,
                EARNINGS_TOTAL_COLUMN,
                "All areas earnings total",
            )

            all_earn_wage = value(
                all_areas_row,
                EARNINGS_WAGE_COLUMN,
                "All areas earnings wage",
            )

            all_earn_self = value(
                all_areas_row,
                EARNINGS_SELF_EMPLOYED_COLUMN,
                "All areas earnings self-employed",
            )

            wa_earn_total = value(
                washington_row,
                EARNINGS_TOTAL_COLUMN,
                "Washington earnings total",
            )

            wa_earn_wage = value(
                washington_row,
                EARNINGS_WAGE_COLUMN,
                "Washington earnings wage",
            )

            wa_earn_self = value(
                washington_row,
                EARNINGS_SELF_EMPLOYED_COLUMN,
                "Washington earnings self-employed",
            )

            all_contrib_total = value(
                all_areas_row,
                CONTRIBUTIONS_TOTAL_COLUMN,
                "All areas contributions total",
            )

            all_contrib_wage = value(
                all_areas_row,
                CONTRIBUTIONS_WAGE_COLUMN,
                "All areas contributions wage",
            )

            all_contrib_self = value(
                all_areas_row,
                CONTRIBUTIONS_SELF_EMPLOYED_COLUMN,
                "All areas contributions self-employed",
            )

            wa_contrib_total = value(
                washington_row,
                CONTRIBUTIONS_TOTAL_COLUMN,
                "Washington contributions total",
            )

            wa_contrib_wage = value(
                washington_row,
                CONTRIBUTIONS_WAGE_COLUMN,
                "Washington contributions wage",
            )

            wa_contrib_self = value(
                washington_row,
                CONTRIBUTIONS_SELF_EMPLOYED_COLUMN,
                "Washington contributions self-employed",
            )

            all_earn_residual, all_earn_status = (
                arithmetic_status(
                    all_earn_total,
                    all_earn_wage,
                    all_earn_self,
                )
            )

            wa_earn_residual, wa_earn_status = (
                arithmetic_status(
                    wa_earn_total,
                    wa_earn_wage,
                    wa_earn_self,
                )
            )

            all_contrib_residual, all_contrib_status = (
                arithmetic_status(
                    all_contrib_total,
                    all_contrib_wage,
                    all_contrib_self,
                )
            )

            wa_contrib_residual, wa_contrib_status = (
                arithmetic_status(
                    wa_contrib_total,
                    wa_contrib_wage,
                    wa_contrib_self,
                )
            )

            arithmetic_statuses = (
                all_earn_status,
                wa_earn_status,
                all_contrib_status,
                wa_contrib_status,
            )

            arithmetic_pass = all(
                status != "FAIL"
                for status in arithmetic_statuses
            )

            if all_earn_total == 0:
                raise RuntimeError(
                    f"{year} {program}: "
                    "All-areas earnings denominator is zero"
                )

            if all_contrib_total == 0:
                raise RuntimeError(
                    f"{year} {program}: "
                    "All-areas contributions denominator is zero"
                )

            earnings_share = (
                wa_earn_total
                / all_earn_total
            )

            contribution_share = (
                wa_contrib_total
                / all_contrib_total
            )

            share_difference = (
                contribution_share
                - earnings_share
            )

            all_earn_dollars = (
                all_earn_total
                * unit_multiplier
            )

            wa_earn_dollars = (
                wa_earn_total
                * unit_multiplier
            )

            all_contrib_dollars = (
                all_contrib_total
                * unit_multiplier
            )

            wa_contrib_dollars = (
                wa_contrib_total
                * unit_multiplier
            )

            source_limitation = ""

            if (
                year == 2015
                and program == "MEDICARE_HI"
            ):
                source_limitation = (
                    "SOURCE_KNOWN_LIMITATION: "
                    "SSA documented a methodological flaw "
                    "affecting 2015 state-level Medicare HI "
                    "self-employed estimates and total columns."
                )

            record_status = (
                "PASS"
                if arithmetic_pass
                else "FAIL"
            )

            if not arithmetic_pass:
                failures.append(
                    f"{year} {program}: "
                    "component arithmetic failed"
                )

            records.append(
                {
                    "data_year": year,
                    "program": program,
                    "source_file": str(
                        path.relative_to(REPO_ROOT)
                    ),
                    "source_sheet": actual_sheet,
                    "all_areas_row": all_areas_row,
                    "washington_row": washington_row,
                    "source_unit": unit_name,
                    "unit_multiplier_to_dollars": (
                        format_decimal(unit_multiplier)
                    ),
                    "all_areas_taxable_earnings_source_units": (
                        format_decimal(all_earn_total)
                    ),
                    "washington_taxable_earnings_source_units": (
                        format_decimal(wa_earn_total)
                    ),
                    "all_areas_taxable_earnings_dollars": (
                        format_decimal(all_earn_dollars)
                    ),
                    "washington_taxable_earnings_dollars": (
                        format_decimal(wa_earn_dollars)
                    ),
                    "earnings_share": (
                        format_decimal(earnings_share)
                    ),
                    "all_areas_contributions_source_units": (
                        format_decimal(all_contrib_total)
                    ),
                    "washington_contributions_source_units": (
                        format_decimal(wa_contrib_total)
                    ),
                    "all_areas_contributions_dollars": (
                        format_decimal(all_contrib_dollars)
                    ),
                    "washington_contributions_dollars": (
                        format_decimal(wa_contrib_dollars)
                    ),
                    "contribution_share": (
                        format_decimal(contribution_share)
                    ),
                    "contribution_minus_earnings_share": (
                        format_decimal(share_difference)
                    ),
                    "all_areas_earnings_residual_source_units": (
                        format_decimal(all_earn_residual)
                    ),
                    "all_areas_earnings_arithmetic": (
                        all_earn_status
                    ),
                    "washington_earnings_residual_source_units": (
                        format_decimal(wa_earn_residual)
                    ),
                    "washington_earnings_arithmetic": (
                        wa_earn_status
                    ),
                    "all_areas_contributions_residual_source_units": (
                        format_decimal(all_contrib_residual)
                    ),
                    "all_areas_contributions_arithmetic": (
                        all_contrib_status
                    ),
                    "washington_contributions_residual_source_units": (
                        format_decimal(wa_contrib_residual)
                    ),
                    "washington_contributions_arithmetic": (
                        wa_contrib_status
                    ),
                    "source_limitation": source_limitation,
                    "status": record_status,
                }
            )

            report_lines.append(
                f"{year} {program}"
            )
            report_lines.append(
                f"  source unit          : {unit_name}"
            )
            report_lines.append(
                f"  All areas row        : {all_areas_row}"
            )
            report_lines.append(
                f"  Washington row       : {washington_row}"
            )
            report_lines.append(
                "  earnings share       : "
                f"{format_decimal(earnings_share)}"
            )
            report_lines.append(
                "  contribution share   : "
                f"{format_decimal(contribution_share)}"
            )
            report_lines.append(
                "  contribution - earn. : "
                f"{format_decimal(share_difference)}"
            )
            report_lines.append(
                "  arithmetic           : "
                f"{record_status}"
            )

            if source_limitation:
                report_lines.append(
                    f"  limitation           : "
                    f"{source_limitation}"
                )

            report_lines.append("")

            print(
                f"{year} {program:<11} | "
                f"{unit_name:<20} | "
                f"earn={float(earnings_share):.8%} | "
                f"contrib={float(contribution_share):.8%} | "
                f"{record_status}"
            )

        wb.close()

    fieldnames = [
        "data_year",
        "program",
        "source_file",
        "source_sheet",
        "all_areas_row",
        "washington_row",
        "source_unit",
        "unit_multiplier_to_dollars",
        "all_areas_taxable_earnings_source_units",
        "washington_taxable_earnings_source_units",
        "all_areas_taxable_earnings_dollars",
        "washington_taxable_earnings_dollars",
        "earnings_share",
        "all_areas_contributions_source_units",
        "washington_contributions_source_units",
        "all_areas_contributions_dollars",
        "washington_contributions_dollars",
        "contribution_share",
        "contribution_minus_earnings_share",
        "all_areas_earnings_residual_source_units",
        "all_areas_earnings_arithmetic",
        "washington_earnings_residual_source_units",
        "washington_earnings_arithmetic",
        "all_areas_contributions_residual_source_units",
        "all_areas_contributions_arithmetic",
        "washington_contributions_residual_source_units",
        "washington_contributions_arithmetic",
        "source_limitation",
        "status",
    ]

    with EXTRACT_CSV.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=fieldnames,
        )
        writer.writeheader()
        writer.writerows(records)

    expected_records = (
        (END_YEAR - START_YEAR + 1)
        * len(TARGETS)
    )

    report_lines.append("=" * 110)
    report_lines.append("CERTIFICATION SUMMARY")
    report_lines.append("=" * 110)
    report_lines.append(
        f"Expected records : {expected_records}"
    )
    report_lines.append(
        f"Extracted records: {len(records)}"
    )
    report_lines.append(
        f"Failures         : {len(failures)}"
    )
    report_lines.append("")

    if failures:
        report_lines.append("FAILURES")
        report_lines.append("-" * 110)

        for failure in failures:
            report_lines.append(
                f"- {failure}"
            )

        report_lines.append("")

    final_result = (
        "PASS"
        if (
            len(records) == expected_records
            and not failures
        )
        else "FAIL"
    )

    report_lines.append(
        f"FINAL RESULT: {final_result}"
    )
    report_lines.append("")
    report_lines.append(
        "PASS certifies source semantics, units, "
        "Washington/All-areas extraction, component "
        "arithmetic, and candidate allocation shares only."
    )
    report_lines.append(
        "PASS does NOT select the Rockefeller allocator, "
        "allocate federal Budget receipts, establish "
        "fiscal-year mapping, carry an allocator into 2024, "
        "or authorize ledger inclusion."
    )

    REPORT_TXT.write_text(
        "\n".join(report_lines) + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 100)
    print("CERTIFICATION SUMMARY")
    print("=" * 100)
    print(f"Expected records : {expected_records}")
    print(f"Extracted records: {len(records)}")
    print(f"Failures         : {len(failures)}")
    print(f"Extract CSV      : {EXTRACT_CSV.relative_to(REPO_ROOT)}")
    print(f"Report           : {REPORT_TXT.relative_to(REPO_ROOT)}")
    print()
    print(f"ETL-172E-5 RESULT: {final_result}")
    print()

    if final_result == "PASS":
        print(
            "PASS certifies source semantics, units, "
            "Washington/All-areas extraction, component "
            "arithmetic, and candidate allocation shares only."
        )
        print(
            "PASS does NOT select the Rockefeller allocator, "
            "allocate federal Budget receipts, establish "
            "fiscal-year mapping, carry an allocator into 2024, "
            "or authorize ledger inclusion."
        )
        return 0

    return 1


if __name__ == "__main__":
    sys.exit(main())