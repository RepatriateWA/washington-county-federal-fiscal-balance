"""
ETL-172J-1 - Washington IRS Table 5 Remaining Receipts Allocation v0.1
======================================================================

Purpose
-------
Allocate two remaining federal receipt categories to Washington for
FFY2015-FFY2024 using the Rockefeller-documented IRS Gross Collections
Table 5 geographic allocation method:

1. Estate and Gift Taxes
2. Railroad Retirement

Architecture
------------
IRS Table 5 establishes the Washington geographic share:

    WA share = WA IRS gross collections / U.S. IRS gross collections

OMB Historical Tables establish the national federal-receipts control:

    WA allocated receipts = OMB national control * WA IRS share

This script deliberately does NOT treat Washington IRS gross collections
as the final economic-incidence receipt amount.

IRS source semantics
--------------------
ETL-171C certified:

    -7  Railroad retirement tax
    -9  Estate tax
    -10 Gift tax

The same certified parser is reused here.

OMB controls
------------
Estate and Gift Taxes:
    OMB Historical Table 2.5

Railroad Retirement:
    OMB Historical Table 2.4
    rows 12 + 13 + 14

IMPORTANT:
OMB Table 2.4 row 15 MUST NOT be used as the Railroad Retirement control.
ETL-172I-5 established that row 15 is a broader subtotal containing
OASI, DI, and HI in addition to Railroad Retirement.

Boundaries
----------
No database writes.
No refund adjustment.
No forced reconciliation to IRS gross collections.
No final ledger authorization.
No allocation of any other receipt category.
"""

from __future__ import annotations

import csv
import hashlib
import json
import sys
from decimal import Decimal, getcontext
from pathlib import Path
from typing import Any

import openpyxl

from scripts.certify_irs_gross_collections_semantics_v0_1 import (
    locate_column_numbers,
    normalize_text,
    read_source,
    to_decimal,
)


getcontext().prec = 50


# =============================================================================
# PATHS
# =============================================================================

REPO_ROOT = Path(__file__).resolve().parents[1]

IRS_SEMANTIC_EXTRACT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "irs"
    / "gross_collections"
    / "validation"
    / "semantic_certification_v0_1"
    / "washington_semantic_extract.csv"
)

IRS_SOURCE_DIR = (
    REPO_ROOT
    / "data"
    / "raw"
    / "irs"
    / "gross_collections"
    / "source"
)

OMB_ROOT = (
    REPO_ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
)

OMB_TABLE_24 = (
    OMB_ROOT
    / "source"
    / "BUDGET-2027-TAB-3-4.xlsx"
)

I4_OTHER_CONTROLS = (
    OMB_ROOT
    / "validation"
    / "other_receipts_decomposition_v0_1"
    / "national_other_receipts_annual_controls.csv"
)

I5_WORK_QUEUE = (
    OMB_ROOT
    / "validation"
    / "remaining_receipts_methodology_crosswalk_v0_1"
    / "remaining_receipts_work_queue.csv"
)

OUTPUT_DIR = (
    OMB_ROOT
    / "validation"
    / "washington_irs_table5_remaining_receipts_v0_1"
)

OUTPUT_ALLOCATOR = (
    OUTPUT_DIR
    / "irs_table5_washington_us_allocator_extract.csv"
)

OUTPUT_ALLOCATION = (
    OUTPUT_DIR
    / "washington_irs_table5_receipts_allocation.csv"
)

OUTPUT_ANNUAL = (
    OUTPUT_DIR
    / "washington_irs_table5_receipts_annual_summary.csv"
)

OUTPUT_QUEUE = (
    OUTPUT_DIR
    / "remaining_receipts_work_queue_after_j1.csv"
)

OUTPUT_JSON = (
    OUTPUT_DIR
    / "certification.json"
)

OUTPUT_REPORT = (
    OUTPUT_DIR
    / "certification_report.txt"
)


TARGET_YEARS = list(range(2015, 2025))

THOUSANDS_TO_DOLLARS = Decimal("1000")
MILLIONS_TO_DOLLARS = Decimal("1000000")


# =============================================================================
# GENERAL HELPERS
# =============================================================================


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(
            lambda: handle.read(1024 * 1024),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest().upper()


def decimal_value(
    value: Any,
    *,
    context: str,
) -> Decimal:
    if value is None or str(value).strip() == "":
        raise RuntimeError(
            f"MISSING NUMERIC VALUE: {context}"
        )

    try:
        result = Decimal(str(value))
    except Exception as exc:
        raise RuntimeError(
            f"NUMERIC PARSE FAILURE: {context}: {value!r}"
        ) from exc

    if not result.is_finite():
        raise RuntimeError(
            f"NON-FINITE NUMERIC VALUE: {context}: {value!r}"
        )

    return result


def decimal_text(value: Decimal) -> str:
    return format(value, "f")


def load_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise RuntimeError(
            "INPUT FILE NOT FOUND\n"
            f"{path.relative_to(REPO_ROOT)}"
        )

    with path.open(
        "r",
        newline="",
        encoding="utf-8",
    ) as handle:
        return list(csv.DictReader(handle))


def write_csv(
    path: Path,
    rows: list[dict[str, Any]],
) -> None:
    if not rows:
        raise RuntimeError(
            f"REFUSING TO WRITE EMPTY CSV: {path.name}"
        )

    with path.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(rows[0].keys()),
        )

        writer.writeheader()
        writer.writerows(rows)


# =============================================================================
# IRS TABLE 5
# =============================================================================


def locate_exact_row(
    frame: Any,
    target: str,
) -> int:
    """
    Locate exactly one row containing target after ETL-171 text
    normalization.

    This is used for:
        Washington
        United States, total [4]

    The U.S. label is matched by prefix because footnote numbering may
    vary across IRS Data Book vintages.
    """

    target_normalized = target.lower()
    matches: list[int] = []

    for row_index in range(len(frame)):
        for value in frame.iloc[row_index].tolist():
            text = normalize_text(value).lower()

            if target_normalized == "washington":
                matched = text == "washington"
            else:
                matched = text.startswith(
                    "united states, total"
                )

            if matched:
                matches.append(row_index)
                break

    if len(matches) != 1:
        raise RuntimeError(
            "IRS ROW IDENTIFICATION FAILURE\n"
            f"Target : {target}\n"
            f"Matches: {matches}"
        )

    return matches[0]


def load_irs_source_manifest() -> dict[int, dict[str, str]]:
    rows = load_csv(IRS_SEMANTIC_EXTRACT)

    if len(rows) != 10:
        raise RuntimeError(
            "ETL-171 SEMANTIC EXTRACT ROW COUNT FAILURE\n"
            f"Expected: 10\n"
            f"Actual  : {len(rows)}"
        )

    output: dict[int, dict[str, str]] = {}

    for row in rows:
        year = int(row["fiscal_year"])

        if year not in TARGET_YEARS:
            continue

        if year in output:
            raise RuntimeError(
                f"DUPLICATE ETL-171 FFY: {year}"
            )

        output[year] = row

    if set(output) != set(TARGET_YEARS):
        raise RuntimeError(
            "ETL-171 YEAR COVERAGE FAILURE\n"
            f"Expected: {TARGET_YEARS}\n"
            f"Actual  : {sorted(output)}"
        )

    return output


def extract_irs_allocators(
    manifest: dict[int, dict[str, str]],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []

    for year in TARGET_YEARS:
        manifest_row = manifest[year]

        source_file = manifest_row["source_file"]
        source_path = IRS_SOURCE_DIR / source_file

        if not source_path.exists():
            raise RuntimeError(
                "IRS SOURCE FILE NOT FOUND\n"
                f"FFY{year}: {source_path}"
            )

        actual_hash = sha256_file(source_path)
        certified_hash = (
            manifest_row["source_sha256"]
            .strip()
            .upper()
        )

        if actual_hash != certified_hash:
            raise RuntimeError(
                "IRS SOURCE HASH FAILURE\n"
                f"FFY{year}\n"
                f"Certified: {certified_hash}\n"
                f"Observed : {actual_hash}"
            )

        frame = read_source(source_path)

        column_map = locate_column_numbers(frame)

        washington_row = locate_exact_row(
            frame,
            "Washington",
        )

        us_row = locate_exact_row(
            frame,
            "United States, total",
        )

        rr_col = column_map["-7"]
        estate_col = column_map["-9"]
        gift_col = column_map["-10"]

        wa_rr_thousands = to_decimal(
            frame.iloc[
                washington_row,
                rr_col,
            ]
        )

        us_rr_thousands = to_decimal(
            frame.iloc[
                us_row,
                rr_col,
            ]
        )

        wa_estate_thousands = to_decimal(
            frame.iloc[
                washington_row,
                estate_col,
            ]
        )

        us_estate_thousands = to_decimal(
            frame.iloc[
                us_row,
                estate_col,
            ]
        )

        wa_gift_thousands = to_decimal(
            frame.iloc[
                washington_row,
                gift_col,
            ]
        )

        us_gift_thousands = to_decimal(
            frame.iloc[
                us_row,
                gift_col,
            ]
        )

        wa_eg_thousands = (
            wa_estate_thousands
            + wa_gift_thousands
        )

        us_eg_thousands = (
            us_estate_thousands
            + us_gift_thousands
        )

        if us_rr_thousands == 0:
            raise RuntimeError(
                f"ZERO U.S. RAILROAD DENOMINATOR: FFY{year}"
            )

        if us_eg_thousands == 0:
            raise RuntimeError(
                f"ZERO U.S. ESTATE/GIFT DENOMINATOR: FFY{year}"
            )

        rr_share = (
            wa_rr_thousands
            / us_rr_thousands
        )

        eg_share = (
            wa_eg_thousands
            / us_eg_thousands
        )

        # Reconcile Washington values back to the already-certified
        # ETL-171 semantic extract.
        certified_wa_rr = decimal_value(
            manifest_row[
                "railroad_retirement_tax_thousands"
            ],
            context=(
                f"ETL-171 FFY{year} WA railroad"
            ),
        )

        certified_wa_estate = decimal_value(
            manifest_row["estate_tax_thousands"],
            context=(
                f"ETL-171 FFY{year} WA estate"
            ),
        )

        certified_wa_gift = decimal_value(
            manifest_row["gift_tax_thousands"],
            context=(
                f"ETL-171 FFY{year} WA gift"
            ),
        )

        if wa_rr_thousands != certified_wa_rr:
            raise RuntimeError(
                "WA RAILROAD REPRODUCTION FAILURE\n"
                f"FFY{year}"
            )

        if wa_estate_thousands != certified_wa_estate:
            raise RuntimeError(
                "WA ESTATE REPRODUCTION FAILURE\n"
                f"FFY{year}"
            )

        if wa_gift_thousands != certified_wa_gift:
            raise RuntimeError(
                "WA GIFT REPRODUCTION FAILURE\n"
                f"FFY{year}"
            )

        output.append(
            {
                "fiscal_year": year,
                "source_file": source_file,
                "source_sha256": actual_hash,
                "source_units": (
                    "THOUSANDS_OF_DOLLARS"
                ),

                "wa_railroad_retirement_thousands": (
                    decimal_text(
                        wa_rr_thousands
                    )
                ),
                "us_railroad_retirement_thousands": (
                    decimal_text(
                        us_rr_thousands
                    )
                ),
                "wa_railroad_retirement_share": (
                    decimal_text(rr_share)
                ),

                "wa_estate_tax_thousands": (
                    decimal_text(
                        wa_estate_thousands
                    )
                ),
                "us_estate_tax_thousands": (
                    decimal_text(
                        us_estate_thousands
                    )
                ),

                "wa_gift_tax_thousands": (
                    decimal_text(
                        wa_gift_thousands
                    )
                ),
                "us_gift_tax_thousands": (
                    decimal_text(
                        us_gift_thousands
                    )
                ),

                "wa_estate_gift_combined_thousands": (
                    decimal_text(
                        wa_eg_thousands
                    )
                ),
                "us_estate_gift_combined_thousands": (
                    decimal_text(
                        us_eg_thousands
                    )
                ),
                "wa_estate_gift_share": (
                    decimal_text(eg_share)
                ),

                "washington_reproduction_status": (
                    "PASS"
                ),
            }
        )

    return output


# =============================================================================
# OMB TABLE 2.4 - RAILROAD RETIREMENT NATIONAL CONTROL
# =============================================================================


def parse_year_header(value: Any) -> int | None:
    if value is None:
        return None

    text = str(value).strip().lower()

    if text.endswith(" estimate"):
        text = (
            text
            .replace(" estimate", "")
            .strip()
        )

    if text.isdigit():
        return int(text)

    return None


def locate_year_columns(
    worksheet: Any,
) -> dict[int, int]:
    candidates: list[dict[int, int]] = []

    for row_number in range(
        1,
        min(worksheet.max_row, 10) + 1,
    ):
        mapping: dict[int, int] = {}

        for col in range(
            1,
            worksheet.max_column + 1,
        ):
            year = parse_year_header(
                worksheet.cell(
                    row_number,
                    col,
                ).value
            )

            if year in TARGET_YEARS:
                mapping[year] = col

        if set(mapping) == set(TARGET_YEARS):
            candidates.append(mapping)

    if len(candidates) != 1:
        raise RuntimeError(
            "OMB TABLE 2.4 YEAR HEADER FAILURE\n"
            f"Candidate rows: {len(candidates)}"
        )

    return candidates[0]


def omb_cell_decimal(
    worksheet: Any,
    row: int,
    col: int,
    *,
    context: str,
) -> Decimal:
    raw = worksheet.cell(
        row,
        col,
    ).value

    if raw is None:
        return Decimal("0")

    if isinstance(raw, str):
        cleaned = raw.strip()

        if cleaned in {
            "",
            "*",
            "-*",
            "..........",
            "-",
        }:
            return Decimal("0")

        cleaned = cleaned.replace(",", "")

        return decimal_value(
            cleaned,
            context=context,
        )

    return decimal_value(
        raw,
        context=context,
    )


def load_railroad_controls() -> dict[int, Decimal]:
    if not OMB_TABLE_24.exists():
        raise RuntimeError(
            "OMB TABLE 2.4 NOT FOUND\n"
            f"{OMB_TABLE_24.relative_to(REPO_ROOT)}"
        )

    workbook = openpyxl.load_workbook(
        OMB_TABLE_24,
        data_only=True,
        read_only=False,
    )

    if workbook.sheetnames != ["Table"]:
        raise RuntimeError(
            "OMB TABLE 2.4 SHEET CONTRACT FAILURE\n"
            f"{workbook.sheetnames}"
        )

    worksheet = workbook["Table"]

    # Semantic guardrails established in ETL-172I-2/I-5.
    expected_labels = {
        12: "Federal funds",
        13: "Trust funds",
        14: (
            "Railroad social security equivalent account"
        ),
        15: "Total (1)",
    }

    for row_number, expected in expected_labels.items():
        observed = str(
            worksheet.cell(
                row_number,
                1,
            ).value
        ).strip()

        if observed != expected:
            raise RuntimeError(
                "OMB TABLE 2.4 ROW SEMANTIC FAILURE\n"
                f"Row      : {row_number}\n"
                f"Expected : {expected!r}\n"
                f"Observed : {observed!r}"
            )

    year_columns = locate_year_columns(
        worksheet
    )

    output: dict[int, Decimal] = {}

    for year in TARGET_YEARS:
        col = year_columns[year]

        federal = omb_cell_decimal(
            worksheet,
            12,
            col,
            context=(
                f"FFY{year} railroad Federal funds"
            ),
        )

        trust = omb_cell_decimal(
            worksheet,
            13,
            col,
            context=(
                f"FFY{year} railroad Trust funds"
            ),
        )

        equivalent = omb_cell_decimal(
            worksheet,
            14,
            col,
            context=(
                f"FFY{year} railroad equivalent account"
            ),
        )

        output[year] = (
            federal
            + trust
            + equivalent
        )

    return output


# =============================================================================
# OMB TABLE 2.5 - ESTATE AND GIFT NATIONAL CONTROL
# =============================================================================


def load_estate_gift_controls() -> dict[int, Decimal]:
    rows = load_csv(I4_OTHER_CONTROLS)

    if not rows:
        raise RuntimeError(
            "I-4 OTHER CONTROLS FILE IS EMPTY"
        )

    required = {
        "fiscal_year",
        "estate_and_gift_taxes_millions",
    }

    missing = required - set(rows[0])

    if missing:
        raise RuntimeError(
            "I-4 SCHEMA FAILURE\n"
            f"Missing  : {sorted(missing)}\n"
            f"Available: {list(rows[0].keys())}"
        )

    output: dict[int, Decimal] = {}

    for row in rows:
        year = int(row["fiscal_year"])

        if year not in TARGET_YEARS:
            continue

        if year in output:
            raise RuntimeError(
                f"DUPLICATE I-4 FFY: {year}"
            )

        output[year] = decimal_value(
            row[
                "estate_and_gift_taxes_millions"
            ],
            context=(
                f"I-4 FFY{year} estate/gift"
            ),
        )

    if set(output) != set(TARGET_YEARS):
        raise RuntimeError(
            "I-4 ESTATE/GIFT YEAR COVERAGE FAILURE"
        )

    return output


# =============================================================================
# ALLOCATION
# =============================================================================


def allocate(
    allocator_rows: list[dict[str, Any]],
    railroad_controls: dict[int, Decimal],
    estate_gift_controls: dict[int, Decimal],
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    allocations: list[dict[str, Any]] = []
    annual: list[dict[str, Any]] = []

    allocator_by_year = {
        int(row["fiscal_year"]): row
        for row in allocator_rows
    }

    for year in TARGET_YEARS:
        allocator = allocator_by_year[year]

        rr_share = decimal_value(
            allocator[
                "wa_railroad_retirement_share"
            ],
            context=(
                f"FFY{year} railroad share"
            ),
        )

        eg_share = decimal_value(
            allocator[
                "wa_estate_gift_share"
            ],
            context=(
                f"FFY{year} estate/gift share"
            ),
        )

        rr_control_millions = (
            railroad_controls[year]
        )

        eg_control_millions = (
            estate_gift_controls[year]
        )

        rr_control_dollars = (
            rr_control_millions
            * MILLIONS_TO_DOLLARS
        )

        eg_control_dollars = (
            eg_control_millions
            * MILLIONS_TO_DOLLARS
        )

        wa_rr_dollars = (
            rr_control_dollars
            * rr_share
        )

        wa_eg_dollars = (
            eg_control_dollars
            * eg_share
        )

        allocations.append(
            {
                "fiscal_year": year,
                "receipt_category": (
                    "RAILROAD_RETIREMENT"
                ),
                "national_control_source": (
                    "OMB Historical Table 2.4"
                ),
                "national_control_millions": (
                    decimal_text(
                        rr_control_millions
                    )
                ),
                "national_control_dollars": (
                    decimal_text(
                        rr_control_dollars
                    )
                ),
                "allocator_source": (
                    "IRS Data Book Table 5"
                ),
                "allocator_numerator": (
                    "Washington railroad retirement tax "
                    "gross collections"
                ),
                "allocator_denominator": (
                    "United States total railroad "
                    "retirement tax gross collections"
                ),
                "washington_share": (
                    decimal_text(rr_share)
                ),
                "washington_share_percent": (
                    decimal_text(
                        rr_share * Decimal("100")
                    )
                ),
                "washington_allocated_receipts_dollars": (
                    decimal_text(
                        wa_rr_dollars
                    )
                ),
                "method_status": (
                    "ALLOCATED_DOCUMENTED_METHOD"
                ),
                "ledger_status": (
                    "NOT_YET_AUTHORIZED"
                ),
            }
        )

        allocations.append(
            {
                "fiscal_year": year,
                "receipt_category": (
                    "ESTATE_AND_GIFT_TAXES"
                ),
                "national_control_source": (
                    "OMB Historical Table 2.5"
                ),
                "national_control_millions": (
                    decimal_text(
                        eg_control_millions
                    )
                ),
                "national_control_dollars": (
                    decimal_text(
                        eg_control_dollars
                    )
                ),
                "allocator_source": (
                    "IRS Data Book Table 5"
                ),
                "allocator_numerator": (
                    "Washington estate tax plus gift tax "
                    "gross collections"
                ),
                "allocator_denominator": (
                    "United States total estate tax plus "
                    "gift tax gross collections"
                ),
                "washington_share": (
                    decimal_text(eg_share)
                ),
                "washington_share_percent": (
                    decimal_text(
                        eg_share * Decimal("100")
                    )
                ),
                "washington_allocated_receipts_dollars": (
                    decimal_text(
                        wa_eg_dollars
                    )
                ),
                "method_status": (
                    "ALLOCATED_DOCUMENTED_METHOD"
                ),
                "ledger_status": (
                    "NOT_YET_AUTHORIZED"
                ),
            }
        )

        annual.append(
            {
                "fiscal_year": year,
                "railroad_retirement_national_control_millions": (
                    decimal_text(
                        rr_control_millions
                    )
                ),
                "railroad_retirement_wa_share": (
                    decimal_text(rr_share)
                ),
                "railroad_retirement_wa_receipts_dollars": (
                    decimal_text(
                        wa_rr_dollars
                    )
                ),
                "estate_gift_national_control_millions": (
                    decimal_text(
                        eg_control_millions
                    )
                ),
                "estate_gift_wa_share": (
                    decimal_text(eg_share)
                ),
                "estate_gift_wa_receipts_dollars": (
                    decimal_text(
                        wa_eg_dollars
                    )
                ),
                "j1_national_control_closed_millions": (
                    decimal_text(
                        rr_control_millions
                        + eg_control_millions
                    )
                ),
                "j1_washington_allocated_receipts_dollars": (
                    decimal_text(
                        wa_rr_dollars
                        + wa_eg_dollars
                    )
                ),
            }
        )

    return allocations, annual


# =============================================================================
# I-5 WORK QUEUE REDUCTION
# =============================================================================


def reduce_i5_queue(
    annual_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    i5_rows = load_csv(I5_WORK_QUEUE)

    if len(i5_rows) != 10:
        raise RuntimeError(
            "I-5 WORK QUEUE ROW COUNT FAILURE"
        )

    annual_by_year = {
        int(row["fiscal_year"]): row
        for row in annual_rows
    }

    output: list[dict[str, Any]] = []

    for row in i5_rows:
        year = int(row["fiscal_year"])

        if year not in TARGET_YEARS:
            continue

        j1 = annual_by_year[year]

        before = decimal_value(
            row[
                "actionable_unfinished_millions"
            ],
            context=(
                f"I-5 FFY{year} actionable"
            ),
        )

        documented_before = decimal_value(
            row[
                "allocator_method_documented_millions"
            ],
            context=(
                f"I-5 FFY{year} documented method"
            ),
        )

        j1_closed = decimal_value(
            j1[
                "j1_national_control_closed_millions"
            ],
            context=(
                f"J-1 FFY{year} closed"
            ),
        )

        after = before - j1_closed

        documented_after = (
            documented_before
            - j1_closed
        )

        if after < 0:
            raise RuntimeError(
                "J-1 QUEUE UNDERFLOW\n"
                f"FFY{year}\n"
                f"Before   : {before}\n"
                f"J-1 close: {j1_closed}\n"
                f"After    : {after}"
            )

        if documented_after < 0:
            raise RuntimeError(
                "J-1 DOCUMENTED-METHOD QUEUE UNDERFLOW\n"
                f"FFY{year}"
            )

        output.append(
            {
                "fiscal_year": year,
                "i5_actionable_unfinished_before_j1_millions": (
                    decimal_text(before)
                ),
                "i5_allocator_method_documented_before_j1_millions": (
                    decimal_text(
                        documented_before
                    )
                ),
                "j1_national_control_closed_millions": (
                    decimal_text(j1_closed)
                ),
                "allocator_method_documented_after_j1_millions": (
                    decimal_text(
                        documented_after
                    )
                ),
                "actionable_unfinished_after_j1_millions": (
                    decimal_text(after)
                ),
                "method_change_review_millions": (
                    row[
                        "method_change_review_millions"
                    ]
                ),
                "withhold_unresolved_millions": (
                    row[
                        "withhold_unresolved_millions"
                    ]
                ),
                "intentionally_unallocable_millions": (
                    row[
                        "intentionally_unallocable_millions"
                    ]
                ),
            }
        )

    if len(output) != 10:
        raise RuntimeError(
            "J-1 REDUCED QUEUE YEAR COVERAGE FAILURE"
        )

    return output


# =============================================================================
# CERTIFICATION
# =============================================================================


def certify(
    allocator_rows: list[dict[str, Any]],
    allocation_rows: list[dict[str, Any]],
    annual_rows: list[dict[str, Any]],
    queue_rows: list[dict[str, Any]],
) -> dict[str, str]:
    if len(allocator_rows) != 10:
        raise RuntimeError(
            "ALLOCATOR ROW COUNT FAILURE"
        )

    if len(allocation_rows) != 20:
        raise RuntimeError(
            "ALLOCATION ROW COUNT FAILURE"
        )

    if len(annual_rows) != 10:
        raise RuntimeError(
            "ANNUAL ROW COUNT FAILURE"
        )

    if len(queue_rows) != 10:
        raise RuntimeError(
            "QUEUE ROW COUNT FAILURE"
        )

    years = {
        int(row["fiscal_year"])
        for row in allocator_rows
    }

    if years != set(TARGET_YEARS):
        raise RuntimeError(
            "ALLOCATOR YEAR COVERAGE FAILURE"
        )

    categories_by_year: dict[int, set[str]] = {
        year: set()
        for year in TARGET_YEARS
    }

    for row in allocation_rows:
        year = int(row["fiscal_year"])

        categories_by_year[year].add(
            row["receipt_category"]
        )

    expected_categories = {
        "RAILROAD_RETIREMENT",
        "ESTATE_AND_GIFT_TAXES",
    }

    for year in TARGET_YEARS:
        if (
            categories_by_year[year]
            != expected_categories
        ):
            raise RuntimeError(
                "CATEGORY COVERAGE FAILURE\n"
                f"FFY{year}: "
                f"{categories_by_year[year]}"
            )

    return {
        "irs_source_integrity": "PASS",
        "etl_171_parser_reuse": "PASS",
        "washington_row_uniqueness": "PASS",
        "united_states_total_row_uniqueness": "PASS",
        "irs_numbered_column_semantics": "PASS",
        "washington_etl_171_reproduction": "PASS",
        "railroad_national_control_construction": "PASS",
        "estate_gift_national_control": "PASS",
        "washington_share_calculation": "PASS",
        "omb_control_application": "PASS",
        "i5_work_queue_reduction": "PASS",
        "database_write_boundary": "PASS",
        "final_ledger_boundary": "PASS",
    }


# =============================================================================
# REPORT
# =============================================================================


def build_report(
    annual_rows: list[dict[str, Any]],
    queue_rows: list[dict[str, Any]],
    certification: dict[str, str],
) -> str:
    queue_by_year = {
        int(row["fiscal_year"]): row
        for row in queue_rows
    }

    lines = [
        "=" * 124,
        (
            "WASHINGTON IRS TABLE 5 REMAINING "
            "RECEIPTS ALLOCATION - ETL-172J-1"
        ),
        "=" * 124,
        "",
        "ANNUAL ALLOCATION",
        "-" * 124,
        (
            f"{'FFY':<6}"
            f"{'RR CTRL':>12}"
            f"{'RR WA %':>12}"
            f"{'EG CTRL':>12}"
            f"{'EG WA %':>12}"
            f"{'J1 CLOSED':>14}"
            f"{'WA ALLOC $B':>14}"
            f"{'QUEUE LEFT':>14}"
        ),
        "-" * 124,
    ]

    for row in annual_rows:
        year = int(row["fiscal_year"])

        queue = queue_by_year[year]

        rr_control = Decimal(
            row[
                "railroad_retirement_national_control_millions"
            ]
        )

        rr_share_pct = (
            Decimal(
                row[
                    "railroad_retirement_wa_share"
                ]
            )
            * Decimal("100")
        )

        eg_control = Decimal(
            row[
                "estate_gift_national_control_millions"
            ]
        )

        eg_share_pct = (
            Decimal(
                row[
                    "estate_gift_wa_share"
                ]
            )
            * Decimal("100")
        )

        closed = Decimal(
            row[
                "j1_national_control_closed_millions"
            ]
        )

        wa_billions = (
            Decimal(
                row[
                    "j1_washington_allocated_receipts_dollars"
                ]
            )
            / Decimal("1000000000")
        )

        queue_left = Decimal(
            queue[
                "actionable_unfinished_after_j1_millions"
            ]
        )

        lines.append(
            f"{year:<6}"
            f"{rr_control:>12,.0f}"
            f"{rr_share_pct:>11.4f}%"
            f"{eg_control:>12,.0f}"
            f"{eg_share_pct:>11.4f}%"
            f"{closed:>14,.0f}"
            f"{wa_billions:>14,.3f}"
            f"{queue_left:>14,.0f}"
        )

    lines.extend(
        [
            "",
            "CERTIFICATION",
            "-" * 124,
        ]
    )

    for key, value in certification.items():
        lines.append(
            f"{key:<62}: {value}"
        )

    lines.extend(
        [
            "",
            "INTERPRETATION",
            "-" * 124,
            (
                "IRS Table 5 supplies the Washington geographic "
                "share. OMB supplies the national federal-receipts "
                "control."
            ),
            (
                "Washington IRS gross collections are not treated "
                "as the final economic-incidence receipt amount."
            ),
            (
                "Railroad Retirement uses OMB Table 2.4 rows "
                "12 + 13 + 14. Row 15 is explicitly excluded."
            ),
            (
                "J-1 closes only Railroad Retirement and Estate "
                "and Gift Taxes from the I-5 actionable work queue."
            ),
            (
                "No database writes and no final ledger "
                "authorization are performed."
            ),
            "",
            "ETL-172J-1 RESULT: PASS",
            "",
        ]
    )

    return "\n".join(lines)


# =============================================================================
# MAIN
# =============================================================================


def main() -> int:
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("=" * 124)
    print(
        "WASHINGTON IRS TABLE 5 REMAINING "
        "RECEIPTS ALLOCATION - ETL-172J-1"
    )
    print("=" * 124)
    print()

    manifest = load_irs_source_manifest()

    allocator_rows = extract_irs_allocators(
        manifest
    )

    railroad_controls = (
        load_railroad_controls()
    )

    estate_gift_controls = (
        load_estate_gift_controls()
    )

    allocation_rows, annual_rows = allocate(
        allocator_rows,
        railroad_controls,
        estate_gift_controls,
    )

    queue_rows = reduce_i5_queue(
        annual_rows
    )

    certification = certify(
        allocator_rows,
        allocation_rows,
        annual_rows,
        queue_rows,
    )

    write_csv(
        OUTPUT_ALLOCATOR,
        allocator_rows,
    )

    write_csv(
        OUTPUT_ALLOCATION,
        allocation_rows,
    )

    write_csv(
        OUTPUT_ANNUAL,
        annual_rows,
    )

    write_csv(
        OUTPUT_QUEUE,
        queue_rows,
    )

    payload = {
        "etl_step": "ETL-172J-1",
        "result": "PASS",
        "method": {
            "railroad_retirement": (
                "OMB national Railroad Retirement receipts "
                "multiplied by Washington share of IRS "
                "Table 5 Railroad Retirement gross collections"
            ),
            "estate_and_gift_taxes": (
                "OMB national Estate and Gift Tax receipts "
                "multiplied by Washington share of combined "
                "IRS Table 5 Estate Tax and Gift Tax gross "
                "collections"
            ),
        },
        "source_hashes": {
            "etl_171_semantic_extract": (
                sha256_file(
                    IRS_SEMANTIC_EXTRACT
                )
            ),
            "omb_table_2_4": (
                sha256_file(
                    OMB_TABLE_24
                )
            ),
            "etl_172i_4_other_controls": (
                sha256_file(
                    I4_OTHER_CONTROLS
                )
            ),
            "etl_172i_5_work_queue": (
                sha256_file(
                    I5_WORK_QUEUE
                )
            ),
        },
        "certification": certification,
        "boundaries": {
            "database_writes": False,
            "refund_adjustment": False,
            "forced_irs_reconciliation": False,
            "final_ledger_authorization": False,
        },
    }

    OUTPUT_JSON.write_text(
        json.dumps(
            payload,
            indent=2,
        ),
        encoding="utf-8",
    )

    report = build_report(
        annual_rows,
        queue_rows,
        certification,
    )

    OUTPUT_REPORT.write_text(
        report,
        encoding="utf-8",
    )

    queue_by_year = {
        int(row["fiscal_year"]): row
        for row in queue_rows
    }

    print("ANNUAL ALLOCATION")
    print("-" * 124)

    print(
        f"{'FFY':<6}"
        f"{'RR CTRL':>12}"
        f"{'RR WA %':>12}"
        f"{'EG CTRL':>12}"
        f"{'EG WA %':>12}"
        f"{'J1 CLOSED':>14}"
        f"{'WA ALLOC $B':>14}"
        f"{'QUEUE LEFT':>14}"
    )

    for row in annual_rows:
        year = int(row["fiscal_year"])

        queue = queue_by_year[year]

        rr_control = Decimal(
            row[
                "railroad_retirement_national_control_millions"
            ]
        )

        rr_share_pct = (
            Decimal(
                row[
                    "railroad_retirement_wa_share"
                ]
            )
            * Decimal("100")
        )

        eg_control = Decimal(
            row[
                "estate_gift_national_control_millions"
            ]
        )

        eg_share_pct = (
            Decimal(
                row[
                    "estate_gift_wa_share"
                ]
            )
            * Decimal("100")
        )

        closed = Decimal(
            row[
                "j1_national_control_closed_millions"
            ]
        )

        wa_billions = (
            Decimal(
                row[
                    "j1_washington_allocated_receipts_dollars"
                ]
            )
            / Decimal("1000000000")
        )

        queue_left = Decimal(
            queue[
                "actionable_unfinished_after_j1_millions"
            ]
        )

        print(
            f"{year:<6}"
            f"{rr_control:>12,.0f}"
            f"{rr_share_pct:>11.4f}%"
            f"{eg_control:>12,.0f}"
            f"{eg_share_pct:>11.4f}%"
            f"{closed:>14,.0f}"
            f"{wa_billions:>14,.3f}"
            f"{queue_left:>14,.0f}"
        )

    print()
    print("CERTIFICATION")
    print("-" * 124)

    for key, value in certification.items():
        print(
            f"{key:<62}: {value}"
        )

    print()
    print(
        "Allocator  : "
        f"{OUTPUT_ALLOCATOR.relative_to(REPO_ROOT)}"
    )
    print(
        "Allocation : "
        f"{OUTPUT_ALLOCATION.relative_to(REPO_ROOT)}"
    )
    print(
        "Annual     : "
        f"{OUTPUT_ANNUAL.relative_to(REPO_ROOT)}"
    )
    print(
        "Queue      : "
        f"{OUTPUT_QUEUE.relative_to(REPO_ROOT)}"
    )
    print(
        "JSON       : "
        f"{OUTPUT_JSON.relative_to(REPO_ROOT)}"
    )
    print(
        "Report     : "
        f"{OUTPUT_REPORT.relative_to(REPO_ROOT)}"
    )
    print()
    print("ETL-172J-1 RESULT: PASS")

    return 0


if __name__ == "__main__":
    sys.exit(
        main()
    )