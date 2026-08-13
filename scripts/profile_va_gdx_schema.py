from __future__ import annotations

import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl


REPO_ROOT = Path(__file__).resolve().parents[1]
RAW_ROOT = REPO_ROOT / "data" / "raw" / "va" / "gdx"
EVIDENCE_ROOT = REPO_ROOT / "evidence" / "va" / "gdx"
OUTPUT_PATH = EVIDENCE_ROOT / "schema_profile.json"

START_FY = 2015
END_FY = 2024

EXPECTED_WA_COUNTIES = 39


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()

def write_json(
    path: Path,
    value: Any,
) -> None:
    path.write_text(
        json.dumps(
            value,
            indent=2,
            sort_keys=True,
        ) + "\n",
        encoding="utf-8",
    )

def find_workbook(fiscal_year: int) -> Path:
    year_dir = RAW_ROOT / f"FY{fiscal_year}"

    if not year_dir.exists():
        raise RuntimeError(
            f"Missing VA GDX directory for FY{fiscal_year}: {year_dir}"
        )

    candidates = sorted(year_dir.glob("*.xlsx"))

    if len(candidates) != 1:
        raise RuntimeError(
            f"Expected exactly one XLSX workbook in {year_dir}; "
            f"found {len(candidates)}."
        )

    return candidates[0]


def normalize_header(value: Any) -> str | None:
    if value is None:
        return None

    text = str(value).strip()

    if not text:
        return None

    return text


def duplicate_values(values: list[str]) -> list[str]:
    counts = Counter(values)
    return sorted(
        value
        for value, count in counts.items()
        if count > 1
    )


def inspect_legacy_year(
    fiscal_year: int,
    workbook_path: Path,
) -> dict[str, Any]:
    wb = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        if "WA" not in wb.sheetnames:
            raise RuntimeError(
                f"FY{fiscal_year} is expected to use the legacy schema "
                "but has no WA worksheet."
            )

        ws = wb["WA"]

        title_row = [
            normalize_header(v)
            for v in next(
                ws.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                )
            )
        ]

        units_row = [
            normalize_header(v)
            for v in next(
                ws.iter_rows(
                    min_row=2,
                    max_row=2,
                    values_only=True,
                )
            )
        ]

        header_values = list(
            next(
                ws.iter_rows(
                    min_row=3,
                    max_row=3,
                    values_only=True,
                )
            )
        )

        headers = [
            normalize_header(v)
            for v in header_values
        ]

        nonblank_headers = [
            h
            for h in headers
            if h is not None
        ]

        county_names: list[str] = []

        for row in ws.iter_rows(
            min_row=4,
            values_only=True,
        ):
            raw_name = row[0]

            if raw_name is None:
                continue

            name = str(raw_name).strip()

            if not name:
                continue

            county_names.append(name)

            if len(county_names) == EXPECTED_WA_COUNTIES:
                break

        if len(county_names) != EXPECTED_WA_COUNTIES:
            raise RuntimeError(
                f"FY{fiscal_year} WA worksheet yielded "
                f"{len(county_names)} county rows; expected 39."
            )

        duplicates = duplicate_values(
            county_names
        )

        return {
            "fiscal_year": fiscal_year,
            "schema_regime": "LEGACY_STATE_SHEETS",
            "workbook_filename": workbook_path.name,
            "sha256": sha256_file(workbook_path),
            "worksheet_names": list(wb.sheetnames),
            "worksheet_count": len(wb.sheetnames),
            "washington_sheet": "WA",
            "washington_sheet_rows": ws.max_row,
            "washington_sheet_columns": ws.max_column,
            "title_row": title_row,
            "units_row": units_row,
            "headers_full": headers,
            "headers_nonblank": nonblank_headers,
            "header_count_nonblank": len(
                nonblank_headers
            ),
            "washington_county_count": len(
                county_names
            ),
            "washington_county_names": (
                county_names
            ),
            "duplicate_county_names": duplicates,
            "county_population_valid": (
                len(county_names)
                == EXPECTED_WA_COUNTIES
                and not duplicates
            ),
        }

    finally:
        wb.close()


def rows_for_state(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    state_name: str,
) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []

    for row in ws.iter_rows(
        min_row=6,
        values_only=True,
    ):
        if row[1] == state_name:
            rows.append(row)

    return rows


def inspect_fy2024(
    workbook_path: Path,
) -> dict[str, Any]:
    wb = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        required = [
            "Data Guide",
            "County Direct Exp",
            "County Indirect Exp",
        ]

        missing = [
            name
            for name in required
            if name not in wb.sheetnames
        ]

        if missing:
            raise RuntimeError(
                "FY2024 missing required revised-schema sheets: "
                + ", ".join(missing)
            )

        direct = wb["County Direct Exp"]
        indirect = wb["County Indirect Exp"]

        direct_headers = [
            normalize_header(v)
            for v in next(
                direct.iter_rows(
                    min_row=3,
                    max_row=3,
                    values_only=True,
                )
            )
        ]

        indirect_headers = [
            normalize_header(v)
            for v in next(
                indirect.iter_rows(
                    min_row=3,
                    max_row=3,
                    values_only=True,
                )
            )
        ]

        direct_units = [
            normalize_header(v)
            for v in next(
                direct.iter_rows(
                    min_row=2,
                    max_row=2,
                    values_only=True,
                )
            )
        ]

        indirect_units = [
            normalize_header(v)
            for v in next(
                indirect.iter_rows(
                    min_row=2,
                    max_row=2,
                    values_only=True,
                )
            )
        ]

        direct_rows = rows_for_state(
            direct,
            "Washington",
        )

        indirect_rows = rows_for_state(
            indirect,
            "Washington",
        )

        direct_counties = [
            str(row[2]).strip()
            for row in direct_rows
        ]

        indirect_counties = [
            str(row[2]).strip()
            for row in indirect_rows
        ]

        direct_duplicates = duplicate_values(
            direct_counties
        )

        indirect_duplicates = duplicate_values(
            indirect_counties
        )

        direct_set = set(
            direct_counties
        )

        indirect_set = set(
            indirect_counties
        )

        return {
            "fiscal_year": 2024,
            "schema_regime": (
                "REVISED_DIRECT_INDIRECT"
            ),
            "workbook_filename": (
                workbook_path.name
            ),
            "sha256": sha256_file(
                workbook_path
            ),
            "worksheet_names": list(
                wb.sheetnames
            ),
            "worksheet_count": len(
                wb.sheetnames
            ),
            "direct": {
                "worksheet": (
                    "County Direct Exp"
                ),
                "rows": direct.max_row,
                "columns": direct.max_column,
                "units_row": direct_units,
                "headers_full": direct_headers,
                "headers_nonblank": [
                    h
                    for h in direct_headers
                    if h is not None
                ],
                "washington_county_count": len(
                    direct_counties
                ),
                "washington_county_names": (
                    direct_counties
                ),
                "duplicate_county_names": (
                    direct_duplicates
                ),
            },
            "indirect": {
                "worksheet": (
                    "County Indirect Exp"
                ),
                "rows": indirect.max_row,
                "columns": (
                    indirect.max_column
                ),
                "units_row": indirect_units,
                "headers_full": (
                    indirect_headers
                ),
                "headers_nonblank": [
                    h
                    for h in indirect_headers
                    if h is not None
                ],
                "washington_county_count": len(
                    indirect_counties
                ),
                "washington_county_names": (
                    indirect_counties
                ),
                "duplicate_county_names": (
                    indirect_duplicates
                ),
            },
            "washington_county_populations_match": (
                direct_set == indirect_set
            ),
            "direct_only_counties": sorted(
                direct_set - indirect_set
            ),
            "indirect_only_counties": sorted(
                indirect_set - direct_set
            ),
            "county_population_valid": (
                len(direct_counties)
                == EXPECTED_WA_COUNTIES
                and len(indirect_counties)
                == EXPECTED_WA_COUNTIES
                and not direct_duplicates
                and not indirect_duplicates
                and direct_set == indirect_set
            ),
            "methodology_break": True,
        }

    finally:
        wb.close()


def compare_legacy_headers(
    profiles: dict[int, dict[str, Any]],
) -> dict[str, Any]:
    base_year = 2015
    base_headers = profiles[
        base_year
    ]["headers_nonblank"]

    by_year: dict[str, Any] = {}

    all_match = True

    for year in range(
        2015,
        2024,
    ):
        headers = profiles[
            year
        ]["headers_nonblank"]

        exact_match = (
            headers == base_headers
        )

        if not exact_match:
            all_match = False

        added = [
            h
            for h in headers
            if h not in base_headers
        ]

        removed = [
            h
            for h in base_headers
            if h not in headers
        ]

        order_changed = (
            set(headers) == set(base_headers)
            and headers != base_headers
        )

        by_year[str(year)] = {
            "matches_fy2015_exactly": (
                exact_match
            ),
            "added_headers_vs_fy2015": (
                added
            ),
            "removed_headers_vs_fy2015": (
                removed
            ),
            "same_set_different_order": (
                order_changed
            ),
        }

    return {
        "baseline_year": base_year,
        "all_legacy_headers_match_exactly": (
            all_match
        ),
        "comparison_by_year": by_year,
    }




def main() -> None:
    print("=" * 88)
    print(
        "WCFFBP — VA GDX SCHEMA PROFILER"
    )
    print("=" * 88)
    print()

    profiles: dict[
        int,
        dict[str, Any],
    ] = {}

    for fiscal_year in range(
        START_FY,
        END_FY + 1,
    ):
        workbook_path = find_workbook(
            fiscal_year
        )

        print(
            f"FY{fiscal_year} "
            f"{workbook_path.name} ... ",
            end="",
        )

        if fiscal_year == 2024:
            profile = inspect_fy2024(
                workbook_path
            )

        else:
            profile = inspect_legacy_year(
                fiscal_year,
                workbook_path,
            )

        profiles[
            fiscal_year
        ] = profile

        if fiscal_year == 2024:
            print(
                "REVISED "
                f"direct={profile['direct']['washington_county_count']} "
                f"indirect={profile['indirect']['washington_county_count']}"
            )

        else:
            print(
                "LEGACY "
                f"counties={profile['washington_county_count']} "
                f"headers={profile['header_count_nonblank']}"
            )

    legacy_comparison = (
        compare_legacy_headers(
            profiles
        )
    )

    county_checks_pass = all(
        profiles[year][
            "county_population_valid"
        ]
        for year in range(
            START_FY,
            END_FY + 1,
        )
    )

    single_legacy_mapping_safe = (
        legacy_comparison[
            "all_legacy_headers_match_exactly"
        ]
        and all(
            profiles[year][
                "county_population_valid"
            ]
            for year in range(
                2015,
                2024,
            )
        )
    )

    conclusion = {
        "all_years_have_valid_washington_county_population": (
            county_checks_pass
        ),
        "single_legacy_etl_mapping_safe": (
            single_legacy_mapping_safe
        ),
        "fy2024_requires_separate_mapping": True,
        "recommended_loader_strategy": (
            "Use one legacy loader mapping for FY2015-FY2023 "
            "only if single_legacy_etl_mapping_safe is true. "
            "Use a distinct FY2024 direct/indirect mapping regardless."
        ),
    }

    output = {
        "dataset": (
            "VA Geographic Distribution "
            "of Expenditures (GDX)"
        ),
        "profile_scope": (
            "FY2015-FY2024 Washington "
            "county expenditure schema"
        ),
        "raw_data_mutated": False,
        "expected_washington_counties": (
            EXPECTED_WA_COUNTIES
        ),
        "year_profiles": {
            str(year): profiles[year]
            for year in sorted(profiles)
        },
        "legacy_schema_comparison": (
            legacy_comparison
        ),
        "conclusion": conclusion,
    }

    EVIDENCE_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    write_json(
        OUTPUT_PATH,
        output,
    )

    print()
    print("-" * 88)
    print("LEGACY SCHEMA COMPARISON")
    print("-" * 88)
    print(
        "FY2015-FY2023 headers identical ..... "
        f"{legacy_comparison['all_legacy_headers_match_exactly']}"
    )
    print(
        "All county population checks pass ... "
        f"{county_checks_pass}"
    )
    print(
        "Single legacy ETL mapping safe ....... "
        f"{single_legacy_mapping_safe}"
    )
    print(
        "FY2024 separate mapping required ..... "
        "True"
    )
    print()
    print(
        f"Evidence written ...................... "
        f"{OUTPUT_PATH}"
    )
    print()
    print("=" * 88)
    print("VA GDX SCHEMA PROFILING COMPLETE")
    print("=" * 88)


if __name__ == "__main__":
    main()