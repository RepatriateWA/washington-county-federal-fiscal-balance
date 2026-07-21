from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path(
    "data/raw/chip/cms21/2015/FY 2015 CHIP FMR NET EXPENDITURES.xlsx"
)

DEFAULT_SHEETS = [
    "CHIP - Washington",
    "MCHIP - Washington",
    "MCHIP 20% - Washington",
]

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}

HEADER_KEYWORDS = {
    "state",
    "quarter",
    "fiscal",
    "year",
    "category",
    "service",
    "description",
    "expenditure",
    "expenditures",
    "federal",
    "state share",
    "total",
    "computable",
    "net",
    "amount",
    "form",
    "line",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Inspect CMS-21 CHIP Excel workbooks and optionally export "
            "worksheet data and metadata."
        )
    )

    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help=f"Workbook path. Default: {DEFAULT_WORKBOOK}",
    )

    parser.add_argument(
        "--sheet",
        action="append",
        dest="sheets",
        help=(
            "Worksheet to inspect. May be supplied more than once. "
            "If omitted, the three Washington worksheets are inspected."
        ),
    )

    parser.add_argument(
        "--rows",
        type=int,
        default=40,
        help="Number of leading rows to print. Default: 40",
    )

    parser.add_argument(
        "--tail-rows",
        type=int,
        default=10,
        help="Number of trailing rows to print. Default: 10",
    )

    parser.add_argument(
        "--list-sheets",
        action="store_true",
        help="List all worksheet names and exit.",
    )

    parser.add_argument(
        "--summary-all",
        action="store_true",
        help="Print a compact structural summary for every worksheet.",
    )

    parser.add_argument(
        "--output",
        type=Path,
        help="Optional text file for the inspection report.",
    )

    parser.add_argument(
        "--json-output",
        type=Path,
        help="Optional JSON file for workbook and worksheet metadata.",
    )

    parser.add_argument(
        "--export-csv-dir",
        type=Path,
        help=(
            "Optional directory for raw CSV exports of inspected worksheets."
        ),
    )

    parser.add_argument(
        "--search",
        action="append",
        dest="search_terms",
        help=(
            "Search inspected worksheets for a case-insensitive term. "
            "May be supplied more than once."
        ),
    )

    return parser.parse_args()


def normalize_json_value(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, float) and pd.isna(value):
        return None

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, TypeError):
            pass

    return value


def safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", value)
    cleaned = re.sub(r"\s+", "_", cleaned.strip())
    return cleaned[:150] or "worksheet"


def find_used_bounds(dataframe: pd.DataFrame) -> dict[str, int | None]:
    nonempty_mask = dataframe.notna()

    occupied_rows = nonempty_mask.any(axis=1)
    occupied_columns = nonempty_mask.any(axis=0)

    row_indexes = dataframe.index[occupied_rows].tolist()
    column_indexes = dataframe.columns[occupied_columns].tolist()

    return {
        "first_nonempty_row_zero_based": row_indexes[0] if row_indexes else None,
        "last_nonempty_row_zero_based": row_indexes[-1] if row_indexes else None,
        "first_nonempty_column_zero_based": (
            int(column_indexes[0]) if column_indexes else None
        ),
        "last_nonempty_column_zero_based": (
            int(column_indexes[-1]) if column_indexes else None
        ),
    }


def classify_cell(value: Any) -> str:
    if value is None or pd.isna(value):
        return "blank"

    if isinstance(value, bool):
        return "boolean"

    if isinstance(value, (int, float)):
        return "numeric"

    if isinstance(value, (datetime, date, pd.Timestamp)):
        return "date_or_datetime"

    if isinstance(value, str):
        stripped = value.strip()

        if not stripped:
            return "blank"

        if stripped.startswith("="):
            return "formula_text"

        return "text"

    return type(value).__name__


def cell_type_counts(dataframe: pd.DataFrame) -> dict[str, int]:
    counts: Counter[str] = Counter()

    for value in dataframe.to_numpy().ravel():
        counts[classify_cell(value)] += 1

    return dict(sorted(counts.items()))


def score_header_row(row: pd.Series) -> tuple[int, list[str]]:
    nonblank_values = [
        str(value).strip()
        for value in row.tolist()
        if value is not None and not pd.isna(value) and str(value).strip()
    ]

    if not nonblank_values:
        return 0, []

    text_values = [
        value for value in nonblank_values if not value.replace(",", "").isdigit()
    ]

    keyword_hits: list[str] = []

    lowered_text = " | ".join(text_values).lower()

    for keyword in HEADER_KEYWORDS:
        if keyword in lowered_text:
            keyword_hits.append(keyword)

    score = len(text_values) + (3 * len(keyword_hits))

    if len(nonblank_values) >= 3:
        score += 2

    if len(text_values) == len(nonblank_values):
        score += 1

    return score, sorted(keyword_hits)


def detect_likely_headers(
    dataframe: pd.DataFrame,
    rows_to_scan: int = 30,
    maximum_results: int = 5,
) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []

    for row_index in range(min(rows_to_scan, len(dataframe))):
        score, keyword_hits = score_header_row(dataframe.iloc[row_index])

        if score <= 0:
            continue

        row_values = [
            normalize_json_value(value)
            for value in dataframe.iloc[row_index].tolist()
        ]

        candidates.append(
            {
                "row_zero_based": row_index,
                "excel_row": row_index + 1,
                "score": score,
                "keyword_hits": keyword_hits,
                "values": row_values,
            }
        )

    candidates.sort(
        key=lambda item: (-item["score"], item["row_zero_based"])
    )

    return candidates[:maximum_results]


def find_search_matches(
    dataframe: pd.DataFrame,
    search_terms: list[str],
    maximum_matches_per_term: int = 50,
) -> dict[str, list[dict[str, Any]]]:
    results: dict[str, list[dict[str, Any]]] = {}

    for term in search_terms:
        matches: list[dict[str, Any]] = []
        term_lower = term.lower()

        for row_index, row in dataframe.iterrows():
            for column_index, value in row.items():
                if value is None or pd.isna(value):
                    continue

                if term_lower in str(value).lower():
                    matches.append(
                        {
                            "row_zero_based": int(row_index),
                            "excel_row": int(row_index) + 1,
                            "column_zero_based": int(column_index),
                            "excel_column_number": int(column_index) + 1,
                            "value": normalize_json_value(value),
                        }
                    )

                    if len(matches) >= maximum_matches_per_term:
                        break

            if len(matches) >= maximum_matches_per_term:
                break

        results[term] = matches

    return results


def get_openpyxl_metadata(
    workbook_path: Path,
    sheet_name: str,
) -> dict[str, Any]:
    if workbook_path.suffix.lower() == ".xls":
        return {
            "available": False,
            "reason": (
                "openpyxl does not support legacy .xls files. "
                "Merged-cell and formula metadata were not inspected."
            ),
        }

    workbook = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
    )

    try:
        worksheet = workbook[sheet_name]

        merged_ranges = [str(item) for item in worksheet.merged_cells.ranges]

        formula_cells: list[str] = []
        hidden_rows: list[int] = []
        hidden_columns: list[str] = []

        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(cell.coordinate)

        for row_number, dimension in worksheet.row_dimensions.items():
            if dimension.hidden:
                hidden_rows.append(int(row_number))

        for column_letter, dimension in worksheet.column_dimensions.items():
            if dimension.hidden:
                hidden_columns.append(str(column_letter))

        return {
            "available": True,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "merged_cell_count": len(merged_ranges),
            "merged_ranges": merged_ranges,
            "formula_cell_count": len(formula_cells),
            "formula_cells": formula_cells[:100],
            "hidden_rows": hidden_rows,
            "hidden_columns": hidden_columns,
            "freeze_panes": (
                str(worksheet.freeze_panes)
                if worksheet.freeze_panes is not None
                else None
            ),
        }
    finally:
        workbook.close()


def load_sheet(
    workbook_path: Path,
    sheet_name: str,
) -> pd.DataFrame:
    return pd.read_excel(
        workbook_path,
        sheet_name=sheet_name,
        header=None,
    )


def build_sheet_metadata(
    workbook_path: Path,
    sheet_name: str,
    dataframe: pd.DataFrame,
    search_terms: list[str],
) -> dict[str, Any]:
    nonempty_rows = dataframe.dropna(how="all")
    nonempty_columns = dataframe.dropna(axis=1, how="all")

    return {
        "sheet_name": sheet_name,
        "shape": {
            "rows": int(dataframe.shape[0]),
            "columns": int(dataframe.shape[1]),
            "nonempty_rows": int(nonempty_rows.shape[0]),
            "nonempty_columns": int(nonempty_columns.shape[1]),
        },
        "used_bounds": find_used_bounds(dataframe),
        "cell_types": cell_type_counts(dataframe),
        "likely_header_rows": detect_likely_headers(dataframe),
        "search_matches": (
            find_search_matches(dataframe, search_terms)
            if search_terms
            else {}
        ),
        "excel_metadata": get_openpyxl_metadata(
            workbook_path,
            sheet_name,
        ),
    }


def format_sheet_report(
    dataframe: pd.DataFrame,
    metadata: dict[str, Any],
    head_rows: int,
    tail_rows: int,
) -> str:
    bounds = metadata["used_bounds"]
    excel_metadata = metadata["excel_metadata"]

    output = [
        "",
        "=" * 100,
        f"SHEET: {metadata['sheet_name']}",
        "=" * 100,
        f"Total rows: {metadata['shape']['rows']}",
        f"Total columns: {metadata['shape']['columns']}",
        f"Nonempty rows: {metadata['shape']['nonempty_rows']}",
        f"Nonempty columns: {metadata['shape']['nonempty_columns']}",
        (
            "Used row range: "
            f"{bounds['first_nonempty_row_zero_based']} to "
            f"{bounds['last_nonempty_row_zero_based']} "
            "(zero-based)"
        ),
        (
            "Used column range: "
            f"{bounds['first_nonempty_column_zero_based']} to "
            f"{bounds['last_nonempty_column_zero_based']} "
            "(zero-based)"
        ),
        "",
        "CELL TYPES",
        "-" * 100,
    ]

    for cell_type, count in metadata["cell_types"].items():
        output.append(f"{cell_type}: {count:,}")

    output.extend(
        [
            "",
            "EXCEL STRUCTURE",
            "-" * 100,
        ]
    )

    if excel_metadata["available"]:
        output.extend(
            [
                f"Merged ranges: {excel_metadata['merged_cell_count']}",
                f"Formula cells: {excel_metadata['formula_cell_count']}",
                f"Hidden rows: {excel_metadata['hidden_rows']}",
                f"Hidden columns: {excel_metadata['hidden_columns']}",
                f"Freeze panes: {excel_metadata['freeze_panes']}",
            ]
        )

        if excel_metadata["merged_ranges"]:
            output.append(
                "Merged cell ranges: "
                + ", ".join(excel_metadata["merged_ranges"])
            )
    else:
        output.append(excel_metadata["reason"])

    output.extend(
        [
            "",
            "LIKELY HEADER ROWS",
            "-" * 100,
        ]
    )

    if metadata["likely_header_rows"]:
        for candidate in metadata["likely_header_rows"]:
            values = [
                str(value)
                for value in candidate["values"]
                if value is not None
            ]

            output.append(
                f"Excel row {candidate['excel_row']} "
                f"(score={candidate['score']}, "
                f"keywords={candidate['keyword_hits']}): "
                + " | ".join(values)
            )
    else:
        output.append("No likely header rows detected.")

    if metadata["search_matches"]:
        output.extend(
            [
                "",
                "SEARCH MATCHES",
                "-" * 100,
            ]
        )

        for term, matches in metadata["search_matches"].items():
            output.append(f"Term: {term}")

            if not matches:
                output.append("  No matches.")
                continue

            for match in matches:
                output.append(
                    "  "
                    f"row={match['excel_row']}, "
                    f"column={match['excel_column_number']}, "
                    f"value={match['value']}"
                )

    output.extend(
        [
            "",
            f"FIRST {head_rows} ROWS",
            "-" * 100,
            dataframe.head(head_rows).to_string(
                index=True,
                header=True,
                max_cols=None,
            ),
        ]
    )

    if tail_rows > 0:
        output.extend(
            [
                "",
                f"LAST {tail_rows} ROWS",
                "-" * 100,
                dataframe.tail(tail_rows).to_string(
                    index=True,
                    header=True,
                    max_cols=None,
                ),
            ]
        )

    output.append("")

    return "\n".join(output)


def export_sheet_csv(
    dataframe: pd.DataFrame,
    output_directory: Path,
    sheet_name: str,
) -> Path:
    output_directory.mkdir(parents=True, exist_ok=True)

    output_path = output_directory / f"{safe_filename(sheet_name)}.csv"

    dataframe.to_csv(
        output_path,
        index=False,
        header=False,
        encoding="utf-8-sig",
    )

    return output_path


def workbook_summary(
    workbook_path: Path,
    worksheet_names: list[str],
) -> dict[str, Any]:
    stat = workbook_path.stat()

    return {
        "workbook_path": str(workbook_path),
        "filename": workbook_path.name,
        "extension": workbook_path.suffix.lower(),
        "file_size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(
            stat.st_mtime
        ).isoformat(timespec="seconds"),
        "worksheet_count": len(worksheet_names),
        "worksheet_names": worksheet_names,
    }


def print_all_sheet_summary(
    workbook_path: Path,
    worksheet_names: list[str],
) -> None:
    print(f"Workbook: {workbook_path}")
    print(f"Worksheet count: {len(worksheet_names)}")
    print()
    print(
        f"{'No.':>4}  {'Rows':>6}  {'Cols':>6}  Worksheet"
    )
    print("-" * 100)

    for index, sheet_name in enumerate(worksheet_names, start=1):
        dataframe = load_sheet(workbook_path, sheet_name)

        print(
            f"{index:>4}  "
            f"{dataframe.shape[0]:>6}  "
            f"{dataframe.shape[1]:>6}  "
            f"{sheet_name}"
        )


def main() -> int:
    args = parse_args()
    workbook_path = args.workbook.resolve()

    if not workbook_path.exists():
        print(
            f"ERROR: Workbook not found:\n{workbook_path}",
            file=sys.stderr,
        )
        return 1

    if workbook_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        print(
            f"ERROR: Unsupported workbook type: "
            f"{workbook_path.suffix}",
            file=sys.stderr,
        )
        return 1

    if args.rows < 0 or args.tail_rows < 0:
        print(
            "ERROR: --rows and --tail-rows must be nonnegative.",
            file=sys.stderr,
        )
        return 1

    try:
        excel_file = pd.ExcelFile(workbook_path)
    except Exception as exc:
        print(
            f"ERROR: Could not open workbook:\n"
            f"{workbook_path}\n\n{exc}",
            file=sys.stderr,
        )
        return 1

    worksheet_names = excel_file.sheet_names
    workbook_metadata = workbook_summary(
        workbook_path,
        worksheet_names,
    )

    if args.list_sheets:
        print(f"Workbook: {workbook_path}")
        print(f"Worksheet count: {len(worksheet_names)}")
        print()

        for index, sheet_name in enumerate(
            worksheet_names,
            start=1,
        ):
            print(f"{index:03d}: {sheet_name}")

        return 0

    if args.summary_all:
        print_all_sheet_summary(
            workbook_path,
            worksheet_names,
        )
        return 0

    requested_sheets = args.sheets or DEFAULT_SHEETS
    search_terms = args.search_terms or []

    output_parts = [
        "CMS-21 CHIP WORKBOOK INSPECTION",
        "=" * 100,
        f"Workbook: {workbook_path}",
        f"File size: {workbook_metadata['file_size_bytes']:,} bytes",
        f"Modified: {workbook_metadata['modified_at']}",
        f"Worksheet count: {len(worksheet_names)}",
        f"Requested sheets: {len(requested_sheets)}",
        "",
    ]

    metadata_document: dict[str, Any] = {
        "workbook": workbook_metadata,
        "inspected_sheets": [],
    }

    for sheet_name in requested_sheets:
        if sheet_name not in worksheet_names:
            message = (
                f"\n{'=' * 100}\n"
                f"SHEET: {sheet_name}\n"
                f"{'=' * 100}\n"
                "ERROR: Worksheet not found.\n"
            )
            output_parts.append(message)

            metadata_document["inspected_sheets"].append(
                {
                    "sheet_name": sheet_name,
                    "error": "Worksheet not found",
                }
            )
            continue

        try:
            dataframe = load_sheet(
                workbook_path,
                sheet_name,
            )

            metadata = build_sheet_metadata(
                workbook_path=workbook_path,
                sheet_name=sheet_name,
                dataframe=dataframe,
                search_terms=search_terms,
            )

            metadata_document["inspected_sheets"].append(metadata)

            output_parts.append(
                format_sheet_report(
                    dataframe=dataframe,
                    metadata=metadata,
                    head_rows=args.rows,
                    tail_rows=args.tail_rows,
                )
            )

            if args.export_csv_dir:
                csv_path = export_sheet_csv(
                    dataframe=dataframe,
                    output_directory=args.export_csv_dir.resolve(),
                    sheet_name=sheet_name,
                )

                output_parts.append(
                    f"CSV exported: {csv_path}\n"
                )

        except Exception as exc:
            message = (
                f"\n{'=' * 100}\n"
                f"SHEET: {sheet_name}\n"
                f"{'=' * 100}\n"
                f"ERROR: {type(exc).__name__}: {exc}\n"
            )
            output_parts.append(message)

            metadata_document["inspected_sheets"].append(
                {
                    "sheet_name": sheet_name,
                    "error": f"{type(exc).__name__}: {exc}",
                }
            )

    final_output = "\n".join(output_parts)
    print(final_output)

    if args.output:
        output_path = args.output.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(
            final_output,
            encoding="utf-8",
        )
        print(f"\nInspection report saved to:\n{output_path}")

    if args.json_output:
        json_path = args.json_output.resolve()
        json_path.parent.mkdir(parents=True, exist_ok=True)
        json_path.write_text(
            json.dumps(
                metadata_document,
                indent=2,
                ensure_ascii=False,
                default=normalize_json_value,
            ),
            encoding="utf-8",
        )
        print(f"\nMetadata JSON saved to:\n{json_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())