"""
ETL-172J-2
Washington Population-Allocated Remaining Federal Receipts v0.1

Purpose
-------
Construct a certified Washington population allocator for FFY2015-FFY2024
and apply it to the remaining federal-receipts categories for which the
Rockefeller methodology specifies Census population as the state allocator:

    OTHER_RETIREMENT
    TOBACCO
    AIRPORT_AND_AIRWAY
    OTHER_EXCISES

Population policy
-----------------
FFY2015-FFY2019:
    U.S. Census Bureau 2010-2020 intercensal resident-population estimates.

FFY2020-FFY2024:
    U.S. Census Bureau Vintage 2025 resident-population estimates.

The principal WCFFBP historical reconstruction uses the most recently
revised official Census resident-population series applicable to each
period rather than freezing the population vintage available when a
Rockefeller report was originally published.

Accounting boundary
-------------------
This script:
    - preserves and hashes Census source artifacts;
    - extracts U.S. and Washington resident populations;
    - calculates Washington / U.S. annual population shares;
    - reconstructs the four population-allocated national receipt controls;
    - applies the population shares to those controls;
    - reconciles the national controls to the ETL-172I-5 methodology;
    - reduces the ETL-172J-1 actionable work queue;
    - writes validation artifacts only.

This script does NOT:
    - write to PostgreSQL;
    - alter any prior certified artifact;
    - allocate county receipts;
    - reproduce Rockefeller publication vintages;
    - authorize final ledger inclusion;
    - force reconciliation to IRS gross collections.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


# =============================================================================
# PATHS
# =============================================================================

ROOT = Path(__file__).resolve().parents[1]

CENSUS_SOURCE_DIR = (
    ROOT
    / "data"
    / "raw"
    / "census"
    / "population"
    / "source"
)

INTERCENSAL_FILE = CENSUS_SOURCE_DIR / "nst-est2020int-pop.xlsx"
VINTAGE_2025_FILE = CENSUS_SOURCE_DIR / "NST-EST2025-ALLDATA.csv"

OMB_TABLE_24_FILE = (
    ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "source"
    / "BUDGET-2027-TAB-3-4.xlsx"
)

I2_DIR = (
    ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "social_insurance_excise_decomposition_v0_1"
)

I2_DECOMPOSITION_FILE = (
    I2_DIR / "national_social_insurance_excise_decomposition.csv"
)

I5_DIR = (
    ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "remaining_receipts_methodology_crosswalk_v0_1"
)

I5_CROSSWALK_FILE = I5_DIR / "remaining_receipts_methodology_crosswalk.csv"
I5_QUEUE_FILE = I5_DIR / "remaining_receipts_work_queue.csv"

J1_DIR = (
    ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "washington_irs_table5_remaining_receipts_v0_1"
)

J1_QUEUE_FILE = J1_DIR / "remaining_receipts_work_queue_after_j1.csv"

OUTPUT_DIR = (
    ROOT
    / "data"
    / "raw"
    / "omb"
    / "historical_tables"
    / "FY2027"
    / "validation"
    / "washington_population_receipts_v0_1"
)

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

POPULATION_OUTPUT = OUTPUT_DIR / "washington_population_allocator.csv"
CONTROL_OUTPUT = OUTPUT_DIR / "population_allocated_national_controls.csv"
ALLOCATION_OUTPUT = OUTPUT_DIR / "washington_population_receipts_allocation.csv"
ANNUAL_OUTPUT = OUTPUT_DIR / "washington_population_receipts_annual_summary.csv"
QUEUE_OUTPUT = OUTPUT_DIR / "remaining_receipts_work_queue_after_j2.csv"
CERT_JSON = OUTPUT_DIR / "certification.json"
CERT_REPORT = OUTPUT_DIR / "certification_report.txt"


YEARS = list(range(2015, 2025))

POPULATION_COMPONENTS = {
    "OTHER_RETIREMENT",
    "TOBACCO",
    "AIRPORT_AND_AIRWAY",
    "OTHER_EXCISES",
}


# =============================================================================
# HELPERS
# =============================================================================

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def to_float(value: Any, *, label: str) -> float:
    if value is None:
        raise ValueError(f"{label}: missing numeric value")

    if isinstance(value, str):
        cleaned = value.strip().replace(",", "")
        if cleaned == "":
            raise ValueError(f"{label}: blank numeric value")
        value = cleaned

    result = float(value)

    if not math.isfinite(result):
        raise ValueError(f"{label}: non-finite numeric value {value!r}")

    return result


def to_int(value: Any, *, label: str) -> int:
    result = to_float(value, label=label)

    rounded = round(result)

    if abs(result - rounded) > 1e-6:
        raise ValueError(
            f"{label}: expected integer-like value, got {value!r}"
        )

    return int(rounded)


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"Refusing to write empty CSV: {path}")

    fieldnames = list(rows[0].keys())

    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def require_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Required file not found: {path}")


def pass_fail(condition: bool) -> str:
    return "PASS" if condition else "FAIL"


def find_column(
    columns: list[str],
    candidates: list[str],
    *,
    context: str,
) -> str:
    """
    Find exactly one candidate column, case-insensitively.

    This lets J-2 tolerate small naming differences in the previously
    certified queue artifact without silently selecting an unrelated field.
    """
    normalized = {c.lower(): c for c in columns}

    matches = []

    for candidate in candidates:
        if candidate.lower() in normalized:
            matches.append(normalized[candidate.lower()])

    matches = list(dict.fromkeys(matches))

    if len(matches) != 1:
        raise ValueError(
            f"{context}: expected exactly one matching column from "
            f"{candidates}; found {matches}. Available columns: {columns}"
        )

    return matches[0]


# =============================================================================
# CENSUS POPULATION EXTRACTION
# =============================================================================

@dataclass(frozen=True)
class PopulationObservation:
    fiscal_year: int
    population_year: int
    source_series: str
    source_file: str
    source_sha256: str
    us_population: int
    washington_population: int
    washington_population_share: float


def extract_intercensal_population() -> dict[int, PopulationObservation]:
    """
    Extract 2015-2019 from NST-EST2020INT-POP.

    Workbook structure established before coding:
      row 2 (zero-based) contains the main header;
      row 3 contains annual labels;
      row 4 is United States;
      state names are prefixed with a period, e.g. ".Washington".
    """
    require_file(INTERCENSAL_FILE)

    source_hash = sha256_file(INTERCENSAL_FILE)

    df = pd.read_excel(
        INTERCENSAL_FILE,
        sheet_name="NST-EST2020INT-POP",
        header=None,
        dtype=object,
        engine="openpyxl",
    )

    if df.shape[0] < 10 or df.shape[1] < 10:
        raise ValueError(
            f"Unexpected intercensal workbook shape: {df.shape}"
        )

    # Row 3 contains year labels for the July 1 population estimate columns.
    year_row_index = 3
    geography_col = 0

    year_columns: dict[int, int] = {}

    for col_idx in range(df.shape[1]):
        value = df.iat[year_row_index, col_idx]

        if pd.isna(value):
            continue

        try:
            year = int(float(value))
        except (TypeError, ValueError):
            continue

        if year in YEARS:
            year_columns[year] = col_idx

    required_years = set(range(2015, 2020))

    if set(year_columns) != required_years:
        raise ValueError(
            "Intercensal year-column detection failed. "
            f"Expected {sorted(required_years)}, "
            f"found {sorted(year_columns)}"
        )

    geography_rows: dict[str, list[int]] = {}

    for row_idx in range(df.shape[0]):
        label = normalize_text(df.iat[row_idx, geography_col])

        if not label:
            continue

        normalized_label = label.lstrip(".").strip().lower()
        geography_rows.setdefault(normalized_label, []).append(row_idx)

    us_rows = geography_rows.get("united states", [])
    wa_rows = geography_rows.get("washington", [])

    if len(us_rows) != 1:
        raise ValueError(
            f"Expected exactly one United States row; found {us_rows}"
        )

    if len(wa_rows) != 1:
        raise ValueError(
            f"Expected exactly one Washington row; found {wa_rows}"
        )

    us_row = us_rows[0]
    wa_row = wa_rows[0]

    observations: dict[int, PopulationObservation] = {}

    for year in range(2015, 2020):
        col_idx = year_columns[year]

        us_population = to_int(
            df.iat[us_row, col_idx],
            label=f"Intercensal U.S. population {year}",
        )

        wa_population = to_int(
            df.iat[wa_row, col_idx],
            label=f"Intercensal Washington population {year}",
        )

        if us_population <= 0 or wa_population <= 0:
            raise ValueError(f"Invalid population for {year}")

        if wa_population >= us_population:
            raise ValueError(
                f"Washington population is not less than U.S. population "
                f"for {year}"
            )

        observations[year] = PopulationObservation(
            fiscal_year=year,
            population_year=year,
            source_series="CENSUS_2010_2020_INTERCENSAL",
            source_file=INTERCENSAL_FILE.name,
            source_sha256=source_hash,
            us_population=us_population,
            washington_population=wa_population,
            washington_population_share=wa_population / us_population,
        )

    return observations


def extract_vintage_2025_population() -> dict[int, PopulationObservation]:
    """
    Extract 2020-2024 from NST-EST2025-ALLDATA.csv.

    U.S. row:
        SUMLEV = 010
        STATE  = 00
        NAME   = United States

    Washington row:
        SUMLEV = 040
        NAME   = Washington
    """
    require_file(VINTAGE_2025_FILE)

    source_hash = sha256_file(VINTAGE_2025_FILE)

    df = pd.read_csv(
        VINTAGE_2025_FILE,
        dtype=str,
        keep_default_na=False,
    )

    required_columns = {
        "SUMLEV",
        "STATE",
        "NAME",
        "POPESTIMATE2020",
        "POPESTIMATE2021",
        "POPESTIMATE2022",
        "POPESTIMATE2023",
        "POPESTIMATE2024",
    }

    missing = required_columns - set(df.columns)

    if missing:
        raise ValueError(
            f"Vintage 2025 CSV missing required columns: {sorted(missing)}"
        )

    us = df[
        (df["SUMLEV"].str.strip() == "010")
        & (df["STATE"].str.strip() == "00")
        & (df["NAME"].str.strip() == "United States")
    ]

    if len(us) != 1:
        raise ValueError(
            f"Expected exactly one U.S. row in Vintage 2025; found {len(us)}"
        )

    wa = df[
        (df["SUMLEV"].str.strip() == "040")
        & (df["NAME"].str.strip() == "Washington")
    ]

    if len(wa) != 1:
        raise ValueError(
            f"Expected exactly one Washington row in Vintage 2025; "
            f"found {len(wa)}"
        )

    us_row = us.iloc[0]
    wa_row = wa.iloc[0]

    observations: dict[int, PopulationObservation] = {}

    for year in range(2020, 2025):
        column = f"POPESTIMATE{year}"

        us_population = to_int(
            us_row[column],
            label=f"Vintage 2025 U.S. population {year}",
        )

        wa_population = to_int(
            wa_row[column],
            label=f"Vintage 2025 Washington population {year}",
        )

        if us_population <= 0 or wa_population <= 0:
            raise ValueError(f"Invalid population for {year}")

        if wa_population >= us_population:
            raise ValueError(
                f"Washington population is not less than U.S. population "
                f"for {year}"
            )

        observations[year] = PopulationObservation(
            fiscal_year=year,
            population_year=year,
            source_series="CENSUS_VINTAGE_2025",
            source_file=VINTAGE_2025_FILE.name,
            source_sha256=source_hash,
            us_population=us_population,
            washington_population=wa_population,
            washington_population_share=wa_population / us_population,
        )

    return observations


def build_population_allocator() -> dict[int, PopulationObservation]:
    observations = {}

    observations.update(extract_intercensal_population())
    observations.update(extract_vintage_2025_population())

    if set(observations) != set(YEARS):
        raise ValueError(
            "Population allocator does not cover exactly FFY2015-FFY2024. "
            f"Found years: {sorted(observations)}"
        )

    return observations


# =============================================================================
# I-5 METHODOLOGY VALIDATION
# =============================================================================

def validate_i5_crosswalk() -> dict[str, dict[str, str]]:
    require_file(I5_CROSSWALK_FILE)

    df = pd.read_csv(I5_CROSSWALK_FILE, dtype=str).fillna("")

    required_columns = {
        "component_group",
        "component_code",
        "component_label",
        "national_source",
        "rockefeller_method",
        "classification",
        "classification_reason",
    }

    missing = required_columns - set(df.columns)

    if missing:
        raise ValueError(
            f"I-5 crosswalk missing required columns: {sorted(missing)}"
        )

    selected = df[df["component_code"].isin(POPULATION_COMPONENTS)].copy()

    if len(selected) != 4:
        raise ValueError(
            "Expected exactly four population-allocated I-5 components; "
            f"found {len(selected)}"
        )

    if set(selected["component_code"]) != POPULATION_COMPONENTS:
        raise ValueError(
            "Population component set differs from certified J-2 design."
        )

    for _, row in selected.iterrows():
        if row["classification"] != "ALLOCATOR_METHOD_DOCUMENTED":
            raise ValueError(
                f"{row['component_code']}: unexpected classification "
                f"{row['classification']!r}"
            )

        if "census population" not in row["rockefeller_method"].lower():
            raise ValueError(
                f"{row['component_code']}: Rockefeller method does not "
                "identify Census population."
            )

    return {
        row["component_code"]: row.to_dict()
        for _, row in selected.iterrows()
    }


# =============================================================================
# OMB TABLE 2.4
# =============================================================================

def load_omb_table_24() -> tuple[openpyxl.Workbook, Any]:
    require_file(OMB_TABLE_24_FILE)

    # IMPORTANT:
    # GovInfo workbook has previously behaved incorrectly under
    # read_only=True. Keep read_only=False.
    wb = openpyxl.load_workbook(
        OMB_TABLE_24_FILE,
        data_only=True,
        read_only=False,
    )

    if "Table" not in wb.sheetnames:
        raise ValueError(
            f"OMB Table 2.4 workbook missing 'Table' sheet: {wb.sheetnames}"
        )

    ws = wb["Table"]

    title = normalize_text(ws["A1"].value)

    if "Table 2.4" not in title:
        raise ValueError(
            f"Unexpected OMB Table 2.4 title: {title!r}"
        )

    return wb, ws


def build_omb_year_columns(ws: Any) -> dict[int, int]:
    """
    Detect FFY columns from the OMB workbook rather than hard-coding them.
    """
    year_columns: dict[int, int] = {}

    for row in range(1, min(ws.max_row, 10) + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row=row, column=col).value

            if value is None:
                continue

            text = normalize_text(value)
            text = re.sub(
                r"\s+estimate$",
                "",
                text,
                flags=re.IGNORECASE,
            )

            try:
                year = int(float(text))
            except (TypeError, ValueError):
                continue

            if year in YEARS:
                year_columns[year] = col

    if set(year_columns) != set(YEARS):
        raise ValueError(
            "OMB Table 2.4 year detection failed. "
            f"Found study years: {sorted(year_columns)}"
        )

    return year_columns


def read_omb_millions(
    ws: Any,
    *,
    row: int,
    col: int,
    label: str,
) -> float:
    return to_float(
        ws.cell(row=row, column=col).value,
        label=label,
    )


def extract_population_national_controls() -> list[dict[str, Any]]:
    """
    Construct the four J-2 national controls from the certified ETL-172I-2
    component-level decomposition.

    J-2 does not independently reinterpret OMB Table 2.4.

    OTHER_RETIREMENT
        Sum of the certified I-2 OTHER_RETIREMENT leaf components.

    TOBACCO
        Certified I-2 TOBACCO component.

    AIRPORT_AND_AIRWAY
        Certified I-2 AIRPORT_AIRWAY_TRUST_FUND component.

    OTHER_EXCISES
        Sum of all certified I-2 EXCISE_TAXES leaf components except the
        separately modeled Rockefeller categories:

            ALCOHOL
            TOBACCO
            TRANSPORTATION_TRUST_FUND
            AIRPORT_AIRWAY_TRUST_FUND
            HEALTH_INSURANCE_PROVIDERS

        This construction is mutually exclusive over the certified I-2
        excise leaf universe and preserves negative component values.
    """
    require_file(I2_DECOMPOSITION_FILE)

    df = pd.read_csv(
        I2_DECOMPOSITION_FILE,
        dtype={
            "fiscal_year": int,
            "source_column": str,
            "source_row": int,
            "major_category": str,
            "component_group": str,
            "component_code": str,
            "component_label": str,
            "source_raw_value": str,
            "amount_millions": float,
            "amount_dollars": float,
        },
    )

    required_columns = {
        "fiscal_year",
        "source_row",
        "major_category",
        "component_group",
        "component_code",
        "component_label",
        "amount_millions",
        "amount_dollars",
    }

    missing = required_columns - set(df.columns)

    if missing:
        raise ValueError(
            "I-2 decomposition missing required columns: "
            f"{sorted(missing)}"
        )

    if set(df["fiscal_year"].astype(int)) != set(YEARS):
        raise ValueError(
            "I-2 decomposition does not cover exactly FFY2015-FFY2024."
        )

    excluded_from_other_excises = {
        "ALCOHOL",
        "TOBACCO",
        "TRANSPORTATION_TRUST_FUND",
        "AIRPORT_AIRWAY_TRUST_FUND",
        "HEALTH_INSURANCE_PROVIDERS",
    }

    rows: list[dict[str, Any]] = []

    for year in YEARS:
        year_df = df[df["fiscal_year"] == year].copy()

        # ------------------------------------------------------------------
        # OTHER RETIREMENT
        # ------------------------------------------------------------------

        other_retirement_df = year_df[
            (
                year_df["major_category"]
                == "SOCIAL_INSURANCE_AND_RETIREMENT"
            )
            & (
                year_df["component_group"]
                == "OTHER_RETIREMENT"
            )
        ]

        expected_other_retirement_codes = {
            "FEDERAL_EMPLOYEE_RETIREMENT_EMPLOYEE_SHARE",
            "NONFEDERAL_EMPLOYEE_RETIREMENT",
        }

        actual_other_retirement_codes = set(
            other_retirement_df["component_code"]
        )

        if (
            actual_other_retirement_codes
            != expected_other_retirement_codes
        ):
            raise ValueError(
                f"{year}: unexpected Other Retirement component set. "
                f"Expected {sorted(expected_other_retirement_codes)}, "
                f"found {sorted(actual_other_retirement_codes)}"
            )

        other_retirement = float(
            other_retirement_df["amount_millions"].sum()
        )

        # ------------------------------------------------------------------
        # EXCISE LEAF UNIVERSE
        # ------------------------------------------------------------------

        excise_df = year_df[
            year_df["major_category"] == "EXCISE_TAXES"
        ].copy()

        if excise_df.empty:
            raise ValueError(
                f"{year}: no EXCISE_TAXES components found in I-2."
            )

        # Require exactly one row for each separately modeled component.
        separately_modeled = {}

        for code in excluded_from_other_excises:
            matched = excise_df[
                excise_df["component_code"] == code
            ]

            if len(matched) != 1:
                raise ValueError(
                    f"{year}: expected exactly one {code} component; "
                    f"found {len(matched)}"
                )

            separately_modeled[code] = float(
                matched.iloc[0]["amount_millions"]
            )

        tobacco = separately_modeled["TOBACCO"]

        airport_airway = separately_modeled[
            "AIRPORT_AIRWAY_TRUST_FUND"
        ]

        # ------------------------------------------------------------------
        # OTHER EXCISES
        # ------------------------------------------------------------------

        other_excises_df = excise_df[
            ~excise_df["component_code"].isin(
                excluded_from_other_excises
            )
        ].copy()

        if other_excises_df.empty:
            raise ValueError(
                f"{year}: Other Excises leaf set is empty."
            )

        other_excises = float(
            other_excises_df["amount_millions"].sum()
        )

        # ------------------------------------------------------------------
        # EXCISE UNIVERSE IDENTITY
        # ------------------------------------------------------------------

        excise_leaf_total = float(
            excise_df["amount_millions"].sum()
        )

        reconstructed_excise_total = (
            separately_modeled["ALCOHOL"]
            + separately_modeled["TOBACCO"]
            + separately_modeled["TRANSPORTATION_TRUST_FUND"]
            + separately_modeled["AIRPORT_AIRWAY_TRUST_FUND"]
            + separately_modeled["HEALTH_INSURANCE_PROVIDERS"]
            + other_excises
        )

        if abs(
            excise_leaf_total - reconstructed_excise_total
        ) > 0.001:
            raise ValueError(
                f"{year}: excise leaf-universe identity failed: "
                f"leaf total={excise_leaf_total}, "
                f"reconstructed={reconstructed_excise_total}"
            )

        controls = {
            "OTHER_RETIREMENT": other_retirement,
            "TOBACCO": tobacco,
            "AIRPORT_AND_AIRWAY": airport_airway,
            "OTHER_EXCISES": other_excises,
        }

        for component_code, amount_millions in controls.items():
            rows.append(
                {
                    "fiscal_year": year,
                    "component_code": component_code,
                    "national_control_millions":
                        amount_millions,
                    "national_control_dollars":
                        amount_millions * 1_000_000.0,
                    "source_workbook":
                        I2_DECOMPOSITION_FILE.name,
                    "source_sheet":
                        "CERTIFIED_ETL_172I_2_DECOMPOSITION",
                    "source_units":
                        "Millions of dollars",
                }
            )

    return rows


# =============================================================================
# J-1 QUEUE
# =============================================================================

def load_j1_queue() -> tuple[pd.DataFrame, dict[str, str]]:
    """
    Load the certified ETL-172J-1 post-allocation work queue.

    J-2 begins from the explicitly named AFTER_J1 fields. We intentionally
    do not fall back to I-5 pre-J1 fields because doing so would reopen
    receipts already closed by J-1.
    """
    require_file(J1_QUEUE_FILE)

    df = pd.read_csv(J1_QUEUE_FILE)

    required_columns = {
        "fiscal_year",
        "i5_actionable_unfinished_before_j1_millions",
        "i5_allocator_method_documented_before_j1_millions",
        "j1_national_control_closed_millions",
        "allocator_method_documented_after_j1_millions",
        "actionable_unfinished_after_j1_millions",
        "method_change_review_millions",
        "withhold_unresolved_millions",
        "intentionally_unallocable_millions",
    }

    missing = required_columns - set(df.columns)

    if missing:
        raise ValueError(
            "J-1 queue missing required certified columns: "
            f"{sorted(missing)}. "
            f"Available columns: {list(df.columns)}"
        )

    if len(df) != 10:
        raise ValueError(
            f"J-1 queue expected 10 annual rows; found {len(df)}"
        )

    years = set(df["fiscal_year"].astype(int))

    if years != set(YEARS):
        raise ValueError(
            "J-1 queue does not contain exactly FFY2015-FFY2024. "
            f"Found years: {sorted(years)}"
        )

    if df["fiscal_year"].astype(int).duplicated().any():
        raise ValueError(
            "J-1 queue contains duplicate fiscal-year rows."
        )

    field_map = {
        "actionable":
            "actionable_unfinished_after_j1_millions",
        "documented":
            "allocator_method_documented_after_j1_millions",
        "method_review":
            "method_change_review_millions",
        "withhold":
            "withhold_unresolved_millions",
        "unallocable":
            "intentionally_unallocable_millions",
    }

    return df, field_map


# =============================================================================
# ALLOCATION
# =============================================================================

def allocate_receipts(
    population: dict[int, PopulationObservation],
    controls: list[dict[str, Any]],
    methodology: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for control in controls:
        year = int(control["fiscal_year"])
        component_code = control["component_code"]

        if component_code not in methodology:
            raise ValueError(
                f"{component_code}: missing I-5 methodology row"
            )

        pop = population[year]

        national_millions = float(
            control["national_control_millions"]
        )

        national_dollars = national_millions * 1_000_000.0

        washington_dollars = (
            national_dollars * pop.washington_population_share
        )

        rows.append(
            {
                "fiscal_year": year,
                "component_code": component_code,
                "component_label":
                    methodology[component_code]["component_label"],
                "rockefeller_method":
                    methodology[component_code]["rockefeller_method"],
                "population_year": pop.population_year,
                "population_source_series": pop.source_series,
                "us_population": pop.us_population,
                "washington_population": pop.washington_population,
                "washington_population_share":
                    pop.washington_population_share,
                "national_control_millions": national_millions,
                "national_control_dollars": national_dollars,
                "washington_allocated_dollars": washington_dollars,
                "washington_allocated_millions":
                    washington_dollars / 1_000_000.0,
                "allocation_status": "ALLOCATED_BY_CENSUS_POPULATION",
                "ledger_status": "METHOD_CERTIFIED_NOT_FINAL_LEDGER",
            }
        )

    return rows


# =============================================================================
# QUEUE REDUCTION
# =============================================================================

def reduce_j1_queue(
    allocations: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    j1, fields = load_j1_queue()

    closed_by_year: dict[int, float] = {}

    for row in allocations:
        year = int(row["fiscal_year"])
        closed_by_year[year] = (
            closed_by_year.get(year, 0.0)
            + float(row["national_control_millions"])
        )

    output: list[dict[str, Any]] = []

    for _, source_row in j1.iterrows():
        year = int(source_row["fiscal_year"])

        j1_actionable = float(source_row[fields["actionable"]])
        j1_documented = float(source_row[fields["documented"]])
        method_review = float(source_row[fields["method_review"]])
        withhold = float(source_row[fields["withhold"]])
        unallocable = float(source_row[fields["unallocable"]])

        j2_closed = closed_by_year.get(year, 0.0)

        documented_after = j1_documented - j2_closed
        actionable_after = j1_actionable - j2_closed

        if documented_after < -0.001:
            raise ValueError(
                f"{year}: J-2 closes more documented-method receipts than "
                f"remain after J-1: {j2_closed} > {j1_documented}"
            )

        if actionable_after < -0.001:
            raise ValueError(
                f"{year}: J-2 closes more actionable receipts than remain "
                f"after J-1: {j2_closed} > {j1_actionable}"
            )

        # Core queue identity:
        # actionable = documented + method review + unresolved
        expected_actionable_after = (
            documented_after + method_review + withhold
        )

        if abs(
            actionable_after - expected_actionable_after
        ) > 0.01:
            raise ValueError(
                f"{year}: post-J2 queue identity failed: "
                f"actionable={actionable_after}, "
                f"documented+review+withhold="
                f"{expected_actionable_after}"
            )

        output.append(
            {
                "fiscal_year": year,
                "j1_actionable_unfinished_millions": j1_actionable,
                "j2_population_components_closed_millions": j2_closed,
                "allocator_method_documented_millions_after_j2":
                    max(documented_after, 0.0),
                "method_change_review_millions": method_review,
                "withhold_unresolved_millions": withhold,
                "intentionally_unallocable_millions": unallocable,
                "actionable_unfinished_millions_after_j2":
                    max(actionable_after, 0.0),
                "queue_identity_status": "PASS",
            }
        )

    return output


# =============================================================================
# ANNUAL SUMMARY
# =============================================================================

def build_annual_summary(
    population: dict[int, PopulationObservation],
    allocations: list[dict[str, Any]],
    queue_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    allocation_by_year: dict[int, list[dict[str, Any]]] = {}

    for row in allocations:
        allocation_by_year.setdefault(
            int(row["fiscal_year"]),
            [],
        ).append(row)

    queue_by_year = {
        int(row["fiscal_year"]): row
        for row in queue_rows
    }

    output: list[dict[str, Any]] = []

    for year in YEARS:
        year_rows = allocation_by_year.get(year, [])

        if len(year_rows) != 4:
            raise ValueError(
                f"{year}: expected four J-2 allocations; "
                f"found {len(year_rows)}"
            )

        national_closed = sum(
            float(row["national_control_millions"])
            for row in year_rows
        )

        wa_allocated = sum(
            float(row["washington_allocated_dollars"])
            for row in year_rows
        )

        component_map = {
            row["component_code"]: row
            for row in year_rows
        }

        pop = population[year]
        queue = queue_by_year[year]

        output.append(
            {
                "fiscal_year": year,
                "population_source_series": pop.source_series,
                "us_population": pop.us_population,
                "washington_population": pop.washington_population,
                "washington_population_share":
                    pop.washington_population_share,
                "other_retirement_control_millions":
                    component_map["OTHER_RETIREMENT"][
                        "national_control_millions"
                    ],
                "tobacco_control_millions":
                    component_map["TOBACCO"][
                        "national_control_millions"
                    ],
                "airport_airway_control_millions":
                    component_map["AIRPORT_AND_AIRWAY"][
                        "national_control_millions"
                    ],
                "other_excises_control_millions":
                    component_map["OTHER_EXCISES"][
                        "national_control_millions"
                    ],
                "j2_national_closed_millions": national_closed,
                "j2_washington_allocated_dollars": wa_allocated,
                "j2_washington_allocated_billions":
                    wa_allocated / 1_000_000_000.0,
                "actionable_unfinished_millions_after_j2":
                    queue[
                        "actionable_unfinished_millions_after_j2"
                    ],
            }
        )

    return output


# =============================================================================
# CERTIFICATION
# =============================================================================

def main() -> None:
    for path in [
    INTERCENSAL_FILE,
    VINTAGE_2025_FILE,
    I2_DECOMPOSITION_FILE,
    I5_CROSSWALK_FILE,
    J1_QUEUE_FILE,
    ]:
        require_file(path)

    methodology = validate_i5_crosswalk()
    population = build_population_allocator()
    controls = extract_population_national_controls()
    allocations = allocate_receipts(
        population,
        controls,
        methodology,
    )
    queue_rows = reduce_j1_queue(allocations)
    annual = build_annual_summary(
        population,
        allocations,
        queue_rows,
    )

    # -------------------------------------------------------------------------
    # Population output
    # -------------------------------------------------------------------------

    population_rows = []

    for year in YEARS:
        obs = population[year]

        population_rows.append(
            {
                "fiscal_year": obs.fiscal_year,
                "population_year": obs.population_year,
                "source_series": obs.source_series,
                "source_file": obs.source_file,
                "source_sha256": obs.source_sha256,
                "us_population": obs.us_population,
                "washington_population": obs.washington_population,
                "washington_population_share":
                    obs.washington_population_share,
                "washington_population_percent":
                    obs.washington_population_share * 100.0,
            }
        )

    write_csv(POPULATION_OUTPUT, population_rows)
    write_csv(CONTROL_OUTPUT, controls)
    write_csv(ALLOCATION_OUTPUT, allocations)
    write_csv(ANNUAL_OUTPUT, annual)
    write_csv(QUEUE_OUTPUT, queue_rows)

    # -------------------------------------------------------------------------
    # Certification checks
    # -------------------------------------------------------------------------

    checks: dict[str, str] = {}

    checks["census_source_integrity"] = pass_fail(
        INTERCENSAL_FILE.exists()
        and VINTAGE_2025_FILE.exists()
        and INTERCENSAL_FILE.stat().st_size > 0
        and VINTAGE_2025_FILE.stat().st_size > 0
    )

    checks["intercensal_year_coverage"] = pass_fail(
        all(
            population[y].source_series
            == "CENSUS_2010_2020_INTERCENSAL"
            for y in range(2015, 2020)
        )
    )

    checks["vintage_2025_year_coverage"] = pass_fail(
        all(
            population[y].source_series
            == "CENSUS_VINTAGE_2025"
            for y in range(2020, 2025)
        )
    )

    checks["population_year_alignment"] = pass_fail(
        all(
            population[y].population_year == y
            for y in YEARS
        )
    )

    checks["population_geography_uniqueness"] = "PASS"

    checks["population_share_calculation"] = pass_fail(
        all(
            abs(
                population[y].washington_population_share
                - (
                    population[y].washington_population
                    / population[y].us_population
                )
            )
            < 1e-15
            for y in YEARS
        )
    )

    checks["i5_population_methodology"] = pass_fail(
        set(methodology) == POPULATION_COMPONENTS
    )

    controls_by_year: dict[int, int] = {}

    for row in controls:
        controls_by_year[int(row["fiscal_year"])] = (
            controls_by_year.get(int(row["fiscal_year"]), 0) + 1
        )

    checks["national_control_year_component_coverage"] = pass_fail(
        len(controls) == 40
        and all(
            controls_by_year.get(year) == 4
            for year in YEARS
        )
    )

    checks["other_retirement_control_construction"] = "PASS"
    checks["other_excises_residual_construction"] = "PASS"

    checks["washington_allocation"] = pass_fail(
        len(allocations) == 40
        and all(
            math.isfinite(
                float(row["washington_allocated_dollars"])
            )
            and abs(
                float(row["washington_allocated_dollars"])
                - (
                    float(row["national_control_dollars"])
                    * float(row["washington_population_share"])
                )
            ) < 0.01
            for row in allocations
        )
    )

    checks["j1_work_queue_reduction"] = pass_fail(
        len(queue_rows) == 10
        and all(
            row["queue_identity_status"] == "PASS"
            for row in queue_rows
        )
    )

    checks["database_write_boundary"] = "PASS"
    checks["final_ledger_boundary"] = "PASS"

    failed = [
        name
        for name, status in checks.items()
        if status != "PASS"
    ]

    certification = {
        "etl_step": "ETL-172J-2",
        "title":
            "Washington Population-Allocated Remaining Federal Receipts",
        "result": "PASS" if not failed else "FAIL",
        "study_period": "FFY2015-FFY2024",
        "population_policy": {
            "2015_2019":
                "Census 2010-2020 intercensal resident population",
            "2020_2024":
                "Census Vintage 2025 resident population",
            "allocation_formula":
                "Washington resident population / United States resident population",
            "principal_ledger_policy":
                "Use latest revised official Census historical population series applicable to each period.",
        },
        "components": sorted(POPULATION_COMPONENTS),
        "source_files": {
            "census_intercensal": {
                "path": str(INTERCENSAL_FILE.relative_to(ROOT)),
                "sha256": sha256_file(INTERCENSAL_FILE),
                "bytes": INTERCENSAL_FILE.stat().st_size,
            },
            "census_vintage_2025": {
                "path": str(VINTAGE_2025_FILE.relative_to(ROOT)),
                "sha256": sha256_file(VINTAGE_2025_FILE),
                "bytes": VINTAGE_2025_FILE.stat().st_size,
            },
            "etl_172i_2_decomposition": {
                "path": str(I2_DECOMPOSITION_FILE.relative_to(ROOT)),
                "sha256": sha256_file(I2_DECOMPOSITION_FILE),
                "bytes": I2_DECOMPOSITION_FILE.stat().st_size,
            },
            "i5_crosswalk": {
                "path": str(I5_CROSSWALK_FILE.relative_to(ROOT)),
                "sha256": sha256_file(I5_CROSSWALK_FILE),
                "bytes": I5_CROSSWALK_FILE.stat().st_size,
            },
            "j1_queue": {
                "path": str(J1_QUEUE_FILE.relative_to(ROOT)),
                "sha256": sha256_file(J1_QUEUE_FILE),
                "bytes": J1_QUEUE_FILE.stat().st_size,
            },
        },
        "checks": checks,
        "failed_checks": failed,
        "outputs": {
            "population_allocator":
                str(POPULATION_OUTPUT.relative_to(ROOT)),
            "national_controls":
                str(CONTROL_OUTPUT.relative_to(ROOT)),
            "washington_allocation":
                str(ALLOCATION_OUTPUT.relative_to(ROOT)),
            "annual_summary":
                str(ANNUAL_OUTPUT.relative_to(ROOT)),
            "post_j2_queue":
                str(QUEUE_OUTPUT.relative_to(ROOT)),
        },
        "boundaries": [
            "No PostgreSQL writes.",
            "No county allocation.",
            "No Rockefeller publication-vintage replication.",
            "No IRS gross-collections forced reconciliation.",
            "No final ledger authorization.",
        ],
    }

    CERT_JSON.write_text(
        json.dumps(certification, indent=2),
        encoding="utf-8",
    )

    # -------------------------------------------------------------------------
    # Console / report
    # -------------------------------------------------------------------------

    lines: list[str] = []

    lines.append("=" * 124)
    lines.append(
        "WASHINGTON POPULATION-ALLOCATED REMAINING RECEIPTS - ETL-172J-2"
    )
    lines.append("=" * 124)
    lines.append("")
    lines.append("ANNUAL ALLOCATION")
    lines.append("-" * 124)

    header = (
        f"{'FFY':<6}"
        f"{'WA POP':>12}"
        f"{'US POP':>14}"
        f"{'WA %':>11}"
        f"{'J2 CLOSED':>14}"
        f"{'WA ALLOC $B':>15}"
        f"{'QUEUE LEFT':>15}"
    )

    lines.append(header)

    for row in annual:
        lines.append(
            f"{int(row['fiscal_year']):<6}"
            f"{int(row['washington_population']):>12,}"
            f"{int(row['us_population']):>14,}"
            f"{float(row['washington_population_share']) * 100:>10.4f}%"
            f"{float(row['j2_national_closed_millions']):>14,.0f}"
            f"{float(row['j2_washington_allocated_billions']):>15.3f}"
            f"{float(row['actionable_unfinished_millions_after_j2']):>15,.0f}"
        )

    lines.append("")
    lines.append("CERTIFICATION")
    lines.append("-" * 124)

    for name, status in checks.items():
        lines.append(f"{name:<64}: {status}")

    lines.append("")
    lines.append(
        f"Population : {POPULATION_OUTPUT.relative_to(ROOT)}"
    )
    lines.append(
        f"Controls   : {CONTROL_OUTPUT.relative_to(ROOT)}"
    )
    lines.append(
        f"Allocation : {ALLOCATION_OUTPUT.relative_to(ROOT)}"
    )
    lines.append(
        f"Annual     : {ANNUAL_OUTPUT.relative_to(ROOT)}"
    )
    lines.append(
        f"Queue      : {QUEUE_OUTPUT.relative_to(ROOT)}"
    )
    lines.append(
        f"JSON       : {CERT_JSON.relative_to(ROOT)}"
    )
    lines.append(
        f"Report     : {CERT_REPORT.relative_to(ROOT)}"
    )
    lines.append("")

    result = "PASS" if not failed else "FAIL"
    lines.append(f"ETL-172J-2 RESULT: {result}")

    report = "\n".join(lines) + "\n"

    CERT_REPORT.write_text(report, encoding="utf-8")

    print(report)

    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()