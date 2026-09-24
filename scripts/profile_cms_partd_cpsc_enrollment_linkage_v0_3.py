#!/usr/bin/env python3
"""Memory-safe CMS Part D CPSC linkage profiler v0.3.

Scope: Washington, Jan 2015-Sep 2024 (117 months).

Key change from v0.2: source indexing retains metadata only. Selected monthly
CPSC files are streamed one at a time and released. No full CSV corpus is held
in memory.

READ ONLY: no payment reconstruction, imputation, DB writes, or ledger inclusion.
RDS remains out of scope.
"""
from __future__ import annotations
import argparse, csv, hashlib, io, re, sys, zipfile
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional
import openpyxl

START=(2015,1); END=(2024,9)
PAYMENT_FIELDS={"CONTRACT_NUMBER","PLAN_BENEFIT_PACKAGE","CONTRACT_NAME",
"ORGANIZATION_TYPE","AVERAGE_PART_D_DIRECT_SUBSIDY_PM_PM_PAYMENT",
"AVERAGE_PART_D_RISK_SCORE","AVERAGE_REINSURANCE_PM_PM_PAYMENT",
"AVERAGE_LOW_INCOME_COST_SHARING_PM_PM_PAYMENT"}
REQ={"contract":{"CONTRACT_NUMBER","CONTRACT_ID","CONTRACT"},
"pbp":{"PLAN_ID","PLAN_BENEFIT_PACKAGE","PBP","PLAN_BENEFIT_PACKAGE_ID"},
"state":{"STATE","STATE_ABBREVIATION","STATE_ABBR"},
"county":{"COUNTY","COUNTY_NAME"},
"enrollment":{"ENROLLMENT","ENROLLMENT_COUNT","ENROLLEES"}}
OPT={"plan_type":{"PLAN_TYPE","PLAN_TYPE_DESCRIPTION"},
"organization_type":{"ORGANIZATION_TYPE","ORG_TYPE"}}
MONTHS={1:("JAN","JANUARY"),2:("FEB","FEBRUARY"),3:("MAR","MARCH"),
4:("APR","APRIL"),5:("MAY",),6:("JUN","JUNE"),7:("JUL","JULY"),
8:("AUG","AUGUST"),9:("SEP","SEPT","SEPTEMBER"),10:("OCT","OCTOBER"),
11:("NOV","NOVEMBER"),12:("DEC","DECEMBER")}

def clean(v): return "" if v is None else str(v).strip()
def norm(v):
    return re.sub(r"[^A-Z0-9]+","_",clean(v).upper().replace("/"," ")).strip("_")
def pbp(v):
    s=clean(v)
    if not s:return ""
    try:return f"{int(float(s)):03d}"
    except:return s.upper()
def enrollment(v):
    s=clean(v).replace(",","")
    if re.fullmatch(r"\d+",s): return int(s),"NUMERIC"
    return None,"NONNUMERIC"
def month_iter():
    y,m=START
    while (y,m)<=END:
        yield y,m
        m+=1
        if m==13:y,m=y+1,1
def sha_file(p):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1048576),b""):h.update(b)
    return h.hexdigest()
def sha_member(p,n):
    h=hashlib.sha256()
    with zipfile.ZipFile(p) as z, z.open(n) as f:
        for b in iter(lambda:f.read(1048576),b""):h.update(b)
    return h.hexdigest()

def infer_date(text):
    u=text.upper(); c=re.sub(r"[^A-Z0-9]","",u)
    a={(int(y),int(m)) for y,m in re.findall(r"(20(?:1[4-9]|2[0-4]))(0[1-9]|1[0-2])",c)}
    if len(a)==1:return (*next(iter(a)),"YYYYMM")
    a={(int(y),int(m)) for m,y in re.findall(r"(0[1-9]|1[0-2])(20(?:1[4-9]|2[0-4]))",c)}
    if len(a)==1:return (*next(iter(a)),"MMYYYY")
    ys={int(x) for x in re.findall(r"20(?:1[4-9]|2[0-4])",u)}
    ms={m for m,names in MONTHS.items() if any(re.search(rf"(?<![A-Z]){n}(?![A-Z])",u) for n in names)}
    if len(ys)==len(ms)==1:return next(iter(ys)),next(iter(ms)),"YEAR_MONTH_NAME"
    return None,None,"UNRESOLVED"

@contextmanager
def text_source(path,member=None):
    if path.suffix.lower()==".zip":
        z=zipfile.ZipFile(path); raw=z.open(member)
        t=io.TextIOWrapper(raw,encoding="utf-8-sig",errors="replace",newline="")
        try:yield t
        finally:t.close(); z.close()
    else:
        f=path.open("r",encoding="utf-8-sig",errors="replace",newline="")
        try:yield f
        finally:f.close()

def header_info(path,member):
    with text_source(path,member) as f:
        sample=f.read(8192)
        delim=max([",","\t","|"],key=lambda x:sample.count(x))
        f.seek(0); r=csv.reader(f,delimiter=delim)
        for rn,row in enumerate(r,1):
            if rn>50:break
            h=[norm(x) for x in row]; mp={}; ok=True
            for k,aliases in REQ.items():
                pos=next((h.index(a) for a in aliases if a in h),None)
                if pos is None:ok=False;break
                mp[k]=pos
            if ok:
                for k,aliases in OPT.items():
                    mp[k]=next((h.index(a) for a in aliases if a in h),None)
                return rn,mp,delim
    return None,None,None

def build_index(root):
    out=[]; scanned=0
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".zip",".csv",".txt"}:continue
        parts={x.lower() for x in path.parts}
        if "validation" in parts or "plan_payment" in parts:continue
        scanned+=1
        if path.suffix.lower()==".zip":
            try:
                with zipfile.ZipFile(path) as z:
                    members=[n for n in z.namelist() if Path(n).suffix.lower() in {".csv",".txt"}]
            except:continue
        else:members=[None]
        for member in members:
            try:rn,mp,d=header_info(path,member)
            except:continue
            if rn is None:continue
            y,m,method=infer_date(f"{path} {member or path.name}")
            out.append({"path":path,"member":member,"year":y,"month":m,
                        "method":method,"header":rn,"map":mp,"delim":d})
    return out,scanned

def choose(idx,y,m):
    a=[x for x in idx if x["year"]==y and x["month"]==m]
    def rank(x):
        s=re.sub(r"[^a-z0-9]","",f"{x['path']} {x['member'] or ''}".lower())
        score=(1000 if "cpsc" in s else 0)+(1000 if "contractplanstatecounty" in s else 0)+(100 if "enrollment" in s else 0)
        if "snp" in s:score-=500
        if "scc" in s and "cpsc" not in s:score-=500
        return -score,str(x["path"]),(x["member"] or "")
    a.sort(key=rank)
    return (a[0] if a else None),a

def wa_rows(src):
    with text_source(src["path"],src["member"]) as f:
        r=csv.reader(f,delimiter=src["delim"]); mp=src["map"]
        for rn,row in enumerate(r,1):
            if rn<=src["header"]:continue
            def g(k):
                i=mp.get(k); return row[i] if i is not None and i<len(row) else ""
            if clean(g("state")).upper() not in {"WA","WASHINGTON"}:continue
            yield {"contract":clean(g("contract")).upper(),"pbp":pbp(g("pbp")),
                   "county":clean(g("county")),"enrollment":clean(g("enrollment")),
                   "plan_type":clean(g("plan_type")),"org":clean(g("organization_type"))}

def payment_keys(root,y):
    found=[]
    for zpath in sorted((root/str(y)).glob("*.zip")):
        try:
            with zipfile.ZipFile(zpath) as z:
                for n in z.namelist():
                    b=Path(n).name.lower().replace(" ","")
                    if b.endswith(".xlsx") and "partdplans" in b and "recon" not in b:
                        data=z.read(n); found.append((zpath,n,data,hashlib.sha256(data).hexdigest()))
        except:pass
    if not found:raise RuntimeError(f"{y}: PartDPlans workbook not found")
    if len({x[3] for x in found})!=1:raise RuntimeError(f"{y}: conflicting PartDPlans workbook bytes")
    zp,n,data,mhash=sorted(found,key=lambda x:(str(x[0]),x[1]))[0]
    wb=openpyxl.load_workbook(io.BytesIO(data),data_only=True,read_only=True)
    try:
        for ws in wb.worksheets:
            for hr,row in enumerate(ws.iter_rows(min_row=1,max_row=30,values_only=True),1):
                h=[norm(x) for x in row]
                if not PAYMENT_FIELDS.issubset(set(h)):continue
                pos={k:h.index(k) for k in PAYMENT_FIELDS}; keys={}; dups=[]
                for vals in ws.iter_rows(min_row=hr+1,values_only=True):
                    def v(k):
                        i=pos[k]; return vals[i] if i<len(vals) else None
                    key=(clean(v("CONTRACT_NUMBER")).upper(),pbp(v("PLAN_BENEFIT_PACKAGE")))
                    if not all(key):continue
                    rec={"name":clean(v("CONTRACT_NAME")),"org":clean(v("ORGANIZATION_TYPE"))}
                    if key in keys:dups.append(key)
                    else:keys[key]=rec
                return keys,dups,{"zip_path":str(zp),"zip_sha256":sha_file(zp),
                    "member_name":n,"member_sha256":mhash,"sheet_name":ws.title,"header_row":hr}
    finally:wb.close()
    raise RuntimeError(f"{y}: PartDPlans semantic table not found")

def write_csv(p,rows,fields):
    with p.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore");w.writeheader();w.writerows(rows)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--cpsc-root",type=Path,default=Path("data/raw/cms/medicare_advantage"))
    ap.add_argument("--plan-payment-root",type=Path,default=Path("data/raw/cms/medicare_advantage/plan_payment"))
    ap.add_argument("--output-dir",type=Path,default=Path("data/raw/cms/medicare_advantage/validation/partd_cpsc_enrollment_linkage_profile_v0_3"))
    a=ap.parse_args();a.output_dir.mkdir(parents=True,exist_ok=True)
    errors=[];warnings=[]
    print("Building metadata-only CPSC source index...",flush=True)
    idx,scanned=build_index(a.cpsc_root)
    print(f"Scanned {scanned:,} candidate files; admitted {len(idx):,} CPSC-like members.",flush=True)

    pkeys={}; pmanifest=[]; dups=[]
    for y in range(2015,2025):
        try:
            k,d,meta=payment_keys(a.plan_payment_root,y);pkeys[y]=k
            pmanifest.append({"payment_year":y,"payment_key_count":len(k),**meta})
            dups += [{"payment_year":y,"contract":x[0],"pbp":x[1]} for x in d]
        except Exception as e:errors.append(str(e));pkeys[y]={}

    monthly=[]; manifest=[]; vocab=Counter(); nomatch=Counter(); absent=[]
    for y,m in month_iter():
        print(f"{y}-{m:02d}: starting...",flush=True)
        src,cands=choose(idx,y,m)
        if not src:
            errors.append(f"{y}-{m:02d}: no indexed CPSC source resolved")
            monthly.append({"year":y,"month":m,"source_status":"MISSING","wa_rows":0,
                "numeric_enrollment":0,"partd_match_numeric_enrollment":0,
                "no_partd_payment_match_numeric_enrollment":0,
                "partd_match_nonnumeric_rows":0,"no_partd_payment_match_nonnumeric_rows":0,
                "numeric_accounting_difference":0})
            print(f"{y}-{m:02d}: MISSING",flush=True);continue
        hashes=[]
        for c in cands:
            try:hashes.append(sha_member(c["path"],c["member"]) if c["path"].suffix.lower()==".zip" else sha_file(c["path"]))
            except Exception as e:warnings.append(f"{y}-{m:02d}: candidate hash error: {e}")
        if len(set(hashes))>1:warnings.append(f"{y}-{m:02d}: {len(cands)} candidates / {len(set(hashes))} distinct hashes")
        sh=sha_member(src["path"],src["member"]) if src["path"].suffix.lower()==".zip" else sha_file(src["path"])
        manifest.append({"year":y,"month":m,"selected_path":str(src["path"]),
            "selected_member":src["member"] or "","selected_source_sha256":sh,
            "date_inference_method":src["method"],"candidate_count":len(cands),
            "distinct_candidate_hashes":len(set(hashes))})
        keys=pkeys[y]; seen=set(); wr=tot=match=no=mn=nn=0
        try:
            for r in wa_rows(src):
                wr+=1; n,status=enrollment(r["enrollment"]); key=(r["contract"],r["pbp"]); hit=key in keys
                if hit:
                    seen.add(key);vocab[(y,r["plan_type"],r["org"],keys[key]["org"])]+=n or 0
                if status=="NUMERIC":
                    tot+=n
                    if hit:match+=n
                    else:
                        no+=n;nomatch[(y,m,key[0],key[1],r["plan_type"],r["org"])]+=n
                elif hit:mn+=1
                else:nn+=1
        except Exception as e:
            errors.append(f"{y}-{m:02d}: streaming parse failed: {e}");continue
        for key,p in keys.items():
            if key not in seen:absent.append({"year":y,"month":m,"contract":key[0],"pbp":key[1],
                "contract_name":p["name"],"payment_organization_type":p["org"],
                "classification":"NO_OBSERVED_WA_CPSC_MATCH_THIS_MONTH"})
        monthly.append({"year":y,"month":m,"source_status":"PARSED","wa_rows":wr,
            "numeric_enrollment":tot,"partd_match_numeric_enrollment":match,
            "no_partd_payment_match_numeric_enrollment":no,"partd_match_nonnumeric_rows":mn,
            "no_partd_payment_match_nonnumeric_rows":nn,
            "numeric_accounting_difference":tot-match-no})
        print(f"{y}-{m:02d}: complete | WA numeric={tot:,} | PartD match={match:,} | no PartD payment match={no:,}",flush=True)

    vocabrows=[{"year":y,"cpsc_plan_type":pt,"cpsc_organization_type":co,
        "payment_organization_type":po,"partd_match_numeric_enrollment":n}
        for (y,pt,co,po),n in sorted(vocab.items())]
    nomatchrows=[{"year":y,"month":m,"contract":c,"pbp":p,"cpsc_plan_type":pt,
        "cpsc_organization_type":o,"numeric_enrollment":n,"classification":"NO_PARTD_PAYMENT_MATCH"}
        for (y,m,c,p,pt,o),n in sorted(nomatch.items())]
    lookup={(r["year"],r["month"]):r for r in monthly};matrix=[]
    for y in range(2015,2025):
        row={"year":y}
        for m in range(1,13):
            row[f"m{m:02d}"]="OUT_OF_SCOPE" if (y,m)<START or (y,m)>END else lookup.get((y,m),{"source_status":"NOT_RECORDED"})["source_status"]
        matrix.append(row)

    write_csv(a.output_dir/"partd_cpsc_monthly_profile.csv",monthly,
        ["year","month","source_status","wa_rows","numeric_enrollment","partd_match_numeric_enrollment",
         "no_partd_payment_match_numeric_enrollment","partd_match_nonnumeric_rows",
         "no_partd_payment_match_nonnumeric_rows","numeric_accounting_difference"])
    write_csv(a.output_dir/"partd_cpsc_source_resolution_matrix.csv",matrix,["year"]+[f"m{x:02d}" for x in range(1,13)])
    write_csv(a.output_dir/"partd_cpsc_source_manifest.csv",manifest,
        ["year","month","selected_path","selected_member","selected_source_sha256","date_inference_method","candidate_count","distinct_candidate_hashes"])
    write_csv(a.output_dir/"partd_cpsc_partd_match_vocabulary.csv",vocabrows,
        ["year","cpsc_plan_type","cpsc_organization_type","payment_organization_type","partd_match_numeric_enrollment"])
    write_csv(a.output_dir/"partd_cpsc_no_partd_payment_match_aggregates.csv",nomatchrows,
        ["year","month","contract","pbp","cpsc_plan_type","cpsc_organization_type","numeric_enrollment","classification"])
    write_csv(a.output_dir/"partd_payment_rows_without_wa_cpsc_match.csv",absent,
        ["year","month","contract","pbp","contract_name","payment_organization_type","classification"])
    write_csv(a.output_dir/"partd_payment_source_manifest.csv",pmanifest,
        ["payment_year","payment_key_count","zip_path","zip_sha256","member_name","member_sha256","sheet_name","header_row"])
    write_csv(a.output_dir/"partd_payment_duplicate_keys.csv",dups,["payment_year","contract","pbp"])

    required=117; parsed=sum(r["source_status"]=="PARSED" for r in monthly)
    total=sum(r["numeric_enrollment"] for r in monthly);matched=sum(r["partd_match_numeric_enrollment"] for r in monthly)
    no=sum(r["no_partd_payment_match_numeric_enrollment"] for r in monthly);diff=sum(r["numeric_accounting_difference"] for r in monthly)
    mn=sum(r["partd_match_nonnumeric_rows"] for r in monthly);nn=sum(r["no_partd_payment_match_nonnumeric_rows"] for r in monthly)
    report=[
"CMS MEDICARE PART D CPSC ENROLLMENT LINKAGE PROFILE v0.3","="*72,
f"Generated UTC: {datetime.now(timezone.utc).isoformat()}","Scope: Washington, Jan 2015-Sep 2024 (117 months)","",
"BOUNDARY","-"*72,"Memory-safe read-only diagnostic profile.","CPSC is broader than the Part D payment universe.",
"NO_PARTD_PAYMENT_MATCH is not automatically a linkage failure.","No payment reconstruction, imputation, database writes, or ledger inclusion.","RDS remains out of scope.","",
"SOURCE INDEX","-"*72,f"Candidate files scanned once: {scanned:,}",f"CPSC-like members admitted by header: {len(idx):,}",
f"Required months: {required}",f"Parsed months: {parsed}","",
"ENROLLMENT ACCOUNTING","-"*72,f"WA numeric enrollment across monthly county rows: {total:,}",
f"PARTD_PAYMENT_MATCH numeric enrollment: {matched:,}",f"NO_PARTD_PAYMENT_MATCH numeric enrollment: {no:,}",
f"PARTD_PAYMENT_MATCH suppressed/nonnumeric rows: {mn:,}",f"NO_PARTD_PAYMENT_MATCH suppressed/nonnumeric rows: {nn:,}",
f"Numeric accounting difference: {diff:,}","",
"OTHER DIAGNOSTICS","-"*72,f"Payment-plan/month rows without observed WA CPSC match: {len(absent):,}",
f"Duplicate annual Part D payment keys: {len(dups):,}",f"Warnings: {len(warnings)}",f"Errors: {len(errors)}","",
"CHECKS","-"*72,f"All 117 months resolved and parsed: {'PASS' if parsed==required else 'FAIL'}",
f"Numeric enrollment accounting identity: {'PASS' if diff==0 else 'FAIL'}",
f"No duplicate annual Part D payment keys: {'PASS' if not dups else 'FAIL'}","",
"MEMORY DISCIPLINE","-"*72,"Source index retains metadata only.","Each selected monthly source is streamed and released before the next month.",
"NO_PARTD_PAYMENT_MATCH detail is aggregated by month/contract/PBP/type.","","WARNINGS","-"*72]
    report += warnings or ["None"]; report += ["","ERRORS","-"*72]; report += errors or ["None"]
    txt="\n".join(report)+"\n";(a.output_dir/"partd_cpsc_enrollment_linkage_profile_report.txt").write_text(txt,encoding="utf-8")
    print("\n"+txt);print(f"Outputs: {a.output_dir}")
    return 0 if parsed==required and diff==0 and not dups and not errors else 1
if __name__=="__main__":sys.exit(main())
