#!/usr/bin/env python3
"""
validate_cms_partd_payment_reconstruction_v0_1.py

CMS Medicare Part D — Read-Only Payment Reconstruction Validator v0.1

PURPOSE
-------
Reconstruct Washington Medicare Part D payment dollars from:
  * exact Contract+PBP monthly Washington CPSC numeric enrollment; and
  * same-year CMS PartDPlans payment inputs.

Certified payment semantics:
    DirectSubsidy_RA_PMPM
        = DirectSubsidy_STD_PMPM * PartD_RiskScore

    PartD_PMPM
        = DirectSubsidy_RA_PMPM + Reinsurance_PMPM + LICS_PMPM

    PartD_Dollars
        = Numeric_Enrollment * PartD_PMPM

PartDReconciliation.xlsx is CONTROL_ONLY and is NOT added to PartDPlans.

BOUNDARY
--------
* CY2015-CY2023 complete monthly source window.
* CY2024 Jan-Sep only, because the certified CPSC archive ends Sep 2024.
* Numeric enrollment only. Suppressed/non-numeric enrollment is NOT imputed.
* Exact same-year Contract+PBP payment match only.
* No sibling-PBP payment propagation.
* No other-year payment propagation.
* Known Part D enrollment with unavailable same-year payment remains explicitly
  SOURCE_PAYMENT_UNAVAILABLE; it is not assigned a zero payment.
* No PostgreSQL writes.
* No final-ledger inclusion.

This validator independently rereads the preserved CPSC and PartDPlans sources.
It does not use PartDReconciliation as an additive payment component.

Run from repository root:
    python scripts/validate_cms_partd_payment_reconstruction_v0_1.py
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]

MA_ROOT = REPO_ROOT / "data" / "raw" / "cms" / "medicare_advantage"
PLAN_PAYMENT_ROOT = MA_ROOT / "plan_payment"
PARTD_ROOT = REPO_ROOT / "data" / "raw" / "cms" / "medicare_partd"
PBP_ROOT = PARTD_ROOT / "pbp_benefits"

OUTPUT_DIR = (
    PARTD_ROOT / "validation" / "partd_payment_reconstruction_v0_1"
)

START_YEAR = 2015
END_YEAR = 2024
REQUIRED_MONTHS = [(y, m) for y in range(2015, 2024) for m in range(1, 13)]
REQUIRED_MONTHS += [(2024, m) for m in range(1, 10)]  # Jan-Sep 2024
EXPECTED_MONTH_COUNT = 117

# Certified exception gate.
AFFIRMATIVE_EXCEPTION_DIR = (
    PBP_ROOT / "validation" / "partd_affirmative_payment_absence_v0_3"
)
AFFIRMATIVE_EXCEPTION_DETAIL = (
    AFFIRMATIVE_EXCEPTION_DIR
    / "cms_partd_affirmative_payment_absence_v0_3_detail.csv"
)
EXPECTED_UNAVAILABLE_RELATIONSHIPS = 3
EXPECTED_UNAVAILABLE_MEMBER_MONTHS = 1472

CONTRACT_ALIASES = {
    "contract", "contract number", "contract_number", "contract id",
    "contract_id", "contractid", "hnumber"
}
PBP_ALIASES = {
    "plan benefit package", "plan benefit package id",
    "plan benefit package number", "plan id", "plan_id",
    "plan identifier", "plan_identifier", "pbp", "pbp id", "pbp_id"
}
ENROLLMENT_ALIASES = {
    "enrollment", "enrollment count", "enrollment_count",
    "total enrollment", "total_enrollment"
}

@dataclass(frozen=True)
class Payment:
    contract: str
    pbp: str
    organization_type: str
    direct_subsidy_std: float
    risk_score: float
    reinsurance: float
    lics: float

    @property
    def direct_subsidy_ra(self) -> float:
        return self.direct_subsidy_std * self.risk_score

    @property
    def total_pmpm(self) -> float:
        return self.direct_subsidy_ra + self.reinsurance + self.lics

def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip())

def low(v) -> str:
    return norm(v).lower()

def norm_contract(v) -> str:
    return norm(v).upper()

def norm_pbp(v) -> str:
    s = norm(v)
    if re.fullmatch(r"\d+(?:\.0+)?", s):
        return str(int(float(s))).zfill(3)
    return s.upper()

def parse_float(v) -> Optional[float]:
    if v is None:
        return None
    s = norm(v).replace(",", "").replace("$", "")
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None

def parse_int(v) -> Optional[int]:
    x = parse_float(v)
    if x is None:
        return None
    if abs(x - round(x)) > 1e-9:
        return None
    return int(round(x))

def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1", errors="replace")

def detect_delimiter(text: str) -> str:
    lines = [x for x in text.splitlines() if x.strip()]
    first = lines[0] if lines else text
    candidates = {"\t": first.count("\t"), ",": first.count(","), "|": first.count("|")}
    delim, count = max(candidates.items(), key=lambda x: x[1])
    if count <= 0:
        raise RuntimeError("Unable to detect delimiter")
    return delim

def find_col(headers: list[str], aliases: set[str]) -> Optional[str]:
    by_low = {low(h): h for h in headers if h is not None}
    for alias in aliases:
        if alias in by_low:
            return by_low[alias]
    return None

def find_contains(headers: list[str], tokens: tuple[str, ...],
                  forbidden: tuple[str, ...] = ()) -> Optional[str]:
    for h in headers:
        lh = low(h)
        if all(t in lh for t in tokens) and not any(t in lh for t in forbidden):
            return h
    return None

def infer_year_month(text: str) -> Optional[tuple[int, int]]:
    """
    Conservative date-token inference supporting YYYYMM, YYYY-MM, YYYY_MM,
    and common month-name + year forms.
    """
    s = text.lower()

    for pat in (
        r"(?<!\d)(20(?:1[5-9]|2[0-4]))[-_]?([01]\d)(?!\d)",
        r"(?<!\d)(20(?:1[5-9]|2[0-4]))[-_]?([1-9])(?!\d)",
    ):
        hits = re.findall(pat, s)
        valid = [(int(y), int(m)) for y, m in hits if 1 <= int(m) <= 12]
        if len(set(valid)) == 1:
            return valid[0]

    month_names = {
        "jan": 1, "january": 1, "feb": 2, "february": 2,
        "mar": 3, "march": 3, "apr": 4, "april": 4,
        "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
        "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
        "oct": 10, "october": 10, "nov": 11, "november": 11,
        "dec": 12, "december": 12,
    }
    years = re.findall(r"(?<!\d)(20(?:1[5-9]|2[0-4]))(?!\d)", s)
    found = []
    for name, month in month_names.items():
        if re.search(rf"(?<![a-z]){re.escape(name)}(?![a-z])", s):
            for y in years:
                found.append((int(y), month))
    if len(set(found)) == 1:
        return found[0]
    return None

def infer_year(text: str) -> Optional[int]:
    vals = {int(x) for x in re.findall(r"(?<!\d)(20(?:1[4-9]|2[0-5]))(?!\d)", text)}
    return next(iter(vals)) if len(vals) == 1 else None

# ----------------------------------------------------------------------
# Certified unavailable exceptions
# ----------------------------------------------------------------------

def load_unavailable_exceptions() -> dict[tuple[int, str, str], int]:
    if not AFFIRMATIVE_EXCEPTION_DETAIL.exists():
        raise FileNotFoundError(
            f"Certified v0.3 exception detail not found: {AFFIRMATIVE_EXCEPTION_DETAIL}"
        )
    out = {}
    with AFFIRMATIVE_EXCEPTION_DETAIL.open("r", encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            key = (
                int(norm(r["year"])),
                norm_contract(r["contract"]),
                norm_pbp(r["pbp"]),
            )
            mm = int(norm(r["member_months_v0_8"]).replace(",", ""))
            if key in out:
                raise RuntimeError(f"Duplicate certified unavailable key: {key}")
            out[key] = mm

    if len(out) != EXPECTED_UNAVAILABLE_RELATIONSHIPS:
        raise RuntimeError(
            f"Unavailable relationship control failed: {len(out)} "
            f"!= {EXPECTED_UNAVAILABLE_RELATIONSHIPS}"
        )
    if sum(out.values()) != EXPECTED_UNAVAILABLE_MEMBER_MONTHS:
        raise RuntimeError(
            f"Unavailable member-month control failed: {sum(out.values())} "
            f"!= {EXPECTED_UNAVAILABLE_MEMBER_MONTHS}"
        )
    return out

# ----------------------------------------------------------------------
# Plan Payment
# ----------------------------------------------------------------------

def discover_plan_payment_zip(year: int) -> Path:
    candidates = []
    for p in PLAN_PAYMENT_ROOT.rglob("*.zip"):
        if "validation" in {x.lower() for x in p.parts}:
            continue
        y = infer_year(p.name) or infer_year(str(p.parent))
        if y == year:
            candidates.append(p.resolve())
    candidates = sorted(set(candidates))
    if not candidates:
        raise FileNotFoundError(f"CY{year}: no Plan Payment ZIP found")

    if len(candidates) > 1:
        hashes = {}
        for p in candidates:
            h = hashlib.sha256()
            with p.open("rb") as f:
                for chunk in iter(lambda: f.read(1024 * 1024), b""):
                    h.update(chunk)
            hashes.setdefault(h.hexdigest(), []).append(p)
        if len(hashes) == 1:
            return sorted(candidates, key=lambda p: (len(str(p)), str(p)))[0]

        scored = []
        for p in candidates:
            s = p.name.lower()
            score = (4 if str(year) in s else 0) + (3 if "payment" in s else 0)
            scored.append((score, p))
        best_score = max(x[0] for x in scored)
        best = [p for score, p in scored if score == best_score]
        if len(best) != 1:
            raise RuntimeError(
                f"CY{year}: ambiguous Plan Payment ZIPs: "
                + " | ".join(map(str, candidates))
            )
        return best[0]
    return candidates[0]

def find_partdplans_member(zf: zipfile.ZipFile) -> str:
    hits = []
    for n in zf.namelist():
        if n.endswith("/"):
            continue
        b = Path(n).name.lower().replace(" ", "").replace("_", "").replace("-", "")
        if b.endswith("partdplans.xlsx"):
            hits.append(n)
    if len(hits) != 1:
        raise RuntimeError(f"Expected exactly one PartDPlans workbook; found {len(hits)}: {hits}")
    return hits[0]

def load_partdplans(year: int) -> tuple[dict[tuple[str, str], Payment], str]:
    zpath = discover_plan_payment_zip(year)
    with zipfile.ZipFile(zpath) as zf:
        member = find_partdplans_member(zf)
        data = zf.read(member)

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    records: dict[tuple[str, str], Payment] = {}
    found_sheet = False

    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)

        header = None
        contract_col = pbp_col = None
        for row_number, candidate in enumerate(rows, start=1):
            if row_number > 25:
                break
            candidate_headers = [norm(x) for x in candidate]
            cc = find_col(candidate_headers, CONTRACT_ALIASES)
            pc = find_col(candidate_headers, PBP_ALIASES)
            if cc and pc:
                header = candidate
                contract_col = cc
                pbp_col = pc
                break

        if header is None:
            continue

        headers = [norm(x) for x in header]
        idx = {h: i for i, h in enumerate(headers)}
        ci = idx[contract_col]
        pi = idx[pbp_col]

        org_col = find_contains(headers, ("organization", "type"))
        direct_col = find_contains(headers, ("direct", "subsidy"))
        risk_col = find_contains(headers, ("risk", "score"))
        reins_col = find_contains(headers, ("reinsurance",))
        lics_col = (
            find_contains(headers, ("low", "income", "cost", "sharing"))
            or find_contains(headers, ("lics",))
        )
        required = {
            "organization type": org_col,
            "direct subsidy": direct_col,
            "risk score": risk_col,
            "reinsurance": reins_col,
            "LICS": lics_col,
        }
        missing = [k for k, v in required.items() if not v]
        if missing:
            raise RuntimeError(
                f"CY{year} {ws.title}: missing required PartDPlans fields: {missing}"
            )

        found_sheet = True
        oi, di, ri, rei, li = (
            idx[org_col], idx[direct_col], idx[risk_col], idx[reins_col], idx[lics_col]
        )

        for row in rows:
            c = norm_contract(row[ci] if ci < len(row) else "")
            p = norm_pbp(row[pi] if pi < len(row) else "")
            if not c or not p:
                continue
            vals = [
                parse_float(row[di] if di < len(row) else None),
                parse_float(row[ri] if ri < len(row) else None),
                parse_float(row[rei] if rei < len(row) else None),
                parse_float(row[li] if li < len(row) else None),
            ]
            if any(v is None for v in vals):
                raise RuntimeError(
                    f"CY{year}: nonnumeric required payment input for {c}-{p}: {vals}"
                )
            key = (c, p)
            if key in records:
                raise RuntimeError(f"CY{year}: duplicate PartDPlans exact key {key}")
            records[key] = Payment(
                contract=c,
                pbp=p,
                organization_type=norm(row[oi] if oi < len(row) else ""),
                direct_subsidy_std=vals[0],
                risk_score=vals[1],
                reinsurance=vals[2],
                lics=vals[3],
            )

    wb.close()
    if not found_sheet:
        raise RuntimeError(f"CY{year}: no Contract+PBP PartDPlans sheet resolved")
    return records, f"{zpath}::{member}"

# ----------------------------------------------------------------------
# CPSC source discovery and parsing
# ----------------------------------------------------------------------

def candidate_cpsc_archives() -> list[Path]:
    """
    Search preserved MA raw archives, excluding Plan Payment, Part D validation,
    and known non-CPSC branches. Header admission below remains authoritative.
    """
    out = []
    for p in MA_ROOT.rglob("*.zip"):
        lower_parts = {x.lower() for x in p.parts}
        if "validation" in lower_parts or "plan_payment" in lower_parts:
            continue
        out.append(p)
    return sorted(set(out))

def tabular_members(zpath: Path):
    with zipfile.ZipFile(zpath) as zf:
        for n in zf.namelist():
            if n.endswith("/") or not n.lower().endswith((".csv", ".txt")):
                continue
            yield n, zf.read(n)

def inspect_header(data: bytes) -> tuple[str, list[str]]:
    text = decode_bytes(data[:256 * 1024])
    delim = detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    for row in reader:
        headers = [norm(x) for x in row]
        if find_col(headers, CONTRACT_ALIASES) and find_col(headers, PBP_ALIASES):
            ec = find_col(headers, ENROLLMENT_ALIASES) or find_contains(headers, ("enrollment",))
            if ec:
                return delim, headers
        # CPSC header should be very near the top.
        if reader.line_num >= 10:
            break
    raise RuntimeError("No CPSC-like Contract+PBP+Enrollment header")

def build_cpsc_month_index() -> tuple[dict[tuple[int, int], tuple[Path, str]], list[dict]]:
    required = set(REQUIRED_MONTHS)
    admitted = defaultdict(list)
    diagnostics = []

    for zpath in candidate_cpsc_archives():
        try:
            with zipfile.ZipFile(zpath) as zf:
                for n in zf.namelist():
                    if n.endswith("/") or not n.lower().endswith((".csv", ".txt")):
                        continue
                    ym = infer_year_month(f"{zpath} {n}")
                    if ym not in required:
                        continue
                    try:
                        sample = zf.read(n)
                        inspect_header(sample)
                    except Exception:
                        continue
                    admitted[ym].append((zpath, n))
        except Exception as exc:
            diagnostics.append({"scope": str(zpath), "message": str(exc)})

    index = {}
    for ym in REQUIRED_MONTHS:
        hits = admitted.get(ym, [])
        # Deduplicate exact byte-identical candidates when packaging duplicates exist.
        if len(hits) > 1:
            by_hash = {}
            for zpath, n in hits:
                with zipfile.ZipFile(zpath) as zf:
                    raw = zf.read(n)
                h = hashlib.sha256(raw).hexdigest()
                by_hash.setdefault(h, []).append((zpath, n))
            if len(by_hash) == 1:
                hits = [sorted(next(iter(by_hash.values())), key=lambda x: (str(x[0]), x[1]))[0]]

        if len(hits) != 1:
            diagnostics.append({
                "scope": f"CPSC {ym[0]}-{ym[1]:02d}",
                "message": f"Expected 1 admitted source member; found {len(hits)}: {hits}",
            })
            continue
        index[ym] = hits[0]
    return index, diagnostics

def read_cpsc_month(zpath: Path, member: str):
    with zipfile.ZipFile(zpath) as zf:
        raw = zf.read(member)
    text = decode_bytes(raw)
    delim = detect_delimiter(text[:256 * 1024])
    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delim)
    headers = reader.fieldnames or []
    cc = find_col(headers, CONTRACT_ALIASES)
    pc = find_col(headers, PBP_ALIASES)
    ec = find_col(headers, ENROLLMENT_ALIASES) or find_contains(headers, ("enrollment",))
    state_col = (
        find_contains(headers, ("state",))
        or find_contains(headers, ("state", "territory"))
    )
    if not cc or not pc or not ec:
        raise RuntimeError(
            f"CPSC member lacks required Contract/PBP/Enrollment fields: {member}"
        )

    for r in reader:
        # If state exists, restrict to WA. Many monthly WA extracts are already scoped.
        if state_col:
            sv = norm(r.get(state_col)).upper()
            if sv and sv not in {"WA", "WASHINGTON"}:
                continue
        yield (
            norm_contract(r.get(cc)),
            norm_pbp(r.get(pc)),
            r.get(ec),
        )

# ----------------------------------------------------------------------
# Output helpers
# ----------------------------------------------------------------------

def write_csv(path: Path, rows: list[dict], fields: list[str]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------

def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 78)
    print("CMS MEDICARE PART D PAYMENT RECONSTRUCTION VALIDATION v0.1")
    print("=" * 78)
    print()
    print("FORMULA:")
    print("  DirectSubsidy_RA_PMPM = DirectSubsidy_STD_PMPM * PartD_RiskScore")
    print("  PartD_PMPM = DirectSubsidy_RA_PMPM + Reinsurance_PMPM + LICS_PMPM")
    print("  PartD_Dollars = Numeric_Enrollment * PartD_PMPM")
    print()
    print("BOUNDARY:")
    print("  Exact same-year Contract+PBP only; numeric enrollment only.")
    print("  Suppressed enrollment and unavailable payment inputs are not imputed.")
    print("  PartDReconciliation is CONTROL_ONLY and is not added.")
    print("  CY2024 is Jan-Sep only.")
    print()

    errors = []

    try:
        unavailable = load_unavailable_exceptions()
    except Exception as exc:
        print(f"FATAL: {exc}")
        return 1

    # Load annual payment tables.
    payment_by_year = {}
    payment_sources = {}
    for year in range(START_YEAR, END_YEAR + 1):
        try:
            payment_by_year[year], payment_sources[year] = load_partdplans(year)
        except Exception as exc:
            errors.append({"scope": f"PartDPlans CY{year}", "message": str(exc)})

    # Build exact monthly CPSC source index.
    cpsc_index, cpsc_diagnostics = build_cpsc_month_index()
    errors.extend(cpsc_diagnostics)

    monthly_rows = []
    annual = defaultdict(lambda: {
        "numeric_cpsc_enrollment": 0,
        "exact_payment_match_enrollment": 0,
        "certified_payment_unavailable_enrollment": 0,
        "other_no_payment_match_enrollment": 0,
        "suppressed_or_nonnumeric_rows": 0,
        "direct_subsidy_ra_dollars": 0.0,
        "reinsurance_dollars": 0.0,
        "lics_dollars": 0.0,
        "total_partd_dollars": 0.0,
    })

    unavailable_observed = defaultdict(int)

    for year, month in REQUIRED_MONTHS:
        month_control = {
            "numeric_cpsc_enrollment": 0,
            "exact_payment_match_enrollment": 0,
            "certified_payment_unavailable_enrollment": 0,
            "other_no_payment_match_enrollment": 0,
            "suppressed_or_nonnumeric_rows": 0,
            "direct_subsidy_ra_dollars": 0.0,
            "reinsurance_dollars": 0.0,
            "lics_dollars": 0.0,
            "total_partd_dollars": 0.0,
        }

        source = cpsc_index.get((year, month))
        payments = payment_by_year.get(year)

        if source is None or payments is None:
            monthly_rows.append({
                "year": year, "month": month,
                **month_control,
                "numeric_accounting_difference": "",
                "dollar_identity_difference": "",
                "source_status": "SOURCE_UNRESOLVED",
            })
            continue

        zpath, member = source
        try:
            for contract, pbp, enrollment_raw in read_cpsc_month(zpath, member):
                enrollment = parse_int(enrollment_raw)
                if enrollment is None:
                    month_control["suppressed_or_nonnumeric_rows"] += 1
                    continue
                if enrollment < 0:
                    errors.append({
                        "scope": f"CPSC {year}-{month:02d} {contract}-{pbp}",
                        "message": f"Negative numeric enrollment: {enrollment}",
                    })
                    continue

                month_control["numeric_cpsc_enrollment"] += enrollment
                payment = payments.get((contract, pbp))

                if payment is not None:
                    month_control["exact_payment_match_enrollment"] += enrollment

                    direct_ra_dollars = enrollment * payment.direct_subsidy_ra
                    reinsurance_dollars = enrollment * payment.reinsurance
                    lics_dollars = enrollment * payment.lics
                    total_dollars = enrollment * payment.total_pmpm

                    month_control["direct_subsidy_ra_dollars"] += direct_ra_dollars
                    month_control["reinsurance_dollars"] += reinsurance_dollars
                    month_control["lics_dollars"] += lics_dollars
                    month_control["total_partd_dollars"] += total_dollars
                else:
                    ukey = (year, contract, pbp)
                    if ukey in unavailable:
                        month_control["certified_payment_unavailable_enrollment"] += enrollment
                        unavailable_observed[ukey] += enrollment
                    else:
                        month_control["other_no_payment_match_enrollment"] += enrollment
        except Exception as exc:
            errors.append({
                "scope": f"CPSC {year}-{month:02d}",
                "message": f"{zpath}::{member}: {exc}",
            })

        numeric_identity = (
            month_control["numeric_cpsc_enrollment"]
            - month_control["exact_payment_match_enrollment"]
            - month_control["certified_payment_unavailable_enrollment"]
            - month_control["other_no_payment_match_enrollment"]
        )
        dollar_identity = (
            month_control["total_partd_dollars"]
            - month_control["direct_subsidy_ra_dollars"]
            - month_control["reinsurance_dollars"]
            - month_control["lics_dollars"]
        )

        monthly_rows.append({
            "year": year, "month": month,
            **month_control,
            "numeric_accounting_difference": numeric_identity,
            "dollar_identity_difference": dollar_identity,
            "source_status": "RESOLVED",
        })

        a = annual[year]
        for k in month_control:
            a[k] += month_control[k]

    annual_rows = []
    for year in range(START_YEAR, END_YEAR + 1):
        a = annual[year]
        numeric_diff = (
            a["numeric_cpsc_enrollment"]
            - a["exact_payment_match_enrollment"]
            - a["certified_payment_unavailable_enrollment"]
            - a["other_no_payment_match_enrollment"]
        )
        dollar_diff = (
            a["total_partd_dollars"]
            - a["direct_subsidy_ra_dollars"]
            - a["reinsurance_dollars"]
            - a["lics_dollars"]
        )
        annual_rows.append({
            "year": year,
            "coverage": "JAN-SEP" if year == 2024 else "JAN-DEC",
            **a,
            "numeric_accounting_difference": numeric_diff,
            "dollar_identity_difference": dollar_diff,
        })

    total = {
        k: sum(r[k] for r in annual_rows)
        for k in (
            "numeric_cpsc_enrollment",
            "exact_payment_match_enrollment",
            "certified_payment_unavailable_enrollment",
            "other_no_payment_match_enrollment",
            "suppressed_or_nonnumeric_rows",
            "direct_subsidy_ra_dollars",
            "reinsurance_dollars",
            "lics_dollars",
            "total_partd_dollars",
        )
    }
    total_numeric_diff = (
        total["numeric_cpsc_enrollment"]
        - total["exact_payment_match_enrollment"]
        - total["certified_payment_unavailable_enrollment"]
        - total["other_no_payment_match_enrollment"]
    )
    total_dollar_diff = (
        total["total_partd_dollars"]
        - total["direct_subsidy_ra_dollars"]
        - total["reinsurance_dollars"]
        - total["lics_dollars"]
    )

    resolved_months = sum(r["source_status"] == "RESOLVED" for r in monthly_rows)

    # IMPORTANT: The 1,472 certified unavailable control applies to the
    # underlying CY2024 relationship population previously certified. This
    # reconstruction scope ends Sep 2024. We require observed Jan-Sep enrollment
    # not to exceed that control, and require every observed unavailable key to
    # be among the certified three. We do not assume all 1,472 necessarily fall
    # within Jan-Sep until source rows demonstrate it.
    unavailable_keys_ok = set(unavailable_observed).issubset(set(unavailable))
    unavailable_not_exceeded = all(
        unavailable_observed[k] <= unavailable[k] for k in unavailable_observed
    )

    checks = [
        ("Required month count is 117",
         len(REQUIRED_MONTHS) == EXPECTED_MONTH_COUNT,
         f"{len(REQUIRED_MONTHS)} == {EXPECTED_MONTH_COUNT}"),
        ("All required CPSC months resolved",
         resolved_months == EXPECTED_MONTH_COUNT,
         f"{resolved_months}/{EXPECTED_MONTH_COUNT}"),
        ("All CY2015-CY2024 PartDPlans sources resolved",
         len(payment_by_year) == 10,
         f"{len(payment_by_year)}/10"),
        ("Certified unavailable relationship control loaded",
         len(unavailable) == EXPECTED_UNAVAILABLE_RELATIONSHIPS,
         f"{len(unavailable)}/{EXPECTED_UNAVAILABLE_RELATIONSHIPS}"),
        ("Certified unavailable member-month control loaded",
         sum(unavailable.values()) == EXPECTED_UNAVAILABLE_MEMBER_MONTHS,
         f"{sum(unavailable.values()):,}/{EXPECTED_UNAVAILABLE_MEMBER_MONTHS:,}"),
        ("Observed unavailable keys are certified exceptions only",
         unavailable_keys_ok,
         f"{len(unavailable_observed)} observed keys"),
        ("Observed unavailable enrollment does not exceed certified controls",
         unavailable_not_exceeded,
         f"{sum(unavailable_observed.values()):,} observed"),
        ("Numeric enrollment accounting identity",
         total_numeric_diff == 0,
         f"difference={total_numeric_diff:,}"),
        ("No other numeric no-payment-match enrollment remains",
         total["other_no_payment_match_enrollment"] == 0,
         f"{total['other_no_payment_match_enrollment']:,}"),
        ("Dollar component identity",
         abs(total_dollar_diff) <= 0.01,
         f"difference=${total_dollar_diff:,.6f}"),
        ("No source/profile errors",
         len(errors) == 0,
         f"errors={len(errors)}"),
    ]
    overall = all(ok for _, ok, _ in checks)

    write_csv(
        OUTPUT_DIR / "cms_partd_payment_reconstruction_monthly_v0_1.csv",
        monthly_rows,
        [
            "year", "month", "numeric_cpsc_enrollment",
            "exact_payment_match_enrollment",
            "certified_payment_unavailable_enrollment",
            "other_no_payment_match_enrollment",
            "suppressed_or_nonnumeric_rows",
            "direct_subsidy_ra_dollars", "reinsurance_dollars", "lics_dollars",
            "total_partd_dollars", "numeric_accounting_difference",
            "dollar_identity_difference", "source_status",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_reconstruction_annual_v0_1.csv",
        annual_rows,
        [
            "year", "coverage", "numeric_cpsc_enrollment",
            "exact_payment_match_enrollment",
            "certified_payment_unavailable_enrollment",
            "other_no_payment_match_enrollment",
            "suppressed_or_nonnumeric_rows",
            "direct_subsidy_ra_dollars", "reinsurance_dollars", "lics_dollars",
            "total_partd_dollars", "numeric_accounting_difference",
            "dollar_identity_difference",
        ],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_reconstruction_checks_v0_1.csv",
        [{"check": n, "status": "PASS" if ok else "FAIL", "detail": d}
         for n, ok, d in checks],
        ["check", "status", "detail"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_reconstruction_errors_v0_1.csv",
        errors,
        ["scope", "message"],
    )
    write_csv(
        OUTPUT_DIR / "cms_partd_payment_unavailable_observed_v0_1.csv",
        [
            {
                "year": k[0], "contract": k[1], "pbp": k[2],
                "certified_member_month_control": unavailable[k],
                "observed_in_reconstruction_scope": unavailable_observed.get(k, 0),
                "remaining_outside_or_unobserved": unavailable[k] - unavailable_observed.get(k, 0),
            }
            for k in sorted(unavailable)
        ],
        [
            "year", "contract", "pbp", "certified_member_month_control",
            "observed_in_reconstruction_scope", "remaining_outside_or_unobserved",
        ],
    )

    print("SOURCE COVERAGE:")
    print(f"  Required months: {EXPECTED_MONTH_COUNT}")
    print(f"  Resolved months: {resolved_months}")
    print(f"  PartDPlans years resolved: {len(payment_by_year)}/10")
    print()
    print("ENROLLMENT ACCOUNTING:")
    print(f"  Numeric CPSC enrollment:                 {total['numeric_cpsc_enrollment']:,}")
    print(f"  Exact payment-match enrollment:          {total['exact_payment_match_enrollment']:,}")
    print(f"  Certified payment-unavailable enrollment:{total['certified_payment_unavailable_enrollment']:>12,}")
    print(f"  Other no-payment-match enrollment:       {total['other_no_payment_match_enrollment']:,}")
    print(f"  Numeric accounting difference:           {total_numeric_diff:,}")
    print(f"  Suppressed/non-numeric rows:             {total['suppressed_or_nonnumeric_rows']:,}")
    print()
    print("RECONSTRUCTED PAYMENT COMPONENTS:")
    print(f"  Risk-adjusted direct subsidy: ${total['direct_subsidy_ra_dollars']:,.2f}")
    print(f"  Reinsurance:                  ${total['reinsurance_dollars']:,.2f}")
    print(f"  LICS:                         ${total['lics_dollars']:,.2f}")
    print(f"  TOTAL PART D:                 ${total['total_partd_dollars']:,.2f}")
    print(f"  Dollar identity difference:  ${total_dollar_diff:,.6f}")
    print()
    print("ANNUAL:")
    for r in annual_rows:
        print(
            f"  {r['year']} {r['coverage']}: "
            f"numeric={r['numeric_cpsc_enrollment']:,} | "
            f"matched={r['exact_payment_match_enrollment']:,} | "
            f"unavailable={r['certified_payment_unavailable_enrollment']:,} | "
            f"other_unmatched={r['other_no_payment_match_enrollment']:,} | "
            f"PartD=${r['total_partd_dollars']:,.2f}"
        )
    print()
    print("CHECKS:")
    for name, ok, detail in checks:
        print(f"  [{'PASS' if ok else 'FAIL'}] {name} — {detail}")
    if errors:
        print()
        print(f"ERRORS/DIAGNOSTICS: {len(errors)}")
        for e in errors[:25]:
            print(f"  {e['scope']}: {e['message']}")
    print()
    print(f"Output directory: {OUTPUT_DIR}")
    print(f"OVERALL: {'PASS' if overall else 'FAIL'}")
    print()
    print("NOTE: PASS certifies deterministic reconstruction mechanics only.")
    print("It does not authorize database writes or final-ledger inclusion.")
    return 0 if overall else 1

if __name__ == "__main__":
    sys.exit(main())
