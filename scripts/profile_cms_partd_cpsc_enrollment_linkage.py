#!/usr/bin/env python3
"""
profile_cms_partd_cpsc_enrollment_linkage.py

Read-only profiler for Washington Medicare Part D CPSC enrollment linkage.

Scope
-----
Jan 2015 through Sep 2024 (117 months), using already-preserved CMS
Monthly Enrollment by Contract/Plan/State/County (CPSC) source files and
annual CMS PartDPlans workbooks.

Questions answered
------------------
1. Which CPSC plan/organization types appear among WA rows linked to Part D?
2. How much WA numeric enrollment links directly by Contract + PBP?
3. How much WA enrollment is unlinked or suppressed/non-numeric?
4. Which annual Part D payment rows have no observed WA CPSC enrollment
   in each month?

Boundary
--------
- READ ONLY.
- No payment reconstruction.
- No enrollment imputation.
- No database writes.
- No ledger inclusion.
- RDS remains out of scope.
- Scope intentionally ends Sep 2024 until Oct-Dec 2024 acquisition is
  separately certified.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl


START = (2015, 1)
END = (2024, 9)

PAYMENT_FIELDS = {
    "CONTRACT_NUMBER",
    "PLAN_BENEFIT_PACKAGE",
    "CONTRACT_NAME",
    "ORGANIZATION_TYPE",
    "AVERAGE_PART_D_DIRECT_SUBSIDY_PM_PM_PAYMENT",
    "AVERAGE_PART_D_RISK_SCORE",
    "AVERAGE_REINSURANCE_PM_PM_PAYMENT",
    "AVERAGE_LOW_INCOME_COST_SHARING_PM_PM_PAYMENT",
}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def norm(value: Any) -> str:
    s = clean(value).upper().replace("/", " ")
    s = re.sub(r"[^A-Z0-9]+", "_", s)
    return s.strip("_")


def canonical_pbp(value: Any) -> str:
    s = clean(value)
    if not s:
        return ""
    try:
        return f"{int(float(s)):03d}"
    except (ValueError, TypeError):
        return s.upper()


def parse_enrollment(value: Any) -> tuple[Optional[int], str]:
    if value is None:
        return None, "BLANK"
    if isinstance(value, bool):
        return None, "NONNUMERIC"
    if isinstance(value, (int, float)):
        x = float(value)
        if x >= 0 and x.is_integer():
            return int(x), "NUMERIC"
        return None, "NONNUMERIC"
    s = str(value).strip().replace(",", "")
    if not s:
        return None, "BLANK"
    if re.fullmatch(r"\d+", s):
        return int(s), "NUMERIC"
    return None, "SUPPRESSED_OR_NONNUMERIC"


def months():
    y, m = START
    while (y, m) <= END:
        yield y, m
        m += 1
        if m == 13:
            y, m = y + 1, 1


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def read_delimited(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:8192]
    delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


CPSC_REQUIRED = {
    "contract": {"CONTRACT_NUMBER", "CONTRACT_ID", "CONTRACT"},
    "pbp": {"PLAN_ID", "PLAN_BENEFIT_PACKAGE", "PBP", "PLAN_BENEFIT_PACKAGE_ID"},
    "state": {"STATE", "STATE_ABBREVIATION", "STATE_ABBR"},
    "county": {"COUNTY", "COUNTY_NAME"},
    "enrollment": {"ENROLLMENT", "ENROLLMENT_COUNT", "ENROLLEES"},
}

CPSC_OPTIONAL = {
    "plan_type": {"PLAN_TYPE", "PLAN_TYPE_DESCRIPTION"},
    "organization_type": {"ORGANIZATION_TYPE", "ORG_TYPE"},
}


def detect_cpsc_header(rows: list[list[str]], max_rows: int = 40):
    for idx, row in enumerate(rows[:max_rows]):
        normalized = [norm(x) for x in row]
        mapping = {}
        ok = True
        for logical, aliases in CPSC_REQUIRED.items():
            col = next((normalized.index(a) for a in aliases if a in normalized), None)
            if col is None:
                ok = False
                break
            mapping[logical] = col
        if not ok:
            continue
        for logical, aliases in CPSC_OPTIONAL.items():
            mapping[logical] = next(
                (normalized.index(a) for a in aliases if a in normalized), None
            )
        return idx, mapping
    return None, {}


def parse_cpsc_container(path: Path) -> tuple[list[dict[str, Any]], str]:
    members: list[tuple[str, bytes]] = []
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if Path(name).suffix.lower() in {".csv", ".txt"}:
                    members.append((name, zf.read(name)))
    elif path.suffix.lower() in {".csv", ".txt"}:
        members.append((path.name, path.read_bytes()))
    else:
        return [], ""

    for member_name, data in members:
        rows = read_delimited(data)
        header_idx, mapping = detect_cpsc_header(rows)
        if header_idx is None:
            continue

        out = []
        for source_row, row in enumerate(rows[header_idx + 1 :], header_idx + 2):
            def get(key: str):
                col = mapping.get(key)
                return row[col] if col is not None and col < len(row) else ""

            out.append(
                {
                    "member_name": member_name,
                    "source_row": source_row,
                    "contract": clean(get("contract")).upper(),
                    "pbp": canonical_pbp(get("pbp")),
                    "state": clean(get("state")).upper(),
                    "county": clean(get("county")),
                    "enrollment_raw": clean(get("enrollment")),
                    "plan_type": clean(get("plan_type")),
                    "organization_type": clean(get("organization_type")),
                }
            )
        return out, member_name

    return [], ""


def date_tokens(year: int, month: int) -> list[str]:
    mon_names = [
        "", "jan", "feb", "mar", "apr", "may", "jun",
        "jul", "aug", "sep", "oct", "nov", "dec",
    ]
    return [
        f"{year}{month:02d}",
        f"{year}_{month:02d}",
        f"{year}-{month:02d}",
        f"{month:02d}_{year}",
        f"{month:02d}-{year}",
        f"{mon_names[month]}{year}",
        f"{mon_names[month]}_{year}",
        f"{mon_names[month]}-{year}",
    ]


def locate_cpsc_candidates(root: Path, year: int, month: int) -> list[Path]:
    tokens = [x.lower() for x in date_tokens(year, month)]
    ranked = []
    for p in root.rglob("*"):
        if not p.is_file() or p.suffix.lower() not in {".zip", ".csv", ".txt"}:
            continue
        name = p.name.lower()
        if not any(t in name for t in tokens):
            continue

        compact = re.sub(r"[^a-z0-9]", "", name)
        score = 0
        if "cpsc" in compact:
            score += 200
        if "contractplanstatecounty" in compact:
            score += 200
        if "enrollment" in compact:
            score += 25
        if "scc" in compact and "cpsc" not in compact:
            score -= 100
        if "snp" in compact:
            score -= 100
        ranked.append((score, str(p).lower(), p))
    return [x[2] for x in sorted(ranked, key=lambda z: (-z[0], z[1]))]


def find_partd_workbook(plan_root: Path, year: int):
    year_dir = plan_root / str(year)
    found = []
    for zp in sorted(year_dir.glob("*.zip")):
        try:
            with zipfile.ZipFile(zp) as zf:
                for member in zf.namelist():
                    base = Path(member).name.lower().replace(" ", "")
                    if (
                        base.endswith(".xlsx")
                        and "partdplans" in base
                        and "recon" not in base
                    ):
                        data = zf.read(member)
                        found.append((zp, member, data, sha256_bytes(data)))
        except Exception:
            continue

    if not found:
        raise RuntimeError(f"{year}: PartDPlans workbook not found")

    hashes = {x[3] for x in found}
    if len(hashes) != 1:
        raise RuntimeError(
            f"{year}: conflicting PartDPlans workbook bytes across annual ZIPs"
        )

    return sorted(found, key=lambda x: (str(x[0]), x[1]))[0]


def load_partd_keys(plan_root: Path, year: int):
    zp, member, data, member_hash = find_partd_workbook(plan_root, year)
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)

    for ws in wb.worksheets:
        for header_row, row in enumerate(
            ws.iter_rows(min_row=1, max_row=30, values_only=True), 1
        ):
            header = [norm(x) for x in row]
            if not PAYMENT_FIELDS.issubset(set(header)):
                continue

            positions = {field: header.index(field) for field in PAYMENT_FIELDS}
            keys = {}
            duplicates = []

            for excel_row, values in enumerate(
                ws.iter_rows(min_row=header_row + 1, values_only=True),
                header_row + 1,
            ):
                def val(field: str):
                    idx = positions[field]
                    return values[idx] if idx < len(values) else None

                contract = clean(val("CONTRACT_NUMBER")).upper()
                pbp = canonical_pbp(val("PLAN_BENEFIT_PACKAGE"))
                if not contract or not pbp:
                    continue

                key = (contract, pbp)
                record = {
                    "contract": contract,
                    "pbp": pbp,
                    "contract_name": clean(val("CONTRACT_NAME")),
                    "payment_organization_type": clean(val("ORGANIZATION_TYPE")),
                    "sheet_name": ws.title,
                    "excel_row": excel_row,
                }
                if key in keys:
                    duplicates.append(key)
                else:
                    keys[key] = record

            return keys, duplicates, {
                "zip_path": str(zp),
                "zip_sha256": sha256_file(zp),
                "member_name": member,
                "member_sha256": member_hash,
                "sheet_name": ws.title,
                "header_row": header_row,
            }

    raise RuntimeError(f"{year}: PartDPlans semantic table not found")


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--cpsc-root",
        type=Path,
        default=Path("data/raw/cms/medicare_advantage"),
    )
    ap.add_argument(
        "--plan-payment-root",
        type=Path,
        default=Path("data/raw/cms/medicare_advantage/plan_payment"),
    )
    ap.add_argument(
        "--output-dir",
        type=Path,
        default=Path(
            "data/raw/cms/medicare_advantage/validation/"
            "partd_cpsc_enrollment_linkage_profile_v0_1"
        ),
    )
    args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    errors: list[str] = []
    annual_keys = {}
    payment_manifest = []
    payment_duplicate_rows = []

    for year in range(2015, 2025):
        try:
            keys, duplicates, meta = load_partd_keys(args.plan_payment_root, year)
            annual_keys[year] = keys
            payment_manifest.append(
                {"payment_year": year, "payment_key_count": len(keys), **meta}
            )
            for contract, pbp in duplicates:
                payment_duplicate_rows.append(
                    {"payment_year": year, "contract": contract, "pbp": pbp}
                )
        except Exception as exc:
            annual_keys[year] = {}
            errors.append(str(exc))

    monthly_rows = []
    source_manifest = []
    vocabulary = Counter()
    unmatched_payment_rows = []
    unlinked_cpsc_rows = []

    for year, month in months():
        candidates = locate_cpsc_candidates(args.cpsc_root, year, month)
        selected = None
        parsed_rows = []
        selected_member = ""

        for candidate in candidates:
            try:
                trial, member = parse_cpsc_container(candidate)
            except Exception:
                continue
            if trial:
                selected = candidate
                parsed_rows = trial
                selected_member = member
                break

        if selected is None:
            errors.append(f"{year}-{month:02d}: no parseable CPSC source found")
            monthly_rows.append(
                {
                    "year": year,
                    "month": month,
                    "source_status": "MISSING_OR_UNPARSED",
                    "wa_rows": 0,
                    "numeric_enrollment": 0,
                    "linked_numeric_enrollment": 0,
                    "unlinked_numeric_enrollment": 0,
                    "linked_suppressed_or_nonnumeric_rows": 0,
                    "unlinked_suppressed_or_nonnumeric_rows": 0,
                    "accounted_numeric_enrollment": 0,
                    "numeric_accounting_difference": 0,
                }
            )
            continue

        source_manifest.append(
            {
                "year": year,
                "month": month,
                "selected_path": str(selected),
                "selected_sha256": sha256_file(selected),
                "selected_member": selected_member,
                "candidate_count": len(candidates),
            }
        )

        wa = [r for r in parsed_rows if r["state"] in {"WA", "WASHINGTON"}]
        keys = annual_keys.get(year, {})
        observed_linked_keys = set()

        numeric_total = 0
        linked_numeric = 0
        unlinked_numeric = 0
        linked_nonnum = 0
        unlinked_nonnum = 0

        for r in wa:
            enrollment, enrollment_status = parse_enrollment(r["enrollment_raw"])
            key = (r["contract"], r["pbp"])
            linked = key in keys

            if linked:
                observed_linked_keys.add(key)
                p = keys[key]
                vocabulary[
                    (
                        year,
                        r["plan_type"],
                        r["organization_type"],
                        p["payment_organization_type"],
                    )
                ] += enrollment or 0

            if enrollment_status == "NUMERIC":
                numeric_total += enrollment or 0
                if linked:
                    linked_numeric += enrollment or 0
                else:
                    unlinked_numeric += enrollment or 0
                    unlinked_cpsc_rows.append(
                        {
                            "year": year,
                            "month": month,
                            "contract": r["contract"],
                            "pbp": r["pbp"],
                            "county": r["county"],
                            "enrollment": enrollment,
                            "plan_type": r["plan_type"],
                            "organization_type": r["organization_type"],
                        }
                    )
            else:
                if linked:
                    linked_nonnum += 1
                else:
                    unlinked_nonnum += 1

        for key, p in keys.items():
            if key not in observed_linked_keys:
                unmatched_payment_rows.append(
                    {
                        "year": year,
                        "month": month,
                        "contract": key[0],
                        "pbp": key[1],
                        "contract_name": p["contract_name"],
                        "payment_organization_type": p[
                            "payment_organization_type"
                        ],
                    }
                )

        accounted = linked_numeric + unlinked_numeric
        monthly_rows.append(
            {
                "year": year,
                "month": month,
                "source_status": "PARSED",
                "wa_rows": len(wa),
                "numeric_enrollment": numeric_total,
                "linked_numeric_enrollment": linked_numeric,
                "unlinked_numeric_enrollment": unlinked_numeric,
                "linked_suppressed_or_nonnumeric_rows": linked_nonnum,
                "unlinked_suppressed_or_nonnumeric_rows": unlinked_nonnum,
                "accounted_numeric_enrollment": accounted,
                "numeric_accounting_difference": numeric_total - accounted,
            }
        )

    vocab_rows = [
        {
            "year": year,
            "cpsc_plan_type": plan_type,
            "cpsc_organization_type": cpsc_org,
            "payment_organization_type": payment_org,
            "linked_numeric_enrollment": enrollment,
        }
        for (year, plan_type, cpsc_org, payment_org), enrollment
        in sorted(vocabulary.items())
    ]

    write_csv(
        args.output_dir / "partd_cpsc_monthly_linkage_profile.csv",
        monthly_rows,
        [
            "year", "month", "source_status", "wa_rows",
            "numeric_enrollment", "linked_numeric_enrollment",
            "unlinked_numeric_enrollment",
            "linked_suppressed_or_nonnumeric_rows",
            "unlinked_suppressed_or_nonnumeric_rows",
            "accounted_numeric_enrollment", "numeric_accounting_difference",
        ],
    )
    write_csv(
        args.output_dir / "partd_cpsc_linked_plan_type_vocabulary.csv",
        vocab_rows,
        [
            "year", "cpsc_plan_type", "cpsc_organization_type",
            "payment_organization_type", "linked_numeric_enrollment",
        ],
    )
    write_csv(
        args.output_dir / "partd_cpsc_unlinked_numeric_rows.csv",
        unlinked_cpsc_rows,
        [
            "year", "month", "contract", "pbp", "county", "enrollment",
            "plan_type", "organization_type",
        ],
    )
    write_csv(
        args.output_dir / "partd_payment_rows_without_wa_monthly_enrollment.csv",
        unmatched_payment_rows,
        [
            "year", "month", "contract", "pbp", "contract_name",
            "payment_organization_type",
        ],
    )
    write_csv(
        args.output_dir / "partd_cpsc_source_manifest.csv",
        source_manifest,
        [
            "year", "month", "selected_path", "selected_sha256",
            "selected_member", "candidate_count",
        ],
    )
    write_csv(
        args.output_dir / "partd_payment_source_manifest.csv",
        payment_manifest,
        [
            "payment_year", "payment_key_count", "zip_path", "zip_sha256",
            "member_name", "member_sha256", "sheet_name", "header_row",
        ],
    )
    write_csv(
        args.output_dir / "partd_payment_duplicate_keys.csv",
        payment_duplicate_rows,
        ["payment_year", "contract", "pbp"],
    )

    required_months = 117
    parsed_months = sum(r["source_status"] == "PARSED" for r in monthly_rows)
    numeric = sum(int(r["numeric_enrollment"]) for r in monthly_rows)
    linked = sum(int(r["linked_numeric_enrollment"]) for r in monthly_rows)
    unlinked = sum(int(r["unlinked_numeric_enrollment"]) for r in monthly_rows)
    accounting_diff = sum(int(r["numeric_accounting_difference"]) for r in monthly_rows)
    linked_nonnum = sum(
        int(r["linked_suppressed_or_nonnumeric_rows"]) for r in monthly_rows
    )
    unlinked_nonnum = sum(
        int(r["unlinked_suppressed_or_nonnumeric_rows"]) for r in monthly_rows
    )

    report = [
        "CMS MEDICARE PART D CPSC ENROLLMENT LINKAGE PROFILE",
        "=" * 72,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        "Scope: Washington, Jan 2015-Sep 2024 (117 months)",
        "",
        "BOUNDARY",
        "-" * 72,
        "Read-only enrollment/linkage profiling.",
        "No payment reconstruction, imputation, database writes, or ledger inclusion.",
        "RDS remains out of scope.",
        "",
        "COUNTS",
        "-" * 72,
        f"Required months: {required_months}",
        f"Parsed months: {parsed_months}",
        f"WA numeric enrollment across monthly county rows: {numeric:,}",
        f"Directly linked numeric enrollment: {linked:,}",
        f"Unlinked numeric enrollment: {unlinked:,}",
        f"Linked suppressed/nonnumeric rows: {linked_nonnum:,}",
        f"Unlinked suppressed/nonnumeric rows: {unlinked_nonnum:,}",
        f"Numeric enrollment accounting difference: {accounting_diff:,}",
        f"Payment-plan/month rows without observed WA linked enrollment: "
        f"{len(unmatched_payment_rows):,}",
        f"Duplicate annual Part D payment keys: {len(payment_duplicate_rows):,}",
        f"Errors: {len(errors)}",
        "",
        "CHECKS",
        "-" * 72,
        f"All 117 months parsed: "
        f"{'PASS' if parsed_months == required_months else 'FAIL'}",
        f"Numeric enrollment accounting identity: "
        f"{'PASS' if accounting_diff == 0 else 'FAIL'}",
        f"No duplicate Part D payment keys: "
        f"{'PASS' if not payment_duplicate_rows else 'FAIL'}",
        f"No source/profile errors: {'PASS' if not errors else 'FAIL'}",
        "",
        "INTERPRETATION",
        "-" * 72,
        "This is a PROFILE, not a reconstruction certification.",
        "Unlinked enrollment is evidence to investigate, not a reason to broaden",
        "the linkage or impute payment fields.",
        "A clean source profile authorizes detailed linkage validation next; it",
        "does not authorize payment multiplication or ledger inclusion.",
        "",
        "ERRORS",
        "-" * 72,
    ]
    report += errors if errors else ["None"]
    report_text = "\n".join(report) + "\n"
    (args.output_dir / "partd_cpsc_enrollment_linkage_profile_report.txt").write_text(
        report_text, encoding="utf-8"
    )

    print(report_text)
    print(f"Outputs: {args.output_dir}")

    clean_source = (
        parsed_months == required_months
        and accounting_diff == 0
        and not payment_duplicate_rows
        and not errors
    )
    return 0 if clean_source else 1


if __name__ == "__main__":
    sys.exit(main())
