#!/usr/bin/env python3
"""
profile_cms_partd_landscape_v0_3.py

Recursive structural profiler for the official CMS CY2006-CY2025 Landscape
archive. Analytical target: CY2015-CY2024.

Why v0.3 exists
---------------
v0.1 assumed target-year tabular files were top-level members.
v0.2 recognized one nesting level but inferred coverage year from full CMS
paths and did not recurse through the additional ZIP layer used in CY2015-23.
Both failures are retained as provenance.

v0.3 uses a deterministic coverage-year rule:
  - identify the target annual OUTER package from the LEAF outer ZIP name;
  - recurse through ZIP-within-ZIP packaging until tabular files are reached;
  - assign all descendants the coverage year of that annual outer package.

Boundary
--------
Structural profiling only. No field or row is interpreted as proving Part D
coverage. No imputation, payment propagation, reconstruction, DB writes, or
ledger inclusion.
"""

from __future__ import annotations
import argparse, csv, hashlib, io, re, sys, zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

DEFAULT_ARCHIVE = Path(
    "data/raw/cms/medicare_partd/landscape/cy2006-cy2025-landscape-files.zip"
)
DEFAULT_OUT = Path(
    "data/raw/cms/medicare_partd/landscape/validation/landscape_profile_v0_3"
)
TARGET_YEARS = tuple(range(2015, 2025))
TABULAR = {".csv", ".txt", ".tsv", ".xlsx", ".xlsm"}
MAX_DEPTH = 5

def clean(v: Any) -> str:
    return "" if v is None else str(v).strip()

def norm(v: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", clean(v).upper()).strip("_")

def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for c in iter(lambda: f.read(1024 * 1024), b""):
            h.update(c)
    return h.hexdigest()

def coverage_year_from_outer_leaf(name: str):
    leaf = Path(name).name
    # Annual target packages begin with 2015..., 2016..., ... or CY2022...
    m = re.match(r"^(?:CY)?(20(?:1[5-9]|2[0-4]))(?:\D|$)", leaf, re.I)
    return int(m.group(1)) if m else None

def candidate_field(s: str) -> bool:
    n = norm(s)
    return any(t in n for t in (
        "CONTRACT","PLAN","PBP","DRUG","PART_D","PARTD","PDP","MA_PD","MAPD",
        "PRESCRIPTION","RX","PLAN_TYPE","PLAN_NAME","ORGANIZATION"
    ))

def decode_delimited(data: bytes):
    text = enc = None
    for e in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text, enc = data.decode(e), e
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise RuntimeError("unable to decode delimited file")
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",|\t;")
    except csv.Error:
        dialect = csv.excel
    r = csv.reader(io.StringIO(text), dialect)
    try:
        header = [clean(x) for x in next(r)]
    except StopIteration:
        return enc, [], 0
    count = sum(1 for row in r if any(clean(x) for x in row))
    return enc, header, count

def workbook_profiles(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            best = None
            for rn, vals in enumerate(
                ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 50), values_only=True), 1
            ):
                vals = list(vals)
                pop = sum(bool(clean(x)) for x in vals)
                textual = sum(bool(clean(x)) and not clean(x).replace(".","",1).isdigit()
                              for x in vals)
                score = (pop, textual)
                if best is None or score > best[0]:
                    best = (score, rn, [clean(x) for x in vals])
            if best is None:
                out.append((ws.title, "", [], 0))
                continue
            _, hr, header = best
            while header and not header[-1]:
                header.pop()
            n = sum(1 for vals in ws.iter_rows(min_row=hr+1, values_only=True)
                    if any(clean(x) for x in vals))
            out.append((ws.title, hr, header, n))
    finally:
        wb.close()
    return out

def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w=csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--archive", type=Path, default=DEFAULT_ARCHIVE)
    ap.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    a=ap.parse_args()
    a.output_dir.mkdir(parents=True, exist_ok=True)

    if not a.archive.exists() or not zipfile.is_zipfile(a.archive):
        print(f"ERROR: archive missing/not ZIP: {a.archive}", file=sys.stderr); return 2

    outer_rows=[]; member_rows=[]; schemas=[]; fields=[]; errors=[]
    annual_count=Counter(); table_count=Counter()
    ahash=sha256_file(a.archive)

    def recurse_zip(data: bytes, year: int, outer_name: str, chain: str, depth: int):
        if depth > MAX_DEPTH:
            errors.append({"year":year,"outer_member":outer_name,"container_chain":chain,
                           "member_name":"","error":f"maximum ZIP recursion depth {MAX_DEPTH} exceeded"})
            return
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                bad=z.testzip()
                if bad:
                    raise RuntimeError(f"ZIP integrity failure at {bad}")
                for info in z.infolist():
                    if info.is_dir(): continue
                    name=info.filename
                    suffix=Path(name).suffix.lower()
                    path=f"{chain} -> {name}"
                    row={"year":year,"outer_member":outer_name,"depth":depth,
                         "container_chain":chain,"member_name":name,"suffix":suffix,
                         "bytes":info.file_size,"crc":info.CRC,
                         "member_kind":"ZIP" if suffix==".zip" else
                                       ("TABULAR" if suffix in TABULAR else "OTHER"),
                         "profile_status":"INVENTORIED"}
                    member_rows.append(row)
                    try:
                        b=z.read(info)
                        if suffix==".zip":
                            row["profile_status"]="RECURSED"
                            recurse_zip(b,year,outer_name,path,depth+1)
                        elif suffix in TABULAR:
                            h=sha256_bytes(b)
                            if suffix in {".csv",".txt",".tsv"}:
                                enc,header,nrows=decode_delimited(b)
                                profiles=[("",1,header,nrows,enc)]
                            else:
                                profiles=[(s,hr,hd,n,"") for s,hr,hd,n in workbook_profiles(b)]
                            for sheet,hr,header,nrows,enc in profiles:
                                schemas.append({
                                    "year":year,"outer_member":outer_name,
                                    "container_chain":chain,"member_name":name,
                                    "member_sha256":h,"sheet_name":sheet,
                                    "header_row":hr,"encoding":enc,"data_rows":nrows,
                                    "column_count":len(header),"columns":" | ".join(header)
                                })
                                for fld in header:
                                    if candidate_field(fld):
                                        fields.append({
                                            "year":year,"outer_member":outer_name,
                                            "container_chain":chain,"member_name":name,
                                            "sheet_name":sheet,"field_name":fld,
                                            "normalized_field":norm(fld)
                                        })
                                table_count[year]+=1
                            row["profile_status"]="PROFILED"
                        del b
                    except Exception as e:
                        row["profile_status"]="ERROR"
                        errors.append({"year":year,"outer_member":outer_name,
                                       "container_chain":chain,"member_name":name,
                                       "error":f"{type(e).__name__}: {e}"})
        except Exception as e:
            errors.append({"year":year,"outer_member":outer_name,
                           "container_chain":chain,"member_name":"",
                           "error":f"{type(e).__name__}: {e}"})

    with zipfile.ZipFile(a.archive) as outer:
        outer_bad=outer.testzip()
        for info in outer.infolist():
            if info.is_dir(): continue
            y=coverage_year_from_outer_leaf(info.filename)
            row={"outer_member":info.filename,"leaf_name":Path(info.filename).name,
                 "coverage_year":y or "","bytes":info.file_size,"crc":info.CRC,
                 "target_annual_package":"YES" if y in TARGET_YEARS else "NO",
                 "status":"INVENTORIED"}
            outer_rows.append(row)
            if y not in TARGET_YEARS:
                continue
            annual_count[y]+=1
            if Path(info.filename).suffix.lower()!=".zip":
                row["status"]="ERROR"
                errors.append({"year":y,"outer_member":info.filename,
                               "container_chain":"","member_name":"",
                               "error":"annual target package is not ZIP"})
                continue
            try:
                b=outer.read(info)
                if not zipfile.is_zipfile(io.BytesIO(b)):
                    raise RuntimeError("annual target package is not valid ZIP")
                row["status"]="RECURSED"
                recurse_zip(b,y,info.filename,Path(info.filename).name,1)
                del b
            except Exception as e:
                row["status"]="ERROR"
                errors.append({"year":y,"outer_member":info.filename,
                               "container_chain":"","member_name":"",
                               "error":f"{type(e).__name__}: {e}"})

    checks=[
        ("outer archive ZIP integrity", outer_bad is None),
        ("exactly one annual outer package for each CY2015-CY2024",
         all(annual_count[y]==1 for y in TARGET_YEARS)),
        ("all CY2015-CY2024 reach >=1 tabular schema/sheet",
         all(table_count[y]>=1 for y in TARGET_YEARS)),
        ("no recursive/profile errors", not errors),
        ("candidate Contract/Plan/PBP/Part-D-like fields observed", bool(fields)),
    ]
    overall=all(x for _,x in checks)

    write_csv(a.output_dir/"cms_partd_landscape_outer_annual_inventory.csv",outer_rows,
              ["outer_member","leaf_name","coverage_year","bytes","crc","target_annual_package","status"])
    write_csv(a.output_dir/"cms_partd_landscape_recursive_member_inventory.csv",member_rows,
              ["year","outer_member","depth","container_chain","member_name","suffix","bytes","crc",
               "member_kind","profile_status"])
    write_csv(a.output_dir/"cms_partd_landscape_schema_profile.csv",schemas,
              ["year","outer_member","container_chain","member_name","member_sha256","sheet_name",
               "header_row","encoding","data_rows","column_count","columns"])
    write_csv(a.output_dir/"cms_partd_landscape_candidate_fields.csv",fields,
              ["year","outer_member","container_chain","member_name","sheet_name",
               "field_name","normalized_field"])
    write_csv(a.output_dir/"cms_partd_landscape_profile_errors.csv",errors,
              ["year","outer_member","container_chain","member_name","error"])

    lines=[
        "CMS MEDICARE PART D LANDSCAPE STRUCTURAL PROFILE v0.3",
        "="*76,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        f"Archive: {a.archive}",f"Archive SHA-256: {ahash}","",
        "BOUNDARY","-"*76,
        "Recursive structural profiling only. Coverage year is assigned from the",
        "leaf name of each annual outer CMS package, then inherited by descendants.",
        "No Part D status semantics, imputation, payment propagation, reconstruction,",
        "database writes, or ledger inclusion.","",
        "INVENTORY","-"*76,
        f"Outer file members: {len(outer_rows):,}",
        f"Recursive descendant members in target packages: {len(member_rows):,}",
        f"Profiled tabular schemas/sheets: {len(schemas):,}",
        f"Candidate semantic fields: {len(fields):,}",
        f"Errors: {len(errors):,}","",
        "TARGET-YEAR COVERAGE","-"*76
    ]
    for y in TARGET_YEARS:
        lines.append(f"CY{y}: annual packages={annual_count[y]}, profiled tabular schemas/sheets={table_count[y]}")
    lines+=["","CHECKS","-"*76]+[f"{n}: {'PASS' if ok else 'FAIL'}" for n,ok in checks]
    lines+=["","ERRORS","-"*76]
    lines += ["None"] if not errors else [
        f"CY{e['year']} | {e['outer_member']} | {e['container_chain']} | "
        f"{e['member_name']} | {e['error']}" for e in errors
    ]
    lines+=["","INTERPRETIVE BOUNDARY","-"*76,
            "A PASS certifies only that the official CMS annual Landscape packages",
            "for CY2015-CY2024 are deterministically identified and recursively readable.",
            "The candidate fields must be reviewed before a Part D status rule is defined.",
            "","OVERALL","-"*76,f"OVERALL: {'PASS' if overall else 'FAIL'}"]
    report="\n".join(lines)+"\n"
    (a.output_dir/"cms_partd_landscape_profile_report.txt").write_text(report,encoding="utf-8")
    print(report); print(f"Outputs: {a.output_dir}")
    return 0 if overall else 1

if __name__=="__main__":
    sys.exit(main())
