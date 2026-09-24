#!/usr/bin/env python3
"""
profile_cms_partd_landscape_v0_2.py

Nested-archive structural profiler for the official CMS CY2006-CY2025
Landscape archive. Analytical target: CY2015-CY2024.

v0.2 change from v0.1
---------------------
CMS packages each target coverage year as a ZIP nested inside the outer
CY2006-CY2025 archive. v0.1 correctly failed because it only admitted
top-level tabular members. v0.2 profiles the actual nested packaging.

Boundary
--------
Structural profiling only. This script does NOT decide whether any field or
row means "offers Part D"; it only inventories nested files, tabular schemas,
and candidate fields. No imputation, payment propagation, reconstruction,
database writes, or ledger inclusion.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

DEFAULT_ARCHIVE = Path(
    "data/raw/cms/medicare_partd/landscape/cy2006-cy2025-landscape-files.zip"
)
DEFAULT_OUT = Path(
    "data/raw/cms/medicare_partd/landscape/validation/landscape_profile_v0_2"
)
TARGET_YEARS = tuple(range(2015, 2025))
TABULAR_SUFFIXES = {".csv", ".txt", ".tsv", ".xlsx", ".xlsm"}

def clean(v: Any) -> str:
    return "" if v is None else str(v).strip()

def norm(v: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", clean(v).upper()).strip("_")

def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def infer_target_year(name: str):
    years = [
        int(x)
        for x in re.findall(r"(?<!\d)(20(?:1[5-9]|2[0-4]))(?!\d)", name)
    ]
    return years[0] if years else None

def is_candidate_field(name: str) -> bool:
    n = norm(name)
    tokens = (
        "CONTRACT", "PLAN", "PBP", "DRUG", "PART_D", "PARTD", "PDP",
        "MA_PD", "MAPD", "PRESCRIPTION", "RX", "PLAN_TYPE", "PLAN_NAME"
    )
    return any(t in n for t in tokens)

def decode_delimited(data: bytes):
    text = None
    encoding = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            encoding = enc
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise RuntimeError("Unable to decode delimited file.")

    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",|\t;")
    except csv.Error:
        # Landscape text files have historically varied; fall back to comma.
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    try:
        header = next(reader)
    except StopIteration:
        return encoding, [], 0
    rows = sum(1 for row in reader if any(clean(x) for x in row))
    return encoding, [clean(x) for x in header], rows

def workbook_profiles(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    results = []
    try:
        for ws in wb.worksheets:
            best = None
            max_scan = min(ws.max_row or 1, 50)
            for rno, vals in enumerate(
                ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), 1
            ):
                vals = list(vals)
                populated = sum(bool(clean(x)) for x in vals)
                textual = sum(
                    bool(clean(x)) and not clean(x).replace(".", "", 1).isdigit()
                    for x in vals
                )
                score = (populated, textual)
                if best is None or score > best[0]:
                    best = (score, rno, [clean(x) for x in vals])

            if best is None:
                results.append((ws.title, "", [], 0))
                continue

            _, header_row, header = best
            while header and not header[-1]:
                header.pop()

            data_rows = 0
            for vals in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if any(clean(x) for x in vals):
                    data_rows += 1
            results.append((ws.title, header_row, header, data_rows))
    finally:
        wb.close()
    return results

def write_csv(path: Path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--archive", type=Path, default=DEFAULT_ARCHIVE)
    ap.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    if not args.archive.exists():
        print(f"ERROR: archive not found: {args.archive}", file=sys.stderr)
        return 2
    if not zipfile.is_zipfile(args.archive):
        print(f"ERROR: not a ZIP archive: {args.archive}", file=sys.stderr)
        return 2

    outer_inventory = []
    nested_inventory = []
    schemas = []
    candidate_fields = []
    errors = []
    year_nested_zip_count = Counter()
    year_profiled_tables = Counter()
    archive_hash = sha256_file(args.archive)

    with zipfile.ZipFile(args.archive) as outer:
        outer_bad = outer.testzip()

        for outer_info in outer.infolist():
            if outer_info.is_dir():
                continue

            outer_name = outer_info.filename
            outer_suffix = Path(outer_name).suffix.lower()
            year = infer_target_year(outer_name)
            outer_row = {
                "outer_member_name": outer_name,
                "inferred_year": year or "",
                "suffix": outer_suffix,
                "bytes": outer_info.file_size,
                "crc": outer_info.CRC,
                "target_year": "YES" if year in TARGET_YEARS else "NO",
                "nested_zip_status": "NOT_TARGET",
            }
            outer_inventory.append(outer_row)

            if year not in TARGET_YEARS:
                continue
            if outer_suffix != ".zip":
                outer_row["nested_zip_status"] = "TARGET_NOT_ZIP"
                errors.append({
                    "year": year,
                    "outer_member_name": outer_name,
                    "inner_member_name": "",
                    "error": "Target-year outer member is not ZIP."
                })
                continue

            try:
                nested_bytes = outer.read(outer_info)
                if not zipfile.is_zipfile(io.BytesIO(nested_bytes)):
                    raise RuntimeError("Outer target member is not a valid nested ZIP.")

                with zipfile.ZipFile(io.BytesIO(nested_bytes)) as nested:
                    nested_bad = nested.testzip()
                    if nested_bad is not None:
                        raise RuntimeError(
                            f"Nested ZIP integrity failure at member: {nested_bad}"
                        )

                    year_nested_zip_count[year] += 1
                    outer_row["nested_zip_status"] = "PROFILED_VALID_ZIP"

                    for inner_info in nested.infolist():
                        if inner_info.is_dir():
                            continue

                        inner_name = inner_info.filename
                        suffix = Path(inner_name).suffix.lower()
                        nrow = {
                            "year": year,
                            "outer_member_name": outer_name,
                            "outer_member_sha256": sha256_bytes(nested_bytes),
                            "inner_member_name": inner_name,
                            "suffix": suffix,
                            "bytes": inner_info.file_size,
                            "crc": inner_info.CRC,
                            "tabular_candidate": "YES" if suffix in TABULAR_SUFFIXES else "NO",
                            "profile_status": "INVENTORIED",
                        }
                        nested_inventory.append(nrow)

                        if suffix not in TABULAR_SUFFIXES:
                            continue

                        try:
                            data = nested.read(inner_info)
                            member_hash = sha256_bytes(data)

                            if suffix in {".csv", ".txt", ".tsv"}:
                                encoding, header, data_rows = decode_delimited(data)
                                schemas.append({
                                    "year": year,
                                    "outer_member_name": outer_name,
                                    "inner_member_name": inner_name,
                                    "inner_member_sha256": member_hash,
                                    "sheet_name": "",
                                    "header_row": 1,
                                    "encoding": encoding,
                                    "data_rows": data_rows,
                                    "column_count": len(header),
                                    "columns": " | ".join(header),
                                })
                                for field in header:
                                    if is_candidate_field(field):
                                        candidate_fields.append({
                                            "year": year,
                                            "outer_member_name": outer_name,
                                            "inner_member_name": inner_name,
                                            "sheet_name": "",
                                            "field_name": field,
                                            "normalized_field": norm(field),
                                        })
                                year_profiled_tables[year] += 1

                            else:
                                profiles = workbook_profiles(data)
                                for sheet, header_row, header, data_rows in profiles:
                                    schemas.append({
                                        "year": year,
                                        "outer_member_name": outer_name,
                                        "inner_member_name": inner_name,
                                        "inner_member_sha256": member_hash,
                                        "sheet_name": sheet,
                                        "header_row": header_row,
                                        "encoding": "",
                                        "data_rows": data_rows,
                                        "column_count": len(header),
                                        "columns": " | ".join(header),
                                    })
                                    for field in header:
                                        if is_candidate_field(field):
                                            candidate_fields.append({
                                                "year": year,
                                                "outer_member_name": outer_name,
                                                "inner_member_name": inner_name,
                                                "sheet_name": sheet,
                                                "field_name": field,
                                                "normalized_field": norm(field),
                                            })
                                    year_profiled_tables[year] += 1

                            nrow["profile_status"] = "PROFILED"
                            del data

                        except Exception as exc:
                            nrow["profile_status"] = "ERROR"
                            errors.append({
                                "year": year,
                                "outer_member_name": outer_name,
                                "inner_member_name": inner_name,
                                "error": f"{type(exc).__name__}: {exc}",
                            })

                del nested_bytes

            except Exception as exc:
                outer_row["nested_zip_status"] = "ERROR"
                errors.append({
                    "year": year,
                    "outer_member_name": outer_name,
                    "inner_member_name": "",
                    "error": f"{type(exc).__name__}: {exc}",
                })

    missing_nested = [y for y in TARGET_YEARS if year_nested_zip_count[y] == 0]
    missing_tables = [y for y in TARGET_YEARS if year_profiled_tables[y] == 0]

    checks = [
        ("outer archive ZIP integrity", outer_bad is None),
        ("exactly one target nested ZIP found for each CY2015-CY2024",
         all(year_nested_zip_count[y] == 1 for y in TARGET_YEARS)),
        ("all CY2015-CY2024 have >=1 profiled tabular member/sheet",
         not missing_tables),
        ("no target nested/profile errors", not errors),
        ("candidate Contract/Plan/PBP/Part-D-like fields observed",
         bool(candidate_fields)),
    ]
    overall = all(ok for _, ok in checks)

    write_csv(
        args.output_dir / "cms_partd_landscape_outer_inventory.csv",
        outer_inventory,
        ["outer_member_name","inferred_year","suffix","bytes","crc",
         "target_year","nested_zip_status"],
    )
    write_csv(
        args.output_dir / "cms_partd_landscape_nested_member_inventory.csv",
        nested_inventory,
        ["year","outer_member_name","outer_member_sha256","inner_member_name",
         "suffix","bytes","crc","tabular_candidate","profile_status"],
    )
    write_csv(
        args.output_dir / "cms_partd_landscape_schema_profile.csv",
        schemas,
        ["year","outer_member_name","inner_member_name","inner_member_sha256",
         "sheet_name","header_row","encoding","data_rows","column_count","columns"],
    )
    write_csv(
        args.output_dir / "cms_partd_landscape_candidate_fields.csv",
        candidate_fields,
        ["year","outer_member_name","inner_member_name","sheet_name",
         "field_name","normalized_field"],
    )
    write_csv(
        args.output_dir / "cms_partd_landscape_profile_errors.csv",
        errors,
        ["year","outer_member_name","inner_member_name","error"],
    )

    lines = [
        "CMS MEDICARE PART D LANDSCAPE STRUCTURAL PROFILE v0.2",
        "=" * 76,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        f"Archive: {args.archive}",
        f"Archive SHA-256: {archive_hash}",
        "",
        "BOUNDARY",
        "-" * 76,
        "Structural profiling only. CMS target years are nested ZIPs.",
        "No Part D status semantics are assigned here.",
        "No imputation, payment propagation, reconstruction, DB writes, or ledger inclusion.",
        "",
        "INVENTORY",
        "-" * 76,
        f"Outer file members: {len(outer_inventory):,}",
        f"Nested file members in target years: {len(nested_inventory):,}",
        f"Profiled target-year tabular schemas/sheets: {len(schemas):,}",
        f"Candidate semantic fields: {len(candidate_fields):,}",
        f"Profile errors: {len(errors):,}",
        "",
        "TARGET-YEAR COVERAGE",
        "-" * 76,
    ]
    for y in TARGET_YEARS:
        lines.append(
            f"CY{y}: nested ZIPs={year_nested_zip_count[y]}, "
            f"profiled tabular member/sheets={year_profiled_tables[y]}"
        )

    lines += ["", "CHECKS", "-" * 76]
    lines += [f"{label}: {'PASS' if ok else 'FAIL'}" for label, ok in checks]
    lines += ["", "ERRORS", "-" * 76]
    if errors:
        lines += [
            f"CY{e['year']} | {e['outer_member_name']} | "
            f"{e['inner_member_name']} | {e['error']}"
            for e in errors
        ]
    else:
        lines.append("None")

    lines += [
        "",
        "INTERPRETIVE BOUNDARY",
        "-" * 76,
        "A structural PASS means the preserved CMS Landscape archive can be read",
        "across CY2015-CY2024 and exposes candidate plan-level fields. It does NOT",
        "establish which rows or fields prove Part D coverage. That requires the",
        "next semantic-status gate.",
        "",
        "OVERALL",
        "-" * 76,
        f"OVERALL: {'PASS' if overall else 'FAIL'}",
    ]

    report = "\n".join(lines) + "\n"
    (args.output_dir / "cms_partd_landscape_profile_report.txt").write_text(
        report, encoding="utf-8"
    )

    print(report)
    print(f"Outputs: {args.output_dir}")
    return 0 if overall else 1

if __name__ == "__main__":
    sys.exit(main())
